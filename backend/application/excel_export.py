from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from database.models.quant_run import QuantRankResult
from database.models.validation import ModelValidationAllocation, ModelValidationOrderPlan, ModelValidationSample, ProCandidateReview, ProResumeRun
from market_review.excel import add_market_review_sheet
from market_review.repository import MarketReviewRepository
from stock_codes import display_stock_code
from entry_timing.excel import latest_entry_timing_rows


class DailyExcelExportService:
    def __init__(self, session, output_root: Path) -> None:
        self.session = session
        self.output_root = output_root

    def export(self, trade_date: date, flash_run_id: str, pro_run_id: str) -> dict[str, Any]:
        pro = self.session.scalar(select(ProResumeRun).where(ProResumeRun.run_id == pro_run_id))
        if pro is None or pro.flash_validation_run_id != flash_run_id:
            raise ValueError("COMPATIBLE_PRO_RUN_REQUIRED")
        quant = list(self.session.scalars(select(QuantRankResult).where(QuantRankResult.quant_run_id == pro.quant_run_id).order_by(QuantRankResult.rank)))
        samples = list(self.session.scalars(select(ModelValidationSample).where(ModelValidationSample.validation_run_id == flash_run_id).order_by(ModelValidationSample.rank)))
        reviews = list(self.session.scalars(select(ProCandidateReview).where(ProCandidateReview.pro_resume_run_id == pro_run_id).order_by(ProCandidateReview.pro_rank)))
        plans = {row.stock_code: row for row in self.session.scalars(select(ModelValidationOrderPlan).where(ModelValidationOrderPlan.validation_run_id == flash_run_id))}
        allocations = {row.stock_code: row for row in self.session.scalars(select(ModelValidationAllocation).where(ModelValidationAllocation.validation_run_id == flash_run_id))}

        workbook = Workbook()
        workbook.remove(workbook.active)
        _write_sheet(workbook, "01_全A量化排名", ({
            "排名": row.rank, "股票代码": display_stock_code(row.stock_code), "总分": row.total_score,
            "技术": row.technical_score, "资金": row.capital_score, "情绪": row.emotion_score,
            "动量": row.momentum_score, "风险": row.risk_score,
        } for row in quant))
        _write_sheet(workbook, "02_Top100_LLM评分", ({
            "排名": row.rank, "股票代码": display_stock_code(row.stock_code), "股票名称": row.stock_name,
            "LLM分数": (row.screening_result or {}).get("llm_score"),
            "筛选结论": (row.screening_result or {}).get("screening_decision"),
            "入选来源": (row.screening_result or {}).get("_trader_demo", {}).get("selection_source"),
            "风险提示": _compact((row.screening_result or {}).get("risk_flags")),
        } for row in samples))
        _write_sheet(workbook, "03_挂单与仓位", (_order_row(review, plans.get(review.stock_code), allocations.get(review.stock_code)) for review in reviews))
        _write_sheet(workbook, "04_重点基本面", ({
            "股票代码": display_stock_code(row.stock_code), "股票名称": row.stock_name,
            "主营业务": _compact((row.fundamental_result or {}).get("main_business_summary")),
            "行业位置": _compact((row.fundamental_result or {}).get("industry_position")),
            "核心优势": _compact((row.fundamental_result or {}).get("competitive_advantage")),
            "主要风险": _compact((row.fundamental_result or {}).get("key_risks")),
            "数据缺失": _compact(row.missing_fields),
        } for row in samples if (row.screening_result or {}).get("_trader_demo", {}).get("selection_source")))
        _write_sheet(workbook, "05_说明与异常", [{
            "交易日期": trade_date, "流水线ID": pro.pipeline_run_id, "量化数量": len(quant),
            "Flash数量": len(samples), "最终候选": len(reviews), "真实交易": "关闭",
            "用途": "交易辅助与模型验证", "说明": "系统不执行真实下单，最终决策由人工负责。",
        }])
        review_row = MarketReviewRepository(self.session).latest_run(trade_date)
        add_market_review_sheet(workbook, MarketReviewRepository(self.session).bundle(review_row) if review_row else None)
        from post_close.excel import daily_sheet_rows
        _write_sheet(workbook, "07_盘后操作建议", daily_sheet_rows(self.session, trade_date))
        admission_run_id, admission_rows = latest_entry_timing_rows(self.session, trade_date)
        if admission_run_id:
            _write_sheet(workbook, "08_买入准入分析", admission_rows)
        workbook.active = workbook.sheetnames.index("03_挂单与仓位")
        sheet_count = len(workbook.sheetnames)

        output_dir = self.output_root / trade_date.isoformat()
        output_dir.mkdir(parents=True, exist_ok=True)
        output = _available_path(output_dir / f"AI交易助手_{trade_date.isoformat()}.xlsx")
        workbook.save(output)
        workbook.close()
        digest = hashlib.sha256(output.read_bytes()).hexdigest()
        audit = output.with_name(f"{output.stem}_audit.json")
        audit.write_text(json.dumps({
            "trade_date": trade_date.isoformat(), "pipeline_run_id": pro.pipeline_run_id,
            "flash_run_id": flash_run_id, "pro_run_id": pro_run_id,
            "candidate_set_hash": pro.candidate_set_hash, "output_path": str(output),
            "sha256": digest, "sheet_count": sheet_count, "advisory_only": True,
            "market_review_run_id": review_row.run_id if review_row else None,
            "entry_timing_run_id": admission_run_id,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        return {"output_path": str(output), "audit_path": str(audit), "sha256": digest, "sheet_count": sheet_count, "status": "SUCCESS"}


def write_tabular_workbook(output: Path, sheets: dict[str, Iterable[dict[str, Any]]]) -> dict[str, Any]:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        _write_sheet(workbook, name, rows)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    content = output.read_bytes()
    return {"output_path": str(output), "size": len(content), "sha256": hashlib.sha256(content).hexdigest(), "sheet_count": len(sheets), "status": "SUCCESS"}


def _write_sheet(workbook: Workbook, title: str, rows: Iterable[dict[str, Any]]) -> None:
    values = list(rows)
    sheet = workbook.create_sheet(title[:31])
    headers = list(dict.fromkeys(key for row in values for key in row)) or ["暂无数据"]
    sheet.append(headers)
    for row in values:
        sheet.append([_cell_value(row.get(key)) for key in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            header = str(sheet.cell(1, cell.column).value)
            if "股票代码" in header or header.lower() in {"stock_code", "code"}:
                cell.number_format = "@"
            if isinstance(cell.value, float) and any(word in header for word in ("收益", "涨跌", "仓位", "置信度", "概率")):
                cell.number_format = "0.00%"
                cell.font = Font(color="C00000" if cell.value > 0 else "008000" if cell.value < 0 else "000000")
            if "建议" in header and isinstance(cell.value, str):
                if any(word in cell.value for word in ("退出", "可交易时")):
                    cell.fill = PatternFill("solid", fgColor="F4CCCC")
                elif "减仓" in cell.value or "收紧止损" in cell.value:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
                elif "继续持有" in cell.value or "准备次日" in cell.value:
                    cell.fill = PatternFill("solid", fgColor="D9EAD3")
                elif "数据不足" in cell.value:
                    cell.fill = PatternFill("solid", fgColor="E7E6E6")
    for index, header in enumerate(headers, start=1):
        values_length = [len(str(sheet.cell(row, index).value or "")) for row in range(2, min(sheet.max_row, 100) + 1)]
        width = min(42, max([12, len(str(header)) * 2, *values_length]))
        sheet.column_dimensions[get_column_letter(index)].width = width


def _order_row(review, plan, allocation) -> dict[str, Any]:
    return {
        "最终排名": review.pro_rank, "股票代码": display_stock_code(review.stock_code),
        "Pro分数": review.pro_score, "优先级": review.priority,
        "建议价": getattr(plan, "recommended_price", None), "最高接受价": getattr(plan, "max_acceptable_price", None),
        "止损价": getattr(plan, "stop_loss_price", None), "止盈一": getattr(plan, "take_profit_1_price", None),
        "止盈二": getattr(plan, "take_profit_2_price", None), "风险收益比": getattr(plan, "active_risk_reward", None),
        "建议仓位": getattr(allocation, "suggested_position_percent", None), "建议股数": getattr(allocation, "suggested_quantity", None),
        "状态": getattr(plan, "status", None), "摘要": review.final_summary,
    }


def _cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return _compact(value)


def _compact(value: Any) -> str:
    if value in (None, "", [], {}):
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, default=str, separators=(",", ":"))


def _available_path(path: Path) -> Path:
    if not path.exists():
        return path
    stamp = datetime.now().strftime("%H%M%S")
    return path.with_name(f"{path.stem}_{stamp}{path.suffix}")
