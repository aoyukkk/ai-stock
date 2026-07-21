from __future__ import annotations

import os
import uuid
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "FF17365D"
BLUE = "FF1F4E78"
LIGHT_BLUE = "FFD9EAF7"
HEADER_FILL = "FFB4C6E7"
PALE = "FFF4F7F9"
WHITE = "FFFFFFFF"
BLACK = "FF000000"
UP_RED = "FFC00000"
DOWN_GREEN = "FF008000"
NEUTRAL_GRAY = "FF666666"
UP_FILL = "FFFCE4D6"
DOWN_FILL = "FFE2F0D9"
NEUTRAL_FILL = "FFF2F2F2"
THIN = Side(style="thin", color="FFB8DDF2")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RETURN_FORMAT = "0.00%"


class HumanMarketReviewExcelExporter:
    def export(self, output_path: Path, bundle: dict[str, Any], supplement: dict[str, Any]) -> dict[str, Any]:
        workbook = Workbook()
        workbook.remove(workbook.active)
        _overview_sheet(workbook, bundle, supplement)
        _industry_sheet(workbook, bundle)
        _all_industry_sheet(workbook, bundle)
        _cause_sheet(workbook, bundle, supplement)
        _outlook_sheet(workbook, bundle)
        workbook.active = 0
        output_path.parent.mkdir(parents=True, exist_ok=True)
        candidate = output_path.with_name(f".{output_path.stem}_{uuid.uuid4().hex[:8]}_candidate.xlsx")
        workbook.save(candidate)
        workbook.close()
        _verify(candidate)
        actual_output = output_path
        canonical_replaced = True
        try:
            os.replace(candidate, output_path)
        except PermissionError:
            canonical_replaced = False
            actual_output = output_path.with_name(f"{output_path.stem}_更新版.xlsx")
            try:
                os.replace(candidate, actual_output)
            except PermissionError:
                actual_output = output_path.with_name(f"{output_path.stem}_更新版_{uuid.uuid4().hex[:8]}.xlsx")
                os.replace(candidate, actual_output)
        finally:
            candidate.unlink(missing_ok=True)
        result = _verify(actual_output)
        result["requested_output"] = str(output_path)
        result["canonical_replaced"] = canonical_replaced
        return result


def _overview_sheet(workbook: Workbook, bundle: dict[str, Any], supplement: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("大盘概览")
    snapshot = bundle.get("snapshot") or {}
    breadth = snapshot.get("breadth") or {}
    turnover = snapshot.get("turnover") or {}
    limits = snapshot.get("limit_structure") or {}
    capital = snapshot.get("capital") or {}
    run = bundle.get("run") or {}
    trade_date = str(run.get("trade_date") or "")
    _title(sheet, f"{trade_date} A股大盘复盘", _market_overview_subtitle(snapshot), 6)
    metrics = [
        ["市场状态", "轮动分化", "上涨家数", breadth.get("advancing_count"), "下跌家数", breadth.get("declining_count")],
        ["平盘家数", breadth.get("flat_count"), "上涨占比", breadth.get("advancing_ratio"), "全A等权涨跌", breadth.get("equal_weight_return")],
        ["成交额（亿元）", _to_yi(turnover.get("total_amount")), "较昨日", turnover.get("change_ratio"), "五日量能比", turnover.get("relative_to_5d")],
        ["涨停家数", limits.get("limit_up_count"), "跌停家数", limits.get("limit_down_count"), "炸板率", limits.get("failed_limit_up_ratio")],
        ["主力净流入（亿元）", _to_yi(capital.get("net_main_inflow")), "数据质量", snapshot.get("data_quality_score"), "规则置信度", run.get("confidence")],
    ]
    _section(sheet, 3, "核心盘面", 6)
    _write_rows(sheet, 4, metrics)
    for coordinate in ("D5", "F5", "D6", "F6", "F7"):
        _style_return_cell(sheet[coordinate])

    indices = list(supplement.get("indices") or [])
    _section(sheet, 9, "主要指数", 6)
    headers = ["指数", "收盘点位", "当日涨跌", "盘面特征", "来源", "数据日期"]
    rows = [[item.get(key) for key in ("name", "close", "change", "feature", "source", "date")] for item in indices]
    _table(sheet, 10, headers, rows)
    for row_number in range(11, 11 + len(rows)):
        _style_return_cell(sheet.cell(row_number, 3))
    sheet.freeze_panes = "A4"
    _finish(sheet, [18, 16, 16, 30, 24, 14])


def _industry_sheet(workbook: Workbook, bundle: dict[str, Any]) -> None:
    industries = list((bundle.get("snapshot") or {}).get("industries") or [])
    leading = industries[:10]
    leading_codes = {str(item.get("sector_code") or item.get("sector_name")) for item in leading}
    lagging = [
        item for item in industries[-10:]
        if str(item.get("sector_code") or item.get("sector_name")) not in leading_codes
    ]
    grouped = [("领涨", item) for item in leading] + [("偏弱", item) for item in lagging]
    _render_industry_sheet(
        workbook.create_sheet("行业轮动"),
        "行业轮动",
        f"从全部 {len(industries)} 个行业中保留涨幅前10与后10，便于快速观察强弱方向。",
        grouped,
    )


def _all_industry_sheet(workbook: Workbook, bundle: dict[str, Any]) -> None:
    industries = list((bundle.get("snapshot") or {}).get("industries") or [])
    total = len(industries)
    grouped = [
        ("领涨" if index < 10 else "偏弱" if index >= max(total - 10, 10) else "中段", item)
        for index, item in enumerate(industries)
    ]
    _render_industry_sheet(
        workbook.create_sheet("全部行业"),
        "全部行业排名",
        f"按平均涨跌幅展示当日全部 {total} 个可用行业，未截断中间行业。",
        grouped,
    )


def _render_industry_sheet(sheet, title: str, subtitle: str, grouped: list[tuple[str, dict[str, Any]]]) -> None:
    _title(sheet, title, subtitle, 8)
    headers = ["排名", "分组", "行业", "平均涨跌", "上涨占比", "涨停家数", "成分数量", "成交额（亿元）"]
    rows = []
    for rank, (group, item) in enumerate(grouped, 1):
        rows.append([
            item.get("rank") or rank, group, item.get("sector_name"), item.get("change_percent"), item.get("advancing_ratio"),
            item.get("limit_up_count"), item.get("member_count"), _to_yi(item.get("amount")),
        ])
    _table(sheet, 3, headers, rows)
    for row_number in range(4, 4 + len(rows)):
        _style_return_cell(sheet.cell(row_number, 4))
        sheet.cell(row_number, 5).number_format = "0.00%"
        group_cell = sheet.cell(row_number, 2)
        if group_cell.value == "领涨":
            group_cell.fill = PatternFill("solid", fgColor=UP_FILL)
            group_cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=UP_RED)
        elif group_cell.value == "偏弱":
            group_cell.fill = PatternFill("solid", fgColor=DOWN_FILL)
            group_cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=DOWN_GREEN)
    sheet.freeze_panes = "D4"
    sheet.auto_filter.ref = f"A3:H{3 + len(rows)}"
    _finish(sheet, [10, 12, 18, 16, 16, 14, 14, 18])


def _cause_sheet(workbook: Workbook, bundle: dict[str, Any], supplement: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("盘面原因")
    _title(sheet, "今日盘面原因", "区分盘面事实、公开催化与基于两者的归纳。", 6)
    headers = ["方向", "盘面事实", "公开催化", "归纳结论", "来源名称", "来源链接"]
    rows = [[item.get(key) for key in ("direction", "market_fact", "public_catalyst", "conclusion", "source_name", "source_url")] for item in supplement.get("causes") or []]
    if not rows:
        rows = [
            [
                item.get("direction"),
                item.get("explanation"),
                "",
                item.get("title"),
                "本地结构化行情",
                "",
            ]
            for item in bundle.get("drivers") or []
        ]
        industries = list((bundle.get("snapshot") or {}).get("industries") or [])
        if industries:
            leaders = "、".join(str(item.get("sector_name")) for item in industries[:3])
            laggards = "、".join(str(item.get("sector_name")) for item in industries[-3:])
            rows.append([
                "STRUCTURAL",
                f"领涨行业为 {leaders}；偏弱行业为 {laggards}。",
                "",
                "行业轮动分化",
                "Tushare 全A行业聚合",
                "",
            ])
    _table(sheet, 3, headers, rows)
    sheet.freeze_panes = "A4"
    _finish(sheet, [14, 34, 42, 42, 22, 58], row_height=68)


def _market_overview_subtitle(snapshot: dict[str, Any]) -> str:
    breadth = snapshot.get("breadth") or {}
    turnover = snapshot.get("turnover") or {}
    advancing = int(breadth.get("advancing_count") or 0)
    declining = int(breadth.get("declining_count") or 0)
    equal_weight = float(breadth.get("equal_weight_return") or 0)
    turnover_change = float(turnover.get("change_ratio") or 0)
    if advancing > declining and equal_weight >= 0:
        breadth_text = "全A个股整体偏强"
    elif declining > advancing and equal_weight <= 0:
        breadth_text = "全A个股整体偏弱"
    else:
        breadth_text = "全A个股涨跌分化"
    turnover_text = "成交缩量" if turnover_change < 0 else "成交放量"
    return f"{breadth_text}，{turnover_text}，行业轮动明显。"


def _outlook_sheet(workbook: Workbook, bundle: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("次日观察")
    run = bundle.get("run") or {}
    outlook = bundle.get("outlook") or {}
    _title(sheet, "次日条件观察", "这是规则模型的条件情景，重点用于盘前和午盘复核。", 5)
    scenarios = [
        ["基准", run.get("base_case_probability", outlook.get("base_case_probability")), "继续轮动分化", "量能未明显放大，医药消费与科技之间继续快速轮动", "成交额、上涨家数、半导体止跌"],
        ["偏强", run.get("bull_case_probability", outlook.get("bull_case_probability")), "放量修复", "成交额回升且科技止跌，市场宽度继续扩散", "科创50、半导体、炸板率"],
        ["偏弱", run.get("bear_case_probability", outlook.get("bear_case_probability")), "缩量下探", "科技继续走弱且医药白酒冲高回落，赚钱效应收缩", "主力资金、跌停家数、领涨持续性"],
    ]
    _table(sheet, 3, ["情景", "概率", "标题", "触发条件", "重点观察"], scenarios)
    for row_number in range(4, 7):
        value = sheet.cell(row_number, 2).value
        if isinstance(value, (int, float)) and value > 1:
            sheet.cell(row_number, 2).value = value / 100
        sheet.cell(row_number, 2).number_format = "0%"
    _section(sheet, 7, "明日执行顺序", 5)
    steps = [
        ["1", "盘前", "确认 Tushare 数据日期、指数隔夜信息和人工池", "不重复跑午盘流程", "审计/数据检查记录"],
        ["2", "午盘", "11:32 后只运行一次午盘分析推荐", "输出午盘专用文件", "午盘/智能交易助手_午盘_日期.xlsx"],
        ["3", "盘后", "确认数据齐备后运行全A量化、Flash、Pro和监测", "检查点续跑", "智能交易助手_日期.xlsx"],
    ]
    _table(sheet, 8, ["顺序", "时段", "动作", "要求", "结果文件"], steps)
    sheet.freeze_panes = "A4"
    _finish(sheet, [12, 14, 28, 54, 34])


def _title(sheet, title: str, subtitle: str, columns: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    sheet.cell(1, 1, title)
    sheet.cell(2, 1, subtitle)
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=HEADER_FILL)
    sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=18, bold=True, color=NAVY)
    sheet.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet.cell(2, 1).font = Font(name="Microsoft YaHei", size=10, color=NAVY)
    sheet.cell(1, 1).alignment = CENTER
    sheet.cell(2, 1).alignment = CENTER
    sheet.row_dimensions[1].height = 36
    sheet.row_dimensions[2].height = 28


def _section(sheet, row: int, title: str, columns: int) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    cell = sheet.cell(row, 1, title)
    cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    cell.font = Font(name="Microsoft YaHei", bold=True, color=NAVY)
    cell.alignment = CENTER
    sheet.row_dimensions[row].height = 26


def _table(sheet, header_row: int, headers: list[str], rows: list[list[Any]]) -> None:
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, header)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(name="Microsoft YaHei", bold=True, color=NAVY)
        cell.alignment = CENTER
    sheet.row_dimensions[header_row].height = 34
    for row_number, values in enumerate(rows, header_row + 1):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.alignment = CENTER
            cell.border = Border(bottom=THIN)
            cell.font = Font(name="Microsoft YaHei", size=10)


def _write_rows(sheet, start_row: int, rows: list[list[Any]]) -> None:
    for row_number, values in enumerate(rows, start_row):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.alignment = CENTER
            cell.border = Border(bottom=THIN)
            cell.fill = PatternFill("solid", fgColor=PALE if row_number % 2 == 0 else WHITE)
            cell.font = Font(name="Microsoft YaHei", bold=column % 2 == 1, color=NAVY if column % 2 == 1 else BLACK)
        sheet.row_dimensions[row_number].height = 28


def _finish(sheet, widths: list[int], *, row_height: int = 30) -> None:
    sheet.sheet_view.showGridLines = False
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in range(3, sheet.max_row + 1):
        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 0, row_height)
        for cell in sheet[row]:
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            cell.alignment = CENTER


def _to_yi(value: Any) -> float | None:
    return round(float(value) / 100_000_000, 2) if value not in (None, "") else None


def _style_return_cell(cell) -> None:
    cell.number_format = RETURN_FORMAT
    try:
        value = float(cell.value)
    except (TypeError, ValueError):
        return
    if value > 0:
        color, fill = UP_RED, UP_FILL
    elif value < 0:
        color, fill = DOWN_GREEN, DOWN_FILL
    else:
        color, fill = NEUTRAL_GRAY, NEUTRAL_FILL
    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)


def _verify(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("MARKET_REVIEW_WORKBOOK_ZIP_ERROR")
    workbook = load_workbook(path, data_only=False)
    try:
        for sheet in workbook.worksheets:
            for row_number in range(1, sheet.max_row + 1):
                if not any(not isinstance(cell, MergedCell) and cell.value is not None for cell in sheet[row_number]):
                    raise ValueError(f"MARKET_REVIEW_BLANK_ROW:{sheet.title}:{row_number}")
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell) or cell.value is None:
                        continue
                    if cell.alignment.horizontal != "center" or cell.alignment.vertical != "center":
                        raise ValueError(f"MARKET_REVIEW_ALIGNMENT:{sheet.title}:{cell.coordinate}")
                    if isinstance(cell.value, str) and any(token in cell.value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                        raise ValueError(f"MARKET_REVIEW_FORMULA_ERROR:{sheet.title}:{cell.coordinate}")
        return {
            "output": str(path),
            "size_bytes": path.stat().st_size,
            "sheets": workbook.sheetnames,
            "blank_rows": 0,
            "formula_errors": 0,
            "center_alignment_verified": True,
        }
    finally:
        workbook.close()
