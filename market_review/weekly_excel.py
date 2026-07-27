from __future__ import annotations

import math
import zipfile
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE = "F4F7F9"
WHITE = "FFFFFF"
UP_RED = "C00000"
DOWN_GREEN = "008000"
THIN = Side(style="thin", color="B8DDF2")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BORDER = Border(bottom=THIN)
PERCENT = "[Red]0.00%;[Green]-0.00%;-"


class WeeklyMarketReviewExcelExporter:
    def export(
        self,
        output_path: Path,
        bundles: list[dict[str, Any]],
        industry_moneyflow: dict[date, dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        ordered = sorted(bundles, key=lambda row: str((row.get("run") or {}).get("trade_date") or ""))
        if not ordered:
            raise ValueError("WEEKLY_MARKET_REVIEW_EMPTY")
        industry_moneyflow = industry_moneyflow or {}
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._overview(workbook, ordered)
        self._daily_review(workbook, ordered)
        self._daily_industry_sheets(workbook, ordered, industry_moneyflow)
        self._capital_flow(workbook, ordered, industry_moneyflow)
        self._all_industries(workbook, ordered)
        workbook.active = 0
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        workbook.close()
        return self._verify(output_path)

    def _overview(self, workbook: Workbook, bundles: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet("本周概览")
        headers = [
            "交易日", "市场状态", "等权涨跌", "中位数涨跌", "上涨家数", "下跌家数",
            "上涨占比", "成交额(亿元)", "较前日变化", "涨停家数", "跌停家数", "数据质量",
        ]
        self._title(sheet, "本周 A 股市场复盘", self._period(bundles), len(headers))
        self._headers(sheet, 3, headers)
        daily_returns: list[float] = []
        for row_number, bundle in enumerate(bundles, 4):
            snapshot = bundle.get("snapshot") or {}
            breadth = snapshot.get("breadth") or {}
            turnover = snapshot.get("turnover") or {}
            limits = snapshot.get("limit_structure") or {}
            regime = bundle.get("regime") or {}
            daily_return = self._number(breadth.get("equal_weight_return"))
            if daily_return is not None:
                daily_returns.append(daily_return)
            values = [
                date.fromisoformat(str((bundle.get("run") or {})["trade_date"])),
                self._regime_label(str(regime.get("market_regime") or "")),
                daily_return,
                self._number(breadth.get("median_return")),
                breadth.get("advancing_count"),
                breadth.get("declining_count"),
                self._number(breadth.get("advancing_ratio")),
                self._number(turnover.get("total_amount")) / 100_000_000 if turnover.get("total_amount") is not None else None,
                self._number(turnover.get("change_ratio")),
                limits.get("limit_up_count"),
                limits.get("limit_down_count"),
                self._number(snapshot.get("data_quality_score")),
            ]
            self._row(sheet, row_number, values)
            sheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
            for column in (3, 4, 7, 9):
                sheet.cell(row_number, column).number_format = PERCENT
            sheet.cell(row_number, 8).number_format = "#,##0.00"
            sheet.cell(row_number, 12).number_format = "0.00"
        total_row = 4 + len(bundles)
        cumulative = math.prod(1 + value for value in daily_returns) - 1 if daily_returns else None
        self._row(sheet, total_row, ["本周合计", "", cumulative, "", "", "", "", "", "", "", "", ""])
        sheet.cell(total_row, 3).number_format = PERCENT
        sheet.cell(total_row, 1).font = Font(name="Microsoft YaHei", bold=True, color=NAVY)
        sheet.cell(total_row, 3).font = Font(name="Microsoft YaHei", bold=True, color=UP_RED if (cumulative or 0) >= 0 else DOWN_GREEN)
        sheet.freeze_panes = "C4"
        sheet.auto_filter.ref = f"A3:L{3 + len(bundles)}"
        self._finish(sheet, [14, 18, 14, 14, 12, 12, 14, 16, 14, 12, 12, 12])

    def _daily_review(self, workbook: Workbook, bundles: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet("每日复盘")
        headers = ["交易日", "复盘结论", "市场宽度", "成交情况", "领涨行业", "偏弱行业", "主要原因"]
        self._title(sheet, "本周每日市场复盘", "原因栏仅汇总已落库的市场结构与成交事实", len(headers))
        self._headers(sheet, 3, headers)
        for row_number, bundle in enumerate(bundles, 4):
            review = bundle.get("review") or {}
            industries = list((bundle.get("snapshot") or {}).get("industries") or [])
            drivers = bundle.get("drivers") or []
            reason = "；".join(str(item.get("explanation") or "") for item in drivers if item.get("explanation")) or "暂无可核验原因"
            values = [
                date.fromisoformat(str((bundle.get("run") or {})["trade_date"])),
                review.get("market_summary") or review.get("headline") or "暂无",
                review.get("breadth_summary") or "暂无",
                review.get("turnover_summary") or "暂无",
                "、".join(str(item.get("sector_name")) for item in industries[:3]) or "暂无",
                "、".join(str(item.get("sector_name")) for item in industries[-3:]) or "暂无",
                reason,
            ]
            self._row(sheet, row_number, values, height=72)
            sheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
        sheet.freeze_panes = "B4"
        sheet.auto_filter.ref = f"A3:G{3 + len(bundles)}"
        self._finish(sheet, [14, 38, 28, 28, 24, 24, 54])

    def _daily_industry_sheets(
        self,
        workbook: Workbook,
        bundles: list[dict[str, Any]],
        industry_moneyflow: dict[date, dict[str, Any]],
    ) -> None:
        for bundle in bundles:
            trade_day = date.fromisoformat(str((bundle.get("run") or {})["trade_date"]))
            industries = list((bundle.get("snapshot") or {}).get("industries") or [])
            flow_bundle = industry_moneyflow.get(trade_day) or {}
            flow_values = flow_bundle.get("values") or {}
            sheet = workbook.create_sheet(f"{trade_day.month}月{trade_day.day}日行业")
            headers = [
                "分组", "当日排名", "行业", "行业涨跌", "上涨占比", "涨停家数",
                "成交额(亿元)", "主力净流入(亿元)", "成分数量",
            ]
            coverage = self._number(flow_bundle.get("coverage"))
            coverage_text = f"；moneyflow 行业映射覆盖率 {coverage:.2%}" if coverage is not None else ""
            self._title(
                sheet,
                f"{trade_day.month}月{trade_day.day}日行业前十后十",
                f"领涨 10 个与偏弱 10 个行业；主力净流入按 Tushare moneyflow 汇总{coverage_text}",
                len(headers),
            )
            self._headers(sheet, 3, headers)
            selected = [("领涨", item) for item in industries[:10]] + [("偏弱", item) for item in industries[-10:]]
            row_number = 4
            for group, item in selected:
                sector = str(item.get("sector_name") or item.get("sector_code") or "")
                values = [
                    group, item.get("rank"), sector,
                    self._number(item.get("change_percent")), self._number(item.get("advancing_ratio")),
                    item.get("limit_up_count"), self._number(item.get("amount")) / 100_000_000 if item.get("amount") is not None else None,
                    self._number(flow_values.get(sector)),
                    item.get("member_count"),
                ]
                self._row(sheet, row_number, values)
                sheet.cell(row_number, 4).number_format = PERCENT
                sheet.cell(row_number, 5).number_format = "0.00%"
                sheet.cell(row_number, 7).number_format = "#,##0.00"
                sheet.cell(row_number, 8).number_format = "[Red]#,##0.00;[Green]-#,##0.00;-"
                row_number += 1
            sheet.freeze_panes = "D4"
            sheet.auto_filter.ref = f"A3:I{row_number - 1}"
            self._finish(sheet, [10, 12, 20, 14, 14, 12, 16, 18, 12])

    def _capital_flow(
        self,
        workbook: Workbook,
        bundles: list[dict[str, Any]],
        industry_moneyflow: dict[date, dict[str, Any]],
    ) -> None:
        sheet = workbook.create_sheet("行业资金流向")
        headers = [
            "交易日", "序号", "主要流出行业", "主力净流出(亿元)", "成交占比变化",
            "主要流入行业", "主力净流入(亿元)", "成交占比变化", "直观判断",
        ]
        self._title(
            sheet,
            "本周行业资金流向变化",
            "主力净流入来自 Tushare moneyflow 行业汇总；成交占比变化用于观察关注度迁移，行业配对仅为直观展示。",
            len(headers),
        )
        self._headers(sheet, 3, headers)
        previous_shares: dict[str, float] = {}
        row_number = 4
        for bundle in bundles:
            trade_day = date.fromisoformat(str((bundle.get("run") or {})["trade_date"]))
            industries = list((bundle.get("snapshot") or {}).get("industries") or [])
            total_amount = sum(self._number(item.get("amount")) or 0 for item in industries)
            current_shares = {
                str(item.get("sector_name") or item.get("sector_code")): (self._number(item.get("amount")) or 0) / total_amount
                for item in industries
            } if total_amount else {}
            share_delta = {
                sector: share - previous_shares.get(sector, 0.0)
                for sector, share in current_shares.items()
            } if previous_shares else {}
            flow_bundle = industry_moneyflow.get(trade_day) or {}
            flow_values = {
                str(sector): self._number(value) or 0.0
                for sector, value in (flow_bundle.get("values") or {}).items()
            }
            if flow_values:
                inflows = sorted(
                    ((sector, value) for sector, value in flow_values.items() if value > 0),
                    key=lambda item: (-item[1], item[0]),
                )[:5]
                outflows = sorted(
                    ((sector, value) for sector, value in flow_values.items() if value < 0),
                    key=lambda item: (item[1], item[0]),
                )[:5]
                flow_basis = "主力资金"
            elif share_delta:
                inflows = [(sector, None) for sector, _ in sorted(
                    ((sector, value) for sector, value in share_delta.items() if value > 0),
                    key=lambda item: (-item[1], item[0]),
                )[:5]]
                outflows = [(sector, None) for sector, _ in sorted(
                    ((sector, value) for sector, value in share_delta.items() if value < 0),
                    key=lambda item: (item[1], item[0]),
                )[:5]]
                flow_basis = "成交关注度"
            else:
                inflows, outflows, flow_basis = [], [], "行业资金"
            pair_count = max(len(inflows), len(outflows), 1)
            for rank in range(pair_count):
                out_sector, out_value = outflows[rank] if rank < len(outflows) else ("暂无", None)
                in_sector, in_value = inflows[rank] if rank < len(inflows) else ("暂无", None)
                if out_value is None and in_value is None:
                    judgement = (
                        f"{flow_basis}由{out_sector}侧下降，向{in_sector}侧提升"
                        if out_sector != "暂无" or in_sector != "暂无"
                        else "行业资金数据不足"
                    )
                else:
                    judgement = f"资金由{out_sector}侧流出，向{in_sector}侧集中"
                values = [
                    trade_day, rank + 1, out_sector, out_value, share_delta.get(out_sector),
                    in_sector, in_value, share_delta.get(in_sector), judgement,
                ]
                self._row(sheet, row_number, values, height=30)
                sheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
                for column in (4, 7):
                    sheet.cell(row_number, column).number_format = "[Red]#,##0.00;[Green]-#,##0.00;-"
                for column in (5, 8):
                    sheet.cell(row_number, column).number_format = "[Red]0.00%;[Green]-0.00%;-"
                row_number += 1
            previous_shares = current_shares
        sheet.freeze_panes = "C4"
        sheet.auto_filter.ref = f"A3:I{row_number - 1}"
        self._finish(sheet, [14, 10, 20, 19, 17, 20, 19, 17, 38])

    def _all_industries(self, workbook: Workbook, bundles: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet("全部行业")
        trade_days = [date.fromisoformat(str((bundle.get("run") or {})["trade_date"])) for bundle in bundles]
        headers = ["本周排名", "行业", *[f"{day.month}月{day.day}日涨跌" for day in trade_days], "本周累计涨跌", "平均上涨占比", "本周成交额(亿元)"]
        self._title(sheet, "本周全部行业表现", "本周累计涨跌按每日行业涨跌复合计算", len(headers))
        self._headers(sheet, 3, headers)
        values_by_sector: dict[str, dict[date, dict[str, Any]]] = defaultdict(dict)
        for bundle, trade_day in zip(bundles, trade_days):
            for item in (bundle.get("snapshot") or {}).get("industries") or []:
                values_by_sector[str(item.get("sector_name") or item.get("sector_code"))][trade_day] = item
        ranked = []
        for sector, by_day in values_by_sector.items():
            returns = [self._number(by_day[day].get("change_percent")) for day in trade_days if day in by_day]
            cumulative = math.prod(1 + value for value in returns if value is not None) - 1 if returns else None
            ranked.append((sector, by_day, cumulative))
        ranked.sort(key=lambda item: (item[2] is None, -(item[2] or 0), item[0]))
        for row_number, (sector, by_day, cumulative) in enumerate(ranked, 4):
            advances = [self._number(item.get("advancing_ratio")) for item in by_day.values() if item.get("advancing_ratio") is not None]
            amount = sum(self._number(item.get("amount")) or 0 for item in by_day.values()) / 100_000_000
            values: list[Any] = [row_number - 3, sector]
            values.extend(self._number(by_day[day].get("change_percent")) if day in by_day else "无数据" for day in trade_days)
            values.extend([cumulative, sum(advances) / len(advances) if advances else None, amount])
            self._row(sheet, row_number, values)
            for column in range(3, 3 + len(trade_days)):
                if isinstance(sheet.cell(row_number, column).value, (int, float)):
                    sheet.cell(row_number, column).number_format = PERCENT
            sheet.cell(row_number, 3 + len(trade_days)).number_format = PERCENT
            sheet.cell(row_number, 4 + len(trade_days)).number_format = "0.00%"
            sheet.cell(row_number, 5 + len(trade_days)).number_format = "#,##0.00"
        sheet.freeze_panes = "C4"
        sheet.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{3 + len(ranked)}"
        self._finish(sheet, [12, 20, *([15] * len(trade_days)), 16, 16, 18])

    @staticmethod
    def _period(bundles: list[dict[str, Any]]) -> str:
        days = [str((bundle.get("run") or {}).get("trade_date")) for bundle in bundles]
        return f"统计区间：{days[0]} 至 {days[-1]}，共 {len(days)} 个交易日"

    @staticmethod
    def _regime_label(value: str) -> str:
        return {
            "BROAD_RALLY": "普涨", "BROAD_SELL_OFF": "普跌", "RISK_ON": "偏强",
            "RISK_OFF": "偏弱", "ROTATION": "轮动", "MIXED_ROTATION": "震荡轮动",
            "RANGE_BOUND": "震荡",
        }.get(value, value or "数据不足")

    @staticmethod
    def _number(value: Any) -> float | None:
        try:
            return float(value) if value is not None else None
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _title(sheet, title: str, subtitle: str, columns: int) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
        sheet.cell(1, 1, title)
        sheet.cell(2, 1, subtitle)
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
        sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
        sheet.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        sheet.cell(2, 1).font = Font(name="Microsoft YaHei", size=10, color=NAVY)
        sheet.cell(1, 1).alignment = CENTER
        sheet.cell(2, 1).alignment = CENTER
        sheet.row_dimensions[1].height = 36
        sheet.row_dimensions[2].height = 30

    @staticmethod
    def _headers(sheet, row: int, headers: list[str]) -> None:
        for column, value in enumerate(headers, 1):
            cell = sheet.cell(row, column, value)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(name="Microsoft YaHei", bold=True, color=WHITE)
            cell.alignment = CENTER
        sheet.row_dimensions[row].height = 32

    @staticmethod
    def _row(sheet, row: int, values: list[Any], height: int = 26) -> None:
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, value)
            cell.alignment = CENTER
            cell.border = BORDER
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.fill = PatternFill("solid", fgColor=WHITE)
        sheet.row_dimensions[row].height = height

    @staticmethod
    def _finish(sheet, widths: list[float]) -> None:
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.auto_filter.ref

    @staticmethod
    def _verify(path: Path) -> dict[str, Any]:
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise ValueError("WEEKLY_WORKBOOK_ZIP_INVALID")
        workbook = load_workbook(path, data_only=False)
        try:
            blank_rows = 0
            for sheet in workbook.worksheets:
                for row in range(1, sheet.max_row + 1):
                    if not any(sheet.cell(row, column).value is not None for column in range(1, sheet.max_column + 1)):
                        blank_rows += 1
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell, MergedCell) or cell.value is None:
                            continue
                        if cell.alignment.horizontal != "center" or cell.alignment.vertical != "center":
                            raise ValueError(f"WEEKLY_WORKBOOK_ALIGNMENT_INVALID:{sheet.title}:{cell.coordinate}")
            if blank_rows:
                raise ValueError(f"WEEKLY_WORKBOOK_BLANK_ROWS:{blank_rows}")
            return {
                "output": str(path), "size_bytes": path.stat().st_size,
                "sheets": workbook.sheetnames, "blank_rows": 0,
                "center_alignment_verified": True,
            }
        finally:
            workbook.close()
