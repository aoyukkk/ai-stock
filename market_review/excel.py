from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEET_NAME = "06_大盘复盘"
DISCLAIMER = "本报告基于已取得的市场数据、公开信息与规则模型生成，仅用于市场复盘和模型验证，不构成投资建议。次日走势为条件情景分析，不是确定性预测。"


def add_market_review_sheet(workbook: Workbook, bundle: dict[str, Any] | None) -> None:
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "A股每日大盘复盘"
    sheet["A1"].font = Font(size=16, bold=True, color="17365D")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")

    if not bundle:
        _section(sheet, 3, "运行状态")
        sheet.append(["状态", "尚未生成当日大盘复盘"])
        sheet.append(["说明", DISCLAIMER])
        _finish(sheet)
        return

    run = bundle.get("run") or {}
    snapshot = bundle.get("snapshot") or {}
    review = bundle.get("review") or {}
    breadth = snapshot.get("breadth") or {}
    turnover = snapshot.get("turnover") or {}

    _section(sheet, 3, "市场概览")
    sheet.append(["交易日", run.get("trade_date"), "市场方向", run.get("market_direction"), "市场状态", run.get("market_regime"), "复盘状态", run.get("status")])
    sheet.append(["标题", review.get("headline"), "置信度", run.get("confidence"), "证据状态", run.get("search_status"), "数据质量", snapshot.get("data_quality_score")])
    sheet.append(["市场总结", review.get("market_summary")])
    sheet.merge_cells(start_row=sheet.max_row, start_column=2, end_row=sheet.max_row, end_column=8)

    _section(sheet, sheet.max_row + 2, "市场宽度与成交")
    sheet.append(["上涨家数", breadth.get("advancing_count"), "下跌家数", breadth.get("declining_count"), "平盘家数", breadth.get("flat_count"), "有效样本", breadth.get("valid_count")])
    sheet.append(["上涨占比", breadth.get("advancing_ratio"), "等权涨跌幅", breadth.get("equal_weight_return"), "中位数涨跌幅", breadth.get("median_return"), "成交额（亿元）", _to_yi(turnover.get("total_amount"))])
    sheet.append(["宽度总结", review.get("breadth_summary"), "成交总结", review.get("turnover_summary")])
    sheet.merge_cells(start_row=sheet.max_row, start_column=4, end_row=sheet.max_row, end_column=8)

    _append_table(sheet, "主要指数", [
        {
            "指数": item.get("index_name"), "代码": item.get("index_code"),
            "收盘": item.get("close"), "涨跌幅": item.get("change_percent"),
            "趋势": item.get("trend_state"), "数据状态": item.get("data_status"),
        }
        for item in snapshot.get("indices") or []
    ])
    _append_table(sheet, "行业表现", [
        {"排名": item.get("rank"), "行业": item.get("sector_name"), "涨跌幅": item.get("change_percent"), "上涨占比": item.get("advancing_ratio"), "涨停数": item.get("limit_up_count"), "成分数": item.get("member_count")}
        for item in snapshot.get("industries") or []
    ])
    _append_table(sheet, "概念表现", [
        {"排名": item.get("rank"), "概念": item.get("sector_name"), "涨跌幅": item.get("change_percent"), "上涨占比": item.get("advancing_ratio"), "涨停数": item.get("limit_up_count"), "成分数": item.get("member_count")}
        for item in snapshot.get("concepts") or []
    ])
    _append_table(sheet, "主要驱动", [
        {"类型": item.get("driver_type"), "方向": item.get("direction"), "标题": item.get("title"), "影响强度": item.get("impact_strength"), "置信度": item.get("confidence"), "说明": item.get("explanation")}
        for item in bundle.get("drivers") or []
    ])
    _append_table(sheet, "次日条件情景", [
        {"情景": item.get("scenario_type"), "概率": item.get("probability"), "标题": item.get("title"), "条件说明": item.get("description"), "观察项": "；".join(item.get("watch_items") or [])}
        for item in bundle.get("scenarios") or []
    ])
    _append_table(sheet, "公开证据", [
        {"证据ID": item.get("evidence_id"), "来源": item.get("source_name"), "标题": item.get("title"), "发布时间": item.get("publish_time"), "证据等级": item.get("source_tier"), "状态": item.get("status"), "摘要": item.get("summary")}
        for item in bundle.get("evidence") or []
    ])
    _section(sheet, sheet.max_row + 2, "说明")
    sheet.append([DISCLAIMER])
    sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=8)
    _finish(sheet)


def _append_table(sheet, title: str, rows: list[dict[str, Any]]) -> None:
    _section(sheet, sheet.max_row + 2, title)
    if not rows:
        sheet.append(["暂无可用数据"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    header_row = sheet.max_row
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")


def _section(sheet, row: int, title: str) -> None:
    sheet.cell(row=row, column=1, value=title)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    cell = sheet.cell(row=row, column=1)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _finish(sheet) -> None:
    percent_headers = {"涨跌幅", "上涨占比", "等权涨跌幅", "中位数涨跌幅", "概率", "置信度"}
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if isinstance(cell.value, float) and _is_percentage_cell(sheet, cell.row, cell.column, percent_headers):
                cell.number_format = "0.00%"
    widths = [18, 28, 18, 28, 18, 28, 18, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 28


def _to_yi(value: Any) -> float | None:
    return round(float(value) / 100_000_000, 2) if value not in (None, "") else None


def _is_percentage_cell(sheet, row: int, column: int, headers: set[str]) -> bool:
    if column > 1 and sheet.cell(row=row, column=column - 1).value in headers:
        return True
    for previous_row in range(row - 1, 0, -1):
        value = sheet.cell(row=previous_row, column=column).value
        if value in headers:
            return True
        row_values = [sheet.cell(row=previous_row, column=index).value for index in range(1, 9)]
        if not any(item not in (None, "") for item in row_values) or (
            row_values[0] not in (None, "") and all(item in (None, "") for item in row_values[1:])
        ):
            break
    return False
