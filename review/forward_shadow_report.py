from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from database.models import (
    FactorPerformanceHistory, ForwardOutcome, GateCounterfactualRun, GateValueEvaluation,
    ModelVersionComparison, StrategyTimingContract,
)
from review.forward_shadow import HORIZONS, clean_code, sample_status


SHEETS = (
    "01_总览", "02_Baseline_V22_V3", "03_逐股结果", "04_V3_REVIEW分层", "05_Gate价值", "06_Gate反事实",
    "07_因子表现", "08_策略概率表现", "09_LLM黑箱表现", "10_待成熟样本", "11_时间契约", "12_数据口径", "13_运行审计",
)


def _value(value: Any) -> Any:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _append_table(ws, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([_value(x) for x in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 200) + 1)]
        ws.column_dimensions[get_column_letter(col)].width = min(42, max(12, max(map(len, values), default=10) + 2))
    code_columns = {i + 1 for i, title in enumerate(headers) if "股票代码" in title or title == "stock_code"}
    for col in code_columns:
        for row in range(2, ws.max_row + 1):
            value = ws.cell(row, col).value
            if value:
                ws.cell(row, col).value = clean_code(str(value))
                ws.cell(row, col).number_format = "@"


class ForwardShadowWorkbookExporter:
    def __init__(self, session: Session) -> None:
        self.session = session

    def export(self, as_of: date, output_root: Path | str = Path("outputs/forward_shadow")) -> tuple[Path, Path]:
        outcomes = list(self.session.scalars(select(ForwardOutcome).where(ForwardOutcome.entry_trade_date <= as_of).order_by(ForwardOutcome.trade_date, ForwardOutcome.route, ForwardOutcome.stock_code)))
        comparisons = list(self.session.scalars(select(ModelVersionComparison).where(ModelVersionComparison.as_of_date == as_of).order_by(ModelVersionComparison.model_route, ModelVersionComparison.segment, ModelVersionComparison.horizon)))
        gates = list(self.session.scalars(select(GateValueEvaluation).where(GateValueEvaluation.as_of_date == as_of).order_by(GateValueEvaluation.gate_name)))
        counterfactuals = list(self.session.scalars(select(GateCounterfactualRun).where(GateCounterfactualRun.as_of_date == as_of).order_by(GateCounterfactualRun.disabled_gate)))
        factors = list(self.session.scalars(select(FactorPerformanceHistory).where(FactorPerformanceHistory.period_end == as_of).order_by(FactorPerformanceHistory.factor_family)))
        contracts = list(self.session.scalars(select(StrategyTimingContract).where(StrategyTimingContract.trade_date <= as_of).order_by(StrategyTimingContract.trade_date, StrategyTimingContract.stock_code)))
        matured = {h: sum((r.horizon_status_json or {}).get(f"d{h}") == "MATURED" for r in outcomes) for h in HORIZONS}
        pending = [(r, h) for r in outcomes for h in HORIZONS if (r.horizon_status_json or {}).get(f"d{h}") == "PENDING"]
        baseline_formal = [r for r in outcomes if r.route == "TUSHARE_BASELINE_V1" and r.baseline_status == "FORMAL_CANDIDATE"]
        baseline_d1 = [float(r.return_d1) for r in baseline_formal if r.return_d1 is not None]
        wb = Workbook(); wb.remove(wb.active)
        for name in SHEETS:
            wb.create_sheet(name)
        _append_table(wb[SHEETS[0]], ["项目", "值"], [
            ("截至日期", as_of), ("结果记录", len(outcomes)), ("D1成熟", matured[1]), ("D3成熟", matured[3]),
            ("D5成熟", matured[5]), ("D10成熟", matured[10]), ("待成熟Horizon", len(pending)),
            ("样本状态(D3)", sample_status(matured[3])), ("晋级建议", "KEEP_V22_V3_SHADOW"),
            ("严格公平A/B成熟样本", 0), ("严格公平A/B状态", "INSUFFICIENT_SAMPLE"),
            ("Baseline正式候选", len(baseline_formal)), ("被V2.2降级候选", len(baseline_formal)),
            ("降级候选D1平均收益", (sum(baseline_d1) / len(baseline_d1) if baseline_d1 else None)),
        ])
        _append_table(wb[SHEETS[1]], ["路线", "分层", "Horizon", "样本", "成交", "正收益率", "平均收益", "中位收益", "MAE", "MFE", "Profit Factor", "最大亏损", "样本状态", "公平样本", "说明"], [
            (x.model_route, x.segment, x.horizon, x.sample_count, x.fill_count, x.positive_rate, x.average_return, x.median_return,
             x.average_mae, x.average_mfe, x.profit_factor, x.maximum_loss, x.sample_status, x.fair_sample, x.details_json) for x in comparisons])
        _append_table(wb[SHEETS[2]], ["来源运行", "路线", "交易日", "股票代码", "股票名称", "Baseline", "V2.2", "V3", "入场日", "入场价", "成交状态", "D1收益", "D3收益", "D5收益", "D10收益", "D1 MAE", "D1 MFE", "错失盈利", "避免亏损", "未成交质量", "成熟状态", "触发门禁", "绑定门禁", "输入哈希"], [
            (r.source_run_id, r.route, r.trade_date, r.stock_code, r.stock_name, r.baseline_status, r.v2_2_status, r.v3_status,
             r.entry_trade_date, r.entry_price, r.entry_status, r.return_d1, r.return_d3, r.return_d5, r.return_d10,
             r.mae_d1, r.mfe_d1, r.missed_opportunity, r.avoided_loss, r.non_fill_quality,
             r.horizon_status_json, r.all_triggered_gates, r.binding_gate, r.input_hash) for r in outcomes])
        strata = [x for x in comparisons if x.model_route == "V3_REVIEW_STRATIFICATION"]
        _append_table(wb[SHEETS[3]], ["分层", "Horizon", "样本", "成交", "正收益率", "平均收益", "中位收益", "MAE", "MFE", "Profit Factor", "最大亏损", "样本状态", "说明"], [
            (x.segment, x.horizon, x.sample_count, x.fill_count, x.positive_rate, x.average_return, x.median_return,
             x.average_mae, x.average_mfe, x.profit_factor, x.maximum_loss, x.sample_status, x.details_json) for x in strata])
        _append_table(wb[SHEETS[4]], ["Gate", "作用域", "到达", "触发", "阻断", "唯一阻断", "共同阻断", "绑定", "避免亏损", "错失盈利", "局部净价值", "组合边际价值", "误杀率", "拒绝精度", "样本状态"], [
            (x.gate_name, x.gate_scope, x.reached_count, x.triggered_count, x.blocked_count, x.unique_blocked_count,
             x.co_blocked_count, x.binding_count, x.avoided_loss, x.missed_gain, x.local_net_gate_value,
             x.portfolio_marginal_value, x.false_negative_rate, x.reject_precision, x.sample_status) for x in gates])
        _append_table(wb[SHEETS[5]], ["禁用Gate", "Horizon", "净收益变化", "胜率变化", "PF变化", "最大回撤变化", "CVaR变化", "候选数变化", "Top20变化", "避免亏损", "错失盈利", "净价值", "重放状态", "Shapley预留"], [
            (x.disabled_gate, x.evaluation_horizon, x.delta_net_return, x.delta_win_rate, x.delta_profit_factor,
             x.delta_max_drawdown, x.delta_cvar, x.delta_candidate_count, x.delta_top20_membership,
             x.avoided_loss, x.missed_gain, x.net_gate_value, x.replay_status, x.shapley_ready_json) for x in counterfactuals])
        _append_table(wb[SHEETS[6]], ["因子族", "样本", "正贡献", "负贡献", "平均贡献", "中位贡献", "贡献收益相关", "Rank IC", "Top收益", "Bottom收益", "命中率", "D1平均", "D3平均", "D5平均", "MAE", "MFE", "Profit Factor", "PASS贡献", "排名贡献"], [
            (x.factor_family, x.sample_count, x.positive_contribution_count, x.negative_contribution_count,
             x.mean_contribution, x.median_contribution, x.contribution_return_correlation, x.contribution_rank_ic,
             x.top_contribution_return, x.bottom_contribution_return, x.contribution_hit_rate, x.avg_return_d1,
             x.avg_return_d3, x.avg_return_d5, x.avg_drawdown, x.average_mfe, x.profit_factor,
             x.pass_contribution_json, x.rank_contribution_json) for x in factors])
        strategies = [x for x in comparisons if x.model_route == "STRATEGY_PROBABILITY"]
        _append_table(wb[SHEETS[7]], ["策略", "Horizon", "样本", "成交", "正收益率", "平均收益", "中位收益", "Profit Factor", "最大亏损", "样本状态"], [
            (x.segment, x.horizon, x.sample_count, x.fill_count, x.positive_rate, x.average_return, x.median_return,
             x.profit_factor, x.maximum_loss, x.sample_status) for x in strategies])
        _append_table(wb[SHEETS[8]], ["标记", "说明", "结论"], [("OPAQUE_LLM_CONTRIBUTION", "Flash/Pro Prompt未修改；不拆入六个底层因子", "仅允许未来按ADVANCE/WATCH/BLOCK分层评价")])
        _append_table(wb[SHEETS[9]], ["来源运行", "股票代码", "Horizon", "入场日", "当前状态"], [(r.source_run_id, r.stock_code, f"D{h}", r.entry_trade_date, (r.horizon_status_json or {}).get(f"d{h}")) for r, h in pending])
        _append_table(wb[SHEETS[10]], ["合同ID", "股票代码", "交易日", "观察结束", "数据可用", "信号生成", "最早下单", "执行策略", "校验"], [
            (x.id, x.stock_code, x.trade_date, x.observation_end_ts, x.available_at_ts, x.signal_generated_at, x.order_eligible_at,
             x.execution_policy, "PASS" if x.observation_end_ts <= x.available_at_ts <= x.signal_generated_at < x.order_eligible_at else "INVALID_TIMING_CONTRACT") for x in contracts])
        _append_table(wb[SHEETS[11]], ["口径", "值"], [
            ("Timing", "EOD_T_TO_NEXT_OPEN"), ("Entry", "T+1 Open × (1 + buy_slippage)"),
            ("Sensitivity", "Open / +10bp / +20bp / +30bp; 09:30-09:35 VWAP仅合法分钟缓存存在时"),
            ("Exit", "买入后第1/3/5/10个交易日Close"), ("缺失收益", "NULL，禁止填0"),
            ("行情来源", "仅data/cache/tushare本地缓存"), ("公平比较", "相同决策时点、入场日、滑点、可成交规则；当前路线时点不同，不强行合并"),
        ])
        _append_table(wb[SHEETS[12]], ["项目", "值"], [
            ("版本", "FORWARD_SHADOW_EVALUATION_V1"), ("外部API调用", 0), ("LLM调用", 0), ("订单", 0),
            ("Scheduler", False), ("历史推荐重跑", False), ("真实/虚拟交易", 0),
            ("Quant/Flash/Pro", "UNCHANGED_BY_THIS_PHASE"), ("公式错误", 0),
        ])
        output_root = Path(output_root); output_root.mkdir(parents=True, exist_ok=True)
        path = output_root / f"前向Shadow评价_截至_{as_of.isoformat()}.xlsx"
        wb.save(path)
        self.validate(path)
        md = output_root / f"前向Shadow评价_截至_{as_of.isoformat()}.md"
        md.write_text(self._markdown(as_of, outcomes, matured, pending), encoding="utf-8")
        return path.resolve(), md.resolve()

    @staticmethod
    def validate(path: Path) -> None:
        wb = load_workbook(path, data_only=False)
        assert tuple(wb.sheetnames) == SHEETS
        errors = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
        for ws in wb.worksheets:
            assert ws.freeze_panes == "A2"
            assert ws.auto_filter.ref
            for row in ws.iter_rows():
                for cell in row:
                    assert cell.alignment.horizontal == "center" and cell.alignment.vertical == "center" and cell.alignment.wrap_text
                    assert not (isinstance(cell.value, str) and any(x in cell.value for x in errors))
        wb.close()

    @staticmethod
    def _markdown(as_of, outcomes, matured, pending) -> str:
        return (f"# 前向 Shadow 评价（截至 {as_of}）\n\n"
                f"- 结果记录：{len(outcomes)}\n- D1/D3/D5/D10 成熟：{matured[1]}/{matured[3]}/{matured[5]}/{matured[10]}\n"
                f"- 待成熟 Horizon：{len(pending)}\n- D3样本状态：{sample_status(matured[3])}\n"
                "- 晋级建议：KEEP_V22_V3_SHADOW\n- 外部API / LLM / 订单：0 / 0 / 0\n\n"
                "> Baseline 与 Shadow 的实际信号生成时间不同；未到共同合法执行时点前，不输出伪公平 A/B 结论。\n")
