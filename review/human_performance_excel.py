from __future__ import annotations

import statistics
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "17365D"
HEADER_BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F4F7F9"
WHITE = "FFFFFF"
THIN_BLUE = Side(style="thin", color="9CD8F4")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_BORDER = Border(bottom=THIN_BLUE)
RETURN_FORMAT = "[Red]0.00%;[Green]-0.00%;-"


class HumanPerformanceExcelExporter:
    def export(self, output_path: Path, payload: dict[str, Any]) -> dict[str, Any]:
        cohorts = sorted(payload.get("cohorts") or [], key=lambda row: _day(row["selection_trade_date"]))
        portfolio_rows = payload.get("daily") or []
        stock_rows = payload.get("stocks") or []
        member_rows = payload.get("members") or []
        midday = payload.get("midday") or {}
        run = payload.get("run") or {}
        report_label = _report_label(run.get("selection_scope"))
        lookback_value = int(run.get("lookback_value") or 3)
        evaluation_end = _day(run.get("evaluation_end_date") or max(
            (_day(row["evaluation_trade_date"]) for row in stock_rows),
            default=date.today(),
        ))
        if not cohorts:
            raise ValueError("HUMAN_PERFORMANCE_COHORTS_EMPTY")

        workbook = Workbook()
        workbook.remove(workbook.active)
        details: dict[date, dict[str, Any]] = {}
        for cohort in cohorts:
            selection_day = _day(cohort["selection_trade_date"])
            cohort_stocks = [row for row in stock_rows if _day(row["selection_trade_date"]) == selection_day]
            cohort_members = [row for row in member_rows if _day(row["selection_trade_date"]) == selection_day]
            details[selection_day] = _build_detail_sheet(
                workbook,
                cohort,
                cohort_stocks,
                evaluation_end,
                cohort_members,
                report_label,
            )
        summary = _build_summary_sheet(
            workbook,
            cohorts,
            portfolio_rows,
            details,
            evaluation_end,
            lookback_value,
            report_label,
        )
        workbook.move_sheet(summary, offset=-len(cohorts))
        if midday.get("rows"):
            midday_sheet = _build_midday_sheet(workbook, midday, evaluation_end)
            workbook.move_sheet(midday_sheet, offset=-len(cohorts))
        workbook.active = 0
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        actual_output = _save_with_update_fallback(workbook, output_path)
        workbook.close()
        return _verify(actual_output)


class WeeklyStockPerformanceExcelExporter:
    """Export one row per recommended stock for compact weekly comparison."""

    def export(self, output_path: Path, payload: dict[str, Any]) -> dict[str, Any]:
        member_rows = list(payload.get("members") or [])
        stock_rows = list(payload.get("stocks") or [])
        run = payload.get("run") or {}
        scope = str(run.get("selection_scope") or "FINAL_LLM_ONLY").upper()
        if scope == "KEY_CANDIDATES":
            sheet_name = "本周重点候选对比"
            title = "本周重点候选表现对比"
            subtitle = "按股票代码去重并保留本周最早入选记录；包含模型选股、人工选股和共同入选。"
            member_rows = _earliest_member_per_stock(member_rows)
        else:
            sheet_name = "本周模型选股对比"
            title = "本周今日推荐模型选股对比"
            subtitle = "每只股票一行；仅统计今日推荐中的模型选股。"
        evaluation_end = _day(run.get("evaluation_end_date") or date.today())
        evaluation_dates = sorted({_day(row["evaluation_trade_date"]) for row in stock_rows})

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        headers = [
            "推荐日期", "观测起始", "原排名", "股票代码", "股票名称", "来源", "复核分",
            *[f"{_month_day(day)}当日涨跌" for day in evaluation_dates],
            f"截至{_month_day(evaluation_end)}总涨跌",
        ]
        _title(
            sheet,
            title,
            f"{subtitle}观测截止 {evaluation_end.isoformat()}。",
            len(headers),
        )
        _write_headers(sheet, 3, headers)

        daily_by_member: dict[tuple[date, str], dict[date, dict[str, Any]]] = defaultdict(dict)
        for row in stock_rows:
            key = (_day(row["selection_trade_date"]), _display_code(row["stock_code"]))
            daily_by_member[key][_day(row["evaluation_trade_date"])] = row

        members = sorted(
            member_rows,
            key=lambda row: (
                _day(row["selection_trade_date"]),
                _display_rank(row) is None,
                _display_rank(row) or 999999,
                _display_code(row["stock_code"]),
            ),
        )
        total_column = len(headers)
        for row_number, member in enumerate(members, 4):
            selection_day = _day(member["selection_trade_date"])
            code = _display_code(member["stock_code"])
            daily_map = daily_by_member.get((selection_day, code), {})
            baseline_day = min(daily_map) if daily_map else None
            values = [
                selection_day,
                baseline_day,
                _display_rank(member),
                code,
                member.get("stock_name") or code,
                _source_label(member.get("selection_source")),
                _number(member.get("pro_score")),
            ]
            for column, value in enumerate(values, 1):
                sheet.cell(row_number, column, value)
            for offset, day in enumerate(evaluation_dates, 8):
                sheet.cell(row_number, offset, _number(daily_map.get(day, {}).get("daily_return")))
                sheet.cell(row_number, offset).number_format = RETURN_FORMAT
            latest = max(daily_map.values(), key=lambda row: _day(row["evaluation_trade_date"]), default=None)
            sheet.cell(
                row_number,
                total_column,
                _number((latest or {}).get("cumulative_return")) if latest else "待观测",
            )
            if latest:
                sheet.cell(row_number, total_column).number_format = RETURN_FORMAT
            for cell in sheet[row_number]:
                cell.alignment = CENTER
                cell.border = BODY_BORDER
                cell.font = Font(name="Microsoft YaHei", size=10)
            sheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
            sheet.cell(row_number, 2).number_format = "yyyy-mm-dd"
            sheet.cell(row_number, 4).number_format = "@"
            sheet.cell(row_number, 7).number_format = "0.00"
            sheet.row_dimensions[row_number].height = 25

        last_row = 3 + len(members)
        sheet.auto_filter.ref = f"A3:{get_column_letter(total_column)}{last_row}"
        sheet.freeze_panes = "H4"
        widths = [14, 14, 10, 15, 16, 14, 12, *([18] * len(evaluation_dates)), 21]
        _finish_sheet(sheet, widths)
        workbook.active = 0
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        actual_output = _save_with_update_fallback(workbook, output_path)
        workbook.close()
        result = _verify(actual_output)
        result["displayed_stock_count"] = len(members)
        result["deduplicated_by_earliest_selection"] = scope == "KEY_CANDIDATES"
        return result


def _earliest_member_per_stock(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    earliest: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = _display_code(row.get("stock_code"))
        existing = earliest.get(code)
        if existing is None or _day(row["selection_trade_date"]) < _day(existing["selection_trade_date"]):
            earliest[code] = row
    return list(earliest.values())


def _save_with_update_fallback(workbook: Workbook, output_path: Path) -> Path:
    try:
        workbook.save(output_path)
        return output_path
    except PermissionError:
        for suffix in ("_更新版", "_更新版_2", "_更新版_3"):
            candidate = output_path.with_name(f"{output_path.stem}{suffix}{output_path.suffix}")
            try:
                workbook.save(candidate)
                return candidate
            except PermissionError:
                continue
        raise


def _build_summary_sheet(
    workbook: Workbook,
    cohorts: list[dict[str, Any]],
    portfolio_rows: list[dict[str, Any]],
    details: dict[date, dict[str, Any]],
    evaluation_end: date,
    lookback_value: int,
    report_label: str,
):
    label = _lookback_label(lookback_value)
    sheet = workbook.create_sheet(f"近{label}日汇总")
    evaluation_dates = sorted({_day(row["evaluation_trade_date"]) for row in portfolio_rows})
    headers = ["选股日", "观测起始", "观测交易日", *[f"{_month_day(day)}组合涨跌" for day in evaluation_dates], f"截至{_month_day(evaluation_end)}组合总涨跌", "上涨家数", "下跌家数", "数据覆盖率"]
    _title(sheet, f"近{label}日{report_label}", f"逐日展示此前{report_label.removesuffix('复盘')}股票的价格表现，累计涨跌按每日收益复合计算，统计截止 {evaluation_end.isoformat()}。", len(headers))
    _write_headers(sheet, 3, headers)
    portfolio_map = {
        (_day(row["selection_trade_date"]), _day(row["evaluation_trade_date"])): row
        for row in portfolio_rows
    }
    for row_number, cohort in enumerate(cohorts, 4):
        selection_day = _day(cohort["selection_trade_date"])
        applicable_dates = [day for day in evaluation_dates if (selection_day, day) in portfolio_map]
        values: list[Any] = [
            selection_day,
            _day(cohort["baseline_trade_date"]) if cohort.get("baseline_trade_date") else None,
            len(applicable_dates),
        ]
        values.extend(
            _number(portfolio_map.get((selection_day, day), {}).get("daily_return"))
            for day in evaluation_dates
        )
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
        total_column = 4 + len(evaluation_dates)
        latest = portfolio_map.get((selection_day, applicable_dates[-1])) if applicable_dates else None
        sheet.cell(row_number, total_column, _number((latest or {}).get("cumulative_return")))
        detail = details[selection_day]
        sheet.cell(row_number, total_column + 1, detail["positive_count"])
        sheet.cell(row_number, total_column + 2, detail["negative_count"])
        sheet.cell(row_number, total_column + 3, _number(cohort.get("coverage_ratio")))
        for cell in sheet[row_number]:
            cell.alignment = CENTER
            cell.border = BODY_BORDER
            cell.font = Font(name="Microsoft YaHei", size=10)
        sheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
        sheet.cell(row_number, 2).number_format = "yyyy-mm-dd"
        for column in range(4, total_column + 1):
            sheet.cell(row_number, column).number_format = RETURN_FORMAT
        sheet.cell(row_number, total_column + 3).number_format = "0.00%"
        sheet.row_dimensions[row_number].height = 24

    last_row = 3 + len(cohorts)
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{last_row}"
    sheet.freeze_panes = "D4"
    _signed_format(sheet, 4, 4 + len(evaluation_dates), 4, last_row)
    widths = [14, 14, 12, *([18] * len(evaluation_dates)), 22, 12, 12, 14]
    _finish_sheet(sheet, widths)
    return sheet


def _build_detail_sheet(
    workbook: Workbook,
    cohort: dict[str, Any],
    stock_rows: list[dict[str, Any]],
    evaluation_end: date,
    member_rows: list[dict[str, Any]] | None = None,
    report_label: str = "推荐复盘",
) -> dict[str, Any]:
    selection_day = _day(cohort["selection_trade_date"])
    sheet_name = f"{selection_day.month}月{selection_day.day}日选股复盘"
    sheet = workbook.create_sheet(sheet_name)
    evaluation_dates = sorted({_day(row["evaluation_trade_date"]) for row in stock_rows})
    headers = ["原排名", "股票代码", "股票名称", "来源", *[f"{_month_day(day)}当日涨跌" for day in evaluation_dates], f"截至{_month_day(evaluation_end)}总涨跌"]
    _title(
        sheet,
        f"{selection_day.month}月{selection_day.day}日{report_label}",
        f"选股日：{selection_day.isoformat()}    观测截止：{evaluation_end.isoformat()}    当前观测：{len(evaluation_dates)}个交易日",
        len(headers),
    )
    stock_map: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in member_rows or []:
        stock_map[_display_code(row["stock_code"])].append(row)
    for row in stock_rows:
        stock_map[_display_code(row["stock_code"])].append(row)
    members = sorted(
        (rows[0] for rows in stock_map.values()),
        key=lambda row: (_display_rank(row) is None, _display_rank(row) or 999999, row["stock_code"]),
    )
    total_column = len(headers)
    last_row = 4 + max(1, len(members))

    sheet.cell(3, 1, "选股数量")
    sheet.cell(3, 2, f"=COUNTA(B5:B{last_row})")
    sheet.cell(3, 3, "截至今日平均")
    latest_returns = []
    for rows in stock_map.values():
        evaluated = [row for row in rows if row.get("evaluation_trade_date") and row.get("cumulative_return") is not None]
        if evaluated:
            latest_returns.append(float(max(evaluated, key=lambda row: _day(row["evaluation_trade_date"]))["cumulative_return"]))
    sheet.cell(3, 4, statistics.fmean(latest_returns) if latest_returns else "等待下一交易日")
    sheet.cell(3, 5, "上涨 / 下跌")
    if total_column > 6:
        sheet.merge_cells(start_row=3, start_column=6, end_row=3, end_column=total_column)
    total_range = f"{get_column_letter(total_column)}5:{get_column_letter(total_column)}{last_row}"
    sheet.cell(3, 6, f'=COUNTIF({total_range},">0")&" / "&COUNTIF({total_range},"<0")')
    for cell in sheet[3]:
        if isinstance(cell, MergedCell):
            continue
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
        cell.alignment = CENTER
        cell.font = Font(name="Microsoft YaHei", bold=cell.column in {1, 3, 5}, color=NAVY if cell.column in {1, 2} else "000000")
    sheet.cell(3, 4).number_format = RETURN_FORMAT
    sheet.row_dimensions[3].height = 30

    _write_headers(sheet, 4, headers)
    if not members:
        sheet.cell(5, 1, "暂无推荐" if report_label == "今日推荐复盘" else "暂无重点候选")
        for cell in sheet[5]:
            cell.alignment = CENTER
            cell.border = BODY_BORDER
            cell.font = Font(name="Microsoft YaHei", size=10)
        sheet.row_dimensions[5].height = 24
    for row_number, member in enumerate(members, 5):
        code = _display_code(member["stock_code"])
        daily_map = {
            _day(row["evaluation_trade_date"]): row
            for row in stock_map[code]
            if row.get("evaluation_trade_date")
        }
        base_values = [
            _display_rank(member),
            code,
            member.get("stock_name") or code,
            _source_label(member.get("selection_source")),
        ]
        for column, value in enumerate(base_values, 1):
            sheet.cell(row_number, column, value)
        for offset, day in enumerate(evaluation_dates, 5):
            sheet.cell(row_number, offset, _number(daily_map.get(day, {}).get("daily_return")))
            sheet.cell(row_number, offset).number_format = RETURN_FORMAT
        latest_row = max(daily_map.values(), key=lambda row: _day(row["evaluation_trade_date"]), default=None)
        sheet.cell(row_number, total_column, _number((latest_row or {}).get("cumulative_return")))
        sheet.cell(row_number, total_column).number_format = RETURN_FORMAT
        for cell in sheet[row_number]:
            cell.alignment = CENTER
            cell.border = BODY_BORDER
            cell.font = Font(name="Microsoft YaHei", size=10)
        sheet.cell(row_number, 2).number_format = "@"
        sheet.row_dimensions[row_number].height = 24

    sheet.auto_filter.ref = f"A4:{get_column_letter(total_column)}{last_row}"
    sheet.freeze_panes = "E5"
    _signed_format(sheet, 5, total_column, 5, last_row)
    widths = [10, 16, 16, 14, *([19] * len(evaluation_dates)), 22]
    _finish_sheet(sheet, widths)
    return {
        "sheet_name": sheet_name,
        "total_column": total_column,
        "last_row": last_row,
        "evaluation_dates": evaluation_dates,
        "positive_count": sum(value > 0 for value in latest_returns),
        "negative_count": sum(value < 0 for value in latest_returns),
    }


def _build_midday_sheet(workbook: Workbook, midday: dict[str, Any], evaluation_end: date):
    sheet = workbook.create_sheet(f"{evaluation_end.month}月{evaluation_end.day}日午盘监测")
    headers = [
        "午盘排名", "股票代码", "股票名称", "午盘结论", "模型复核分",
        "午盘参考价", "当日收盘价", "午盘至收盘涨跌", "观察状态",
    ]
    _title(
        sheet,
        f"{evaluation_end.month}月{evaluation_end.day}日午盘推荐监测",
        f"午盘运行：{midday.get('decision_time') or '11:30后'}    监测截止：{evaluation_end.isoformat()}",
        len(headers),
    )
    _write_headers(sheet, 3, headers)
    rows = list(midday.get("rows") or [])
    for row_number, row in enumerate(rows, 4):
        values = [
            row.get("rank"), row.get("stock_code"), row.get("stock_name"),
            row.get("action"), _number(row.get("pro_score")),
            _number(row.get("reference_price")), _number(row.get("close_price")),
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
        if row.get("reference_price") and row.get("close_price"):
            sheet.cell(row_number, 8, float(row["close_price"]) / float(row["reference_price"]) - 1)
            status = "收盘强于午盘参考" if float(row["close_price"]) >= float(row["reference_price"]) else "收盘弱于午盘参考"
        else:
            status = "行情数据待补"
        sheet.cell(row_number, 9, status)
        for cell in sheet[row_number]:
            cell.alignment = CENTER
            cell.border = BODY_BORDER
            cell.font = Font(name="Microsoft YaHei", size=10)
        sheet.cell(row_number, 2).number_format = "@"
        sheet.cell(row_number, 5).number_format = "0.00"
        sheet.cell(row_number, 6).number_format = "0.00"
        sheet.cell(row_number, 7).number_format = "0.00"
        sheet.cell(row_number, 8).number_format = RETURN_FORMAT
        sheet.row_dimensions[row_number].height = 26

    last_row = 3 + len(rows)
    sheet.auto_filter.ref = f"A3:I{last_row}"
    sheet.freeze_panes = "D4"
    _signed_format(sheet, 8, 8, 4, last_row)
    _finish_sheet(sheet, [11, 16, 16, 18, 14, 15, 15, 19, 22])
    return sheet


def _title(sheet, title: str, subtitle: str, last_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(1, 1, title)
    sheet.cell(2, 1, subtitle)
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
    sheet.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet.cell(2, 1).font = Font(name="Microsoft YaHei", size=10, color=NAVY)
    sheet.cell(1, 1).alignment = CENTER
    sheet.cell(2, 1).alignment = CENTER
    sheet.row_dimensions[1].height = 34
    sheet.row_dimensions[2].height = 28


def _write_headers(sheet, row_number: int, headers: list[str]) -> None:
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row_number, column, header)
        cell.fill = PatternFill("solid", fgColor=HEADER_BLUE)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = CENTER
    sheet.row_dimensions[row_number].height = 34


def _signed_format(sheet, first_column: int, last_column: int, first_row: int, last_row: int) -> None:
    # RETURN_FORMAT already renders A-share gains red and losses green.
    # Avoid extra conditional rules because some Excel-compatible renderers paint the whole cell.
    return None


def _finish_sheet(sheet, widths: list[int]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _compound_formula(cells: list[str]) -> str:
    if not cells:
        return ""
    if len(cells) == 1:
        return f"={cells[0]}"
    return "=" + "*".join(f"(1+{cell})" for cell in cells) + "-1"


def _source_label(value: Any) -> str:
    return {"LLM": "模型选股", "MANUAL": "人工选股", "BOTH": "模型+人工"}.get(str(value or "").upper(), str(value or "未知"))


def _display_rank(row: dict[str, Any]) -> Any:
    return row.get("pro_rank") if row.get("pro_rank") is not None else row.get("quant_rank")


def _month_day(value: date) -> str:
    return f"{value.month}月{value.day}日"


def _lookback_label(value: int) -> str:
    return {3: "三", 5: "五", 7: "七"}.get(value, str(value))


def _report_label(selection_scope: Any) -> str:
    scope = str(selection_scope or "").upper()
    if scope == "KEY_CANDIDATES":
        return "重点候选复盘"
    if scope == "FINAL_LLM_ONLY":
        return "今日推荐模型选股复盘"
    return "今日推荐复盘"


def _day(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value))


def _number(value: Any) -> float | None:
    return float(value) if value is not None else None


def _display_code(value: Any) -> str:
    return str(value or "").split(".", 1)[0].zfill(6)


def _verify(output_path: Path) -> dict[str, Any]:
    workbook = load_workbook(output_path, data_only=False)
    formula_count = 0
    for sheet in workbook.worksheets:
        if sheet.tables:
            raise ValueError(f"HUMAN_PERFORMANCE_TABLE_OBJECT_FOUND:{sheet.title}")
        for row_number in range(1, sheet.max_row + 1):
            if not any(not isinstance(cell, MergedCell) and cell.value is not None for cell in sheet[row_number]):
                raise ValueError(f"HUMAN_PERFORMANCE_BLANK_ROW:{sheet.title}:{row_number}")
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                if cell.alignment.horizontal != "center" or cell.alignment.vertical != "center":
                    raise ValueError(f"HUMAN_PERFORMANCE_ALIGNMENT:{sheet.title}:{cell.coordinate}")
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                if isinstance(cell.value, str) and any(token in cell.value for token in ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?")):
                    raise ValueError(f"HUMAN_PERFORMANCE_FORMULA_ERROR:{sheet.title}:{cell.coordinate}")
    result = {
        "output": str(output_path),
        "size_bytes": output_path.stat().st_size,
        "sheets": workbook.sheetnames,
        "formula_count": formula_count,
        "blank_rows": 0,
        "table_objects": 0,
        "center_alignment_verified": True,
    }
    workbook.close()
    return result
