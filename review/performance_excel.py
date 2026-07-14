from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


DISCLAIMER = "本统计反映选股后的市场价格表现，不代表实际成交收益或交易建议。"


class PerformanceExcelExporter:
    def export(self, output_path: Path, payload: dict[str, Any]) -> dict[str, Any]:
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets = (
            ("01_选股日汇总", payload.get("cohorts") or []),
            ("02_组合每日涨跌", payload.get("daily") or []),
            ("03_个股收益明细", payload.get("stocks") or []),
            ("04_口径与异常", payload.get("methodology") or {}),
        )
        for index, (name, raw_rows) in enumerate(sheets, start=1):
            rows = _normalize_rows(raw_rows)
            _write_sheet(workbook, name, rows, index)
        workbook.active = 0
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        workbook.close()
        content = output_path.read_bytes()
        return {
            "output_path": str(output_path),
            "size": len(content),
            "sha256": hashlib.sha256(content).hexdigest(),
            "sheet_count": 4,
            "status": "SUCCESS",
        }


def _normalize_rows(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, list):
        return [row for row in value if isinstance(row, dict)]
    if isinstance(value, dict):
        return [{"item": key, "value": item} for key, item in value.items()]
    return []


def _write_sheet(workbook: Workbook, name: str, rows: list[dict[str, Any]], index: int) -> None:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    headers = list(rows[0]) if rows else ["状态"]
    last_column = get_column_letter(len(headers))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A1"] = name
    sheet["A2"] = DISCLAIMER
    sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    sheet["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["A2"].font = Font(color="7F6000")
    for cell in (sheet["A1"], sheet["A2"]):
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column, header in enumerate(headers, start=1):
        sheet.cell(4, column, header)
    body = rows or [{headers[0]: "暂无数据"}]
    for row_index, row in enumerate(body, start=5):
        for column, header in enumerate(headers, start=1):
            sheet.cell(row_index, column, _cell_value(row.get(header)))
    last_row = 4 + len(body)
    for row in sheet.iter_rows(min_row=4, max_row=last_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True, color="17365D")

    table = Table(displayName=f"SelectionPerformanceTable{index}", ref=f"A4:{last_column}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A5"
    for column, header in enumerate(headers, start=1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = 28 if any(word in header.lower() for word in ("summary", "status", "说明", "口径", "异常")) else 16
        if "stock_code" in header.lower():
            for row_index in range(5, last_row + 1):
                sheet.cell(row_index, column).number_format = "@"
        if any(word in header.lower() for word in ("return", "drawdown", "win_rate", "coverage_ratio", "position_percent")):
            for row_index in range(5, last_row + 1):
                sheet.cell(row_index, column).number_format = "0.00%"
            if "return" in header.lower() or "drawdown" in header.lower():
                target = f"{letter}5:{letter}{last_row}"
                sheet.conditional_formatting.add(target, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="C00000")))
                sheet.conditional_formatting.add(target, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="008000")))


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, default=str)
