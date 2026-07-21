from __future__ import annotations

import hashlib
import json
import statistics
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from database.models.entry_timing_v2 import EntryTimingV2Result
from database.models.performance import SelectionCohort
from database.models.quant_run import QuantRankResult, QuantRun
from database.models.validation import ModelValidationSample, ProCandidateReview, ProResumeRun
from entry_timing.service import _rows_by_code
from entry_timing.service_v2 import EntryTimingV2ShadowService, _result_dict


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
STRATEGIES = ("TREND_BREAKOUT", "STRONG_PULLBACK", "SECTOR_RESONANCE", "OVERSOLD_REBOUND", "UNCLASSIFIED", "DATA_INSUFFICIENT")


class HistoricalEntryTimingV2Validator:
    def __init__(self, session, *, cache_root: Path) -> None:
        self.session = session
        self.cache_root = Path(cache_root)

    def run(self, start_date: date, end_date: date, output: Path) -> dict[str, Any]:
        service = EntryTimingV2ShadowService(self.session, cache_root=self.cache_root)
        cohorts = list(self.session.scalars(select(SelectionCohort).where(
            SelectionCohort.selection_trade_date >= start_date,
            SelectionCohort.selection_trade_date <= end_date,
        ).order_by(SelectionCohort.selection_trade_date)))
        if not cohorts: raise ValueError("V2_HISTORICAL_COHORTS_NOT_FOUND")
        hashes_before = self._global_source_hashes(cohorts)
        run_ids = [service.run(day.selection_trade_date, quant_run_id=day.quant_run_id, candidate_mode="HISTORICAL_CANDIDATES", force_shadow=True)["run_id"] for day in cohorts]
        rows = list(self.session.scalars(select(EntryTimingV2Result).where(EntryTimingV2Result.run_id.in_(run_ids)).order_by(EntryTimingV2Result.trade_date, EntryTimingV2Result.quant_rank)))
        records = []
        for row in rows:
            item = _result_dict(row)
            item.update(self._outcome(row.trade_date, row.stock_code, end_date))
            records.append(item)
        hashes_after = self._global_source_hashes(cohorts)
        if hashes_before != hashes_after: raise RuntimeError("V2_HISTORICAL_SOURCE_HASH_CHANGED")
        ai_all = [row for row in records if row["pool_type"] == "AI_POOL"]
        manual_all = [row for row in records if row["pool_type"] != "AI_POOL"]
        missing_forward = [row for row in ai_all if row.get("cumulative_return") is None]
        missing_features = [row for row in ai_all if row["classification_status"] == "DATA_INSUFFICIENT"]
        eligible = [row for row in ai_all if row.get("cumulative_return") is not None and row["classification_status"] != "DATA_INSUFFICIENT"]
        original = eligible
        v1 = [row for row in eligible if row["admission_status_v1"] == "PASS"]
        v2 = [row for row in eligible if row["admission_status_v2"] == "PASS"]
        versions = {"ORIGINAL": _all_metrics(original), "ENTRY_TIMING_V1": _all_metrics(v1), "ENTRY_TIMING_V2_1": _all_metrics(v2)}
        removed_winners = [row for row in eligible if row["admission_status_v2"] != "PASS" and row["cumulative_return"] > 0]
        removed_losers = [row for row in eligible if row["admission_status_v2"] != "PASS" and row["cumulative_return"] <= 0]
        retained_winners = [row for row in v2 if row["cumulative_return"] > 0]
        retained_losers = [row for row in v2 if row["cumulative_return"] <= 0]
        daily = self._daily(records, start_date, end_date)
        strategy_groups = _groups(ai_all, "strategy_id", STRATEGIES)
        emotion_groups = _groups(ai_all, "market_emotion_state", ("GREEN", "YELLOW", "RED", "DATA_INSUFFICIENT"))
        regime_values = sorted({row["market_regime"] for row in ai_all})
        regime_groups = _groups(ai_all, "market_regime", regime_values)
        comparison = {
            "historical_period": f"{start_date.isoformat()} to {end_date.isoformat()}",
            "versions": versions,
            "fair_sample_intersection": {
                "original_total": len(ai_all), "v1_total": sum(row["admission_status_v1"] == "PASS" for row in ai_all),
                "v2_total": sum(row["admission_status_v2"] == "PASS" for row in ai_all),
                "common_eligible_sample": len(eligible), "excluded_due_missing_forward_data": len(missing_forward),
                "excluded_due_missing_features": len(missing_features), "comparison_uses_identical_input_rows": True,
                "return_definition": "close-to-close compounded pct_chg", "fee_basis": "gross return; identical for all versions",
            },
            "market_emotion_distribution": dict(Counter(row["market_emotion_state"] for row in ai_all)),
            "market_regime_distribution": dict(Counter(row["market_regime"] for row in ai_all)),
            "strategy_distribution": {name: sum((row["strategy_id"] if row["classification_status"] != "DATA_INSUFFICIENT" else "DATA_INSUFFICIENT") == name for row in ai_all) for name in STRATEGIES},
            "classification_coverage": (len(ai_all) - len(missing_features)) / len(ai_all) if ai_all else 0,
            "admission_distribution_v1": dict(Counter(row["admission_status_v1"] for row in ai_all)),
            "admission_distribution_v2": dict(Counter(row["admission_status_v2"] for row in ai_all)),
            "daily_candidate_counts": daily,
            "empty_pool_days": [row["trade_date"] for row in daily if row["v2_pass"] == 0],
            "removed_losing_stocks": removed_losers, "removed_winning_stocks": removed_winners,
            "retained_losing_stocks": retained_losers, "retained_winning_stocks": retained_winners,
            "performance_by_strategy": strategy_groups, "performance_by_emotion_state": emotion_groups,
            "performance_by_market_regime": regime_groups,
            "manual_challenge_pool": {"distribution": dict(Counter(row["admission_status_v2"] for row in manual_all)), "strategy_distribution": dict(Counter(row["strategy_id"] for row in manual_all)), "performance": _all_metrics([row for row in manual_all if row.get("cumulative_return") is not None])},
            "source_hashes_before": hashes_before, "source_hashes_after": hashes_after,
            "quant_hash_unchanged": hashes_before["quant"] == hashes_after["quant"],
            "flash_hash_unchanged": hashes_before["flash"] == hashes_after["flash"],
            "pro_hash_unchanged": hashes_before["pro"] == hashes_after["pro"],
            "llm_calls": 0, "external_api_calls": 0, "order_creation_count": 0,
            "grid_search_performed": False, "production_profile_changed": False,
        }
        comparison["promotion_recommendation"] = _promotion(comparison)
        comparison["reproducible_hash"] = _hash({key: value for key, value in comparison.items() if key not in {"removed_losing_stocks", "removed_winning_stocks", "retained_losing_stocks", "retained_winning_stocks"}})
        excel = EntryTimingV2Excel().export(output, records, comparison)
        return {**comparison, "run_ids": run_ids, "excel": excel}

    def _outcome(self, selection_date: date, code: str, end_date: date) -> dict[str, Any]:
        daily_dir = self.cache_root / "trade_date" / "daily"
        paths = [path for path in sorted(daily_dir.glob("*.json")) if f"{selection_date:%Y%m%d}" < path.stem <= f"{end_date:%Y%m%d}"]
        observations = []
        wealth = 1.0; peak = 1.0; max_drawdown = 0.0; mae = 0.0; mfe = 0.0
        for path in paths:
            row = _rows_by_code(path).get(code)
            if not row or row.get("pct_chg") is None: continue
            pre_close = _number(row.get("pre_close")); high = _number(row.get("high")); low = _number(row.get("low"))
            high_excursion = wealth * (high / pre_close) - 1 if pre_close and high is not None else None
            low_excursion = wealth * (low / pre_close) - 1 if pre_close and low is not None else None
            if high_excursion is not None: mfe = max(mfe, high_excursion)
            if low_excursion is not None: mae = min(mae, low_excursion)
            daily_return = float(row["pct_chg"]) / 100
            wealth *= 1 + daily_return; peak = max(peak, wealth); max_drawdown = min(max_drawdown, wealth / peak - 1)
            observations.append({"trade_date": path.stem, "daily_return": daily_return, "cumulative_return": wealth - 1})
        return {
            "d1_return": observations[0]["cumulative_return"] if len(observations) >= 1 else None,
            "d3_return": observations[2]["cumulative_return"] if len(observations) >= 3 else None,
            "d5_return": observations[4]["cumulative_return"] if len(observations) >= 5 else None,
            "cumulative_return": wealth - 1 if observations else None, "holding_days": len(observations),
            "mae": mae if observations else None, "mfe": mfe if observations else None,
            "max_drawdown": max_drawdown if observations else None, "daily_observations": observations,
        }

    @staticmethod
    def _daily(records, start_date, end_date):
        output = []
        for day in sorted({row["trade_date"] for row in records}):
            ai = [row for row in records if row["trade_date"] == day and row["pool_type"] == "AI_POOL"]
            output.append({"trade_date": day, "original": len(ai), "v1_pass": sum(row["admission_status_v1"] == "PASS" for row in ai), "v2_pass": sum(row["admission_status_v2"] == "PASS" for row in ai), "emotion_state": ai[0]["market_emotion_state"] if ai else None, "market_regime": ai[0]["market_regime"] if ai else None})
        return output

    def _global_source_hashes(self, cohorts):
        quant=[]; flash=[]; pro=[]
        for cohort in cohorts:
            run=self.session.scalar(select(QuantRun).where(QuantRun.run_id==cohort.quant_run_id)); rows=list(self.session.scalars(select(QuantRankResult).where(QuantRankResult.quant_run_id==cohort.quant_run_id).order_by(QuantRankResult.rank)))
            quant.append([cohort.selection_trade_date, cohort.quant_run_id, run.request_hash, [[r.stock_code,r.rank,str(r.total_score),str(r.technical_score),str(r.capital_score),str(r.emotion_score),str(r.momentum_score),str(r.risk_score)] for r in rows]])
            fs=list(self.session.scalars(select(ModelValidationSample).where(ModelValidationSample.validation_run_id==cohort.flash_run_id).order_by(ModelValidationSample.rank))); flash.append([cohort.selection_trade_date,cohort.flash_run_id,[[r.stock_code,r.rank,r.screening_result] for r in fs]])
            pr=self.session.scalar(select(ProResumeRun).where(ProResumeRun.run_id==cohort.pro_run_id)); ps=list(self.session.scalars(select(ProCandidateReview).where(ProCandidateReview.pro_resume_run_id==cohort.pro_run_id).order_by(ProCandidateReview.pro_rank))); pro.append([cohort.selection_trade_date,cohort.pro_run_id,getattr(pr,"candidate_set_hash",None),[[r.stock_code,r.pro_rank,str(r.pro_score),r.priority,r.final_summary] for r in ps]])
        return {"quant":_baseline_hash(quant),"flash":_baseline_hash(flash),"pro":_baseline_hash(pro)}


class EntryTimingV2Excel:
    SHEETS = ("01_总体对比", "02_逐股V1_V2明细", "03_策略分组", "04_市场情绪分组", "05_被过滤股票", "06_保留股票", "07_人工挑战池", "08_口径与异常")

    def export(self, output: Path, records: list[dict[str, Any]], report: dict[str, Any]) -> dict[str, Any]:
        workbook = Workbook(); workbook.remove(workbook.active)
        comparison_rows=[]
        metric_labels = [("count","候选数"),("observed_count","完整观测数"),("win_rate","正收益比例"),("average_return","平均收益"),("median_return","中位收益"),("average_win","平均盈利"),("average_loss","平均亏损"),("profit_loss_ratio","盈亏比"),("profit_factor","Profit Factor"),("maximum_loss","最大亏损"),("mae","MAE"),("mfe","MFE"),("max_drawdown","最大回撤")]
        for key,label in metric_labels:
            comparison_rows.append({"指标":label,"Original":report["versions"]["ORIGINAL"].get(key),"Entry Timing V1":report["versions"]["ENTRY_TIMING_V1"].get(key),"Entry Timing V2.1":report["versions"]["ENTRY_TIMING_V2_1"].get(key)})
        _sheet(workbook,self.SHEETS[0],comparison_rows)
        _sheet(workbook,self.SHEETS[1],[self._detail(row) for row in records])
        _sheet(workbook,self.SHEETS[2],_group_rows(report["performance_by_strategy"],"策略"))
        emotion_rows = _group_rows(report["performance_by_emotion_state"], "分组", "情绪状态")
        regime_rows = _group_rows(report["performance_by_market_regime"], "分组", "市场状态")
        _sheet(workbook,self.SHEETS[3],emotion_rows+regime_rows)
        _sheet(workbook,self.SHEETS[4],[self._detail(row) for row in records if row["pool_type"]=="AI_POOL" and row["admission_status_v2"]!="PASS"])
        _sheet(workbook,self.SHEETS[5],[self._detail(row) for row in records if row["pool_type"]=="AI_POOL" and row["admission_status_v2"]=="PASS"])
        _sheet(workbook,self.SHEETS[6],[self._detail(row) for row in records if row["pool_type"]!="AI_POOL"])
        fair=report["fair_sample_intersection"]
        note_labels = {
            "original_total": "原始候选总数", "v1_total": "V1通过数", "v2_total": "V2.1通过数",
            "common_eligible_sample": "共同有效样本", "excluded_due_missing_forward_data": "因后续行情缺失排除",
            "excluded_due_missing_features": "因特征缺失排除", "comparison_uses_identical_input_rows": "比较使用相同行",
            "return_definition": "收益口径", "fee_basis": "费用口径",
        }
        notes=[{"项目":note_labels.get(key,key),"值":_display_note(value)} for key,value in fair.items()]
        for version, metrics in report["versions"].items():
            for horizon in ("d1", "d3", "d5"):
                notes.append({
                    "项目": f"{version} {horizon.upper()}覆盖",
                    "值": f'{metrics[horizon]["observed_count"]}/{metrics[horizon]["count"]}',
                })
        notes += [{"项目":"LLM调用","值":0},{"项目":"外部API调用","值":0},{"项目":"订单创建","值":0},{"项目":"网格搜索","值":"否"},{"项目":"Quant Hash未变化","值":_display_note(report["quant_hash_unchanged"])},{"项目":"Flash Hash未变化","值":_display_note(report["flash_hash_unchanged"])},{"项目":"Pro Hash未变化","值":_display_note(report["pro_hash_unchanged"])},{"项目":"Promotion","值":report["promotion_recommendation"]},{"项目":"可复现Hash","值":report["reproducible_hash"]}]
        _sheet(workbook,self.SHEETS[7],notes)
        output.parent.mkdir(parents=True,exist_ok=True); workbook.save(output); workbook.close()
        return _verify(output,self.SHEETS)

    @staticmethod
    def _detail(row):
        return {"选股日期":row["trade_date"],"股票代码":row["stock_code"].split(".",1)[0],"股票名称":row.get("stock_name"),"候选来源":row["pool_type"],"Quant排名":row.get("quant_rank"),"Quant分":row.get("quant_score"),"Strategy ID":row["strategy_id"],"Strategy Fit":row["strategy_fit_score"],"分类状态":row["classification_status"],"Market Emotion":row.get("market_emotion_score"),"Emotion State":row["market_emotion_state"],"Market Regime":row["market_regime"],"Entry Timing V1":row["entry_timing_v1_score"],"Entry Timing V2":row.get("entry_timing_v2_score"),"Admission Ranking V2":row.get("admission_ranking_score_v2"),"Admission V1":row["admission_status_v1"],"Admission V2":row["admission_status_v2"],"市场门禁":row["market_gate_status"],"高位风险":"、".join(row["risk_flags"]) or "无","Block原因":"、".join(row["block_reasons"]) or "无","Review原因":"、".join(row["review_reasons"]) or "无","D+1收益":row.get("d1_return"),"D+3收益":row.get("d3_return"),"D+5收益":row.get("d5_return"),"截至期末收益":row.get("cumulative_return"),"MAE":row.get("mae"),"MFE":row.get("mfe"),"最大回撤":row.get("max_drawdown"),"数据覆盖":json.dumps(row.get("data_coverage"),ensure_ascii=False),"版本":row["version"]}


def _metrics(rows, key="cumulative_return"):
    values=[float(row[key]) for row in rows if row.get(key) is not None]
    if not values:return {"count":len(rows),"observed_count":0,"win_rate":None,"average_return":None,"median_return":None,"average_win":None,"average_loss":None,"profit_loss_ratio":None,"profit_factor":None,"maximum_loss":None}
    wins=[v for v in values if v>0]; losses=[v for v in values if v<0]
    avg_win=statistics.fmean(wins) if wins else None; avg_loss=statistics.fmean(losses) if losses else None
    return {"count":len(rows),"observed_count":len(values),"win_rate":len(wins)/len(values),"average_return":statistics.fmean(values),"median_return":statistics.median(values),"average_win":avg_win,"average_loss":avg_loss,"profit_loss_ratio":avg_win/abs(avg_loss) if avg_win is not None and avg_loss else None,"profit_factor":sum(wins)/abs(sum(losses)) if wins and losses else None,"maximum_loss":min(values)}


def _all_metrics(rows):
    result=_metrics(rows)
    result.update({"d1":_metrics(rows,"d1_return"),"d3":_metrics(rows,"d3_return"),"d5":_metrics(rows,"d5_return")})
    for key in ("mae", "mfe"):
        values=[float(row[key]) for row in rows if row.get(key) is not None]
        result[key]=statistics.fmean(values) if values else None
    drawdowns=[float(row["max_drawdown"]) for row in rows if row.get("max_drawdown") is not None]
    result["max_drawdown"]=min(drawdowns) if drawdowns else None
    result["tail_losses"]={threshold:sum(row.get("cumulative_return") is not None and row["cumulative_return"]<threshold for row in rows) for threshold in (-.05,-.10,-.15,-.20)}
    result["coverage_ratio"]=result["observed_count"]/len(rows) if rows else 0
    return result


def _groups(rows,key,values):
    output={}
    for value in values:
        group=[row for row in rows if ("DATA_INSUFFICIENT" if key=="strategy_id" and row["classification_status"]=="DATA_INSUFFICIENT" else row.get(key))==value]
        metrics=_all_metrics([row for row in group if row.get("cumulative_return") is not None]); output[value]={"candidate_count":len(group),"PASS":sum(row["admission_status_v2"]=="PASS" for row in group),"REVIEW":sum(row["admission_status_v2"]=="REVIEW" for row in group),"BLOCK":sum(row["admission_status_v2"]=="BLOCK" for row in group),**metrics,"sample_status":"SUFFICIENT" if metrics["observed_count"]>=5 else "INSUFFICIENT_SAMPLE"}
    return output


def _group_rows(groups,label,dimension=None):
    rows=[]
    for key,data in groups.items():
        prefix={label:key} if dimension is None else {"分组维度":dimension,label:key}
        rows.append({**prefix,**{name:value for name,value in data.items() if name not in {"d1","d3","d5","tail_losses"}}})
    return rows


def _promotion(report):
    v1=report["versions"]["ENTRY_TIMING_V1"]; v2=report["versions"]["ENTRY_TIMING_V2_1"]
    if not v2["observed_count"]: return "KEEP_V2_SHADOW"
    risk_better=v2["maximum_loss"] is not None and v1["maximum_loss"] is not None and v2["maximum_loss"]>=v1["maximum_loss"]
    return "READY_FOR_FORWARD_SHADOW" if risk_better and v2["observed_count"]>=max(10,int(v1["observed_count"]*.25)) and (v2["average_return"] or -1)>=(v1["average_return"] or -1) and (v2["profit_factor"] or 0)>=(v1["profit_factor"] or 0)*.85 else "KEEP_V2_SHADOW"


def _sheet(workbook,title,rows):
    sheet=workbook.create_sheet(title); headers=list(rows[0]) if rows else ["暂无数据"]; sheet.append(headers)
    for row in rows: sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions; sheet.sheet_view.showGridLines=False
    for cell in sheet[1]: cell.fill=PatternFill("solid",fgColor="1F4E78"); cell.font=Font(name="Microsoft YaHei",bold=True,color="FFFFFF")
    percent_headers={"正收益比例","平均收益","中位收益","平均盈利","平均亏损","最大亏损","MAE","MFE","最大回撤","D+1收益","D+3收益","D+5收益","截至期末收益","win_rate","average_return","median_return","average_win","average_loss","maximum_loss","mae","mfe","max_drawdown","coverage_ratio"}
    for row in sheet.iter_rows():
        has_long_text=False
        for cell in row:
            cell.alignment=CENTER; header=str(sheet.cell(1,cell.column).value)
            has_long_text = has_long_text or (isinstance(cell.value,str) and len(cell.value)>28)
            if "股票代码" in header: cell.number_format="@"
            metric=str(sheet.cell(cell.row,1).value or "")
            if isinstance(cell.value,(int,float)) and (header in percent_headers or metric in percent_headers): cell.number_format="[Red]0.00%;[Green]-0.00%;-"
            if header in {"Admission V2","Admission V1"} and cell.row>1: cell.fill=PatternFill("solid",fgColor={"PASS":"D9EAD3","REVIEW":"FFF2CC","BLOCK":"F4CCCC"}.get(str(cell.value),"FFFFFF"))
        sheet.row_dimensions[row[0].row].height=38 if has_long_text else 26
    for index,header in enumerate(headers,1):
        values=[str(sheet.cell(row,index).value or "") for row in range(1,min(sheet.max_row,200)+1)]
        longest=max((len(value) for value in values),default=len(str(header)))
        sheet.column_dimensions[get_column_letter(index)].width=min(32,max(12,longest*1.45+2))


def _verify(path,sheets):
    workbook=load_workbook(path,data_only=False)
    try:
        assert workbook.sheetnames==list(sheets)
        formula_count=0
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None and (cell.alignment.horizontal!="center" or cell.alignment.vertical!="center" or not cell.alignment.wrap_text): raise ValueError(f"V2_EXCEL_ALIGNMENT:{sheet.title}:{cell.coordinate}")
                    if isinstance(cell.value,str) and cell.value.startswith("="): formula_count+=1
                    if str(cell.value).startswith(("#REF!","#DIV/0!","#VALUE!","#NAME?","#N/A")): raise ValueError(f"V2_EXCEL_FORMULA_ERROR:{sheet.title}:{cell.coordinate}")
        return {"output":str(path),"size_bytes":path.stat().st_size,"sheets":workbook.sheetnames,"centered":True,"formula_error_scan":"PASSED","formula_count":formula_count}
    finally: workbook.close()


def _number(value):
    try:return float(value) if value is not None else None
    except (TypeError,ValueError):return None


def _display_note(value):
    if isinstance(value,bool): return "是" if value else "否"
    return value


def _hash(value):return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False,default=str,separators=(",", ":")).encode()).hexdigest()


def _baseline_hash(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, default=str, separators=(",", ":")).encode()).hexdigest()
