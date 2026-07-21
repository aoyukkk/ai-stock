from __future__ import annotations

import math
import statistics
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from database.models.entry_timing import EntryTimingResult
from database.models.performance import SelectionCohort
from entry_timing.service import EntryTimingShadowService, _result_dict, _rows_by_code


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
STATUS_FILL = {
    "PASS": "D9EAD3", "REVIEW": "FFF2CC", "BLOCK": "F4CCCC", "DATA_INSUFFICIENT": "E7E6E6",
}


class HistoricalEntryTimingValidator:
    def __init__(self, session, *, cache_root: Path) -> None:
        self.session = session
        self.cache_root = cache_root

    def run(self, start_date: date, end_date: date, output: Path) -> dict[str, Any]:
        service = EntryTimingShadowService(self.session, cache_root=self.cache_root)
        cohorts = list(self.session.scalars(
            select(SelectionCohort).where(
                SelectionCohort.selection_trade_date >= start_date,
                SelectionCohort.selection_trade_date <= end_date,
            ).order_by(SelectionCohort.selection_trade_date)
        ))
        if not cohorts:
            raise ValueError("HISTORICAL_SELECTION_COHORTS_NOT_FOUND")
        run_ids = []
        for cohort in cohorts:
            summary = service.run(
                cohort.selection_trade_date,
                quant_run_id=cohort.quant_run_id,
                candidate_mode="HISTORICAL_CANDIDATES",
                force_shadow=True,
            )
            run_ids.append(summary["run_id"])
        rows = list(self.session.scalars(
            select(EntryTimingResult).where(EntryTimingResult.admission_run_id.in_(run_ids))
            .order_by(EntryTimingResult.trade_date, EntryTimingResult.quant_rank)
        ))
        outcome_map = self._outcomes(rows, end_date)
        records = []
        for row in rows:
            item = _result_dict(row)
            item.update(outcome_map.get((row.trade_date, row.stock_code), {}))
            records.append(item)
        ai = [row for row in records if row["pool_type"] == "AI_POOL"]
        manual = [row for row in records if row["pool_type"] != "AI_POOL"]
        admitted = []
        for selection_day in sorted({row["trade_date"] for row in ai}):
            daily = sorted(
                (row for row in ai if row["trade_date"] == selection_day and row["admission_status"] == "PASS"),
                key=lambda row: (row["quant_rank"] is None, row["quant_rank"] or 999999),
            )
            admitted.extend(daily[:20])
        original_metrics = _metrics(ai)
        admitted_metrics = _metrics(admitted)
        manual_metrics = _metrics(manual)
        removed = [
            row for row in ai
            if row.get("cumulative_return") is not None
            and row["cumulative_return"] <= -0.20
            and row["admission_status"] != "PASS"
        ]
        comparison = {
            "historical_period": f"{start_date.isoformat()} to {end_date.isoformat()}",
            "original": original_metrics,
            "entry_filtered": admitted_metrics,
            "manual_challenge": manual_metrics,
            "removed_high_risk_stocks": removed,
            "admission_distribution": _distribution(ai),
            "manual_admission_distribution": _distribution(manual),
            "source_candidate_count": len(ai),
            "entry_filtered_count": len(admitted),
            "manual_count": len(manual),
            "quant_hash_unchanged": all(service.summary(run_id)["quant_hash_unchanged"] for run_id in run_ids),
            "llm_calls": 0,
            "external_api_calls": 0,
            "production_recommendation_changed": False,
        }
        excel = EntryTimingHistoricalExcel().export(output, records, comparison)
        return {**comparison, "admission_run_ids": run_ids, "excel": excel}

    def _outcomes(self, rows: list[EntryTimingResult], end_date: date) -> dict[tuple[date, str], dict[str, Any]]:
        daily_dir = self.cache_root / "trade_date" / "daily"
        paths = [path for path in sorted(daily_dir.glob("*.json")) if path.stem <= f"{end_date:%Y%m%d}"]
        by_date = {date.fromisoformat(f"{path.stem[:4]}-{path.stem[4:6]}-{path.stem[6:]}"): _rows_by_code(path) for path in paths}
        output = {}
        for row in rows:
            daily_returns = []
            cumulative_curve = []
            cumulative = 1.0
            for trade_day in sorted(day for day in by_date if row.trade_date < day <= end_date):
                bar = by_date[trade_day].get(row.stock_code)
                if not bar or bar.get("pct_chg") is None:
                    continue
                daily_return = float(bar["pct_chg"]) / 100
                daily_returns.append((trade_day, daily_return))
                cumulative *= 1 + daily_return
                cumulative_curve.append(cumulative - 1)
            max_drawdown = _curve_drawdown(cumulative_curve)
            output[(row.trade_date, row.stock_code)] = {
                "cumulative_return": cumulative - 1 if daily_returns else None,
                "holding_days": len(daily_returns),
                "max_drawdown": max_drawdown,
                "daily_returns": [{"trade_date": day.isoformat(), "return": value} for day, value in daily_returns],
            }
        return output


class EntryTimingHistoricalExcel:
    def export(self, output: Path, records: list[dict[str, Any]], comparison: dict[str, Any]) -> dict[str, Any]:
        workbook = Workbook()
        workbook.remove(workbook.active)
        _sheet(workbook, "08_买入准入分析", [self._analysis_row(row) for row in records])
        _sheet(workbook, "历史对比", [
            {"指标": label, "原重点候选": comparison["original"].get(key), "Entry过滤后": comparison["entry_filtered"].get(key), "人工挑战池": comparison["manual_challenge"].get(key)}
            for key, label in (
                ("count", "数量"), ("win_rate", "胜率"), ("average_return", "平均收益"),
                ("median_return", "中位收益"), ("max_loss", "最大亏损"),
                ("max_drawdown", "最大回撤"), ("profit_loss_ratio", "盈亏比"),
            )
        ])
        _sheet(workbook, "高风险过滤", [self._analysis_row(row) for row in comparison["removed_high_risk_stocks"]])
        _sheet(workbook, "人工挑战池", [self._analysis_row(row) for row in records if row["pool_type"] != "AI_POOL"])
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
        workbook.close()
        return _verify_excel(output)

    @staticmethod
    def _analysis_row(row: dict[str, Any]) -> dict[str, Any]:
        return {
            "选股日期": row["trade_date"], "股票代码": row["stock_code"].split(".", 1)[0],
            "股票名称": row.get("stock_name"), "候选池": row["pool_type"],
            "Quant排名": row.get("quant_rank"), "Quant分": row.get("quant_score"),
            "Flash分": row.get("flash_score"), "Entry Timing分": row.get("entry_timing_score"),
            "价格位置": row.get("position_score"), "回撤质量": row.get("pullback_score"),
            "量价结构": row.get("volume_price_score"), "板块共振": row.get("sector_score"),
            "市场适配": row.get("market_score"), "流动性": row.get("liquidity_score"),
            "风险标签": "、".join(row.get("risk_flags") or []) or "无",
            "Admission状态": row.get("admission_status"),
            "阻断原因": "、".join(row.get("block_reasons") or []) or "无",
            "截至期末总涨跌": row.get("cumulative_return"), "最大回撤": row.get("max_drawdown"),
            "观测交易日": row.get("holding_days"), "人工分": row.get("manual_score"),
            "AI分": row.get("ai_score"), "分差": row.get("score_difference"),
        }


def _sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    headers = list(rows[0]) if rows else ["暂无数据"]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = CENTER
            header = str(sheet.cell(1, cell.column).value)
            if "股票代码" in header:
                cell.number_format = "@"
            metric_label = str(sheet.cell(cell.row, 1).value or "")
            percentage_metric = metric_label in {"胜率", "平均收益", "中位收益", "最大亏损", "最大回撤"} and cell.column > 1
            percentage_column = header in {"截至期末总涨跌", "最大回撤"}
            if (percentage_metric or percentage_column) and isinstance(cell.value, (int, float)):
                cell.number_format = "[Red]0.00%;[Green]-0.00%;-"
            if header == "Admission状态" and cell.row > 1:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILL.get(str(cell.value), "FFFFFF"))
        sheet.row_dimensions[row[0].row].height = 25
    for column, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(column)].width = min(28, max(12, len(str(header)) * 2 + 2))
    sheet.sheet_view.showGridLines = False


def _metrics(rows: list[dict[str, Any]]) -> dict[str, Any]:
    values = [float(row["cumulative_return"]) for row in rows if row.get("cumulative_return") is not None]
    if not values:
        return {"count": len(rows), "observed_count": 0, "win_rate": None, "average_return": None, "median_return": None, "max_loss": None, "max_drawdown": None, "profit_loss_ratio": None}
    positives = [value for value in values if value > 0]
    negatives = [value for value in values if value < 0]
    loss_average = abs(statistics.fmean(negatives)) if negatives else None
    ratio = statistics.fmean(positives) / loss_average if positives and loss_average else None
    drawdowns = [float(row.get("max_drawdown") or 0) for row in rows]
    return {
        "count": len(rows), "observed_count": len(values), "win_rate": len(positives) / len(values),
        "average_return": statistics.fmean(values), "median_return": statistics.median(values),
        "max_loss": min(values), "max_drawdown": min(drawdowns) if drawdowns else None,
        "profit_loss_ratio": ratio,
    }


def _curve_drawdown(cumulative_curve: list[float]) -> float | None:
    if not cumulative_curve:
        return None
    wealth = [1 + value for value in cumulative_curve]
    peak = 1.0
    worst = 0.0
    for value in wealth:
        peak = max(peak, value)
        worst = min(worst, value / peak - 1)
    return worst


def _distribution(rows: list[dict[str, Any]]) -> dict[str, int]:
    result: dict[str, int] = defaultdict(int)
    for row in rows:
        result[row["admission_status"]] += 1
    return dict(result)


def _verify_excel(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    if cell.alignment.horizontal != "center" or cell.alignment.vertical != "center" or not cell.alignment.wrap_text:
                        raise ValueError(f"ENTRY_TIMING_EXCEL_ALIGNMENT:{sheet.title}:{cell.coordinate}")
        return {"output": str(path), "size_bytes": path.stat().st_size, "sheets": workbook.sheetnames, "centered": True}
    finally:
        workbook.close()
