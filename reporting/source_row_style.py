from __future__ import annotations

from copy import copy
from typing import Any

from openpyxl.styles import PatternFill


MANUAL_SOURCE_COLOR = "FFF2CC"
MODEL_SOURCE_COLOR = "DCEAF5"
NEUTRAL_ROW_COLOR = "FFFFFF"

MANUAL_SOURCE_FILL = PatternFill("solid", fgColor=MANUAL_SOURCE_COLOR)
MODEL_SOURCE_FILL = PatternFill("solid", fgColor=MODEL_SOURCE_COLOR)
NEUTRAL_ROW_FILL = PatternFill("solid", fgColor=NEUTRAL_ROW_COLOR)

SELECTION_SOURCE_HEADERS = {
    "入选来源",
    "来源",
    "选择来源",
    "交易候选来源",
    "selection_source",
    "pool_type",
}


def source_fill(value: Any) -> PatternFill:
    text = str(value or "").strip().upper()
    if not text:
        return NEUTRAL_ROW_FILL
    if any(token in text for token in ("人工", "共同", "MANUAL", "BOTH", "HUMAN")):
        return MANUAL_SOURCE_FILL
    return MODEL_SOURCE_FILL


def apply_selection_source_rows(
    worksheet,
    *,
    header_row: int | None = None,
    first_data_row: int | None = None,
    last_data_row: int | None = None,
) -> bool:
    resolved_header_row = header_row or _find_header_row(worksheet)
    if resolved_header_row is None:
        return False
    source_column = next(
        (
            cell.column
            for cell in worksheet[resolved_header_row]
            if str(cell.value or "").strip() in SELECTION_SOURCE_HEADERS
        ),
        None,
    )
    if source_column is None:
        return False
    start = first_data_row or resolved_header_row + 1
    end = last_data_row or worksheet.max_row
    for row_number in range(start, end + 1):
        fill = source_fill(worksheet.cell(row_number, source_column).value)
        for cell in worksheet[row_number]:
            cell.fill = copy(fill)
    return True


def disable_table_banding(worksheet) -> None:
    for table in worksheet.tables.values():
        if table.tableStyleInfo is not None:
            table.tableStyleInfo.showRowStripes = False
            table.tableStyleInfo.showColumnStripes = False


def _find_header_row(worksheet) -> int | None:
    for row_number in range(1, min(12, worksheet.max_row) + 1):
        headers = {
            str(cell.value or "").strip()
            for cell in worksheet[row_number]
        }
        if headers & SELECTION_SOURCE_HEADERS:
            return row_number
    return None
