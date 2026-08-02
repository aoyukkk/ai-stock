from __future__ import annotations

import hashlib
import json
import os
import re
import uuid
import zipfile
from copy import copy
from dataclasses import asdict, dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.views import Selection

from reporting.workbook_standard import DAILY_SHEETS, LEGACY_DAILY_SHEETS
from reporting.source_row_style import (
    NEUTRAL_ROW_FILL,
    SELECTION_SOURCE_HEADERS,
    disable_table_banding,
    source_fill,
)


FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
FAIL_FILL = PatternFill("solid", fgColor="F4CCCC")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
PASS_FILL = PatternFill("solid", fgColor="D9EAD3")


@dataclass(frozen=True)
class ReferenceWorkbookMetadata:
    template_path: str
    template_file_hash: str
    sheet_names: list[str]
    sheet_order: list[str]
    style_signature: str
    style_profile: dict[str, Any]


class WorkbookStyleService:
    """Apply the established human daily-workbook visual language to other exports."""

    def __init__(self, reference_path: Path | str) -> None:
        self.reference_path = Path(reference_path).resolve()
        if not self.reference_path.is_file():
            raise FileNotFoundError(f"REFERENCE_WORKBOOK_MISSING:{self.reference_path}")

    @classmethod
    def resolve_recent_successful_reference(
        cls,
        output_root: Path | str,
        *,
        start: date,
        end: date,
    ) -> Path:
        root = Path(output_root)
        excluded = ("午盘", "历史", "审计", "失败", "radar", "shadow", "复盘")
        current = end
        while current >= start:
            directory = root / current.isoformat()
            if directory.is_dir():
                candidates = sorted(
                    (
                        path for path in directory.rglob("*.xlsx")
                        if not any(
                            token.lower() in str(path.relative_to(directory)).lower()
                            for token in excluded
                        )
                        and not path.name.startswith("~$")
                    ),
                    key=lambda path: (
                        path.name.startswith(f"智能交易助手_{current.isoformat()}"),
                        path.stat().st_mtime,
                        path.name,
                    ),
                    reverse=True,
                )
                for candidate in candidates:
                    if cls._is_complete_daily_workbook(candidate):
                        return candidate.resolve()
            current -= timedelta(days=1)
        raise FileNotFoundError("SUCCESSFUL_REFERENCE_WORKBOOK_NOT_FOUND")

    @staticmethod
    def _is_complete_daily_workbook(path: Path) -> bool:
        try:
            workbook = load_workbook(path, read_only=False, data_only=False)
        except Exception:
            return False
        try:
            names = tuple(workbook.sheetnames)
            return names in {DAILY_SHEETS, LEGACY_DAILY_SHEETS}
        finally:
            workbook.close()

    def metadata(self) -> ReferenceWorkbookMetadata:
        workbook = load_workbook(self.reference_path, data_only=False)
        try:
            profile = {
                "sheet_count": len(workbook.sheetnames),
                "sheets": [self._sheet_profile(sheet) for sheet in workbook.worksheets],
            }
        finally:
            workbook.close()
        material = json.dumps(profile, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        return ReferenceWorkbookMetadata(
            template_path=str(self.reference_path),
            template_file_hash=hashlib.sha256(self.reference_path.read_bytes()).hexdigest(),
            sheet_names=list(profile_sheet["name"] for profile_sheet in profile["sheets"]),
            sheet_order=list(profile_sheet["name"] for profile_sheet in profile["sheets"]),
            style_signature=hashlib.sha256(material.encode("utf-8")).hexdigest(),
            style_profile=profile,
        )

    def reformat_midday_radar(
        self,
        source_path: Path | str,
        output_path: Path | str,
        *,
        final_status: str,
        result_nature: str,
    ) -> dict[str, Any]:
        source = Path(source_path).resolve()
        output = Path(output_path).resolve()
        if output.exists():
            raise FileExistsError(f"OUTPUT_EXISTS:{output.name}")
        output.parent.mkdir(parents=True, exist_ok=True)
        reference = load_workbook(self.reference_path, data_only=False)
        workbook = load_workbook(source, data_only=False)
        try:
            original_values = self._value_snapshot(workbook)
            self.apply(workbook, reference)
            home = workbook[workbook.sheetnames[0]]
            self._write_home_notice(home, final_status, result_nature, reference[reference.sheetnames[0]])
            temporary = output.with_name(f".{output.stem}_{uuid.uuid4().hex[:8]}.tmp.xlsx")
            workbook.save(temporary)
        finally:
            workbook.close()
            reference.close()
        os.replace(temporary, output)
        validation = self.validate_midday_radar(
            source,
            output,
            original_values=original_values,
            final_status=final_status,
            result_nature=result_nature,
        )
        return {
            "output_path": str(output),
            "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
            "reference": asdict(self.metadata()),
            "validation": validation,
        }

    def align_existing_workbook(self, workbook) -> None:
        reference = load_workbook(self.reference_path, data_only=False)
        try:
            self.apply(workbook, reference)
        finally:
            reference.close()

    def align_midday_workbook(self, workbook, *, final_status: str, result_nature: str) -> None:
        reference = load_workbook(self.reference_path, data_only=False)
        try:
            self.apply(workbook, reference)
            self._write_home_notice(
                workbook[workbook.sheetnames[0]],
                final_status,
                result_nature,
                reference[reference.sheetnames[0]],
            )
        finally:
            reference.close()

    def annotate_official_status(self, path: Path | str, status: str) -> None:
        output = Path(path)
        workbook = load_workbook(output, data_only=False)
        reference = load_workbook(self.reference_path, data_only=False)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            existing = str(worksheet["A3"].value or "").strip()
            notice = f"运行状态：{status}"
            worksheet["A3"] = f"{existing}  {notice}" if existing else notice
            prototype = reference[reference.sheetnames[0]]["A3"]
            self._copy_cell_style(prototype, worksheet["A3"], include_number_format=True)
            worksheet["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            worksheet.row_dimensions[3].height = max(30, worksheet.row_dimensions[3].height or 0)
            temporary = output.with_name(f".{output.stem}_{uuid.uuid4().hex[:8]}.tmp.xlsx")
            workbook.save(temporary)
        finally:
            workbook.close()
            reference.close()
        os.replace(temporary, output)

    def apply(self, workbook, reference) -> None:
        summary = reference[reference.sheetnames[0]]
        table_source = reference["量化前100"] if "量化前100" in reference.sheetnames else reference.worksheets[-2]
        header_widths, header_formats = self._reference_header_maps(reference)
        source_header_row = self._header_row(table_source)
        for worksheet in workbook.worksheets:
            header_row = self._header_row(worksheet)
            if header_row is None:
                continue
            self._copy_title_style(summary, worksheet)
            self._copy_sheet_settings(table_source, worksheet)
            self._style_table(
                worksheet,
                table_source,
                header_row,
                source_header_row or 4,
                header_widths,
                header_formats,
            )
        self._normalize_all_alignments(workbook)
        self.ensure_excel_compatibility(workbook)

    @staticmethod
    def set_freeze_panes_safely(worksheet, coordinate: str | None) -> None:
        """Set frozen panes without accumulating invalid duplicate selections.

        openpyxl inserts additional ``selection`` records every time a two-axis
        freeze pane is assigned.  Re-styling an existing workbook can therefore
        create a file that openpyxl can read but desktop Excel must repair.
        """
        worksheet.sheet_view.pane = None
        worksheet.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
        worksheet.freeze_panes = coordinate

    @staticmethod
    def ensure_excel_compatibility(workbook) -> None:
        """Normalize Excel-strict view and filter structures before saving."""
        for worksheet in workbook.worksheets:
            freeze = str(worksheet.freeze_panes) if worksheet.freeze_panes else None
            WorkbookStyleService.set_freeze_panes_safely(worksheet, freeze)
            sheet_filter = worksheet.auto_filter.ref
            if sheet_filter and any(
                WorkbookStyleService._ranges_overlap(sheet_filter, table.ref)
                for table in worksheet.tables.values()
            ):
                # A table already owns its AutoFilter.  A second worksheet-level
                # AutoFilter over the same range makes Excel report corruption.
                worksheet.auto_filter.ref = None

    @staticmethod
    def validate_excel_compatibility(path: Path | str) -> dict[str, Any]:
        """Reject the two Open XML patterns known to trigger Excel repair."""
        target = Path(path).resolve()
        with zipfile.ZipFile(target) as archive:
            bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f"XLSX_ZIP_CRC_ERROR:{bad_member}")
        workbook = load_workbook(target, data_only=False)
        try:
            selection_records = 0
            table_count = 0
            for worksheet in workbook.worksheets:
                panes = [item.pane for item in worksheet.sheet_view.selection if item.pane]
                selection_records += len(worksheet.sheet_view.selection)
                if len(panes) != len(set(panes)):
                    raise ValueError(f"DUPLICATE_SHEET_VIEW_SELECTION:{worksheet.title}:{panes}")
                table_count += len(worksheet.tables)
                sheet_filter = worksheet.auto_filter.ref
                if sheet_filter and any(
                    WorkbookStyleService._ranges_overlap(sheet_filter, table.ref)
                    for table in worksheet.tables.values()
                ):
                    raise ValueError(f"OVERLAPPING_TABLE_AND_SHEET_FILTER:{worksheet.title}:{sheet_filter}")
            return {
                "status": "PASS",
                "zip_crc": "PASS",
                "sheet_count": len(workbook.sheetnames),
                "table_count": table_count,
                "selection_records": selection_records,
                "duplicate_view_selections": 0,
                "overlapping_table_sheet_filters": 0,
            }
        finally:
            workbook.close()

    @staticmethod
    def _ranges_overlap(left: str, right: str) -> bool:
        left_min_col, left_min_row, left_max_col, left_max_row = range_boundaries(left)
        right_min_col, right_min_row, right_max_col, right_max_row = range_boundaries(right)
        return not (
            left_max_col < right_min_col
            or right_max_col < left_min_col
            or left_max_row < right_min_row
            or right_max_row < left_min_row
        )

    @staticmethod
    def _normalize_all_alignments(workbook) -> None:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def validate_midday_radar(
        self,
        source_path: Path,
        output_path: Path,
        *,
        original_values: dict[str, dict[str, Any]],
        final_status: str,
        result_nature: str,
    ) -> dict[str, Any]:
        workbook = load_workbook(output_path, data_only=False)
        try:
            current = self._value_snapshot(workbook)
            changed = []
            for sheet, cells in original_values.items():
                for coordinate, value in cells.items():
                    if current.get(sheet, {}).get(coordinate) != value:
                        changed.append(f"{sheet}!{coordinate}")
            if changed:
                raise ValueError(f"BUSINESS_VALUE_CHANGED:{changed[:10]}")
            home_text = str(workbook[workbook.sheetnames[0]]["A2"].value or "")
            if final_status not in home_text or result_nature not in home_text:
                raise ValueError("MIDDAY_HOME_NOTICE_MISSING")
            formula_errors = 0
            centered = wrapped = 0
            stock_code_cells = 0
            for worksheet in workbook.worksheets:
                header_row = self._header_row(worksheet)
                headers = {
                    cell.column: str(cell.value or "")
                    for cell in worksheet[header_row]
                } if header_row else {}
                for row in worksheet.iter_rows():
                    for cell in row:
                        if isinstance(cell, MergedCell) or cell.value is None:
                            continue
                        if cell.alignment.horizontal != "center" or cell.alignment.vertical != "center":
                            raise ValueError(f"CELL_NOT_CENTERED:{worksheet.title}!{cell.coordinate}")
                        centered += 1
                        if not cell.alignment.wrap_text:
                            raise ValueError(f"CELL_NOT_WRAPPED:{worksheet.title}!{cell.coordinate}")
                        wrapped += 1
                        if isinstance(cell.value, str) and any(error in cell.value.upper() for error in FORMULA_ERRORS):
                            formula_errors += 1
                        if "股票代码" in headers.get(cell.column, "") and cell.row > int(header_row or 0):
                            stock_code_cells += 1
                            if cell.number_format != "@" or not isinstance(cell.value, str):
                                raise ValueError(f"STOCK_CODE_FORMAT_ERROR:{worksheet.title}!{cell.coordinate}")
            if formula_errors:
                raise ValueError(f"FORMULA_ERRORS:{formula_errors}")
            return {
                "status": "PASS",
                "source_path": str(source_path),
                "source_sha256": hashlib.sha256(source_path.read_bytes()).hexdigest(),
                "sheet_names": list(workbook.sheetnames),
                "sheet_count": len(workbook.sheetnames),
                "business_values_unchanged": True,
                "centered_cells": centered,
                "wrapped_cells": wrapped,
                "stock_code_cells": stock_code_cells,
                "formula_errors": 0,
                "final_status": final_status,
                "result_nature": result_nature,
            }
        finally:
            workbook.close()

    @staticmethod
    def compare_official_structure(reference_path: Path | str, output_path: Path | str) -> dict[str, Any]:
        reference = load_workbook(reference_path, data_only=False)
        output = load_workbook(output_path, data_only=False)
        try:
            reference_headers = WorkbookStyleService._all_headers(reference)
            output_headers = WorkbookStyleService._all_headers(output)
            return {
                "sheet_names_match": reference.sheetnames == output.sheetnames,
                "sheet_order_match": reference.sheetnames == output.sheetnames,
                "headers_match": reference_headers == output_headers,
                "reference_sheets": list(reference.sheetnames),
                "output_sheets": list(output.sheetnames),
            }
        finally:
            reference.close()
            output.close()

    @staticmethod
    def _all_headers(workbook) -> dict[str, list[Any]]:
        result = {}
        for worksheet in workbook.worksheets:
            row = WorkbookStyleService._header_row(worksheet)
            result[worksheet.title] = [cell.value for cell in worksheet[row]] if row else []
        return result

    @staticmethod
    def _write_home_notice(worksheet, status: str, nature: str, reference_summary) -> None:
        max_col = max(2, worksheet.max_column)
        for merged in list(worksheet.merged_cells.ranges):
            if merged.min_row <= 2 <= merged.max_row:
                worksheet.unmerge_cells(str(merged))
        worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        cell = worksheet.cell(2, 1)
        cell.value = f"运行状态：{status}\n结果性质：{nature}"
        prototype = reference_summary.cell(3, 1)
        WorkbookStyleService._copy_cell_style(prototype, cell, include_number_format=True)
        cell.fill = FAIL_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[2].height = 42

    @staticmethod
    def _copy_title_style(reference_summary, worksheet) -> None:
        source = reference_summary.cell(1, 1)
        for cell in worksheet[1]:
            if isinstance(cell, MergedCell):
                continue
            WorkbookStyleService._copy_cell_style(source, cell, include_number_format=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = reference_summary.row_dimensions[1].height or 30

    @staticmethod
    def _copy_sheet_settings(source, target) -> None:
        target.sheet_view.showGridLines = source.sheet_view.showGridLines
        target.sheet_properties.pageSetUpPr = copy(source.sheet_properties.pageSetUpPr)
        target.page_margins = copy(source.page_margins)
        target.page_setup = copy(source.page_setup)
        target.print_options = copy(source.print_options)
        target.sheet_properties.pageSetUpPr.fitToPage = True
        target.page_setup.orientation = source.page_setup.orientation or "landscape"

    @staticmethod
    def _style_table(worksheet, source, header_row, source_header_row, header_widths, header_formats) -> None:
        source_headers = {str(cell.value or ""): cell.column for cell in source[source_header_row]}
        source_max_col = max(1, source.max_column)
        source_body_rows = [min(source_header_row + 1, source.max_row)]
        selection_source_column = next(
            (
                cell.column
                for cell in worksheet[header_row]
                if str(cell.value or "").strip() in SELECTION_SOURCE_HEADERS
            ),
            None,
        )
        disable_table_banding(worksheet)
        for column in range(1, worksheet.max_column + 1):
            header = worksheet.cell(header_row, column)
            prototype = source.cell(source_header_row, source_headers.get(str(header.value or ""), min(column, source_max_col)))
            WorkbookStyleService._copy_cell_style(prototype, header, include_number_format=True)
            header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            width = WorkbookStyleService._column_width(str(header.value or ""), header_widths, worksheet.column_dimensions[get_column_letter(column)].width)
            worksheet.column_dimensions[get_column_letter(column)].width = width
            for row in range(header_row + 1, worksheet.max_row + 1):
                cell = worksheet.cell(row, column)
                body_source_row = source_body_rows[(row - header_row - 1) % len(source_body_rows)]
                body_source_col = source_headers.get(str(header.value or ""), min(column, source_max_col))
                body = source.cell(body_source_row, body_source_col)
                current_format = cell.number_format
                WorkbookStyleService._copy_cell_style(body, cell, include_number_format=False)
                cell.fill = copy(
                    source_fill(worksheet.cell(row, selection_source_column).value)
                    if selection_source_column is not None
                    else NEUTRAL_ROW_FILL
                )
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.number_format = current_format if current_format and current_format != "General" else header_formats.get(str(header.value or ""), body.number_format)
                if "股票代码" in str(header.value or "") and cell.value not in (None, ""):
                    cell.value = str(cell.value).split(".")[0].zfill(6)
                    cell.number_format = "@"
                WorkbookStyleService._semantic_format(cell, str(header.value or ""))
        worksheet.row_dimensions[header_row].height = source.row_dimensions[source_header_row].height or 24
        body_height = max(source.row_dimensions[row].height or 20 for row in source_body_rows)
        for row in range(header_row + 1, worksheet.max_row + 1):
            worksheet.row_dimensions[row].height = max(body_height, min(60, WorkbookStyleService._estimated_height(worksheet, row)))
        WorkbookStyleService.set_freeze_panes_safely(
            worksheet,
            f"{'D' if worksheet.max_column >= 12 else 'A'}{header_row + 1}",
        )
        if worksheet.auto_filter.ref:
            _, _, max_col, max_row = range_boundaries(worksheet.auto_filter.ref)
            worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{max(max_row, worksheet.max_row)}"

    @staticmethod
    def _semantic_format(cell, header: str) -> None:
        value = cell.value
        upper = str(value or "").upper()
        if any(token in upper for token in ("FAILED", "FAIL", "BLOCKED")):
            cell.fill = copy(FAIL_FILL)
        elif upper in {"PASS", "SUCCESS", "TRUE", "是"}:
            cell.fill = copy(PASS_FILL)
        if not isinstance(value, (int, float)):
            return
        if "覆盖率" in header:
            cell.fill = copy(PASS_FILL if value >= 0.99 else WARN_FILL if value >= 0.98 else FAIL_FILL)
        elif header.endswith("分") or "评分" in header or "得分" in header:
            cell.fill = copy(PASS_FILL if value >= 70 else WARN_FILL if value >= 55 else FAIL_FILL)

    @staticmethod
    def _column_width(header: str, widths: dict[str, float], current: float | None) -> float:
        if header in widths:
            return widths[header]
        if "股票代码" in header:
            return widths.get("股票代码", 12)
        if "股票名称" in header:
            return widths.get("股票名称", 14)
        if any(token in header for token in ("说明", "原因", "风险", "明细", "审计", "逻辑", "条件")):
            return max(widths.get("说明", 28), 28)
        if any(token in header for token in ("日期", "时间", "状态", "来源", "行业", "能力", "用途")):
            return 16
        return min(20, max(10, float(current or 12)))

    @staticmethod
    def _estimated_height(worksheet, row: int) -> float:
        lines = 1
        for cell in worksheet[row]:
            if cell.value in (None, ""):
                continue
            width = worksheet.column_dimensions[get_column_letter(cell.column)].width or 12
            text = str(cell.value)
            lines = max(lines, text.count("\n") + 1, (len(text) + max(8, int(width)) - 1) // max(8, int(width)))
        return 18 * min(lines, 3)

    @staticmethod
    def _reference_header_maps(reference) -> tuple[dict[str, float], dict[str, str]]:
        widths: dict[str, float] = {}
        formats: dict[str, str] = {}
        for worksheet in reference.worksheets:
            row = WorkbookStyleService._header_row(worksheet)
            if not row:
                continue
            for cell in worksheet[row]:
                header = str(cell.value or "")
                if not header:
                    continue
                width = worksheet.column_dimensions[get_column_letter(cell.column)].width or 12
                widths.setdefault(header, float(width))
                if worksheet.max_row > row:
                    formats.setdefault(header, worksheet.cell(row + 1, cell.column).number_format)
        return widths, formats

    @staticmethod
    def _header_row(worksheet) -> int | None:
        for table in worksheet.tables.values():
            _, min_row, _, _ = range_boundaries(table.ref)
            return min_row
        if worksheet.auto_filter.ref:
            _, min_row, _, _ = range_boundaries(worksheet.auto_filter.ref)
            return min_row
        best: tuple[int, int] | None = None
        for row in range(1, min(12, worksheet.max_row) + 1):
            count = sum(worksheet.cell(row, col).value not in (None, "") for col in range(1, worksheet.max_column + 1))
            if count >= 2 and (best is None or count > best[1]):
                best = (row, count)
        return best[0] if best else None

    @staticmethod
    def _copy_cell_style(source, target, *, include_number_format: bool) -> None:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        if include_number_format:
            target.number_format = source.number_format

    @staticmethod
    def _value_snapshot(workbook) -> dict[str, dict[str, Any]]:
        result: dict[str, dict[str, Any]] = {}
        for worksheet in workbook.worksheets:
            result[worksheet.title] = {
                cell.coordinate: cell.value
                for row in worksheet.iter_rows()
                for cell in row
                if not isinstance(cell, MergedCell) and cell.value is not None
            }
        return result

    @staticmethod
    def _sheet_profile(worksheet) -> dict[str, Any]:
        header_row = WorkbookStyleService._header_row(worksheet)
        title = worksheet.cell(1, 1)
        header = worksheet.cell(header_row or 1, 1)
        body = worksheet.cell(min((header_row or 1) + 1, worksheet.max_row), 1)
        widths = {
            key: value.width
            for key, value in worksheet.column_dimensions.items()
            if value.width is not None
        }
        return {
            "name": worksheet.title,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "header_row": header_row,
            "freeze_panes": str(worksheet.freeze_panes or ""),
            "auto_filter": str(worksheet.auto_filter.ref or ""),
            "table_count": len(worksheet.tables),
            "conditional_formatting_count": len(worksheet.conditional_formatting),
            "title_style": WorkbookStyleService._cell_signature(title),
            "header_style": WorkbookStyleService._cell_signature(header),
            "body_style": WorkbookStyleService._cell_signature(body),
            "row_heights": {str(key): value.height for key, value in worksheet.row_dimensions.items() if value.height is not None},
            "column_widths": widths,
            "orientation": worksheet.page_setup.orientation,
        }

    @staticmethod
    def _cell_signature(cell) -> dict[str, Any]:
        color = getattr(cell.fill.fgColor, "rgb", None) or getattr(cell.fill.fgColor, "indexed", None)
        font_color = getattr(cell.font.color, "rgb", None) if cell.font.color else None
        return {
            "style_id": cell.style_id,
            "font": cell.font.name,
            "font_size": cell.font.sz,
            "bold": cell.font.bold,
            "font_color": font_color,
            "fill": color,
            "number_format": cell.number_format,
            "horizontal": cell.alignment.horizontal,
            "vertical": cell.alignment.vertical,
            "wrap_text": cell.alignment.wrap_text,
        }


def normalized_stock_code(value: Any) -> str:
    text = str(value or "").strip()
    match = re.search(r"(\d{6})", text)
    return match.group(1) if match else text
