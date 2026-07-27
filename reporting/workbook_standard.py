from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import range_boundaries


TRADING_ASSISTANT_PROFILE = "trading_assistant_human_v1"

LEGACY_DAILY_SHEETS = (
    "今日概览", "重点候选", "挂单与仓位", "基本面摘要", "量化前100", "当前问题",
)

LEGACY_MIDDAY_SHEETS = (
    "今日概览", "重点候选", "价格与权重", "复核依据", "量化前100", "当前问题",
)

DAILY_SHEETS = (
    "今日概览",
    "今日推荐",
    "重点候选",
    "挂单与仓位",
    "基本面摘要",
    "量化前100",
    "当前问题",
)

MIDDAY_SHEETS = (
    "今日概览",
    "今日推荐",
    "重点候选",
    "价格与权重",
    "复核依据",
    "量化前100",
    "当前问题",
)

FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
FORBIDDEN_MACHINE_TEXT = (
    "MODEL_VALIDATION",
    "MODEL_UNVERIFIED",
    "LLM_UNVERIFIED",
    "DEEPSEEK_UNVERIFIED",
    "NON_ACTIONABLE",
    "WATCH_ONLY",
    "ADVANCE",
    "HEALTHY",
    "UNCLEAR",
    "receivable_risk",
    "inventory_risk",
    "goodwill_risk",
    "shareholder_action_risk",
    "request_hash",
    "reasoning_content",
)


def validate_trading_assistant_workbook(path: Path, *, profile: str = TRADING_ASSISTANT_PROFILE) -> dict[str, Any]:
    if profile != TRADING_ASSISTANT_PROFILE:
        raise ValueError(f"UNSUPPORTED_WORKBOOK_PROFILE:{profile}")
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError("WORKBOOK_MISSING")
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("WORKBOOK_ZIP_ERROR")

    workbook = load_workbook(path, data_only=False)
    try:
        sheet_names = tuple(workbook.sheetnames)
        if sheet_names not in {DAILY_SHEETS, MIDDAY_SHEETS, LEGACY_DAILY_SHEETS, LEGACY_MIDDAY_SHEETS}:
            raise ValueError(f"WORKBOOK_SHEET_PROFILE_MISMATCH:{sheet_names}")

        table_names: set[str] = set()
        table_count = 0
        for worksheet in workbook.worksheets:
            if worksheet.freeze_panes is None:
                raise ValueError(f"WORKBOOK_FREEZE_PANES_MISSING:{worksheet.title}")
            for table in worksheet.tables.values():
                table_count += 1
                if table.name in table_names:
                    raise ValueError(f"WORKBOOK_DUPLICATE_TABLE_NAME:{table.name}")
                table_names.add(table.name)
                _validate_table(worksheet, table)
            _validate_cells(worksheet)
        if table_count != len(sheet_names):
            raise ValueError(f"WORKBOOK_TABLE_COUNT_MISMATCH:{table_count}")
    finally:
        workbook.close()

    return {
        "profile": profile,
        "status": "PASS",
        "sheet_count": len(sheet_names),
        "table_count": table_count,
        "table_metadata": "PASS",
        "formula_errors": 0,
        "leading_zero_codes": "PRESERVED",
        "centered_tables": "PASS",
    }


def _validate_table(worksheet, table) -> None:
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = [worksheet.cell(min_row, column).value for column in range(min_col, max_col + 1)]
    metadata_headers = [column.name for column in table.tableColumns]
    if headers != metadata_headers:
        raise ValueError(
            f"WORKBOOK_TABLE_HEADER_METADATA_MISMATCH:{worksheet.title}:{table.name}"
        )
    if len(headers) != len(set(headers)) or any(value in (None, "") for value in headers):
        raise ValueError(f"WORKBOOK_INVALID_TABLE_HEADERS:{worksheet.title}:{table.name}")

    for row_number in range(min_row + 1, max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(min_col, max_col + 1)]
        if all(value in (None, "") for value in values):
            raise ValueError(f"WORKBOOK_BLANK_DATA_ROW:{worksheet.title}:{row_number}")
        for column in range(min_col, max_col + 1):
            cell = worksheet.cell(row_number, column)
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            if cell.alignment.horizontal != "center" or cell.alignment.vertical != "center":
                raise ValueError(f"WORKBOOK_TABLE_ALIGNMENT_ERROR:{worksheet.title}:{cell.coordinate}")

    for index, header in enumerate(headers, start=min_col):
        if "股票代码" not in str(header):
            continue
        for row_number in range(min_row + 1, max_row + 1):
            cell = worksheet.cell(row_number, index)
            if cell.value in (None, ""):
                continue
            code = str(cell.value).strip()
            if code == "无待处理事项":
                continue
            text_code = re.fullmatch(r"\d{6}", code) and cell.number_format == "@"
            numeric_code = (
                isinstance(cell.value, int)
                and 0 <= cell.value <= 999999
                and cell.number_format == "000000"
            )
            if not (text_code or numeric_code):
                raise ValueError(f"WORKBOOK_STOCK_CODE_FORMAT_ERROR:{worksheet.title}:{cell.coordinate}")

    if worksheet.title == "基本面摘要":
        header_map = {str(value): index for index, value in enumerate(headers, start=min_col)}
        concept_col = header_map.get("概念标签")
        review_col = header_map.get("人工复核")
        source_col = header_map.get("核验来源")
        for row_number in range(min_row + 1, max_row + 1):
            if concept_col:
                value = str(worksheet.cell(row_number, concept_col).value or "")
                concepts = [item for item in re.split(r"[；;]", value) if item.strip()]
                if len(concepts) > 8:
                    raise ValueError(
                        f"WORKBOOK_CONCEPT_TAG_OVERFLOW:{worksheet.title}:{row_number}"
                    )
            if review_col and source_col:
                review = str(worksheet.cell(row_number, review_col).value or "")
                source = str(worksheet.cell(row_number, source_col).value or "")
                if review == "已联网核验" and not re.search(r"https?://", source):
                    raise ValueError(
                        f"WORKBOOK_VERIFICATION_SOURCE_MISSING:{worksheet.title}:{row_number}"
                    )


def _validate_cells(worksheet) -> None:
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if not isinstance(value, str):
                continue
            upper = value.upper()
            if any(error in upper for error in FORMULA_ERRORS):
                raise ValueError(f"WORKBOOK_FORMULA_ERROR:{worksheet.title}:{cell.coordinate}")
            if not value.startswith("=") and any(text.lower() in value.lower() for text in FORBIDDEN_MACHINE_TEXT):
                raise ValueError(f"WORKBOOK_MACHINE_TEXT_LEAK:{worksheet.title}:{cell.coordinate}")
