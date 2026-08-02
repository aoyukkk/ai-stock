from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import shutil
import sqlite3
import statistics
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_EVEN
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB = ROOT / "data" / "ai_trader_dev.db"
DEFAULT_OUTPUT_ROOT = ROOT / "outputs" / "quant_factor_audit"
TARGET_CODES = (
    "000862",
    "300896",
    "300335",
    "600644",
    "600797",
    "001258",
    "300040",
    "600744",
)
FACTOR_COLUMNS = {
    "technical": "technical_score",
    "capital": "capital_score",
    "emotion": "emotion_score",
    "momentum": "momentum_score",
    "risk": "risk_score",
}
WEIGHTS = {
    "technical": Decimal("0.25"),
    "capital": Decimal("0.25"),
    "emotion": Decimal("0.20"),
    "momentum": Decimal("0.15"),
    "risk": Decimal("0.15"),
}
SCORING_SUBFACTORS = {
    "technical": (
        "ma_trend_score",
        "rsi_score",
        "atr_stability_score",
        "vwap_position_score",
    ),
    "capital": (
        "amount_score",
        "volume_ratio_score",
        "turnover_score",
        "main_inflow_score",
    ),
    "emotion": (
        "limit_up_environment_score",
        "limit_down_control_score",
        "market_heat_score",
    ),
    "momentum": ("return_5d_score", "return_20d_score"),
    "risk": (
        "volatility_risk_score",
        "drawdown_risk_score",
        "liquidity_risk_score",
    ),
}
FORMULAS = {
    "ma_trend_score": "clip(((MA5-MA20)/MA20*100-(-5))/(5-(-5))*100,0,100)",
    "rsi_score": "clip(100-abs(55-RSI14)*2,0,100)",
    "atr_stability_score": "clip((ATR14/close*100-8)/(1-8)*100,0,100)",
    "vwap_position_score": "clip(((close-VWAP20)/VWAP20*100-(-5))/(5-(-5))*100,0,100)",
    "amount_score": "clip((daily.amount-50000000)/(300000000-50000000)*100,0,100)",
    "volume_ratio_score": "clip((volume_ratio-0.8)/(2.5-0.8)*100,0,100)",
    "turnover_score": "clip((turnover_rate-0.5)/(8-0.5)*100,0,100)",
    "main_inflow_score": "clip((net_mf_amount*10000-(-5000000))/(20000000-(-5000000))*100,0,100)",
    "limit_up_environment_score": "clip((limit_up_count-10)/(80-10)*100,0,100)",
    "limit_down_control_score": "clip((limit_down_count-30)/(0-30)*100,0,100)",
    "market_heat_score": "clip(50+limit_up_count-limit_down_count,0,100)",
    "return_5d_score": "clip((return_5d-(-8))/(12-(-8))*100,0,100)",
    "return_20d_score": "clip((return_20d-(-15))/(25-(-15))*100,0,100)",
    "volatility_risk_score": "clip((volatility20-8)/(1-8)*100,0,100)",
    "drawdown_risk_score": "clip((max_drawdown20-20)/(1-20)*100,0,100)",
    "liquidity_risk_score": "clip((daily.amount-20000000)/(200000000-20000000)*100,0,100)",
}
DATASET_BY_SUBFACTOR = {
    "ma_trend_score": "daily.close",
    "rsi_score": "daily.close",
    "atr_stability_score": "daily.high,daily.low,daily.pre_close,daily.close",
    "vwap_position_score": "daily.amount,daily.vol,daily.close",
    "amount_score": "daily.amount",
    "volume_ratio_score": "HARD_CODED_DEFAULT(1.0); daily_basic.volume_ratio ignored",
    "turnover_score": "daily_basic.turnover_rate",
    "main_inflow_score": "moneyflow.net_mf_amount",
    "limit_up_environment_score": "TushareMarketDataProvider hard-coded 0",
    "limit_down_control_score": "TushareMarketDataProvider hard-coded 0",
    "market_heat_score": "derived from hard-coded limit counts",
    "return_5d_score": "daily.close",
    "return_20d_score": "daily.close",
    "volatility_risk_score": "daily.close",
    "drawdown_risk_score": "daily.close",
    "liquidity_risk_score": "daily.amount",
}
RAW_UNIT_BY_SUBFACTOR = {
    "ma_trend_score": "CNY/share",
    "rsi_score": "ratio",
    "atr_stability_score": "CNY/share",
    "vwap_position_score": "percent",
    "amount_score": "Tushare thousand CNY passed as CNY",
    "volume_ratio_score": "hard-coded ratio",
    "turnover_score": "percentage points",
    "main_inflow_score": "Tushare 10k CNY converted to CNY",
    "limit_up_environment_score": "count",
    "limit_down_control_score": "count",
    "market_heat_score": "score",
    "return_5d_score": "percent",
    "return_20d_score": "percent",
    "volatility_risk_score": "percent",
    "drawdown_risk_score": "percent",
    "liquidity_risk_score": "Tushare thousand CNY passed as CNY",
}


def _json(value: Any) -> Any:
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return value
    return value


def canonical_hash(value: Any) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_code(value: Any) -> str:
    text = str(value or "").strip().upper()
    return text.split(".", 1)[0].zfill(6) if text else ""


def detect_mapping_join_failure(
    stock_codes: Iterable[Any], mapping_codes: Iterable[Any]
) -> dict[str, Any]:
    stocks = {normalize_code(value) for value in stock_codes if normalize_code(value)}
    mappings = {
        normalize_code(value) for value in mapping_codes if normalize_code(value)
    }
    missing = sorted(stocks - mappings)
    return {
        "stock_count": len(stocks),
        "mapping_count": len(mappings),
        "matched_count": len(stocks & mappings),
        "missing_count": len(missing),
        "missing_codes": missing,
        "join_failed": bool(stocks) and len(missing) == len(stocks),
    }


def percentile(sorted_values: list[float], probability: float) -> float | None:
    if not sorted_values:
        return None
    position = (len(sorted_values) - 1) * probability
    lower = int(math.floor(position))
    upper = int(math.ceil(position))
    if lower == upper:
        return sorted_values[lower]
    fraction = position - lower
    return sorted_values[lower] * (1 - fraction) + sorted_values[upper] * fraction


def distribution(values: Iterable[Any]) -> dict[str, Any]:
    clean = [float(value) for value in values if value is not None]
    if not clean:
        return {
            "sample_count": 0,
            "non_null_count": 0,
            "null_count": 0,
            "unique_value_count": 0,
        }
    ordered = sorted(clean)
    count = len(ordered)
    mean = statistics.fmean(ordered)
    variance = statistics.pvariance(ordered)
    std = math.sqrt(variance)
    if std:
        skewness = statistics.fmean(((value - mean) / std) ** 3 for value in ordered)
        kurtosis = statistics.fmean(((value - mean) / std) ** 4 for value in ordered) - 3
    else:
        skewness = 0.0
        kurtosis = 0.0
    result = {
        "sample_count": count,
        "non_null_count": count,
        "null_count": 0,
        "unique_value_count": len(set(ordered)),
        "min": ordered[0],
        "max": ordered[-1],
        "mean": mean,
        "median": percentile(ordered, 0.5),
        "std": std,
        "variance": variance,
        "skewness": skewness,
        "kurtosis": kurtosis,
        "count_equal_50": sum(math.isclose(value, 50.0, abs_tol=1e-9) for value in ordered),
    }
    for label, probability in (
        ("p01", 0.01),
        ("p05", 0.05),
        ("p10", 0.10),
        ("p25", 0.25),
        ("p50", 0.50),
        ("p75", 0.75),
        ("p90", 0.90),
        ("p95", 0.95),
        ("p99", 0.99),
    ):
        result[label] = percentile(ordered, probability)
    result["ratio_equal_50"] = result["count_equal_50"] / count
    return result


def recompute_total(row: dict[str, Any]) -> tuple[Decimal, Decimal]:
    exact = sum(
        Decimal(str(row[FACTOR_COLUMNS[factor]])) * weight
        for factor, weight in WEIGHTS.items()
    )
    return exact, exact.quantize(Decimal("0.0001"), rounding=ROUND_HALF_EVEN)


def _descending_ranks(rows: list[dict[str, Any]], values: list[float]) -> list[int]:
    order = sorted(
        range(len(rows)),
        key=lambda index: (-values[index], normalize_code(rows[index]["stock_code"])),
    )
    ranks = [0] * len(rows)
    for rank, index in enumerate(order, start=1):
        ranks[index] = rank
    return ranks


def _spearman(first: list[int], second: list[int]) -> float:
    if not first:
        return 1.0
    mean = (len(first) + 1) / 2
    numerator = sum((x - mean) * (y - mean) for x, y in zip(first, second))
    denominator = math.sqrt(
        sum((x - mean) ** 2 for x in first)
        * sum((y - mean) ** 2 for y in second)
    )
    return numerator / denominator if denominator else 1.0


def leave_one_factor_out(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    original_ranks = [int(row["rank"]) for row in rows]
    original_top20 = {normalize_code(row["stock_code"]) for row in rows[:20]}
    original_top100 = {normalize_code(row["stock_code"]) for row in rows[:100]}
    results: list[dict[str, Any]] = []
    for factor, column in FACTOR_COLUMNS.items():
        values = [
            float(row["total_score"]) - float(row[column]) * float(WEIGHTS[factor])
            for row in rows
        ]
        new_ranks = _descending_ranks(rows, values)
        order = sorted(range(len(rows)), key=lambda index: new_ranks[index])
        new_top20 = {normalize_code(rows[index]["stock_code"]) for index in order[:20]}
        new_top100 = {normalize_code(rows[index]["stock_code"]) for index in order[:100]}
        deltas = [
            abs(original_ranks[index] - new_ranks[index])
            for index in range(len(rows))
        ]
        correlation = _spearman(original_ranks, new_ranks)
        results.append(
            {
                "factor": factor,
                "rank_correlation_without_factor": correlation,
                "top20_membership_change": len(original_top20 ^ new_top20) // 2,
                "top100_membership_change": len(original_top100 ^ new_top100) // 2,
                "mean_rank_delta": statistics.fmean(deltas),
                "max_rank_delta": max(deltas),
                "rank_disruption": 1 - correlation,
            }
        )
    denominator = sum(row["rank_disruption"] for row in results)
    for row in results:
        row["factor_effective_rank_share"] = (
            row["rank_disruption"] / denominator if denominator else 0.0
        )
    return results


def variance_contributions(
    rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    results = []
    for factor, column in FACTOR_COLUMNS.items():
        values = [float(row[column]) * float(WEIGHTS[factor]) for row in rows]
        results.append(
            {
                "factor": factor,
                "nominal_weight": float(WEIGHTS[factor]),
                "mean_nominal_contribution": statistics.fmean(values),
                "variance_contribution": statistics.pvariance(values),
            }
        )
    total_variance = sum(row["variance_contribution"] for row in results)
    for row in results:
        row["variance_share"] = (
            row["variance_contribution"] / total_variance if total_variance else 0.0
        )
    return results


def detect_global_broadcast(values: Iterable[Any]) -> bool:
    clean = [value for value in values if value is not None]
    return bool(clean) and len(set(clean)) == 1


def classify_score_origin(
    factor_group: str,
    factor_name: str,
    raw_value: Any,
    score: Any,
) -> str:
    if factor_group == "emotion":
        return "HARD_CODED_DEFAULT"
    if factor_name == "volume_ratio_score":
        return "HARD_CODED_DEFAULT"
    if factor_name in FORMULAS:
        if raw_value is None and float(score or 0) == 50:
            return "MISSING_NEUTRAL_FALLBACK"
        return "CLIP_SCORE"
    if raw_value is None and float(score or 0) == 50:
        return "MISSING_NEUTRAL_FALLBACK"
    return "UNKNOWN"


def _ro_connection(path: Path) -> sqlite3.Connection:
    connection = sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    return connection


def _rows(connection: sqlite3.Connection, sql: str, params: tuple = ()) -> list[dict[str, Any]]:
    return [dict(row) for row in connection.execute(sql, params)]


def _load_postclose_rows(connection: sqlite3.Connection) -> list[dict[str, Any]]:
    rows = _rows(connection, "SELECT * FROM postclose_official_run ORDER BY id DESC")
    for row in rows:
        row["report_json"] = _json(row.get("report_json")) or {}
        row["output_paths_json"] = _json(row.get("output_paths_json")) or {}
    return rows


def identify_target(
    connection: sqlite3.Connection,
    *,
    run_id: str | None,
    trade_date: str | None,
    workbook: Path | None,
) -> dict[str, Any]:
    run_candidates = _rows(
        connection,
        """
        SELECT q.*, r.stock_code AS first_stock_code, r.total_score AS first_total_score
        FROM quant_run q
        JOIN quant_rank_result r
          ON r.quant_run_id=q.run_id AND r.rank=1
        WHERE q.status='COMPLETED'
        ORDER BY q.id DESC
        """,
    )
    if run_id:
        run_candidates = [row for row in run_candidates if row["run_id"] == run_id]
    elif trade_date:
        run_candidates = [
            row
            for row in run_candidates
            if str(row["base_market_trade_date"]) == trade_date
        ]
    else:
        run_candidates = [
            row
            for row in run_candidates
            if normalize_code(row["first_stock_code"]) == "000862"
            and abs(float(row["first_total_score"]) - 63.16) < 0.05
        ]
    matched_runs = []
    for row in run_candidates:
        leaders = _rows(
            connection,
            """
            SELECT stock_code,total_score,emotion_score,capital_score
            FROM quant_rank_result
            WHERE quant_run_id=?
            ORDER BY rank LIMIT 5
            """,
            (row["run_id"],),
        )
        if [normalize_code(item["stock_code"]) for item in leaders] == [
            "000862",
            "300896",
            "300335",
            "600644",
            "600797",
        ]:
            matched_runs.append(row)
    if len(matched_runs) != 1:
        raise RuntimeError("AUDIT_RUN_NOT_IDENTIFIED")
    run = matched_runs[0]

    official_matches = []
    for official in _load_postclose_rows(connection):
        report = official["report_json"]
        quant = report.get("quant") or {}
        if quant.get("run_id") == run["run_id"]:
            official_matches.append(official)
    official = official_matches[0] if len(official_matches) == 1 else None
    if workbook is None and official:
        candidate = (
            (official.get("output_paths_json") or {}).get("main_workbook")
            or ((official.get("report_json") or {}).get("output_paths") or {}).get(
                "main_workbook"
            )
        )
        workbook = Path(candidate) if candidate else None
    if workbook is None or not workbook.exists():
        raise RuntimeError("AUDIT_RUN_NOT_IDENTIFIED")
    return {"run": run, "official": official, "workbook": workbook.resolve()}


def locate_quant_report(
    connection: sqlite3.Connection, run: dict[str, Any]
) -> tuple[Path, dict[str, Any]]:
    jobs = _rows(
        connection,
        "SELECT * FROM pipeline_job WHERE job_type='QUANT' AND trade_date=? ORDER BY id DESC",
        (str(run["base_market_trade_date"]),),
    )
    report_root = ROOT / "data" / "reports"
    candidates: list[Path] = []
    for job in jobs:
        ids = _json(job.get("run_ids")) or {}
        if ids.get("quant_run_id") == run["run_id"]:
            candidates.append(
                report_root
                / f"quant_{str(run['base_market_trade_date']).replace('-', '')}_{job['job_id']}.json"
            )
    candidates.extend(
        report_root.glob(
            f"quant_{str(run['base_market_trade_date']).replace('-', '')}_*.json"
        )
    )
    seen: set[Path] = set()
    matched: list[tuple[Path, dict[str, Any]]] = []
    for path in candidates:
        path = path.resolve()
        if path in seen or not path.exists():
            continue
        seen.add(path)
        payload = json.loads(path.read_text(encoding="utf-8"))
        scored = payload.get("all_scored_stocks") or []
        if (
            payload.get("factor_version") == run["factor_version"]
            and int(payload.get("scored_count") or 0) == int(run["scored_count"])
            and scored
            and normalize_code(scored[0].get("stock_code")) == "000862"
            and abs(float(scored[0].get("total_score") or 0) - 63.1629) < 0.0001
        ):
            matched.append((path, payload))
    unique = {path: payload for path, payload in matched}
    if len(unique) != 1:
        raise RuntimeError("RAW_INPUT_CACHE_MISSING")
    return next(iter(unique.items()))


def read_workbook_quant(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "量化前100" not in workbook.sheetnames:
            raise RuntimeError("AUDIT_RUN_NOT_IDENTIFIED")
        sheet = workbook["量化前100"]
        header_row = None
        headers: list[str] = []
        for row_index, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True),
            start=1,
        ):
            values = [str(value).strip() if value is not None else "" for value in row]
            if "股票代码" in values and "量化总分" in values:
                header_row = row_index
                headers = values
                break
        if header_row is None:
            raise RuntimeError("AUDIT_RUN_NOT_IDENTIFIED")
        rows = []
        for row in sheet.iter_rows(
            min_row=header_row + 1,
            max_row=sheet.max_row,
            values_only=True,
        ):
            record = {
                headers[index]: row[index] if index < len(row) else None
                for index in range(len(headers))
                if headers[index]
            }
            if not record.get("股票代码"):
                continue
            record["股票代码"] = normalize_code(record["股票代码"])
            rows.append(record)
        return {
            "sheet_name": sheet.title,
            "header_row": header_row,
            "rows": rows,
            "sheet_names": workbook.sheetnames,
        }
    finally:
        workbook.close()


def database_excel_reconciliation(
    db_rows: list[dict[str, Any]], workbook_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    db_by_code = {normalize_code(row["stock_code"]): row for row in db_rows}
    mapping = {
        "量化总分": "total_score",
        "技术得分": "technical_score",
        "资金得分": "capital_score",
        "情绪得分": "emotion_score",
        "动量得分": "momentum_score",
        "风险得分": "risk_score",
    }
    result = []
    for row in workbook_rows:
        code = normalize_code(row.get("股票代码"))
        database = db_by_code.get(code)
        if not database:
            result.append(
                {
                    "stock_code": code,
                    "field": "ROW",
                    "database_value": None,
                    "excel_value": None,
                    "difference": None,
                    "display_difference": None,
                    "status": "DATABASE_ROW_MISSING",
                }
            )
            continue
        for excel_field, db_field in mapping.items():
            db_value = float(database[db_field])
            excel_value = float(row.get(excel_field))
            difference = excel_value - db_value
            display_difference = round(excel_value, 2) - round(db_value, 2)
            result.append(
                {
                    "stock_code": code,
                    "field": db_field,
                    "database_value": db_value,
                    "excel_value": excel_value,
                    "difference": difference,
                    "display_difference": display_difference,
                    "status": "PASS" if abs(difference) <= 1e-8 else "FAIL",
                }
            )
    return result


def _manifest_dataset_map(connection: sqlite3.Connection, manifest_id: str) -> dict[str, dict[str, Any]]:
    row = connection.execute(
        "SELECT * FROM run_data_manifest WHERE manifest_id=?", (manifest_id,)
    ).fetchone()
    if row is None:
        return {}
    record = dict(row)
    datasets = (_json(record.get("required_dataset_watermarks")) or []) + (
        _json(record.get("optional_dataset_watermarks")) or []
    )
    return {item["dataset_name"]: item for item in datasets}


def _cache_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return payload
    for key in ("records", "data", "items"):
        if isinstance(payload.get(key), list):
            return payload[key]
    return []


def raw_coverage_audit(
    connection: sqlite3.Connection, run: dict[str, Any], scored_codes: set[str]
) -> list[dict[str, Any]]:
    manifest = _manifest_dataset_map(connection, str(run["data_manifest_id"]))
    trade_key = str(run["base_market_trade_date"]).replace("-", "")
    field_map = {
        "daily": (
            "ts_code",
            "trade_date",
            "open",
            "high",
            "low",
            "close",
            "pre_close",
            "change",
            "pct_chg",
            "vol",
            "amount",
        ),
        "daily_basic": (
            "turnover_rate",
            "turnover_rate_f",
            "volume_ratio",
            "total_share",
            "float_share",
            "free_share",
            "total_mv",
            "circ_mv",
            "pe",
            "pb",
        ),
        "moneyflow": (
            "buy_sm_amount",
            "sell_sm_amount",
            "buy_md_amount",
            "sell_md_amount",
            "buy_lg_amount",
            "sell_lg_amount",
            "buy_elg_amount",
            "sell_elg_amount",
            "net_mf_amount",
            "net_mf_vol",
        ),
        "stk_limit": ("pre_close", "up_limit", "down_limit"),
        "adj_factor": ("adj_factor",),
    }
    units = {
        ("daily", "amount"): ("THOUSAND_CNY", 1, "RealtimeQuote.amount"),
        ("daily", "vol"): ("TUSHARE_LOT", 1, "KLineBar.volume"),
        ("daily", "pct_chg"): ("PERCENTAGE_POINTS", 1, "change_percent"),
        ("daily_basic", "turnover_rate"): (
            "PERCENTAGE_POINTS",
            1,
            "CapitalFlowSnapshot.turnover_rate",
        ),
        ("daily_basic", "volume_ratio"): (
            "RATIO",
            0,
            "IGNORED; hard-coded 1.0",
        ),
        ("moneyflow", "net_mf_amount"): (
            "TEN_THOUSAND_CNY",
            10000,
            "CapitalFlowSnapshot.main_net_inflow",
        ),
    }
    used = {
        ("daily", "open"): "technical/risk diagnostics",
        ("daily", "high"): "technical ATR/VWAP",
        ("daily", "low"): "technical ATR/VWAP",
        ("daily", "close"): "technical/momentum/risk",
        ("daily", "pre_close"): "technical ATR",
        ("daily", "vol"): "technical VWAP",
        ("daily", "amount"): "technical VWAP/capital/risk",
        ("daily_basic", "turnover_rate"): "capital",
        ("moneyflow", "net_mf_amount"): "capital",
    }
    result = []
    for dataset, fields in field_map.items():
        watermark = manifest.get(dataset, {})
        cache_path = Path(
            watermark.get("cache_key")
            or ROOT
            / "data"
            / "cache"
            / "tushare"
            / "trade_date"
            / dataset
            / f"{trade_key}.json"
        )
        rows = _cache_rows(cache_path)
        keys = [
            (normalize_code(row.get("ts_code")), str(row.get("trade_date") or trade_key))
            for row in rows
        ]
        duplicate_count = len(keys) - len(set(keys))
        unique_codes = {code for code, _ in keys if code}
        dates = sorted({str(row.get("trade_date")) for row in rows if row.get("trade_date")})
        for field in fields:
            null_count = sum(row.get(field) is None for row in rows)
            raw_unit, multiplier, standardized = units.get(
                (dataset, field), ("UNDECLARED_BY_ADAPTER", 1, field)
            )
            actual_used = (dataset, field) in used
            unit_status = "PASS"
            if dataset == "daily" and field == "amount":
                unit_status = "FAIL_UNIT_MULTIPLIER_1000_MISSING"
            elif dataset == "daily_basic" and field == "volume_ratio":
                unit_status = "FIELD_PRESENT_BUT_IGNORED"
            result.append(
                {
                    "dataset": dataset,
                    "raw_field": field,
                    "raw_unit_from_adapter": raw_unit,
                    "unit_multiplier": multiplier,
                    "standardized_field": standardized,
                    "expected_trade_date": str(run["base_market_trade_date"]),
                    "actual_trade_date": ",".join(dates) if dates else None,
                    "available_at": watermark.get("fetched_at"),
                    "cache_path": str(cache_path),
                    "row_count": len(rows),
                    "unique_stock_count": len(unique_codes),
                    "duplicate_count": duplicate_count,
                    "null_count": null_count,
                    "null_rate": null_count / len(rows) if rows else None,
                    "coverage": len(unique_codes & scored_codes) / len(scored_codes)
                    if scored_codes
                    else None,
                    "future_data_check": "PASS"
                    if not dates or max(dates) <= trade_key
                    else "FAIL",
                    "used_by_factor": used.get((dataset, field), "NOT_USED"),
                    "actually_used_in_run": actual_used,
                    "unit_validation": unit_status,
                }
            )
    return result


def _detail_map(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    return {
        normalize_code(row["stock_code"]): row.get("factor_detail") or []
        for row in report.get("top_stocks") or []
    }


def missing_data_audit(
    rows: list[dict[str, Any]],
    detail_by_code: dict[str, list[dict[str, Any]]],
    coverage_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    dataset_codes: dict[str, set[str]] = {}
    for item in coverage_rows:
        dataset = item["dataset"]
        if dataset in dataset_codes:
            continue
        dataset_codes[dataset] = {
            normalize_code(row.get("ts_code"))
            for row in _cache_rows(Path(item["cache_path"]))
            if row.get("ts_code")
        }
    results = []
    for row in rows:
        code = normalize_code(row["stock_code"])
        details = {
            (item["factor_group"], item["factor_name"]): item
            for item in detail_by_code.get(code, [])
        }
        for factor, expected_names in SCORING_SUBFACTORS.items():
            missing: list[str] = []
            fallback: list[str] = []
            reason = ""
            if factor == "capital":
                if code not in dataset_codes.get("moneyflow", set()):
                    missing.extend(
                        ["volume_ratio_score", "turnover_score", "main_inflow_score"]
                    )
                    fallback.extend(missing)
                    reason = "MONEYFLOW_NOT_COVERED"
                else:
                    missing.append("volume_ratio_score")
                    fallback.append("volume_ratio_score")
                    reason = "DAILY_BASIC_VOLUME_RATIO_IGNORED_AND_HARD_CODED_1"
            elif factor == "emotion":
                missing.extend(expected_names)
                fallback.extend(expected_names)
                reason = "TUSHARE_MARKET_EMOTION_PLACEHOLDER_0_0_50"
            elif details:
                for name in expected_names:
                    item = details.get((factor, name))
                    if not item or item.get("raw_value") is None:
                        missing.append(name)
                        if item and float(item.get("score") or 0) == 50:
                            fallback.append(name)
                reason = "INSUFFICIENT_HISTORY" if missing else ""
            else:
                reason = "RAW_FACTOR_DETAIL_NOT_PERSISTED_FOR_BOTTOM_310"
            available = len(expected_names) - len(missing)
            results.append(
                {
                    "stock_code": code,
                    "factor": factor,
                    "available_subfactor_count": available
                    if reason != "RAW_FACTOR_DETAIL_NOT_PERSISTED_FOR_BOTTOM_310"
                    else None,
                    "expected_subfactor_count": len(expected_names),
                    "missing_subfactors": ",".join(missing),
                    "missing_reason": reason,
                    "missing_is_random": False if reason else None,
                    "fallback_used": bool(fallback),
                    "fallback_subfactors": ",".join(fallback),
                    "weight_before": float(WEIGHTS[factor]),
                    "effective_weight_after_reweight": float(WEIGHTS[factor]),
                    "confidence_penalty": 0,
                    "data_quality_status": "DEGRADED" if reason else "COMPLETE",
                }
            )
    return results


def capital_collision_audit(
    rows: list[dict[str, Any]],
    detail_by_code: dict[str, list[dict[str, Any]]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    full_counter = Counter(round(float(row["capital_score"]), 4) for row in rows)
    top100_counter = Counter(
        round(float(row["capital_score"]), 4) for row in rows[:100]
    )
    result = []
    for score, count in sorted(full_counter.items(), key=lambda item: (-item[1], item[0])):
        if count < 2:
            continue
        representative = next(
            normalize_code(row["stock_code"])
            for row in rows
            if round(float(row["capital_score"]), 4) == score
        )
        source = "FIXED_RANGE_SCORE_COLLISION"
        if math.isclose(score, 52.9412, abs_tol=0.00005):
            source = (
                "(amount=0 + hard-coded volume-ratio=11.7647 + "
                "turnover=100 + inflow=100)/4"
            )
        elif math.isclose(score, 37.5, abs_tol=0.00005):
            source = "missing moneyflow: (amount≈0 + 50 + 50 + 50)/4"
        result.append(
            {
                "factor": "capital",
                "score": score,
                "full_a_count": count,
                "top100_count": top100_counter.get(score, 0),
                "representative_stock": representative,
                "collision_source": source,
                "rank_denominator": "N/A",
                "tie_method": "N/A",
                "reweighted": False,
                "fallback": math.isclose(score, 37.5, abs_tol=0.00005),
                "database_precision": 4,
                "export_precision": 2,
            }
        )
    target_codes = [
        normalize_code(row["stock_code"])
        for row in rows
        if math.isclose(float(row["capital_score"]), 52.9412, abs_tol=0.00005)
    ]
    verified = 0
    for code in target_codes:
        details = {
            item["factor_name"]: item
            for item in detail_by_code.get(code, [])
            if item["factor_group"] == "capital"
        }
        scores = [float(details.get(name, {}).get("score", math.nan)) for name in SCORING_SUBFACTORS["capital"]]
        if all(
            math.isclose(actual, expected, abs_tol=0.0001)
            for actual, expected in zip(scores, (0.0, 11.7647, 100.0, 100.0))
        ):
            verified += 1
    summary = {
        "unique_capital_scores": len(full_counter),
        "duplicated_score_groups": sum(count > 1 for count in full_counter.values()),
        "largest_tie_group": max(full_counter.values()),
        "largest_tie_score": full_counter.most_common(1)[0][0],
        "count_equal_52_94": len(target_codes),
        "count_equal_52_94_formula_verified": verified,
        "count_equal_50": sum(
            math.isclose(float(row["capital_score"]), 50, abs_tol=0.00005)
            for row in rows
        ),
        "decimal_precision": 4,
        "calculation_precision": "Decimal quantized to 4 decimals per subfactor/group",
        "export_precision": 2,
    }
    return result, summary


def normalization_universe_audit(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for factor, names in SCORING_SUBFACTORS.items():
        for name in names:
            result.append(
                {
                    "factor_name": name,
                    "expected_universe": "FULL_A per YAML normalize.method=percentile",
                    "actual_universe": "PER_STOCK_FIXED_RANGE",
                    "expected_count": len(rows),
                    "actual_count": "N/A",
                    "exclusions": "N/A",
                    "missing_exclusions": "N/A",
                    "rank_denominator": "N/A",
                    "tie_method": "N/A",
                    "normalization_scope": "STOCK_LEVEL_FIXED_THRESHOLDS"
                    if factor != "emotion"
                    else "GLOBAL_MARKET_FIXED_THRESHOLDS",
                    "universe_mismatch": True,
                    "finding": "DOCUMENT_CODE_MISMATCH; percentile_score is not called",
                }
            )
    return result


def lineage_rows(
    rows: list[dict[str, Any]],
    workbook_rows: list[dict[str, Any]],
    detail_by_code: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    db_by_code = {normalize_code(row["stock_code"]): row for row in rows}
    excel_by_code = {
        normalize_code(row.get("股票代码")): row for row in workbook_rows
    }
    result = []
    subfactor_weights = {
        "technical": Decimal("0.25"),
        "capital": Decimal("0.25"),
        "emotion": Decimal("0.3333333333333333"),
        "momentum": None,
        "risk": Decimal("0.3333333333333333"),
    }
    for code in TARGET_CODES:
        db_row = db_by_code.get(code)
        if not db_row:
            continue
        for detail in detail_by_code.get(code, []):
            factor = detail["factor_group"]
            name = detail["factor_name"]
            is_scoring = name in SCORING_SUBFACTORS.get(factor, ())
            weight = detail.get("weight")
            if weight is None and is_scoring:
                weight = float(subfactor_weights.get(factor) or 0)
            contribution = (
                float(detail.get("score") or 0) * float(weight)
                if is_scoring and weight is not None
                else 0
            )
            result.append(
                {
                    "stock_code": code,
                    "stock_name": db_row.get("stock_name"),
                    "level": "subfactor" if is_scoring else "diagnostic",
                    "factor_group": factor,
                    "factor_name": name,
                    "dataset": DATASET_BY_SUBFACTOR.get(name, "diagnostic/report"),
                    "field": DATASET_BY_SUBFACTOR.get(name, ""),
                    "trade_date": db_row.get("base_market_trade_date"),
                    "raw_value": detail.get("raw_value"),
                    "raw_unit": RAW_UNIT_BY_SUBFACTOR.get(name, "diagnostic"),
                    "standardized_value": detail.get("raw_value"),
                    "unit_conversion": "see raw coverage audit",
                    "formula": FORMULAS.get(name, detail.get("explain_text")),
                    "lookback": _lookback(name),
                    "derived_value": detail.get("normalized_value"),
                    "method": "FIXED_RANGE_CLIP" if name in FORMULAS else "DIAGNOSTIC",
                    "universe": "N/A",
                    "rank": None,
                    "percentile": None,
                    "tie_method": "N/A",
                    "missing": detail.get("raw_value") is None,
                    "reason": ""
                    if detail.get("raw_value") is not None
                    else "RAW_VALUE_NOT_AVAILABLE",
                    "fallback": classify_score_origin(
                        factor, name, detail.get("raw_value"), detail.get("score")
                    ),
                    "reweight": False,
                    "subfactor_score": detail.get("score"),
                    "weight": weight,
                    "contribution": contribution,
                    "group_score": db_row.get(FACTOR_COLUMNS.get(factor, "")),
                    "effective_weight": float(WEIGHTS.get(factor, 0)),
                    "total_score": db_row["total_score"],
                    "recomputed_total": recompute_total(db_row)[1],
                    "database_value": db_row.get(FACTOR_COLUMNS.get(factor, "")),
                    "excel_value": (excel_by_code.get(code) or {}).get(
                        _excel_factor_name(factor)
                    ),
                    "difference": _difference(
                        db_row.get(FACTOR_COLUMNS.get(factor, "")),
                        (excel_by_code.get(code) or {}).get(_excel_factor_name(factor)),
                    ),
                    "score_origin": classify_score_origin(
                        factor, name, detail.get("raw_value"), detail.get("score")
                    ),
                }
            )
        exact, stored = recompute_total(db_row)
        result.append(
            {
                "stock_code": code,
                "stock_name": db_row.get("stock_name"),
                "level": "total",
                "factor_group": "quant",
                "factor_name": "total_score",
                "dataset": "quant_rank_result",
                "field": "technical/capital/emotion/momentum/risk",
                "trade_date": db_row.get("base_market_trade_date"),
                "raw_value": None,
                "raw_unit": "score",
                "standardized_value": None,
                "unit_conversion": "N/A",
                "formula": "0.25*T+0.25*C+0.20*E+0.15*M+0.15*R",
                "lookback": "N/A",
                "derived_value": float(exact),
                "method": "WEIGHTED_SUM_THEN_QUANTIZE_4DP",
                "universe": "N/A",
                "rank": db_row["rank"],
                "percentile": None,
                "tie_method": "stock_code deterministic secondary key",
                "missing": False,
                "reason": "",
                "fallback": "",
                "reweight": False,
                "subfactor_score": None,
                "weight": 1.0,
                "contribution": None,
                "group_score": None,
                "effective_weight": 1.0,
                "total_score": db_row["total_score"],
                "recomputed_total": float(stored),
                "database_value": db_row["total_score"],
                "excel_value": (excel_by_code.get(code) or {}).get("量化总分"),
                "difference": _difference(
                    db_row["total_score"],
                    (excel_by_code.get(code) or {}).get("量化总分"),
                ),
                "score_origin": "WEIGHTED_SUM",
            }
        )
    return result


def _lookback(name: str) -> str:
    if "5d" in name or name == "ma_trend_score":
        return "5/20 sessions"
    if "20d" in name or name.startswith(("volatility", "drawdown", "vwap")):
        return "20 sessions"
    if name.startswith(("rsi", "atr")):
        return "14 sessions"
    return "current trade date"


def _excel_factor_name(factor: str) -> str:
    return {
        "technical": "技术得分",
        "capital": "资金得分",
        "emotion": "情绪得分",
        "momentum": "动量得分",
        "risk": "风险得分",
    }.get(factor, "")


def _difference(first: Any, second: Any) -> float | None:
    if first is None or second is None:
        return None
    return float(second) - float(first)


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    fields: list[str] = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _effective_config(
    connection: sqlite3.Connection, run: dict[str, Any]
) -> dict[str, Any]:
    import yaml

    yaml_payload = yaml.safe_load((ROOT / "config" / "quant_factor.yaml").read_text("utf-8"))
    yaml_weights = yaml_payload["quant_factor"]["weights"]
    db_rows = _rows(
        connection,
        """
        SELECT * FROM system_config
        WHERE config_key LIKE 'quant_factor.weights.%'
        """,
    )
    history_rows = _rows(
        connection,
        """
        SELECT * FROM config_history
        WHERE config_key LIKE 'quant_factor.weights.%' AND time <= ?
        ORDER BY time DESC
        """,
        (str(run["created_at"]),),
    )
    runtime_options = {}
    jobs = _rows(
        connection,
        "SELECT checkpoint FROM pipeline_job WHERE job_type='QUANT' AND trade_date=?",
        (str(run["base_market_trade_date"]),),
    )
    for job in jobs:
        checkpoint = _json(job.get("checkpoint")) or {}
        options = checkpoint.get("options") or {}
        runtime_options.update(
            {key: value for key, value in options.items() if "weight" in key}
        )
    database_override = {
        row["config_key"]: _json(row["config_value"]) for row in db_rows
    }
    effective = dict(yaml_weights)
    for factor in WEIGHTS:
        key = f"quant_factor.weights.{factor}"
        if key in database_override:
            effective[factor] = database_override[key]
    return {
        "weights": {key: float(value) for key, value in effective.items()},
        "weight_source": "YAML"
        if not database_override and not runtime_options
        else "DATABASE_OR_RUNTIME",
        "runtime_override": runtime_options,
        "database_override": database_override,
        "yaml_value": yaml_weights,
        "code_default": {key: float(value) for key, value in WEIGHTS.items()},
        "config_history_match": not history_rows,
        "effective_config_hash": canonical_hash(yaml_payload),
        "normalize": yaml_payload["quant_factor"]["normalize"],
        "raw": yaml_payload,
    }


def _workbook_payload(
    audit: dict[str, Any],
    distributions: dict[str, dict[str, Any]],
    raw_coverage: list[dict[str, Any]],
    collision_rows: list[dict[str, Any]],
    missing_rows: list[dict[str, Any]],
    normalization_rows: list[dict[str, Any]],
    contribution_rows: list[dict[str, Any]],
    loo_rows: list[dict[str, Any]],
    lineage: list[dict[str, Any]],
    reconciliation: list[dict[str, Any]],
) -> dict[str, Any]:
    summary = [
        {"项目": key, "值": value}
        for key, value in (
            ("Phase", audit["phase"]),
            ("Audit run", audit["audit_run"]),
            ("Trade date", audit["trade_date"]),
            ("Workbook", audit["workbook"]["path"]),
            ("Quant version", audit["quant_version"]),
            ("Effective universe size", audit["effective_universe_size"]),
            ("Emotion finding", audit["emotion"]["finding"]),
            ("Capital finding", audit["capital"]["finding"]),
            ("Score compression", audit["score_compression_diagnosis"]),
            ("Final status", audit["final_status"]),
        )
    ]
    config_rows = [
        {
            "因子": factor,
            "生效权重": audit["config"]["weights"][factor],
            "YAML值": audit["config"]["yaml_value"][factor],
            "代码默认": audit["config"]["code_default"][factor],
            "来源": audit["config"]["weight_source"],
            "运行时覆盖": json.dumps(audit["config"]["runtime_override"], ensure_ascii=False),
            "数据库覆盖": json.dumps(audit["config"]["database_override"], ensure_ascii=False),
        }
        for factor in WEIGHTS
    ]
    distribution_sheets = {}
    for key, title in (
        ("technical_score", "04_技术分布"),
        ("capital_score", "05_资金分布"),
        ("emotion_score", "06_情绪分布"),
        ("momentum_score", "07_动量分布"),
        ("risk_score", "08_风险分布"),
        ("total_score", "09_总分分布"),
    ):
        distribution_sheets[title] = [
            {"统计项": name, "值": value}
            for name, value in distributions[key].items()
        ]
    emotion_rows = [
        {"项目": key, "值": value}
        for key, value in audit["emotion"].items()
    ]
    findings = [
        {"类型": item["type"], "状态": item["status"], "说明": item["detail"]}
        for item in audit["findings"]
    ]
    runtime_rows = [
        {"检查项": key, "结果": value}
        for key, value in audit["runtime_audit"].items()
    ]
    contribution_combined = [
        {**row, "analysis": "variance"} for row in contribution_rows
    ] + [{**row, "analysis": "leave_one_factor_out"} for row in loo_rows]
    sheets = {
        "01_审计总览": summary,
        "02_运行与配置": config_rows,
        "03_Tushare原始覆盖": raw_coverage,
        **distribution_sheets,
        "10_资金同分审计": collision_rows,
        "11_情绪50审计": emotion_rows,
        "12_缺失值与回退": missing_rows,
        "13_归一化Universe": normalization_rows,
        "14_有效因子贡献": contribution_combined,
        "15_逐股Lineage": lineage,
        "16_DB与Excel对账": reconciliation,
        "17_问题与建议": findings,
        "18_运行审计": runtime_rows,
    }
    return {"sheets": sheets, "meta": {"final_status": audit["final_status"]}}


def _build_excel(payload: dict[str, Any], output_path: Path) -> dict[str, Any]:
    dependency_root = (
        Path.home()
        / ".cache"
        / "codex-runtimes"
        / "codex-primary-runtime"
        / "dependencies"
        / "node"
        / "node_modules"
    )
    if not (dependency_root / "@oai" / "artifact-tool").exists():
        raise RuntimeError("ARTIFACT_TOOL_DEPENDENCY_NOT_FOUND")
    builder_source = ROOT / "scripts" / "build_quant_factor_audit_excel.mjs"
    with tempfile.TemporaryDirectory(prefix="quant-factor-audit-") as tmp_name:
        build_dir = Path(tmp_name)
        builder = build_dir / builder_source.name
        shutil.copy2(builder_source, builder)
        payload_path = build_dir / "payload.json"
        payload_path.write_text(
            json.dumps(payload, ensure_ascii=False, default=str),
            encoding="utf-8",
        )
        link = build_dir / "node_modules"
        result = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(dependency_root)],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            raise RuntimeError("ARTIFACT_TOOL_JUNCTION_FAILED")
        completed = subprocess.run(
            ["node", str(builder), str(payload_path), str(output_path)],
            cwd=build_dir,
            capture_output=True,
            text=True,
            check=False,
            timeout=300,
        )
        validation_path = output_path.with_suffix(".validation.json")
        teardown_fault = completed.returncode in {3221226505, -1073740791}
        complete_artifacts = output_path.exists() and validation_path.exists()
        if (completed.returncode != 0 and not (teardown_fault and complete_artifacts)) or (
            not complete_artifacts
        ):
            raise RuntimeError(
                f"AUDIT_WORKBOOK_EXPORT_FAILED:{completed.returncode}:"
                f"{completed.stderr[-1000:]}"
            )
        inspect_sidecar = Path(f"{output_path}.inspect.ndjson")
        inspect_sidecar.unlink(missing_ok=True)
        validation = (
            json.loads(validation_path.read_text(encoding="utf-8"))
            if validation_path.exists()
            else {}
        )
        validation["artifact_tool_process_exit_code"] = completed.returncode
        validation["native_teardown_warning"] = teardown_fault
        validation_path.unlink(missing_ok=True)
        preview_dir = output_path.parent / "preview"
        return {**validation, "preview_dir": str(preview_dir)}


def _markdown(audit: dict[str, Any]) -> str:
    emotion = audit["emotion"]
    capital = audit["capital"]
    score = audit["distributions"]["total_score"]
    lines = [
        "# TUSHARE_BASELINE_V1 Quant因子审计",
        "",
        f"- Audit run: `{audit['audit_run']}`",
        f"- Trade date: `{audit['trade_date']}`",
        f"- Workbook: `{audit['workbook']['path']}`",
        f"- Final status: `{audit['final_status']}`",
        "",
        "## 核心结论",
        "",
        f"1. Emotion全部50是生产数据接线问题：{emotion['finding']}",
        f"2. Emotion横截面排名贡献为 `{emotion['effective_rank_contribution']:.6f}`。",
        f"3. Capital 52.94共 `{capital['count_equal_52_94']}` 只，来源为：{capital['collision_source']}",
        f"4. daily.amount单位校验失败，Tushare千元值未乘1000即进入人民币阈值。",
        f"5. 量化总分范围 `{score['min']:.4f}`—`{score['max']:.4f}`，均值 `{score['mean']:.4f}`。",
        f"6. 当前Baseline不宜原样继续作为已校准绝对分；可保留历史结果，但应先做最小数据接线修复和前向Shadow验证。",
        "",
        "## 证据",
        "",
        "- 情绪输入由Tushare provider硬编码为涨停0、跌停0，转换后市场热度50；三项平均为50。",
        "- 同日全市场情绪本来就是GLOBAL_MARKET值，因此即便接入真实值，也不会区分个股；其横截面方差贡献应为0。",
        "- 资金量比在正式转换中硬编码为1.0，得分固定11.7647。",
        "- daily.amount未执行千元到元的1000倍转换，导致成交额与流动性子因子大量被clip到0。",
        "- 52.9412精确等于 `(0 + 11.7647 + 100 + 100) / 4`，不是百分位，也不是Excel模板默认值。",
        "- YAML声明normalize.method=percentile，但生产因子调用固定区间min-max/clip；标记DOCUMENT_CODE_MISMATCH。",
        "",
        "## 建议的最小修复范围（本阶段未实施）",
        "",
        "1. 仅修复Tushare适配层单位：daily.amount乘1000后再进入绝对金额阈值。",
        "2. 将daily_basic.volume_ratio真实字段传入CapitalFlowSnapshot，移除1.0占位。",
        "3. 将同交易日只读market_emotion_snapshot接入Quant的大盘情绪输入，并明确GLOBAL_MARKET语义。",
        "4. 不回写历史run；以新factor_version做Shadow重算、分布与前向收益验证后再决定晋级。",
        "",
        "## 运行边界",
        "",
        "- External API calls: 0",
        "- LLM calls: 0",
        "- Orders: 0",
        "- Scheduler: off",
        "- Production database writes: 0",
        "- Git commit: not executed",
    ]
    return "\n".join(lines) + "\n"


def run_audit(args: argparse.Namespace) -> dict[str, Any]:
    if not args.read_only or not args.no_external_api or not args.no_llm:
        raise RuntimeError("AUDIT_SAFETY_FLAGS_REQUIRED")
    database_path = Path(args.database).resolve()
    database_before = {
        "size": database_path.stat().st_size,
        "mtime_ns": database_path.stat().st_mtime_ns,
    }
    connection = _ro_connection(database_path)
    try:
        target = identify_target(
            connection,
            run_id=args.run_id or None,
            trade_date=args.trade_date or None,
            workbook=Path(args.workbook).resolve() if args.workbook else None,
        )
        run = target["run"]
        official = target["official"]
        workbook_path = target["workbook"]
        report_path, quant_report = locate_quant_report(connection, run)
        workbook = read_workbook_quant(workbook_path)
        db_rows = _rows(
            connection,
            """
            SELECT r.*,
                   json_extract(r.factor_detail_reference,'$.stock_name') AS stock_name,
                   json_extract(r.factor_detail_reference,'$.level_one_sector') AS level_one_sector
            FROM quant_rank_result r
            WHERE r.quant_run_id=?
            ORDER BY r.rank
            """,
            (run["run_id"],),
        )
        for row in db_rows:
            row["base_market_trade_date"] = str(run["base_market_trade_date"])
        if len(db_rows) != int(run["scored_count"]):
            raise RuntimeError("DATABASE_RANK_COUNT_MISMATCH")
        if len(workbook["rows"]) != 100:
            raise RuntimeError("WORKBOOK_TOP100_COUNT_MISMATCH")
        detail_by_code = _detail_map(quant_report)
        scored_codes = {normalize_code(row["stock_code"]) for row in db_rows}
        config = _effective_config(connection, run)
        raw_coverage = raw_coverage_audit(connection, run, scored_codes)
        missing_rows = missing_data_audit(
            db_rows, detail_by_code, raw_coverage
        )
        collision_rows, capital_summary = capital_collision_audit(
            db_rows, detail_by_code
        )
        normalization_rows = normalization_universe_audit(db_rows)
        reconciliation = database_excel_reconciliation(
            db_rows, workbook["rows"]
        )
        lineage = lineage_rows(db_rows, workbook["rows"], detail_by_code)
        distributions = {
            column: distribution(row[column] for row in db_rows)
            for column in (*FACTOR_COLUMNS.values(), "total_score")
        }
        full_contributions = variance_contributions(db_rows)
        top100_contributions = variance_contributions(db_rows[:100])
        loo = leave_one_factor_out(db_rows)
        contribution_rows = [
            {**row, "scope": "FULL_A"} for row in full_contributions
        ] + [{**row, "scope": "TOP100"} for row in top100_contributions]
        emotion_loo = next(row for row in loo if row["factor"] == "emotion")
        manifest_map = _manifest_dataset_map(
            connection, str(run["data_manifest_id"])
        )
        universe_hash = canonical_hash(sorted(scored_codes))
        official_report = (official or {}).get("report_json") or {}
        official_outputs = (official or {}).get("output_paths_json") or {}
        official_recorded_hash = official_outputs.get("main_workbook_sha256")
        current_workbook_hash = file_sha256(workbook_path)
        exact_reconciliation_max = max(
            abs(float(row["difference"] or 0)) for row in reconciliation
        )
        display_reconciliation_max = max(
            abs(float(row["display_difference"] or 0)) for row in reconciliation
        )
        unit_failures = [
            row
            for row in raw_coverage
            if str(row["unit_validation"]).startswith("FAIL")
        ]
        data_coverage_failures = [
            row
            for row in raw_coverage
            if row["actually_used_in_run"]
            and row["coverage"] is not None
            and float(row["coverage"]) < 0.99
        ]
        findings = [
            {
                "type": "DOCUMENT_CODE_MISMATCH",
                "status": "CONFIRMED",
                "detail": "YAML declares percentile normalization; production uses fixed-range min-max/clip.",
            },
            {
                "type": "IMPLEMENTATION_ISSUE",
                "status": "CONFIRMED",
                "detail": "Tushare market emotion returns hard-coded 0/0 and becomes global constant 50.",
            },
            {
                "type": "UNIT_ERROR",
                "status": "CONFIRMED",
                "detail": "daily.amount thousand-CNY is passed without x1000 into CNY thresholds.",
            },
            {
                "type": "IMPLEMENTATION_ISSUE",
                "status": "CONFIRMED",
                "detail": "daily_basic.volume_ratio is available but production passes hard-coded 1.0.",
            },
            {
                "type": "DATABASE_EXPORT_MISMATCH",
                "status": "NOT_FOUND"
                if exact_reconciliation_max <= 1e-8
                else "CONFIRMED",
                "detail": f"DB/Excel numeric max difference={exact_reconciliation_max:.12g}.",
            },
            {
                "type": "WORKBOOK_HASH_DRIFT",
                "status": "CONFIRMED"
                if official_recorded_hash
                and official_recorded_hash.lower() != current_workbook_hash.lower()
                else "NOT_FOUND",
                "detail": (
                    f"official_recorded={official_recorded_hash}; current={current_workbook_hash}"
                ),
            },
        ]
        quant_hash_before = canonical_hash(
            {
                str(path.relative_to(ROOT)): file_sha256(path)
                for path in (
                    ROOT / "quant" / "config.py",
                    ROOT / "quant" / "factors.py",
                    ROOT / "quant" / "normalizer.py",
                    ROOT / "quant" / "ranking.py",
                    ROOT / "config" / "quant_factor.yaml",
                )
            }
        )
        audit = {
            "phase": "TUSHARE_BASELINE_V1 Quant Factor Audit",
            "audit_run": run["run_id"],
            "trade_date": str(run["base_market_trade_date"]),
            "workbook": {
                "path": str(workbook_path),
                "sha256": current_workbook_hash,
                "official_recorded_sha256": official_recorded_hash,
                "sheet_name": workbook["sheet_name"],
                "source_trade_date": str(run["base_market_trade_date"]),
            },
            "quant_version": run["factor_version"],
            "config_hash": official_report.get("config_hash_before"),
            "input_hash": run["request_hash"],
            "pipeline_request_hash": _pipeline_request_hash(connection, run),
            "data_snapshot_id": run["data_manifest_id"],
            "universe_snapshot_id": f"AUDIT_DERIVED_SHA256:{universe_hash}",
            "universe_snapshot_persisted": False,
            "created_at": run["created_at"],
            "source_database": str(database_path),
            "source_table": "quant_rank_result",
            "effective_universe_size": len(db_rows),
            "quant_report": str(report_path),
            "config": config,
            "formula_sources": {
                "technical": "quant/factors.py::TechnicalFactorCalculator",
                "capital": "quant/factors.py::CapitalFactorCalculator",
                "emotion": "quant/factors.py::EmotionFactorCalculator + datasource/tushare_provider.py::get_market_emotion",
                "momentum": "quant/factors.py::MomentumFactorCalculator",
                "risk": "quant/factors.py::RiskFactorCalculator",
                "total": "quant/ranking.py::QuantRankingEngine.calculate_stock_score",
                "normalization": "quant/normalizer.py::min_max_normalize/clamp_score",
            },
            "tushare_datasets": sorted(manifest_map),
            "raw_coverage": raw_coverage,
            "unit_validation": {
                "status": "FAIL" if unit_failures else "PASS",
                "failures": unit_failures,
            },
            "point_in_time_validation": {
                "status": "PASS"
                if all(row["future_data_check"] == "PASS" for row in raw_coverage)
                else "FAIL",
                "manifest_temporal_status": run["temporal_status"],
            },
            "emotion": {
                **distributions["emotion_score"],
                "source_type": "GLOBAL_MARKET_VALUE_FROM_HARD_CODED_DEFAULT",
                "fallback_reason": "Tushare provider hard-codes limit_up=0, limit_down=0; converter derives heat=50",
                "effective_rank_contribution": emotion_loo[
                    "factor_effective_rank_share"
                ],
                "variance_contribution": 0,
                "finding": "EMOTION_PIPELINE_NOT_DIFFERENTIATED; global factor is valid in scope, but its 0/0/50 source is a placeholder implementation bug",
            },
            "capital": {
                **capital_summary,
                "rank_denominator": "N/A; no percentile rank used",
                "normalization_scope": "PER_STOCK_FIXED_RANGE",
                "fallback_count": sum(
                    row["factor"] == "capital" and row["fallback_used"]
                    for row in missing_rows
                ),
                "collision_source": "(0 + 11.7647 + 100 + 100)/4=52.941175→52.9412",
                "finding": "CAPITAL_SCORE_COLLISION_CONFIRMED; driven by daily.amount unit error, volume-ratio hard-code, and clip saturation",
            },
            "missing_data": {
                "neutral_fallback_count": sum(
                    len(str(row["fallback_subfactors"]).split(","))
                    for row in missing_rows
                    if row["fallback_subfactors"]
                    and row["factor"] != "emotion"
                ),
                "reweighted_rows": 0,
                "non_random_missing_rows": sum(
                    bool(row["missing_reason"]) for row in missing_rows
                ),
                "factor_detail_not_persisted_rows": sum(
                    row["missing_reason"]
                    == "RAW_FACTOR_DETAIL_NOT_PERSISTED_FOR_BOTTOM_310"
                    for row in missing_rows
                ),
            },
            "distributions": distributions,
            "top100_score_range": {
                "min": min(float(row["total_score"]) for row in db_rows[:100]),
                "max": max(float(row["total_score"]) for row in db_rows[:100]),
            },
            "score_compression_diagnosis": (
                "ABNORMAL_IMPLEMENTATION_COMPRESSION: emotion is constant; "
                "daily.amount unit mismatch zeros liquidity scores; volume ratio is hard-coded; "
                "fixed clip ranges and missing fallbacks compress the upper tail"
            ),
            "variance_contributions": {
                "full_a": full_contributions,
                "top100": top100_contributions,
            },
            "leave_one_factor_out": loo,
            "top100_dominant_factors": [
                row["factor"]
                for row in sorted(
                    top100_contributions,
                    key=lambda item: item["variance_share"],
                    reverse=True,
                )
            ],
            "sample_lineage_count": len(lineage),
            "reconciliation_tolerance": {
                "database": 1e-8,
                "excel_display": 0.005,
            },
            "database_reconciliation": {
                "status": "PASS" if exact_reconciliation_max <= 1e-8 else "FAIL",
                "max_difference": exact_reconciliation_max,
            },
            "excel_reconciliation": {
                "status": "PASS"
                if display_reconciliation_max <= 0.005
                else "FAIL",
                "max_display_difference": display_reconciliation_max,
            },
            "findings": findings,
            "implementation_issue": True,
            "data_coverage_issue": bool(data_coverage_failures),
            "runtime_audit": {
                "external_api_calls": 0,
                "llm_calls": 0,
                "orders": 0,
                "scheduler": "OFF",
                "real_orders": 0,
                "virtual_orders": 0,
                "database_mode": "READ_ONLY",
                "production_tables_written": 0,
                "historical_results_overwritten": 0,
                "git_commit": "NOT_EXECUTED",
                "quant_hash_before": quant_hash_before,
                "flash_hash": official_report.get("flash", {}).get("hash"),
                "pro_hash": official_report.get("pro", {}).get("hash"),
            },
            "problems": [
                "Emotion production input is a hard-coded placeholder.",
                "daily.amount has a confirmed x1000 unit mismatch.",
                "daily_basic.volume_ratio is ignored and replaced with 1.0.",
                "YAML percentile declaration does not match production fixed-range code.",
                "Raw factor_detail is present only for report top5000; bottom310 has group scores only.",
                "Historical quant_run does not persist an exact universe_snapshot_id.",
            ],
            "known_limitations": [
                "The audit does not invent missing bottom310 subfactor details.",
                "The current official workbook bytes differ from the hash recorded at initial official export, although score cells reconcile to DB.",
                "No external source was queried to reconstruct unavailable historical metadata.",
            ],
            "recommended_remediation": [
                "Create a new shadow factor_version; do not mutate or overwrite this run.",
                "Fix daily.amount adapter multiplier to x1000 before absolute-liquidity scoring.",
                "Pass daily_basic.volume_ratio through CapitalFlowSnapshot instead of hard-coded 1.0.",
                "Bind Quant to the same-date read-only market_emotion_snapshot and label the factor GLOBAL_MARKET.",
                "Persist per-run effective config, exact universe snapshot hash, score origin and fallback flags.",
                "Run forward shadow evaluation before production promotion.",
            ],
            "final_status": "MULTIPLE_QUANT_ISSUES_FOUND",
            "suggested_commit": "audit(quant): trace tushare inputs and diagnose factor score compression",
        }

        output_dir = (
            Path(args.output_dir).resolve()
            if args.output_dir
            else (DEFAULT_OUTPUT_ROOT / run["run_id"]).resolve()
        )
        output_dir.mkdir(parents=True, exist_ok=True)
        json_path = output_dir / "quant_factor_audit.json"
        md_path = output_dir / "quant_factor_audit.md"
        lineage_path = output_dir / "factor_lineage.csv"
        missing_path = output_dir / "missing_data_audit.csv"
        collision_path = output_dir / "score_collision_audit.csv"
        normalization_path = output_dir / "normalization_universe_audit.csv"
        reconciliation_path = output_dir / "database_excel_reconciliation.csv"
        excel_path = output_dir / "quant_factor_audit.xlsx"

        _write_csv(lineage_path, lineage)
        _write_csv(missing_path, missing_rows)
        _write_csv(collision_path, collision_rows)
        _write_csv(normalization_path, normalization_rows)
        _write_csv(reconciliation_path, reconciliation)
        json_path.write_text(
            json.dumps(audit, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
        md_path.write_text(_markdown(audit), encoding="utf-8")
        workbook_payload = _workbook_payload(
            audit,
            distributions,
            raw_coverage,
            collision_rows,
            missing_rows,
            normalization_rows,
            contribution_rows,
            loo,
            lineage,
            reconciliation,
        )
        excel_validation = _build_excel(workbook_payload, excel_path)
        audit["reports"] = [
            str(json_path),
            str(md_path),
            str(lineage_path),
            str(missing_path),
            str(collision_path),
            str(normalization_path),
            str(reconciliation_path),
        ]
        audit["excel"] = {
            "path": str(excel_path),
            "sha256": file_sha256(excel_path),
            "validation": excel_validation,
        }
        audit["runtime_audit"]["quant_hash_after"] = canonical_hash(
            {
                str(path.relative_to(ROOT)): file_sha256(path)
                for path in (
                    ROOT / "quant" / "config.py",
                    ROOT / "quant" / "factors.py",
                    ROOT / "quant" / "normalizer.py",
                    ROOT / "quant" / "ranking.py",
                    ROOT / "config" / "quant_factor.yaml",
                )
            }
        )
        audit["runtime_audit"]["quant_hash_unchanged"] = (
            audit["runtime_audit"]["quant_hash_before"]
            == audit["runtime_audit"]["quant_hash_after"]
        )
        json_path.write_text(
            json.dumps(audit, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
        md_path.write_text(_markdown(audit), encoding="utf-8")
    finally:
        connection.close()
    database_after = {
        "size": database_path.stat().st_size,
        "mtime_ns": database_path.stat().st_mtime_ns,
    }
    if database_before != database_after:
        raise RuntimeError("HISTORICAL_DATABASE_CHANGED_DURING_READ_ONLY_AUDIT")
    return audit


def _pipeline_request_hash(
    connection: sqlite3.Connection, run: dict[str, Any]
) -> str | None:
    for row in _rows(
        connection,
        "SELECT run_ids,checkpoint FROM pipeline_job WHERE job_type='QUANT' AND trade_date=?",
        (str(run["base_market_trade_date"]),),
    ):
        ids = _json(row.get("run_ids")) or {}
        if ids.get("quant_run_id") == run["run_id"]:
            return (_json(row.get("checkpoint")) or {}).get("request_hash")
    return None


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Read-only TUSHARE_BASELINE_V1 Quant factor audit"
    )
    parser.add_argument("--run-id", default="")
    parser.add_argument("--trade-date", default="")
    parser.add_argument("--workbook", default="")
    parser.add_argument("--database", default=str(DEFAULT_DB))
    parser.add_argument("--no-external-api", action="store_true", default=True)
    parser.add_argument("--no-llm", action="store_true", default=True)
    parser.add_argument("--read-only", action="store_true", default=True)
    parser.add_argument("--output-dir", default="")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        audit = run_audit(args)
    except RuntimeError as exc:
        print(json.dumps({"final_status": "AUDIT_BLOCKED", "error": str(exc)}, ensure_ascii=False))
        return 2
    print(
        json.dumps(
            {
                "phase": audit["phase"],
                "audit_run": audit["audit_run"],
                "trade_date": audit["trade_date"],
                "final_status": audit["final_status"],
                "reports": audit["reports"],
                "excel": audit["excel"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
