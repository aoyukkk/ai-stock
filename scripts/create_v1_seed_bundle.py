from __future__ import annotations

import argparse
import json
import os
import shutil
import sqlite3
import tempfile
import uuid
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT_BOOTSTRAP = Path(__file__).resolve().parents[1]
if str(ROOT_BOOTSTRAP) not in sys.path:
    sys.path.insert(0, str(ROOT_BOOTSTRAP))

from dotenv import dotenv_values
from openpyxl import load_workbook

from backend.version import APP_VERSION, SCHEMA_VERSION
from scripts.v1_release_common import json_dump, sanitize_value, scan_text, sha256_file


ROOT = ROOT_BOOTSTRAP
DEFAULT_SOURCE = ROOT / "data" / "ai_trader_dev.db"
DEFAULT_DESTINATION = ROOT / "build" / "seed"
ACCEPTED_RUN_STATUSES = {"COMPLETED", "SUCCESS", "PARTIAL_SUCCESS"}
PRESERVED_TABLES = {
    "desktop_schema_version", "stock_master", "run_data_manifest", "quant_run", "quant_rank_result",
    "model_validation_run", "model_validation_sample", "model_validation_llm_audit",
    "model_validation_order_plan", "model_validation_allocation", "validation_account_snapshot",
    "pro_resume_run", "pro_candidate_review", "llm_usage", "selection_cohort",
    "selection_cohort_member", "selection_performance_run", "selection_performance_daily",
    "selection_portfolio_daily", "workbench_run_registry", "pipeline_job", "system_config",
}


def create_seed_bundle(source: Path, destination: Path) -> dict[str, Any]:
    source = source.resolve()
    destination = destination.resolve()
    if not source.is_file():
        raise FileNotFoundError(f"SOURCE_DATABASE_NOT_FOUND:{source}")
    secrets = _known_secrets()
    temporary_parent = destination.parent
    temporary_parent.mkdir(parents=True, exist_ok=True)
    temp_root = Path(tempfile.mkdtemp(prefix=f".{destination.name}-", dir=temporary_parent))
    stats = {
        "removed_secrets": 0,
        "rewritten_paths": 0,
        "excluded_prompts_responses": 0,
        "excluded_logs": sum(1 for path in (ROOT / "logs").rglob("*") if path.is_file()) if (ROOT / "logs").is_dir() else 0,
        "excluded_mock_rows": 0,
        "excluded_non_seed_rows": 0,
    }
    try:
        seed_database = temp_root / "data" / "ai_trader_seed.db"
        seed_database.parent.mkdir(parents=True, exist_ok=True)
        _sqlite_backup(source, seed_database)
        pipeline_runs, table_counts = _sanitize_database(seed_database, secrets, stats)
        included_files = _copy_allowlisted_files(temp_root, pipeline_runs, secrets, stats)
        _copy_allowlisted_cache(temp_root, pipeline_runs, secrets, included_files)

        database_hash = sha256_file(seed_database)
        manifest = {
            "seed_version": APP_VERSION,
            "source_database_sha256": sha256_file(source),
            "sanitized_database_sha256": database_hash,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "schema_version": SCHEMA_VERSION,
            "latest_trade_date": max((item["trade_date"] for item in pipeline_runs), default=None),
            "available_trade_dates": sorted({item["trade_date"] for item in pipeline_runs}),
            "pipeline_runs": pipeline_runs,
            "table_row_counts": table_counts,
            "included_files": included_files,
            "excluded_file_counts": {"logs": stats["excluded_logs"], "temporary_or_unregistered": 0},
            "ambiguous_records": [],
            "sanitization": stats,
            "secret_scan_status": "PENDING",
            "absolute_path_scan_status": "PENDING",
            "checksum_status": "PENDING",
            "total_size_bytes": 0,
        }
        manifest_path = temp_root / "SEED_DATA_MANIFEST_1.0.0.json"
        json_dump(manifest_path, manifest)

        from scripts.verify_v1_seed_bundle import verify_seed_bundle

        verification = verify_seed_bundle(temp_root, known_secrets=secrets, update_manifest=False)
        if not verification["passed"]:
            raise RuntimeError(f"SEED_VERIFICATION_FAILED:{','.join(verification['problems'])}")
        manifest.update({
            "secret_scan_status": "PASS",
            "absolute_path_scan_status": "PASS",
            "checksum_status": "PASS",
            "total_size_bytes": sum(path.stat().st_size for path in temp_root.rglob("*") if path.is_file()),
        })
        json_dump(manifest_path, manifest)
        staging = destination.with_name(f".{destination.name}-ready-{uuid.uuid4().hex[:8]}")
        shutil.copytree(temp_root, staging)
        if destination.exists():
            shutil.rmtree(destination)
        os.replace(staging, destination)
        shutil.rmtree(temp_root, ignore_errors=True)
        return manifest
    except Exception:
        shutil.rmtree(temp_root, ignore_errors=True)
        raise


def _sanitize_database(database: Path, secrets: list[str], stats: dict[str, int]) -> tuple[list[dict[str, Any]], dict[str, int]]:
    with sqlite3.connect(database) as connection:
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys=OFF")
        table_names = [row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
        initial_rows = sum(int(connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]) for table in table_names)
        stats["excluded_mock_rows"] = _explicit_mock_row_count(connection)
        registry = [dict(row) for row in connection.execute(
            "SELECT * FROM workbench_run_registry ORDER BY trade_date, reconciled_at"
        ) if str(row["final_status"]).upper() in ACCEPTED_RUN_STATUSES and str(row["source"]).upper() != "MOCK"]
        if not registry:
            raise RuntimeError("NO_AUDITABLE_PIPELINE_RUNS")

        quant_ids = [row["quant_run_id"] for row in registry]
        manifest_ids = [row["manifest_id"] for row in registry]
        flash_ids = [row["flash_run_id"] for row in registry]
        pro_ids = [row["pro_run_id"] for row in registry]
        pipeline_ids = [row["pipeline_run_id"] for row in registry]
        _delete_not_in(connection, "workbench_run_registry", "pipeline_run_id", pipeline_ids)
        _delete_not_in(connection, "quant_run", "run_id", quant_ids)
        _delete_not_in(connection, "quant_rank_result", "quant_run_id", quant_ids)
        _delete_not_in(connection, "run_data_manifest", "manifest_id", manifest_ids)
        for table in ("model_validation_run", "model_validation_sample", "model_validation_llm_audit", "model_validation_order_plan", "model_validation_allocation", "validation_account_snapshot"):
            _delete_not_in(connection, table, "run_id" if table == "model_validation_run" else "validation_run_id", flash_ids)
        _delete_not_in(connection, "pro_resume_run", "run_id", pro_ids)
        _delete_not_in(connection, "pro_candidate_review", "pro_resume_run_id", pro_ids)
        placeholders = ",".join("?" for _ in pipeline_ids + flash_ids + pro_ids)
        connection.execute(
            f"DELETE FROM llm_usage WHERE COALESCE(pipeline_run_id,'') NOT IN ({','.join('?' for _ in pipeline_ids)}) "
            f"AND COALESCE(validation_run_id,'') NOT IN ({','.join('?' for _ in flash_ids)}) "
            f"AND COALESCE(pro_resume_run_id,'') NOT IN ({','.join('?' for _ in pro_ids)})",
            [*pipeline_ids, *flash_ids, *pro_ids],
        )
        _delete_not_in(connection, "selection_cohort", "pipeline_run_id", pipeline_ids)
        cohort_ids = [row[0] for row in connection.execute("SELECT id FROM selection_cohort")]
        _delete_not_in(connection, "selection_cohort_member", "cohort_id", cohort_ids)

        for table in table_names:
            if table not in PRESERVED_TABLES:
                connection.execute(f'DELETE FROM "{table}"')

        connection.execute("DELETE FROM pipeline_job")
        for row in registry:
            connection.execute(
                "INSERT INTO pipeline_job(job_id,job_type,trade_date,status,stage,progress_current,progress_total,success_count,failure_count,token_usage,cost_usd,started_at,finished_at,run_ids,checkpoint,created_at,updated_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    f"seed-{uuid.uuid4().hex[:20]}", "HISTORICAL", row["trade_date"], "SUCCESS", "SEEDED_HISTORICAL",
                    1, 1, 1, 0, 0, 0.0, row["created_at"], row["updated_at"],
                    json.dumps({"pipeline_run_id": row["pipeline_run_id"], "quant_run_id": row["quant_run_id"], "flash_run_id": row["flash_run_id"], "pro_run_id": row["pro_run_id"]}),
                    json.dumps({"mode": "SEED_HISTORICAL", "source": row["source"]}), row["created_at"], row["updated_at"],
                ),
            )

        connection.execute(
            "DELETE FROM system_config WHERE is_sensitive = 1 OR lower(config_key) LIKE '%api_key%' "
            "OR lower(config_key) LIKE '%password%' OR lower(config_key) LIKE '%credential%'"
        )
        if _table_has_column(connection, "llm_usage", "response_metadata"):
            stats["excluded_prompts_responses"] += int(connection.execute(
                "SELECT COUNT(*) FROM llm_usage WHERE response_metadata IS NOT NULL AND trim(response_metadata) NOT IN ('', '{}', 'null')"
            ).fetchone()[0])
            connection.execute("UPDATE llm_usage SET response_metadata='{}', provider_request_id=NULL, error_message=NULL")
        if _table_has_column(connection, "model_validation_llm_audit", "diagnostics"):
            stats["excluded_prompts_responses"] += int(connection.execute(
                "SELECT COUNT(*) FROM model_validation_llm_audit WHERE diagnostics IS NOT NULL AND trim(diagnostics) NOT IN ('', '{}', 'null')"
            ).fetchone()[0])
            connection.execute("UPDATE model_validation_llm_audit SET diagnostics='{}', error_message=NULL")
        if _table_has_column(connection, "workbench_run_registry", "export_path"):
            for row in registry:
                relative = _logical_output_path(row["export_path"], row["trade_date"])
                connection.execute("UPDATE workbench_run_registry SET export_path=? WHERE id=?", (relative, row["id"]))

        secret_hits, path_hits = _sanitize_text_columns(connection, ROOT, secrets)
        stats["removed_secrets"] += secret_hits
        stats["rewritten_paths"] += path_hits
        connection.execute("PRAGMA foreign_keys=ON")
        violations = list(connection.execute("PRAGMA foreign_key_check"))
        if violations:
            raise RuntimeError(f"SEED_FOREIGN_KEY_CHECK_FAILED:{len(violations)}")
        if connection.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("SEED_SQLITE_INTEGRITY_FAILED")
        connection.commit()
        connection.execute("VACUUM")
        counts = {table: int(connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]) for table in table_names}
        stats["excluded_non_seed_rows"] = max(0, initial_rows - sum(counts.values()) + len(registry))
        pipelines = [_pipeline_manifest(connection, row) for row in registry]
        return pipelines, counts


def _explicit_mock_row_count(connection: sqlite3.Connection) -> int:
    count = int(connection.execute(
        "SELECT COUNT(*) FROM pipeline_job WHERE upper(COALESCE(stage,'')) LIKE '%MOCK%' OR upper(COALESCE(checkpoint,'')) LIKE '%MOCK%'"
    ).fetchone()[0])
    count += int(connection.execute(
        "SELECT COUNT(*) FROM workbench_run_registry WHERE upper(COALESCE(source,''))='MOCK'"
    ).fetchone()[0])
    return count


def _pipeline_manifest(connection: sqlite3.Connection, row: dict[str, Any]) -> dict[str, Any]:
    counts = _json_value(row.get("counts"))
    return {
        "trade_date": str(row["trade_date"]), "pipeline_run_id": row["pipeline_run_id"], "status": row["final_status"],
        "quant_run_id": row["quant_run_id"], "flash_run_id": row["flash_run_id"], "pro_run_id": row["pro_run_id"],
        "candidate_set_hash": row["candidate_set_hash"],
        "quant_count": int(connection.execute("SELECT COUNT(*) FROM quant_rank_result WHERE quant_run_id=?", (row["quant_run_id"],)).fetchone()[0]),
        "flash_count": int(connection.execute("SELECT COUNT(*) FROM model_validation_sample WHERE validation_run_id=?", (row["flash_run_id"],)).fetchone()[0]),
        "manual_count": int(counts.get("manual", 0)), "candidate_count": int(counts.get("candidate", 0)),
        "order_count": int(connection.execute("SELECT COUNT(*) FROM model_validation_order_plan WHERE validation_run_id=?", (row["flash_run_id"],)).fetchone()[0]),
        "position_count": int(connection.execute("SELECT COUNT(*) FROM model_validation_allocation WHERE validation_run_id=?", (row["flash_run_id"],)).fetchone()[0]),
        "fundamental_count": int(counts.get("fundamental", counts.get("candidate", 0))),
        "export_files": [_logical_output_path(row.get("export_path"), str(row["trade_date"]))],
    }


def _copy_allowlisted_files(temp_root: Path, runs: list[dict[str, Any]], secrets: list[str], stats: dict[str, int]) -> list[dict[str, Any]]:
    included: list[dict[str, Any]] = []
    for run in runs:
        date_dir = ROOT / "outputs" / run["trade_date"]
        candidates: list[Path] = []
        for logical in run["export_files"]:
            source = ROOT / logical
            if source.is_file(): candidates.append(source)
        if date_dir.is_dir():
            candidates.extend(path for path in date_dir.glob("*.xlsx") if path.is_file())
            audit_dir = date_dir / "审计"
            if audit_dir.is_dir():
                candidates.extend(path for path in audit_dir.glob("*.json") if path.is_file() and "机器版" in path.name)
        for source in dict.fromkeys(candidates):
            logical = Path("outputs") / run["trade_date"] / source.name
            destination = temp_root / logical
            destination.parent.mkdir(parents=True, exist_ok=True)
            if source.suffix.lower() == ".xlsx":
                secret_hits, path_hits = _sanitize_workbook(source, destination, secrets)
            else:
                raw = json.loads(source.read_text(encoding="utf-8"))
                sanitized, secret_hits, path_hits = sanitize_value(raw, source_root=ROOT, secrets=secrets)
                json_dump(destination, sanitized)
            stats["removed_secrets"] += secret_hits
            stats["rewritten_paths"] += path_hits
            included.append(_file_manifest(source, destination, logical, run))
    return included


def _copy_allowlisted_cache(temp_root: Path, runs: list[dict[str, Any]], secrets: list[str], included: list[dict[str, Any]]) -> None:
    trade_dates = {run["trade_date"].replace("-", "") for run in runs}
    cache_root = ROOT / "data" / "cache" / "tushare"
    candidates: list[Path] = []
    for dataset in ("daily", "adj_factor"):
        files = sorted((cache_root / "trade_date" / dataset).glob("*.json"))
        eligible = [path for path in files if any(path.stem <= trade_date for trade_date in trade_dates)]
        candidates.extend(eligible[-30:])
    for dataset in ("daily_basic", "moneyflow", "stk_limit"):
        for trade_date in trade_dates:
            path = cache_root / "trade_date" / dataset / f"{trade_date}.json"
            if path.is_file(): candidates.append(path)
    fundamental = cache_root / "fundamental"
    if fundamental.is_dir():
        candidates.extend(path for path in fundamental.rglob("*.json") if path.is_file())

    for source in dict.fromkeys(candidates):
        relative = source.relative_to(cache_root)
        logical = Path("cache") / "tushare" / relative
        destination = temp_root / logical
        destination.parent.mkdir(parents=True, exist_ok=True)
        raw = json.loads(source.read_text(encoding="utf-8"))
        sanitized, secret_hits, path_hits = sanitize_value(raw, source_root=ROOT, secrets=secrets)
        if secret_hits or path_hits:
            json_dump(destination, sanitized)
        else:
            shutil.copy2(source, destination)
        included.append({
            "logical_path": logical.as_posix(), "source_path": source.name, "destination_path": logical.as_posix(),
            "file_size": destination.stat().st_size, "sha256": sha256_file(destination), "related_trade_date": None,
            "related_pipeline_run_id": None, "file_type": "TUSHARE_BATCH_CACHE", "sanitization_status": "PASS",
        })


def _sanitize_workbook(source: Path, destination: Path, secrets: list[str]) -> tuple[int, int]:
    workbook = load_workbook(source)
    secret_hits = path_hits = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    value, removed, rewritten = sanitize_value(cell.value, source_root=ROOT, secrets=secrets)
                    cell.value = value
                    secret_hits += removed
                    path_hits += rewritten
    workbook.save(destination)
    workbook.close()
    return secret_hits, path_hits


def _sanitize_text_columns(connection: sqlite3.Connection, source_root: Path, secrets: list[str]) -> tuple[int, int]:
    secret_hits = path_hits = 0
    tables = [row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
    for table in tables:
        columns = [row[1] for row in connection.execute(f'PRAGMA table_info("{table}")') if str(row[2]).upper() in {"TEXT", "JSON"}]
        if not columns or not _table_has_column(connection, table, "id"): continue
        for row in connection.execute(f'SELECT id,{",".join(chr(34)+column+chr(34) for column in columns)} FROM "{table}"').fetchall():
            updates: dict[str, Any] = {}
            for index, column in enumerate(columns, start=1):
                original = row[index]
                if original is None: continue
                parsed = _json_value(original)
                value, removed, rewritten = sanitize_value(parsed, source_root=source_root, secrets=secrets)
                encoded = json.dumps(value, ensure_ascii=False, separators=(",", ":")) if not isinstance(parsed, str) else value
                if encoded != original: updates[column] = encoded
                secret_hits += removed
                path_hits += rewritten
            if updates:
                assignment = ",".join(f'"{key}"=?' for key in updates)
                connection.execute(f'UPDATE "{table}" SET {assignment} WHERE id=?', [*updates.values(), row[0]])
    return secret_hits, path_hits


def _known_secrets() -> list[str]:
    values = dotenv_values(ROOT / ".env") if (ROOT / ".env").exists() else {}
    names = ("TUSHARE_TOKEN", "DEEPSEEK_API_KEY", "OPENAI_API_KEY", "GITHUB_TOKEN")
    return list(dict.fromkeys(str(os.getenv(name) or values.get(name) or "").strip() for name in names if os.getenv(name) or values.get(name)))


def _delete_not_in(connection: sqlite3.Connection, table: str, column: str, values: list[Any]) -> None:
    if not values:
        connection.execute(f'DELETE FROM "{table}"')
        return
    connection.execute(f'DELETE FROM "{table}" WHERE "{column}" NOT IN ({",".join("?" for _ in values)})', values)


def _sqlite_backup(source: Path, destination: Path) -> None:
    with sqlite3.connect(source) as source_connection, sqlite3.connect(destination) as destination_connection:
        source_connection.backup(destination_connection)


def _table_has_column(connection: sqlite3.Connection, table: str, column: str) -> bool:
    return column in {row[1] for row in connection.execute(f'PRAGMA table_info("{table}")')}


def _json_value(value: Any) -> Any:
    if not isinstance(value, str): return value or {}
    try: return json.loads(value)
    except (json.JSONDecodeError, TypeError): return value


def _logical_output_path(value: Any, trade_date: str) -> str:
    if value:
        path = Path(str(value))
        if path.is_absolute():
            try: return path.resolve().relative_to(ROOT).as_posix()
            except ValueError: pass
        if str(value).replace("\\", "/").startswith("outputs/"): return str(value).replace("\\", "/")
    return f"outputs/{trade_date}/"


def _file_manifest(source: Path, destination: Path, logical: Path, run: dict[str, Any]) -> dict[str, Any]:
    return {
        "logical_path": logical.as_posix(), "source_path": source.name, "destination_path": logical.as_posix(),
        "file_size": destination.stat().st_size, "sha256": sha256_file(destination),
        "related_trade_date": run["trade_date"], "related_pipeline_run_id": run["pipeline_run_id"],
        "file_type": source.suffix.lstrip(".").upper(), "sanitization_status": "PASS",
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_DESTINATION)
    args = parser.parse_args()
    manifest = create_seed_bundle(args.source, args.output)
    print(json.dumps({"status": "PASS", "output": str(args.output), "pipeline_runs": len(manifest["pipeline_runs"]), "total_size_bytes": manifest["total_size_bytes"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
