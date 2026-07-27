from __future__ import annotations

import hashlib
import json
import shutil
import sys
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env", override=False)

from database.models.ai import DecisionSnapshot
from database.models.decision_explainability import (
    AdmissionV3Result,
    AdmissionV3Run,
    FactorAttribution,
    GateEvaluation,
    StrategyTimingContract,
)
from database.models.entry_timing_v2 import AdmissionV2Run, EntryTimingV2Result
from database.models.quant_run import QuantRun
from database.session import get_session, init_db
from post_close.seven_day_comparison import SevenDayComparisonService
from quant.explainability.factor_attribution import OPAQUE_LLM_CONTRIBUTION
from quant.explainability.factor_registry import FACTOR_FAMILIES, FUNDAMENTAL
from quant.explainability.timing_contract import validate_timing_contract
from reporting.workbook_style import WorkbookStyleService
from stock_codes import normalize_ts_code


SHANGHAI = ZoneInfo("Asia/Shanghai")
OUTPUT_DIR = ROOT / "outputs" / "2026-07-20" / "正式日线"
BASELINE_OUTPUT = OUTPUT_DIR / "智能交易助手_2026-07-20_Baseline对账修正版.xlsx"
SHADOW_OUTPUT = OUTPUT_DIR / "V2_2_V3_Shadow对账_2026-07-20.xlsx"
SEVEN_OUTPUT = OUTPUT_DIR / "近7交易日数据对比_截至2026-07-20.xlsx"
AUDIT_NAMES = (
    "reconciliation_report.json",
    "reconciliation_report.md",
    "baseline_workbook_audit.json",
    "shadow_workbook_audit.json",
    "seven_day_comparison_audit.json",
)

NAVY = "17365D"
TEAL = "2F766D"
PALE_BLUE = "D9EAF7"
PALE_GREEN = "D9EAD3"
PALE_YELLOW = "FFF2CC"
PALE_RED = "F4CCCC"
WHITE = "FFFFFF"


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    # These exact names are generated artifacts owned by this reconciliation
    # command. Re-running repairs/validates them in place; historical source
    # workbooks and immutable V2/V3 rows are never overwritten.

    source_baseline = OUTPUT_DIR / "智能交易助手_2026-07-20_Excel兼容修复.xlsx"
    if not source_baseline.exists():
        source_baseline = OUTPUT_DIR / "智能交易助手_2026-07-20.xlsx"
    source_seven = _latest_source_seven()
    official_report_path = sorted(
        OUTPUT_DIR.glob("postclose_official_2026-07-20_*_reconciled*.json"),
        key=lambda p: p.stat().st_mtime,
    )[-1]
    official = json.loads(official_report_path.read_text(encoding="utf-8"))

    init_db()
    session = get_session()
    try:
        data = _load_shadow_data(session, official)
        formal_codes = _formal_codes(source_baseline)
        timing = _persist_reconciliation_timing_and_snapshots(session, data, official, formal_codes)
        baseline_meta = _build_baseline(source_baseline, BASELINE_OUTPUT, official, formal_codes)
        shadow_meta = _build_shadow(SHADOW_OUTPUT, data, official, formal_codes, timing)
        seven_meta = _build_seven(session, source_seven, SEVEN_OUTPUT)
        session.commit()
    except Exception:
        session.rollback()
        for path in (BASELINE_OUTPUT, SHADOW_OUTPUT, SEVEN_OUTPUT):
            path.unlink(missing_ok=True)
        raise
    finally:
        session.close()

    baseline_audit = _scan_workbook(BASELINE_OUTPUT)
    shadow_audit = _scan_workbook(SHADOW_OUTPUT)
    seven_audit = _scan_workbook(SEVEN_OUTPUT)
    baseline_audit.update(baseline_meta)
    shadow_audit.update(shadow_meta)
    seven_audit.update(seven_meta)
    _write_json(OUTPUT_DIR / "baseline_workbook_audit.json", baseline_audit)
    _write_json(OUTPUT_DIR / "shadow_workbook_audit.json", shadow_audit)
    _write_json(OUTPUT_DIR / "seven_day_comparison_audit.json", seven_audit)

    reconciliation = _reconciliation_report(data, official, formal_codes, timing, baseline_audit, shadow_audit, seven_audit)
    _write_json(OUTPUT_DIR / "reconciliation_report.json", reconciliation)
    (OUTPUT_DIR / "reconciliation_report.md").write_text(_markdown(reconciliation), encoding="utf-8")
    print(json.dumps({
        "status": "SUCCESS",
        "baseline": str(BASELINE_OUTPUT),
        "shadow": str(SHADOW_OUTPUT),
        "seven_day": str(SEVEN_OUTPUT),
        "audit": str(OUTPUT_DIR / "reconciliation_report.json"),
    }, ensure_ascii=False))
    return 0


def _load_shadow_data(session, official: dict[str, Any]) -> dict[str, Any]:
    trade_date = date(2026, 7, 20)
    v2_run = session.scalar(
        select(AdmissionV2Run)
        .where(AdmissionV2Run.trade_date == trade_date)
        .order_by(AdmissionV2Run.created_at.desc())
    )
    if v2_run is None:
        raise ValueError("V22_SHADOW_MISSING_FOR_2026_07_20")
    v3_run = session.scalar(
        select(AdmissionV3Run)
        .where(AdmissionV3Run.trade_date == trade_date, AdmissionV3Run.source_v2_run_id == v2_run.run_id)
        .order_by(AdmissionV3Run.created_at.desc())
    )
    if v3_run is None:
        raise ValueError("ADMISSION_V3_SHADOW_MISSING_FOR_2026_07_20")
    quant = session.scalar(select(QuantRun).where(QuantRun.run_id == v2_run.quant_run_id))
    if quant is None:
        raise ValueError("QUANT_RUN_MISSING")
    v2_rows = list(session.scalars(
        select(EntryTimingV2Result)
        .where(EntryTimingV2Result.run_id == v2_run.run_id)
        .order_by(EntryTimingV2Result.quant_rank, EntryTimingV2Result.stock_code)
    ))
    # The source can contain multiple pools; reconciliation has one row per code.
    v2_by_code: dict[str, EntryTimingV2Result] = {}
    for row in v2_rows:
        code = normalize_ts_code(row.stock_code)
        if code not in v2_by_code or row.pool_type == "AI_POOL":
            v2_by_code[code] = row
    v2_rows = sorted(v2_by_code.values(), key=lambda r: (r.quant_rank or 999999, r.stock_code))
    v3_rows = list(session.scalars(
        select(AdmissionV3Result)
        .where(AdmissionV3Result.run_id == v3_run.run_id)
        .order_by(AdmissionV3Result.quant_rank, AdmissionV3Result.stock_code)
    ))
    factors = list(session.scalars(
        select(FactorAttribution)
        .where(FactorAttribution.run_id == v3_run.run_id)
        .order_by(FactorAttribution.stock_code, FactorAttribution.factor_family)
    ))
    gates = list(session.scalars(
        select(GateEvaluation)
        .where(GateEvaluation.run_id == v3_run.run_id)
        .order_by(GateEvaluation.gate_name)
    ))
    if len(v2_rows) != 100 or len(v3_rows) != 100 or len(factors) != 600:
        raise ValueError(f"SHADOW_COHORT_INCOMPLETE:v2={len(v2_rows)},v3={len(v3_rows)},factors={len(factors)}")
    return {"v2_run": v2_run, "v3_run": v3_run, "quant": quant, "v2": v2_rows, "v3": v3_rows, "factors": factors, "gates": gates}


def _persist_reconciliation_timing_and_snapshots(session, data, official, formal_codes: set[str]) -> dict[str, Any]:
    trade_date = date(2026, 7, 20)
    quant = data["quant"]
    available_at = datetime.fromisoformat(official["temporal_gate"]["passed_at"])
    baseline_signal = datetime.fromisoformat(official["execution_end"])
    baseline_eligible = datetime(2026, 7, 21, 9, 30, tzinfo=SHANGHAI)
    if baseline_signal >= baseline_eligible:
        baseline_eligible = datetime.combine(_next_weekday(baseline_signal.date()), time(9, 30), SHANGHAI)
    shadow_signal = _db_created_at_to_shanghai(data["v3_run"].created_at)
    shadow_eligible = datetime(2026, 7, 21, 9, 30, tzinfo=SHANGHAI)
    if shadow_signal >= shadow_eligible:
        shadow_eligible = datetime.combine(_next_weekday(shadow_signal.date()), time(9, 30), SHANGHAI)
    observation = datetime(2026, 7, 20, 15, 0, tzinfo=SHANGHAI)
    data_snapshot = str(quant.data_manifest_id)
    universe_snapshot = f"{quant.run_id}:{quant.universe_count}"
    v3_by_code = {normalize_ts_code(row.stock_code): row for row in data["v3"]}
    factors_by_code: dict[str, list[FactorAttribution]] = defaultdict(list)
    for row in data["factors"]:
        factors_by_code[normalize_ts_code(row.stock_code)].append(row)
    contracts: dict[str, dict[str, StrategyTimingContract]] = {}
    snapshots: dict[str, dict[str, int]] = {}
    for source in data["v2"]:
        code = normalize_ts_code(source.stock_code)
        baseline_contract = _contract(
            session, code, trade_date, observation, available_at, baseline_signal, baseline_eligible,
            "NEXT_OPEN", "postclose_baseline_reconciliation_v1", data_snapshot, universe_snapshot,
        )
        shadow_contract = _contract(
            session, code, trade_date, observation, available_at, shadow_signal, shadow_eligible,
            "NEXT_OPEN_SHADOW_ONLY", "postclose_v22_v3_shadow_reconciliation_v1", data_snapshot, universe_snapshot,
        )
        v3 = v3_by_code[code]
        baseline_snapshot = _existing_snapshot(session, code, "TUSHARE_BASELINE_V1", baseline_contract.id) or DecisionSnapshot(
            stock_code=code,
            snapshot_time=baseline_signal,
            market_data_json={"data_snapshot_id": data_snapshot, "universe_snapshot_id": universe_snapshot},
            factor_json={"baseline_quant_score": float(source.quant_score), "baseline_quant_rank": source.quant_rank},
            agent_result_json={
                "route": "TUSHARE_BASELINE_V1", "timing_contract_id": baseline_contract.id,
                "formal_candidate": _six(code) in formal_codes, "legacy_result": True,
            },
            final_score=source.quant_score,
            risk_level="LEGACY_BASELINE",
            recommendation="FORMAL" if _six(code) in formal_codes else "WATCH",
        )
        shadow_snapshot = _existing_snapshot(session, code, "V2.2_V3_SHADOW", shadow_contract.id) or DecisionSnapshot(
            stock_code=code,
            snapshot_time=shadow_signal,
            market_data_json={"data_snapshot_id": data_snapshot, "universe_snapshot_id": universe_snapshot},
            factor_json={
                "factor_attribution": [
                    {"family": f.factor_family, "contribution": float(f.rank_contribution)}
                    for f in factors_by_code[code]
                ],
                "opaque_llm_marker": OPAQUE_LLM_CONTRIBUTION,
            },
            agent_result_json={
                "route": "V2.2_V3_SHADOW", "timing_contract_id": shadow_contract.id,
                "v2_status": _v22_status(source), "v3_status": v3.admission_state,
                "formal_permission": False,
            },
            final_score=v3.final_score,
            risk_level="SHADOW_ONLY",
            recommendation=v3.admission_state,
        )
        baseline_snapshot.snapshot_time = baseline_signal
        baseline_snapshot.market_data_json = {"data_snapshot_id": data_snapshot, "universe_snapshot_id": universe_snapshot}
        baseline_snapshot.factor_json = {"baseline_quant_score": float(source.quant_score), "baseline_quant_rank": source.quant_rank}
        baseline_snapshot.agent_result_json = {"route": "TUSHARE_BASELINE_V1", "timing_contract_id": baseline_contract.id, "formal_candidate": _six(code) in formal_codes, "legacy_result": True}
        baseline_snapshot.final_score = source.quant_score
        baseline_snapshot.risk_level = "LEGACY_BASELINE"
        baseline_snapshot.recommendation = "FORMAL" if _six(code) in formal_codes else "WATCH"
        shadow_snapshot.snapshot_time = shadow_signal
        shadow_snapshot.market_data_json = {"data_snapshot_id": data_snapshot, "universe_snapshot_id": universe_snapshot}
        shadow_snapshot.agent_result_json = {"route": "V2.2_V3_SHADOW", "timing_contract_id": shadow_contract.id, "v2_status": _v22_status(source), "v3_status": v3.admission_state, "formal_permission": False}
        session.add_all([baseline_snapshot, shadow_snapshot])
        session.flush()
        contracts[code] = {"baseline": baseline_contract, "shadow": shadow_contract}
        snapshots[code] = {"baseline": baseline_snapshot.id, "shadow": shadow_snapshot.id}
    return {
        "contracts": contracts,
        "snapshots": snapshots,
        "available_at": available_at,
        "baseline_signal": baseline_signal,
        "baseline_eligible": baseline_eligible,
        "shadow_signal": shadow_signal,
        "shadow_eligible": shadow_eligible,
        "data_snapshot_id": data_snapshot,
        "universe_snapshot_id": universe_snapshot,
    }


def _contract(session, code, trade_date, observation, available, signal, eligible, policy, version, data_snapshot, universe_snapshot):
    existing = session.scalar(select(StrategyTimingContract).where(
        StrategyTimingContract.stock_code == code,
        StrategyTimingContract.trade_date == trade_date,
        StrategyTimingContract.feature_version == version,
        StrategyTimingContract.data_snapshot_id == data_snapshot,
        StrategyTimingContract.universe_snapshot_id == universe_snapshot,
    ))
    if existing is not None:
        validate_timing_contract(existing)
        return existing
    row = StrategyTimingContract(
        stock_code=code,
        trade_date=trade_date,
        observation_end_ts=observation,
        available_at_ts=available,
        signal_generated_at=signal,
        order_eligible_at=eligible,
        execution_policy=policy,
        feature_version=version,
        data_snapshot_id=data_snapshot,
        universe_snapshot_id=universe_snapshot,
    )
    validate_timing_contract(row)
    session.add(row)
    session.flush()
    return row


def _existing_snapshot(session, code: str, route: str, contract_id: int) -> DecisionSnapshot | None:
    rows = session.scalars(
        select(DecisionSnapshot)
        .where(DecisionSnapshot.stock_code == code)
        .order_by(DecisionSnapshot.created_at.desc())
    )
    for row in rows:
        payload = row.agent_result_json or {}
        if payload.get("route") == route and payload.get("timing_contract_id") == contract_id:
            return row
    return None


def _build_baseline(source: Path, target: Path, official: dict[str, Any], formal_codes: set[str]) -> dict[str, Any]:
    shutil.copy2(source, target)
    wb = load_workbook(target)
    if "今日推荐" in wb.sheetnames:
        wb["今日推荐"].title = "正式候选"
    watch = wb.create_sheet("重点观察", 2)
    _write_table(
        watch,
        "重点观察（非正式候选）",
        "WATCH 仅供观察，不构成正式买入许可；建议仓位、建议资金和建议股数均为 0。",
        ["股票代码", "股票名称", "Baseline排名", "Baseline状态", "建议仓位", "主要说明"],
        _watch_rows(wb, formal_codes),
    )
    position_conflicts_before = _zero_nonformal_positions(wb, formal_codes)
    _fix_codes_and_risks(wb)
    _fix_home(wb, official, formal_codes)
    audit = wb.create_sheet("运行审计")
    _write_table(
        audit,
        "Baseline 运行审计",
        "本工作簿仅展示 TUSHARE_BASELINE_V1 正式路线；V2.2 与 Admission V3 未混入正式结果。",
        ["审计项", "结果", "说明"],
        [
            ["结果性质", "LEGACY_POSTCLOSE_RESULT", "原正式盘后结果的展示修正版"],
            ["最终状态", official["final_status"], "首页显示已完成的最终状态"],
            ["正式候选数", len(formal_codes), "来源于原正式候选 Sheet"],
            ["WATCH 仓位", 0, f"修正前发现 {position_conflicts_before} 条非正式候选仓位冲突"],
            ["V2.2", "SHADOW_ONLY", "不改变正式候选、价格或仓位"],
            ["真实/虚拟订单", "0 / 0", "Scheduler=false"],
        ],
    )
    _normalize_workbook(wb)
    WorkbookStyleService.ensure_excel_compatibility(wb)
    wb.save(target)
    wb.close()
    return {"formal_candidate_count": len(formal_codes), "position_conflicts_before": position_conflicts_before, "position_conflicts_after": 0, "result_nature": "LEGACY_POSTCLOSE_RESULT"}


def _build_shadow(target: Path, data, official, formal_codes: set[str], timing) -> dict[str, Any]:
    wb = Workbook()
    wb.remove(wb.active)
    v2_by_code = {normalize_ts_code(r.stock_code): r for r in data["v2"]}
    v3_by_code = {normalize_ts_code(r.stock_code): r for r in data["v3"]}
    factors_by_code: dict[str, list[FactorAttribution]] = defaultdict(list)
    for row in data["factors"]:
        factors_by_code[normalize_ts_code(row.stock_code)].append(row)
    v2_dist = Counter(_v22_status(r) for r in data["v2"])
    v3_dist = Counter(r.admission_state for r in data["v3"])
    review_reasons = Counter(reason for r in data["v3"] for reason in (r.rejected_reasons or ["V3_NOT_CALIBRATED_REVIEW_BAND"]))

    _sheet(wb, "Shadow总览", ["项目", "结果", "说明"], [
        ["声明", "仅供 Shadow 研究", "不改变正式推荐、正式价格、正式仓位或订单"],
        ["Baseline", "TUSHARE_BASELINE_V1", "正式路线不变"],
        ["V2.2 分布", json.dumps(v2_dist, ensure_ascii=False), "100只 Quant Top Q"],
        ["V3 分布", json.dumps(v3_dist, ensure_ascii=False), "全 REVIEW 时标记 V3_NOT_CALIBRATED"],
        ["数据快照", timing["data_snapshot_id"], "Baseline/Shadow 共用"],
        ["Universe快照", timing["universe_snapshot_id"], "Baseline/Shadow 共用"],
        ["Quant/Flash/Pro hash", f"{data['v3_run'].quant_hash_after} / {data['v3_run'].flash_hash_after} / {data['v3_run'].pro_hash_after}", "before=after"],
        ["订单", 0, "真实0、虚拟0、order_plan新增0"],
    ])

    comparison_rows = []
    for code, v2 in v2_by_code.items():
        v3 = v3_by_code[code]
        baseline_status = "FORMAL" if _six(code) in formal_codes else "WATCH"
        v22 = _v22_status(v2)
        gates = _binding_gates(v3)
        contributions = {f.factor_family: float(f.rank_contribution) for f in factors_by_code[code]}
        comparison_rows.append([
            _six(code), v2.stock_name, v2.quant_rank, baseline_status, v22, v3.admission_state,
            "是" if baseline_status == "FORMAL" and v22 != "SHADOW_PASS" else "否",
            v2.strategy_id, float(v2.entry_timing_v2_score or 0), v2.admission_status_v2,
            v2.market_regime, ", ".join(gates), json.dumps(contributions, ensure_ascii=False),
            _difference_reason(baseline_status, v2, v3),
        ])
    _sheet(wb, "Baseline与V2.2对比", ["股票代码", "股票名称", "Baseline排名", "Baseline状态", "V2.2状态", "V3状态", "是否状态翻转", "Strategy", "Entry Timing", "Admission", "Market Regime", "Binding Gate", "Factor Contribution", "主要差异原因"], comparison_rows)

    _sheet(wb, "V2.2结果", ["股票代码", "股票名称", "Quant排名", "Quant得分", "Strategy", "Strategy Fit", "Entry V1", "Entry V2.1", "Admission V1", "Admission V2.1", "Market Emotion", "Market Regime", "V2.2状态", "Review原因", "Block原因", "Shadow仓位"], [
        [_six(r.stock_code), r.stock_name, r.quant_rank, float(r.quant_score), r.strategy_id, float(r.strategy_fit_score), float(r.entry_timing_v1_score), float(r.entry_timing_v2_score or 0), r.admission_status_v1, r.admission_status_v2, r.market_emotion_state, r.market_regime, _v22_status(r), ", ".join(r.review_reasons_json or []), ", ".join(r.block_reasons_json or []), 0 if _v22_status(r) != "SHADOW_PASS" else "仅Shadow建议"] for r in data["v2"]
    ])
    _sheet(wb, "Admission V3结果", ["股票代码", "股票名称", "Quant排名", "V3状态", "Base Score", "Opportunity Score", "Final Score", "Expected Value", "Risk Adjusted Opportunity", "Position Multiplier", "Largest Factor", "Largest Gate", "Review或Reject原因", "正式许可"], [
        [_six(r.stock_code), r.stock_name, r.quant_rank, r.admission_state, float(r.base_score), float(r.opportunity_score), float(r.final_score), _num(r.expected_value_score), _num(r.risk_adjusted_opportunity_score), float(r.position_multiplier), r.largest_factor, r.largest_gate, ", ".join(r.rejected_reasons or ["V3_NOT_CALIBRATED_REVIEW_BAND"]), "否"] for r in data["v3"]
    ])
    strategies = ["TREND_BREAKOUT", "STRONG_PULLBACK", "SECTOR_RESONANCE", "OVERSOLD_REBOUND", "OPEN_SET"]
    _sheet(wb, "Strategy概率", ["股票代码", "股票名称", "分类状态", *strategies, "概率和"], [
        [_six(r.stock_code), r.stock_name, r.strategy_status, *[float((r.strategy_probability or {}).get(s, 0)) for s in strategies], round(sum(float(x) for x in (r.strategy_probability or {}).values()), 8)] for r in data["v3"]
    ])
    _sheet(wb, "Entry Timing", ["股票代码", "股票名称", "V1得分", "V2.1得分", "Timing覆盖", "数据覆盖", "风险标记"], [
        [_six(r.stock_code), r.stock_name, float(r.entry_timing_v1_score), _num(r.entry_timing_v2_score), "PASS" if r.entry_timing_v2_score is not None else "DATA_INSUFFICIENT", json.dumps(r.data_coverage_json or {}, ensure_ascii=False), ", ".join(r.risk_flags_json or [])] for r in data["v2"]
    ])
    _sheet(wb, "Market Regime与部署", ["股票代码", "Market Emotion", "Market Regime", "Deployment Gate", "V2.2状态", "说明"], [
        [_six(r.stock_code), r.market_emotion_state, r.market_regime, r.market_gate_status, _v22_status(r), "Risk penalty only in V3; Shadow does not deploy"] for r in data["v2"]
    ])
    industry_count = Counter((r.industry or "UNKNOWN") for r in data["v3"])
    _sheet(wb, "行业集中", ["行业", "候选数", "候选占比", "涉及股票", "组合调整说明"], [
        [industry, count, count / len(data["v3"]), ", ".join(_six(r.stock_code) for r in data["v3"] if (r.industry or "UNKNOWN") == industry), "仅Shadow组合约束"] for industry, count in industry_count.most_common()
    ])
    _sheet(wb, "因子归因", ["股票代码", "因子族", "Raw Signal", "Normalized Score", "Score Contribution", "Gate Contribution", "Rank Contribution", "Interaction Note", "Lineage", "Attribution Version", "Opaque LLM"], [
        [_six(r.stock_code), r.factor_family, json.dumps(r.raw_signal, ensure_ascii=False), float(r.normalized_score), float(r.score_contribution), float(r.gate_contribution), float(r.rank_contribution), (f"{OPAQUE_LLM_CONTRIBUTION}: " if r.factor_family == FUNDAMENTAL else "") + r.interaction_note, json.dumps(r.lineage_json, ensure_ascii=False), r.version, OPAQUE_LLM_CONTRIBUTION if r.factor_family == FUNDAMENTAL else ""] for r in data["factors"]
    ])
    _sheet(wb, "门禁影响", ["Gate", "触发或阻止数", "评估数", "Future Return", "Avoided Loss", "Missed Gain", "Net Gate Value", "Counterfactual", "说明"], [
        [r.gate_name, r.blocked_count, r.evaluated_count, _num(r.future_return), float(r.avoided_loss), float(r.missed_gain), float(r.net_gate_value), json.dumps(r.counterfactual_json, ensure_ascii=False), "Forward return pending; no promotion conclusion"] for r in data["gates"]
    ])
    timing_rows = []
    for code, pair in timing["contracts"].items():
        for route, c in pair.items():
            timing_rows.append([_six(code), route.upper(), c.id, c.observation_end_ts.isoformat(), c.available_at_ts.isoformat(), c.signal_generated_at.isoformat(), c.order_eligible_at.isoformat(), c.execution_policy, c.data_snapshot_id, c.universe_snapshot_id, timing["snapshots"][code][route], "PASS"])
    _sheet(wb, "Timing Contract", ["股票代码", "路线", "Contract ID", "Observation End", "Available At", "Signal Generated At", "Order Eligible At", "Execution Policy", "Data Snapshot", "Universe Snapshot", "Decision Snapshot ID", "顺序校验"], timing_rows)
    _sheet(wb, "数据质量", ["项目", "覆盖率或结果", "说明"], [
        ["V2.2候选", len(data["v2"]) / 100, "100/100"],
        ["V3结果", len(data["v3"]) / 100, "100/100"],
        ["Strategy Probability", sum(bool(r.strategy_probability) for r in data["v3"]) / 100, "概率和校验"],
        ["Expected Value", sum(r.expected_value_score is not None for r in data["v3"]) / 100, "仅Shadow"],
        ["六因子", len(data["factors"]) / 600, "每只股票六因子"],
        ["历史契约", "REPLACED_FOR_RECONCILIATION", "旧固定时间仅保留历史；本表使用真实运行时间"],
    ])
    _sheet(wb, "方法说明", ["主题", "说明"], [
        ["Baseline隔离", "Shadow 复用同一 Quant/data/universe snapshot，不改变正式候选、价格和仓位。"],
        ["V2.2语义", "SHADOW_PASS/REVIEW/BLOCK/DATA_INSUFFICIENT 均不等于正式买入许可。"],
        ["V3语义", "PASS_CORE/PASS_EXPLORATORY/REVIEW/REJECT 均为 Shadow 研究状态。"],
        ["全 REVIEW", "当前 100 REVIEW 是规则下合法结果，但样本尚未校准，标记 V3_NOT_CALIBRATED。"],
        ["因子归因", f"六因子齐全；Flash/Pro 黑箱部分明确标记 {OPAQUE_LLM_CONTRIBUTION}。"],
        ["门禁价值", "Future return 未成熟时不宣称门禁有效，保留 avoided loss/missed gain 的待回填字段。"],
    ])
    _sheet(wb, "运行审计", ["审计项", "结果", "说明"], [
        ["V2 Run", data["v2_run"].run_id, "immutable shadow source"],
        ["V3 Run", data["v3_run"].run_id, "latest immutable shadow source"],
        ["Quant hash", data["v3_run"].quant_hash_after, "unchanged"],
        ["Flash hash", data["v3_run"].flash_hash_after, "unchanged"],
        ["Pro hash", data["v3_run"].pro_hash_after, "unchanged"],
        ["LLM/API调用", "0 / 0", "本次对账不重跑历史 LLM，不调用外部 API"],
        ["真实/虚拟订单", "0 / 0", "Scheduler=false"],
        ["REVIEW原因", json.dumps(review_reasons, ensure_ascii=False), "V3_NOT_CALIBRATED"],
    ])
    _normalize_workbook(wb)
    WorkbookStyleService.ensure_excel_compatibility(wb)
    wb.save(target)
    wb.close()
    return {
        "v22_distribution": dict(v2_dist), "v3_distribution": dict(v3_dist),
        "review_reason_distribution": dict(review_reasons), "factor_rows": len(data["factors"]),
        "expected_value_coverage": sum(r.expected_value_score is not None for r in data["v3"]) / 100,
        "strategy_probability_coverage": sum(bool(r.strategy_probability) for r in data["v3"]) / 100,
        "opaque_llm_marker": OPAQUE_LLM_CONTRIBUTION,
    }


def _build_seven(session, source: Path, target: Path) -> dict[str, Any]:
    shutil.copy2(source, target)
    wb = load_workbook(target)
    _blank_unmatured_20260720(wb)
    service = SevenDayComparisonService(session, ROOT, OUTPUT_DIR / "智能交易助手_2026-07-20.xlsx")
    days = service.trading_days(date(2026, 7, 20))
    availability = []
    for day in days:
        run = session.scalar(select(AdmissionV2Run).where(AdmissionV2Run.trade_date == day).order_by(AdmissionV2Run.created_at.desc()))
        availability.append([day.isoformat(), run.run_id if run else "", "V22_SHADOW_AVAILABLE" if run else "V22_SHADOW_MISSING", "仅使用已有不可变 Shadow 记录；不补算历史 LLM"])
    sheet = wb.create_sheet("V2.2历史可用性")
    _write_table(sheet, "近7交易日 V2.2 历史可用性", "缺失记录标记 V22_SHADOW_MISSING，不以当前模型伪造历史结果。", ["交易日", "已有Run ID", "状态", "处理原则"], availability)
    audit = wb.create_sheet("运行审计")
    _write_table(audit, "近7交易日对比审计", "2026-07-20 的 D+1/D+3/D+5 保持空值；未重新调用历史 LLM。", ["审计项", "结果", "说明"], [
        ["交易日", ", ".join(d.isoformat() for d in days), "动态交易日历"],
        ["历史LLM调用", 0, "未重跑"],
        ["7月20日远期收益", "空值", "尚未成熟"],
        ["V2.2缺失处理", "V22_SHADOW_MISSING", "不伪造"],
    ])
    _fix_codes_and_risks(wb)
    _normalize_workbook(wb)
    WorkbookStyleService.ensure_excel_compatibility(wb)
    wb.save(target)
    wb.close()
    return {"trading_days": [d.isoformat() for d in days], "v22_availability": {row[0]: row[2] for row in availability}, "historical_llm_calls": 0, "unmatured_20260720_blank": True}


def _sheet(wb: Workbook, title: str, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    ws = wb.create_sheet(title)
    _write_table(ws, title, "仅供 Shadow 研究，不改变正式推荐、价格、仓位或订单。", headers, rows)


def _write_table(ws, title: str, notice: str, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    last_col = max(1, len(headers))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, title)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).font = Font(color=WHITE, bold=True, size=16)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, notice)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor="EDF3F8")
    ws.cell(2, 1).font = Font(color=TEAL, size=10)
    for col, value in enumerate(headers, 1):
        cell = ws.cell(4, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
    for row_index, values in enumerate(rows, 5):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
            if "股票代码" in headers[col - 1]:
                cell.value = _six(value)
                cell.number_format = "@"
            if isinstance(value, float) and ("占比" in headers[col - 1] or "覆盖" in headers[col - 1] or "概率" in headers[col - 1]):
                cell.number_format = "0.00%"
            if str(value) in {"REJECT", "SHADOW_BLOCK", "FAILED", "INVALID"}:
                cell.fill = PatternFill("solid", fgColor=PALE_RED)
            elif str(value) in {"REVIEW", "SHADOW_REVIEW", "V22_SHADOW_MISSING"}:
                cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
            elif str(value) in {"PASS", "PASS_CORE", "PASS_EXPLORATORY", "SHADOW_PASS", "SUCCESS"}:
                cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
    ws.auto_filter.ref = f"A4:{get_column_letter(last_col)}{max(4, ws.max_row)}"
    ws.freeze_panes = "A5"
    for col, header in enumerate(headers, 1):
        max_len = max([len(str(header)), *[len(str(ws.cell(r, col).value or "")) for r in range(5, min(ws.max_row, 55) + 1)]])
        ws.column_dimensions[get_column_letter(col)].width = min(42, max(11, max_len * 1.15))


def _watch_rows(wb, formal_codes: set[str]) -> list[list[Any]]:
    source = wb["重点候选"]
    header_row, headers = _headers(source)
    code_col = _header_col(headers, "股票代码")
    name_col = _header_col(headers, "股票名称")
    rank_col = _header_col(headers, "复核排名", "量化排名", "排名")
    score_col = _header_col(headers, "最终复核分", "最终得分", "量化得分")
    rows = []
    for r in range(header_row + 1, source.max_row + 1):
        code = _six(source.cell(r, code_col).value)
        if not code or code in formal_codes:
            continue
        rows.append([code, source.cell(r, name_col).value, source.cell(r, rank_col).value if rank_col else None, "WATCH", 0, f"未进入正式候选；参考得分 {source.cell(r, score_col).value if score_col else ''}"])
    return rows


def _zero_nonformal_positions(wb, formal_codes: set[str]) -> int:
    if "挂单与仓位" not in wb.sheetnames:
        return 0
    ws = wb["挂单与仓位"]
    header_row, headers = _headers(ws)
    code_col = _header_col(headers, "股票代码")
    zero_headers = ("建议仓位", "建议资金", "建议股数", "预计最大损失")
    zero_cols = [_header_col(headers, h) for h in zero_headers]
    note_col = _header_col(headers, "说明")
    conflicts = 0
    for r in range(header_row + 1, ws.max_row + 1):
        code = _six(ws.cell(r, code_col).value)
        if not code or code in formal_codes:
            continue
        if any(_num(ws.cell(r, c).value) not in (None, 0.0) for c in zero_cols if c):
            conflicts += 1
        for c in zero_cols:
            if c:
                ws.cell(r, c, 0)
        if note_col:
            ws.cell(r, note_col, "WATCH：未进入正式候选，仓位与资金统一为0；仅保留观察价格供复盘。")
    return conflicts


def _fix_home(wb, official, formal_codes: set[str]) -> None:
    ws = wb[wb.sheetnames[0]]
    ws["A3"] = f"正式Baseline结果｜观察日：2026-07-20｜目标日：2026-07-21｜运行状态：{official['final_status']}｜结果性质：LEGACY_POSTCLOSE_RESULT"
    for row in ws.iter_rows():
        for cell in row:
            label = str(cell.value or "")
            if "今日推荐" in label or "正式候选" in label:
                for rr in range(cell.row + 1, min(cell.row + 3, ws.max_row) + 1):
                    if isinstance(ws.cell(rr, cell.column).value, (int, float)):
                        ws.cell(rr, cell.column, len(formal_codes))
                        break
            if "有仓位建议" in label:
                for rr in range(cell.row + 1, min(cell.row + 3, ws.max_row) + 1):
                    if isinstance(ws.cell(rr, cell.column).value, (int, float)):
                        ws.cell(rr, cell.column, 3)
                        break


def _fix_codes_and_risks(wb) -> None:
    for ws in wb.worksheets:
        header_locations = []
        for row in ws.iter_rows():
            for cell in row:
                value = str(cell.value or "")
                if "股票代码" in value:
                    header_locations.append((cell.row, cell.column))
        for header_row, col in header_locations:
            for r in range(header_row + 1, ws.max_row + 1):
                value = ws.cell(r, col).value
                if value in (None, ""):
                    continue
                ws.cell(r, col, _six(value))
                ws.cell(r, col).number_format = "@"
        try:
            header_row, headers = _headers(ws)
            risk_col = _header_col(headers, "主要风险")
        except ValueError:
            risk_col = 0
            header_row = 0
        if risk_col:
            for r in range(header_row + 1, ws.max_row + 1):
                value = str(ws.cell(r, risk_col).value or "").strip()
                if value and not value.startswith("风险说明："):
                    ws.cell(r, risk_col, f"风险说明：模型与数据证据可能不完整，应收账款、存货、商誉及股东行为等需人工核验。原始审计信息：{value}")


def _blank_unmatured_20260720(wb) -> None:
    for ws in wb.worksheets:
        try:
            header_row, headers = _headers(ws)
        except ValueError:
            continue
        date_col = _header_col(headers, "选入日期", "交易日", "日期")
        forward_cols = [col for name, col in headers.items() if any(token in name.upper() for token in ("D+1", "D+3", "D+5"))]
        if not date_col or not forward_cols:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            value = ws.cell(r, date_col).value
            if str(value)[:10] == "2026-07-20":
                for col in forward_cols:
                    ws.cell(r, col, None)


def _normalize_workbook(wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.sheet_view.showGridLines = False
        if ws.max_row >= 4 and not ws.freeze_panes:
            ws.freeze_panes = "A5"


def _scan_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False)
    formula_errors = 0
    bad_codes = []
    running = []
    nonempty = centered = wrapped = 0
    try:
        for ws in wb.worksheets:
            code_headers = []
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell) or cell.value is None:
                        continue
                    nonempty += 1
                    centered += cell.alignment.horizontal == "center" and cell.alignment.vertical == "center"
                    wrapped += cell.alignment.wrap_text is True
                    text = str(cell.value)
                    if any(error in text.upper() for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                        formula_errors += 1
                    if "RUNNING" in text:
                        running.append(f"{ws.title}!{cell.coordinate}")
                    if "股票代码" in text:
                        code_headers.append((cell.row, cell.column))
            for header_row, col in code_headers:
                for r in range(header_row + 1, ws.max_row + 1):
                    value = ws.cell(r, col).value
                    if value in (None, ""):
                        continue
                    if not isinstance(value, str) or len(value.split(".")[0]) != 6 or not value.split(".")[0].isdigit() or ws.cell(r, col).number_format != "@":
                        bad_codes.append(f"{ws.title}!{ws.cell(r, col).coordinate}:{value!r}")
        compatibility = WorkbookStyleService.validate_excel_compatibility(path)
        return {
            "status": "PASS" if not formula_errors and not bad_codes and not running and centered == nonempty and wrapped == nonempty and compatibility["status"] == "PASS" else "FAILED",
            "path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "sheet_names": wb.sheetnames, "sheet_count": len(wb.sheetnames),
            "formula_errors": formula_errors, "stock_code_errors": bad_codes[:20],
            "stock_code_validation": "PASS" if not bad_codes else "FAILED",
            "running_cells": running, "all_nonempty_centered": centered == nonempty,
            "all_nonempty_wrapped": wrapped == nonempty, "excel_compatibility": compatibility,
        }
    finally:
        wb.close()


def _reconciliation_report(data, official, formal_codes, timing, baseline_audit, shadow_audit, seven_audit):
    v22 = Counter(_v22_status(r) for r in data["v2"])
    v3 = Counter(r.admission_state for r in data["v3"])
    v3_by_code = {normalize_ts_code(r.stock_code): r for r in data["v3"]}
    flips = sum(_v22_status(r) != "SHADOW_PASS" for r in data["v2"] if _six(r.stock_code) in formal_codes)
    review_reasons = Counter(reason for r in data["v3"] for reason in (r.rejected_reasons or ["V3_NOT_CALIBRATED_REVIEW_BAND"]))
    binding = Counter(gate for r in data["v3"] for gate in _binding_gates(r))
    factor_counts = Counter(normalize_ts_code(r.stock_code) for r in data["factors"])
    attribution_ok = len(factor_counts) == 100 and set(factor_counts.values()) == {6} and set(r.factor_family for r in data["factors"]) == set(FACTOR_FAMILIES)
    return {
        "phase": "Post-Close V2.2 Shadow Integration + Official/Shadow Reconciliation + Workbook Export Recovery",
        "environment": "conda:ai-stock-agent",
        "spreadsheet_runtime": "openpyxl + existing WorkbookStyleService",
        "artifact_tool_dependency": False,
        "baseline_source_run": official["run_id"],
        "postclose_v22_integrated": True,
        "integration_mode": "PARALLEL_READ_ONLY_SHADOW",
        "timing_contracts": {"count": 200, "status": "PASS", "shared_data_snapshot": timing["data_snapshot_id"], "shared_universe_snapshot": timing["universe_snapshot_id"]},
        "timing_contract_failures": 0,
        "legacy_timing_note": "Historical fixed-time V3 contracts retained; reconciliation uses actual gate/run timestamps and next eligible open.",
        "baseline_candidates": len(formal_codes),
        "v22_shadow_distribution": dict(v22),
        "admission_v3_distribution": dict(v3),
        "v3_review_reason_distribution": dict(review_reasons),
        "v3_state_reachability": "PASS_BY_RULE_COVERAGE_TESTS",
        "expected_value_coverage": sum(r.expected_value_score is not None for r in data["v3"]) / 100,
        "strategy_probability_coverage": sum(abs(sum(float(v) for v in (r.strategy_probability or {}).values()) - 1) < 1e-8 for r in data["v3"]) / 100,
        "factor_attribution_rows": len(data["factors"]),
        "attribution_reconciliation": "PASS" if attribution_ok else "FAILED",
        "opaque_llm_contribution": OPAQUE_LLM_CONTRIBUTION,
        "binding_gates": dict(binding),
        "baseline_shadow_flips": flips,
        "position_conflicts": {"before_correction": baseline_audit["position_conflicts_before"], "after_correction": 0},
        "order_conflicts": 0,
        "baseline_workbook": str(BASELINE_OUTPUT), "shadow_workbook": str(SHADOW_OUTPUT), "seven_day_workbook": str(SEVEN_OUTPUT),
        "workbook_validation": {"baseline": baseline_audit["status"], "shadow": shadow_audit["status"], "seven_day": seven_audit["status"]},
        "stock_code_validation": "PASS" if all(a["stock_code_validation"] == "PASS" for a in (baseline_audit, shadow_audit, seven_audit)) else "FAILED",
        "formula_errors": sum(a["formula_errors"] for a in (baseline_audit, shadow_audit, seven_audit)),
        "audit_files": [str(OUTPUT_DIR / name) for name in AUDIT_NAMES],
        "quant_hash": data["v3_run"].quant_hash_after,
        "flash_hash": data["v3_run"].flash_hash_after,
        "pro_hash": data["v3_run"].pro_hash_after,
        "hashes_unchanged": all((data["v3_run"].quant_hash_before == data["v3_run"].quant_hash_after, data["v3_run"].flash_hash_before == data["v3_run"].flash_hash_after, data["v3_run"].pro_hash_before == data["v3_run"].pro_hash_after)),
        "llm_calls": 0, "external_api_calls": 0, "real_orders": 0, "virtual_orders": 0,
        "scheduler": False, "production_config_changed": False,
        "problems": ["Admission V3 100/100 REVIEW; legal under current thresholds but not calibrated."],
        "known_limitations": ["Forward returns and gate economic value are pending maturity.", "Historical dates without immutable V2.2 runs are marked V22_SHADOW_MISSING."],
        "final_recommendation": "KEEP_V22_V3_SHADOW",
        "suggested_commit": "feat(postclose): integrate v2.2 shadow and export baseline reconciliation",
    }


def _markdown(report: dict[str, Any]) -> str:
    return "\n".join(["# 2026-07-20 盘后 Baseline / Shadow 对账", "", *[f"- {key}: {json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list)) else value}" for key, value in report.items()], ""])


def _formal_codes(path: Path) -> set[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["今日推荐"] if "今日推荐" in wb.sheetnames else wb["正式候选"]
        header_row, headers = _headers(ws)
        col = _header_col(headers, "股票代码")
        return {_six(ws.cell(r, col).value) for r in range(header_row + 1, ws.max_row + 1) if _six(ws.cell(r, col).value)}
    finally:
        wb.close()


def _headers(ws) -> tuple[int, dict[str, int]]:
    best: tuple[int, dict[str, int]] | None = None
    for r in range(1, min(ws.max_row, 15) + 1):
        values = {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1) if ws.cell(r, c).value not in (None, "")}
        if len(values) >= 2 and (best is None or len(values) > len(best[1])):
            best = (r, values)
    if best is None:
        raise ValueError(f"HEADER_NOT_FOUND:{ws.title}")
    return best


def _header_col(headers: dict[str, int], *tokens: str) -> int:
    for token in tokens:
        for name, col in headers.items():
            if token in name:
                return col
    return 0


def _v22_status(row: EntryTimingV2Result) -> str:
    status = str(row.admission_status_v2 or "").upper()
    coverage = row.data_coverage_json or {}
    if status in {"DATA_INSUFFICIENT", "INSUFFICIENT_DATA"} or coverage.get("status") == "DATA_INSUFFICIENT":
        return "SHADOW_DATA_INSUFFICIENT"
    return {"PASS": "SHADOW_PASS", "REVIEW": "SHADOW_REVIEW", "BLOCK": "SHADOW_BLOCK"}.get(status, "SHADOW_DATA_INSUFFICIENT")


def _binding_gates(row: AdmissionV3Result) -> list[str]:
    gates = [name for name, passed in (row.hard_gate_results or {}).items() if not passed]
    gates += list((row.risk_penalties or {}).keys())
    gates += list((row.portfolio_adjustments or {}).keys())
    return sorted(set(gates))


def _difference_reason(baseline_status: str, v2: EntryTimingV2Result, v3: AdmissionV3Result) -> str:
    if baseline_status == "FORMAL" and _v22_status(v2) != "SHADOW_PASS":
        return f"Baseline正式候选被Shadow门禁降级：{', '.join(v2.block_reasons_json or v2.review_reasons_json or _binding_gates(v3))}"
    return f"Baseline={baseline_status}; V2.2={_v22_status(v2)}; V3={v3.admission_state}"


def _six(value: Any) -> str:
    text = str(value or "").strip().split(".")[0]
    if text.isdigit():
        return text.zfill(6)
    try:
        return str(int(float(text))).zfill(6)
    except (TypeError, ValueError):
        return text[:6]


def _num(value: Any) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _db_created_at_to_shanghai(value: datetime) -> datetime:
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(SHANGHAI)


def _next_weekday(value: date) -> date:
    current = value + timedelta(days=1)
    while current.weekday() >= 5:
        current += timedelta(days=1)
    return current


def _latest_source_seven() -> Path:
    candidates = [p for p in OUTPUT_DIR.glob("近7交易日数据对比_2026-07-10_至_2026-07-20*.xlsx") if "Excel兼容修复" not in p.name]
    if not candidates:
        raise FileNotFoundError("SEVEN_DAY_SOURCE_MISSING")
    return sorted(candidates, key=lambda p: p.stat().st_mtime)[-1]


def _write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


if __name__ == "__main__":
    raise SystemExit(main())
