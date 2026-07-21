from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from sqlalchemy import select


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env", override=False)

from database.models.ifind_shadow import ExternalProviderUsage
from database.models.postclose_official import PostCloseOfficialRun
from database.session import get_session
from post_close.seven_day_comparison import SevenDayComparisonService
from reporting.workbook_style import WorkbookStyleService


ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")


def _scan(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    try:
        nonempty = centered = wrapped = formula_errors = stock_codes = 0
        bad_alignment: list[str] = []
        bad_stock_codes: list[str] = []
        for worksheet in workbook.worksheets:
            code_columns: set[int] = set()
            header_rows: dict[int, int] = {}
            for row in worksheet.iter_rows():
                for cell in row:
                    if "股票代码" in str(cell.value or ""):
                        code_columns.add(cell.column)
                        header_rows[cell.column] = cell.row
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell) or cell.value is None:
                        continue
                    nonempty += 1
                    if cell.alignment.horizontal == "center" and cell.alignment.vertical == "center":
                        centered += 1
                    else:
                        bad_alignment.append(f"{worksheet.title}!{cell.coordinate}:center")
                    if cell.alignment.wrap_text:
                        wrapped += 1
                    else:
                        bad_alignment.append(f"{worksheet.title}!{cell.coordinate}:wrap")
                    if isinstance(cell.value, str) and any(error in cell.value.upper() for error in ERRORS):
                        formula_errors += 1
                    if cell.column in code_columns and cell.row > header_rows[cell.column] and cell.value not in (None, ""):
                        stock_codes += 1
                        code = str(cell.value).split(".")[0]
                        if cell.number_format != "@" or not isinstance(cell.value, str) or len(code) != 6 or not code.isdigit():
                            bad_stock_codes.append(f"{worksheet.title}!{cell.coordinate}:{cell.value!r}:{cell.number_format}")
        return {
            "path": str(path),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "sheet_names": list(workbook.sheetnames),
            "nonempty_cells": nonempty,
            "centered_cells": centered,
            "wrapped_cells": wrapped,
            "all_nonempty_centered": centered == nonempty,
            "all_nonempty_wrapped": wrapped == nonempty,
            "alignment_errors": bad_alignment[:20],
            "stock_code_cells": stock_codes,
            "stock_code_format_pass": not bad_stock_codes,
            "stock_code_errors": bad_stock_codes[:20],
            "formula_errors": formula_errors,
            "home_status": str(workbook[workbook.sheetnames[0]]["A3"].value or ""),
        }
    finally:
        workbook.close()


def main() -> int:
    output_dir = ROOT / "outputs" / "2026-07-20" / "正式日线"
    reference = ROOT / "outputs" / "2026-07-17" / "智能交易助手_2026-07-17.xlsx"
    main_workbook = sorted(
        output_dir.glob("智能交易助手_2026-07-20*.xlsx"),
        key=lambda path: path.stat().st_mtime,
    )[-1]
    seven_workbook = sorted(
        output_dir.glob("近7交易日数据对比_*_至_2026-07-20*.xlsx"),
        key=lambda path: path.stat().st_mtime,
    )[-1]
    report_path = sorted(
        output_dir.glob("postclose_official_2026-07-20_*_reconciled*.json"),
        key=lambda path: path.stat().st_mtime,
    )[-1]
    report = json.loads(report_path.read_text(encoding="utf-8"))

    session = get_session()
    try:
        row = session.scalar(select(PostCloseOfficialRun).where(
            PostCloseOfficialRun.run_id == report["run_id"]
        ))
        usage = list(session.scalars(select(ExternalProviderUsage).where(
            ExternalProviderUsage.provider.ilike("%ifind%")
        )))
        days_service = SevenDayComparisonService(session, ROOT, reference)
        days = days_service.trading_days(date(2026, 7, 20))
        seven_validation = days_service.validate(seven_workbook, days)
    finally:
        session.close()

    task_query = subprocess.run(
        ["schtasks.exe", "/Query", "/TN", "AITrader_PostClose_20260720_1700"],
        capture_output=True,
        text=True,
        check=False,
    )
    main_scan = _scan(main_workbook)
    seven_scan = _scan(seven_workbook)
    structure = WorkbookStyleService.compare_official_structure(reference, main_workbook)
    main_excel_compatibility = WorkbookStyleService.validate_excel_compatibility(main_workbook)
    result = {
        "status": "PASS",
        "run_id": report["run_id"],
        "final_status": report["final_status"],
        "database_status": row.status if row else None,
        "database_llm_calls": row.llm_calls if row else None,
        "database_total_tokens": row.total_tokens if row else None,
        "report_llm_calls": report["llm_calls"],
        "report_total_tokens": report["total_token_usage"],
        "main_workbook": main_scan,
        "main_structure": structure,
        "main_excel_compatibility": main_excel_compatibility,
        "seven_day_workbook": seven_scan,
        "seven_day_validation": seven_validation,
        "seven_trading_days": [value.isoformat() for value in days],
        "ifind_usage_rows_after_official_run": len([
            item for item in usage
            if item.requested_at and item.requested_at.replace(tzinfo=None) >= datetime(2026, 7, 20, 17, 0)
        ]),
        "real_orders": report["real_orders"],
        "virtual_orders": report["virtual_orders"],
        "scheduler": report["scheduler"],
        "production_config_changed": report["production_config_changed"],
        "one_shot_task_deleted": task_query.returncode != 0,
        "report_path": str(report_path),
    }
    required = (
        result["final_status"] == "POSTCLOSE_FULL_A_PARTIAL_SUCCESS",
        result["database_status"] == result["final_status"],
        result["database_llm_calls"] == result["report_llm_calls"] == 222,
        result["database_total_tokens"] == result["report_total_tokens"] == 987761,
        all(structure[key] for key in ("sheet_names_match", "sheet_order_match", "headers_match")),
        main_excel_compatibility["status"] == "PASS",
        main_scan["all_nonempty_centered"], main_scan["all_nonempty_wrapped"],
        main_scan["stock_code_format_pass"], main_scan["formula_errors"] == 0,
        "POSTCLOSE_FULL_A_PARTIAL_SUCCESS" in main_scan["home_status"],
        seven_scan["all_nonempty_centered"], seven_scan["all_nonempty_wrapped"],
        seven_scan["stock_code_format_pass"], seven_scan["formula_errors"] == 0,
        seven_validation["status"] == "PASS",
        seven_validation["excel_compatibility"]["status"] == "PASS",
        result["ifind_usage_rows_after_official_run"] == 0,
        result["real_orders"] == result["virtual_orders"] == 0,
        result["scheduler"] is False,
        result["production_config_changed"] is False,
        result["one_shot_task_deleted"],
    )
    if not all(required):
        result["status"] = "FAILED"
    print(json.dumps(result, ensure_ascii=True, default=str))
    return 0 if result["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
