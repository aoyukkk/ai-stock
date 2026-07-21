from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def main() -> int:
    parser = argparse.ArgumentParser(description="Render every workbook sheet with artifact-tool.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--sheet", action="append", dest="selected_sheets")
    parser.add_argument("--range", dest="render_range", default=None)
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    output_dir = args.output_dir.resolve()

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        all_sheet_names = list(workbook.sheetnames)
    finally:
        workbook.close()
    sheet_names = args.selected_sheets or all_sheet_names
    unknown = sorted(set(sheet_names) - set(all_sheet_names))
    if unknown:
        raise ValueError(f"UNKNOWN_WORKBOOK_SHEETS:{unknown}")

    dependency_root = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules"
    if not (dependency_root / "@oai" / "artifact-tool").exists():
        raise RuntimeError("ARTIFACT_TOOL_DEPENDENCY_NOT_FOUND")

    output_dir.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="artifact-render-", dir=output_dir.parent) as temp_value:
        temp_dir = Path(temp_value)
        script = temp_dir / "render_workbook_sheets.mjs"
        shutil.copy2(ROOT / "scripts" / "render_workbook_sheets.mjs", script)
        link = temp_dir / "node_modules"
        linked = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(dependency_root)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        if linked.returncode != 0:
            raise RuntimeError("ARTIFACT_TOOL_JUNCTION_FAILED")
        output_dir.mkdir(parents=True, exist_ok=True)
        render_env = os.environ.copy()
        if args.render_range:
            render_env["ARTIFACT_RENDER_RANGE"] = args.render_range
        completed = subprocess.run(
            ["node", str(script), str(workbook_path), str(output_dir), *sheet_names],
            cwd=temp_dir,
            env=render_env,
            check=False,
        )
        if link.exists():
            os.rmdir(link)
        expected = [output_dir / f"{name}.png" for name in sheet_names]
        rendered_completely = all(path.is_file() and path.stat().st_size > 0 for path in expected)
        formula_check = output_dir / "公式检查.ndjson"
        if completed.returncode != 0 and not (rendered_completely and formula_check.is_file()):
            raise RuntimeError(f"ARTIFACT_TOOL_RENDER_FAILED:{completed.returncode}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
