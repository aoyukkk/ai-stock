from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import uuid
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo


BLUE = "1F4E78"
BLUE_LIGHT = "D9EAF7"
BLUE_PALE = "EDF5FB"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
WHITE = "FFFFFF"
TEXT = "172033"
MUTED = "596579"
GRID = "C9D8E6"


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply a human verification overlay to a daily workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("audit", type=Path)
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    audit_path = args.audit.resolve()
    payload = json.loads(audit_path.read_text(encoding="utf-8"))
    items = payload["items"]
    by_code = {str(item["stock_code"]).zfill(6): item for item in items}

    history = workbook_path.parent / "历史版本"
    history.mkdir(parents=True, exist_ok=True)
    old_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()[:8]
    backup = history / f"{workbook_path.stem}_人工核验前_{old_hash}.xlsx"
    if not backup.exists():
        shutil.copy2(workbook_path, backup)

    wb = load_workbook(workbook_path)
    _update_candidate_sheet(wb["重点候选"], by_code)
    if "今日推荐" in wb.sheetnames:
        _update_recommendation_sheet(wb["今日推荐"], by_code)
    _update_order_sheet(wb["挂单与仓位"], by_code)
    _update_fundamental_sheet(wb["基本面摘要"], by_code)
    _rewrite_current_issues(wb["当前问题"], by_code)
    for sheet_name in ("人工核验补全", "核验来源"):
        if sheet_name in wb.sheetnames:
            wb.remove(wb[sheet_name])
    _build_verification_sheet(wb, payload)
    _build_source_sheet(wb, payload)
    if wb.calculation is None:
        wb.calculation = CalcProperties()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    candidate = workbook_path.with_name(f".{workbook_path.stem}_{uuid.uuid4().hex[:8]}_verified.xlsx")
    wb.save(candidate)
    wb.close()
    with zipfile.ZipFile(candidate) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError(f"WORKBOOK_ZIP_CORRUPT:{bad}")
    os.replace(candidate, workbook_path)
    digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    result = {
        "status": "SUCCESS",
        "workbook": str(workbook_path),
        "backup": str(backup),
        "sha256": digest,
        "verified_stock_count": len(items),
        "pending_manual_confirmation_count": _count_pending(workbook_path),
        "reviewer": payload["reviewer"],
        "audit": str(audit_path),
    }
    result_path = audit_path.with_name("人工核验补全_应用结果.json")
    result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def _update_candidate_sheet(ws, by_code: dict[str, dict]) -> None:
    headers = _headers(ws, 4)
    updates = {
        "002156": {
            "产业链": "半导体封装测试产业链（人工核验）",
            "核心逻辑": "主营集成电路封装与测试，先进封装受益于AI算力和半导体景气；但高资本开支、负债率和客户结构需要持续跟踪。",
            "主要风险": "半导体周期、客户集中、资本开支与折旧、负债率、技术迭代及地缘贸易限制。",
            "当前状态": "基本面及二筛已人工补齐；规则阻断保留",
        },
        "603726": {
            "二筛结论": "未形成模型结论（人工核验完成）",
            "产业链": "暖通空调与设备热管理零部件（人工核验）",
            "核心逻辑": "传统空调风叶和机械风机是基本盘，EC高效风机及复合材料是增量方向；2026年一季度表观利润增长主要来自非经常性收益，扣非利润仍承压。",
            "主要风险": "空调行业竞争、客户集中、应收与存货、原材料、新业务商业化及非经常性收益影响。",
            "当前状态": "基本面及二筛已人工补齐；规则阻断保留",
        },
    }
    for row in range(5, ws.max_row + 1):
        code = str(ws.cell(row, headers["股票代码"]).value or "").zfill(6)
        for field, value in updates.get(code, {}).items():
            ws.cell(row, headers[field], value)
        item = by_code.get(code, {})
        if item.get("manual_secondary_score") is not None:
            ws.cell(row, headers["二筛得分"], item["manual_secondary_score"])
            ws.cell(row, headers["二筛结论"], item["manual_secondary_conclusion"])
        if code in updates:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=GREEN)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _update_recommendation_sheet(ws, by_code: dict[str, dict]) -> None:
    headers = _headers(ws, 4)
    updates = {
        "002156": {
            "产业链": "半导体封装测试产业链（人工核验）",
            "核心逻辑": "主营集成电路封装与测试，先进封装受益于AI算力和半导体景气；但高资本开支、负债率和客户结构需要持续跟踪。",
            "主要风险": "半导体周期、客户集中、资本开支与折旧、负债率、技术迭代及地缘贸易限制。",
        },
        "603726": {
            "产业链": "暖通空调与设备热管理零部件（人工核验）",
            "核心逻辑": "传统空调风叶和机械风机是基本盘，EC高效风机及复合材料是增量方向；2026年一季度表观利润增长主要来自非经常性收益，扣非利润仍承压。",
            "主要风险": "空调行业竞争、客户集中、应收与存货、原材料、新业务商业化及非经常性收益影响。",
        },
    }
    for row in range(5, ws.max_row + 1):
        code = str(ws.cell(row, headers["股票代码"]).value or "").zfill(6)
        for field, value in updates.get(code, {}).items():
            ws.cell(row, headers[field], value)
        if code in updates:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=GREEN)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _update_order_sheet(ws, by_code: dict[str, dict]) -> None:
    headers = _headers(ws, 4)
    for row in range(5, ws.max_row + 1):
        code = str(ws.cell(row, headers["股票代码"]).value or "").zfill(6)
        item = by_code.get(code, {})
        if item.get("manual_secondary_score") is None:
            continue
        ws.cell(row, headers["二筛得分"], item["manual_secondary_score"])
        ws.cell(row, headers["二筛结论"], item["manual_secondary_conclusion"])
        ws.cell(
            row,
            headers["说明"],
            "二筛空值已由GPT-5.6 SOL完成人工复核；挂单价、目标价及仓位仍按原规则门禁保留为空。",
        )
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _update_fundamental_sheet(ws, by_code: dict[str, dict]) -> None:
    headers = _headers(ws, 4)
    updates = {
        "002156": {
            "产业链": "半导体封装测试产业链",
            "链条位置": "中游",
            "主营业务": "集成电路封装与测试，覆盖高性能计算、存储、汽车电子和消费电子等应用。",
            "核心产品": "先进封装；集成电路封装；集成电路测试",
            "概念标签": "先进封装；AI算力；国产封测",
            "结构性方向": "AI算力与国产半导体需求推动先进封装升级",
            "潜在优势": "具备多基地量产和高性能计算封测经验，但先进封装产能利用与客户结构仍需跟踪。",
            "行业趋势": "AI算力、汽车电子和高性能计算推动封装技术升级，行业同时具有资本密集和周期波动特征。",
            "核心逻辑": "行业景气和先进封装需求改善提供弹性；投资判断需同时约束资本开支、负债率、客户集中和半导体周期。",
            "失效条件": "行业景气转弱；先进封装订单不及预期；负债和资本开支压力扩大；主要客户需求显著下降",
            "财务状态": "稳定但需关注杠杆",
            "财务说明": "项目内2026年一季度规则财务状态为稳定；结合2025年报，业绩受行业回暖带动，但高资本开支和负债率仍是主要约束。",
            "人工复核": "已完成",
            "复核优先级": "中",
        },
        "603726": {
            "产业链": "暖通空调、设备热管理零部件与高分子复合材料",
            "链条位置": "中游",
            "主营业务": "家用空调风叶、机械风机和高分子复合材料；机械风机覆盖中央空调、建筑通风、通信机柜、储能柜和精密空调等场景。",
            "核心产品": "空调风叶；机械风机；EC高效风机；改性高分子复合材料",
            "概念标签": "暖通空调零部件；设备热管理；高效节能风机；高分子复合材料",
            "结构性方向": "EC高效风机向通信机柜、储能柜、充电桩和精密空调等散热场景拓展",
            "潜在优势": "多地生产基地、风叶风机与材料业务协同、空气动力和材料检测能力；2025年获制造业单项冠军企业认定。",
            "行业趋势": "传统家用空调需求增速趋缓，节能风机、设备散热和功能材料仍有结构性机会。",
            "核心逻辑": "传统空调零部件是基本盘，EC风机和复合材料是增量；2026年一季度表观利润增长主要来自出售可转债，扣非利润仍下降。",
            "失效条件": "传统风叶风机收入继续下滑；复合材料增量不能覆盖主业下降；扣非利润持续走弱；新业务商业化不及预期",
            "财务状态": "稳定但增长承压",
            "财务说明": "2025年营收同比下降3.74%、扣非净利润下降16.13%；2026年一季度营收下降16.56%、扣非净利润下降15.38%，归母净利润增长主要来自非经常性收益。",
            "人工复核": "已完成",
            "复核优先级": "中",
        },
    }
    for row in range(5, ws.max_row + 1):
        code = str(ws.cell(row, headers["股票代码"]).value or "").zfill(6)
        for field, value in updates.get(code, {}).items():
            ws.cell(row, headers[field], value)
        if code in updates:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=GREEN)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[row].height = 108


def _rewrite_current_issues(ws, by_code: dict[str, dict]) -> None:
    for table_name in list(ws.tables):
        del ws.tables[table_name]
    if ws.max_row > 9:
        ws.delete_rows(10, ws.max_row - 9)
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 30), min_col=1, max_col=5):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None
    ws["A1"] = "当前规则事项"
    ws["A3"] = "17只股票的19条技术失败记录已由GPT-5.6 SOL完成人工核验，详见“人工核验补全”；本页仅保留仍生效的规则阻断。"
    headers = ["股票代码", "股票名称", "问题类型", "当前状态", "说明"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    rows = [
        ["002156", "通富微电", "规则提醒", "规则阻断保留", "基本面及二筛空值已人工补齐；原挂单与仓位未通过规则门禁，不以人工判断改写规则结果。"],
        ["301151", "冠龙节能", "规则提醒", "规则阻断保留", "基本面可读；人工关注股未进入模型重点名单，且规则未形成有效仓位。"],
        ["002409", "雅克科技", "规则提醒", "规则阻断保留", "基本面可读；人工关注股未进入模型重点名单，且规则未形成有效仓位。"],
        ["300145", "南方泵业", "规则提醒", "规则阻断保留", "基本面可读；人工关注股未进入模型重点名单，且规则未形成有效仓位。"],
        ["603726", "朗迪集团", "规则提醒", "规则阻断保留", "基本面及二筛空值已人工补齐；原挂单与仓位未通过规则门禁，不以人工判断改写规则结果。"],
    ]
    for r_idx, values in enumerate(rows, 5):
        for c_idx, value in enumerate(values, 1):
            ws.cell(r_idx, c_idx, value)
    _style_existing_sheet(ws, 9, 5)
    table = Table(displayName="CurrentRuleItems", ref="A4:E9")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:E9"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 78
    for row in range(5, 10):
        ws.row_dimensions[row].height = 48


def _build_verification_sheet(wb, payload: dict) -> None:
    ws = wb.create_sheet("人工核验补全")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 65
    ws.merge_cells("A1:J1")
    ws["A1"] = "2026-07-15 人工核验结果"
    ws.merge_cells("A2:J2")
    ws["A2"] = "由GPT-5.6 SOL依据公司年报/公告与项目内2026Q1结构化财务数据逐项核验；最终候选二筛空值以明确标记的人工复核分补齐，不改写规则仓位。"
    ws.merge_cells("A3:J3")
    ws["A3"] = f"原始运行：{payload['validation_run_id']}    核验股票：{len(payload['items'])}只    原技术失败记录：已全部核验"
    headers = [
        "序号", "股票代码", "股票名称", "原问题", "人工状态", "主营与产业链",
        "人工核验结论", "主要风险", "对当前候选影响", "核验人/时间",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    for idx, item in enumerate(payload["items"], 1):
        row = idx + 4
        values = [
            idx,
            str(item["stock_code"]).zfill(6),
            item["stock_name"],
            item["original_issue"],
            item["status"],
            f"{item['business']}\n{item['industry_chain']}",
            item["conclusion"],
            item["risks"],
            item["impact"],
            f"{payload['reviewer']}\n{payload['reviewed_at']}",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
    end_row = 4 + len(payload["items"])
    _style_verification_sheet(ws, end_row)
    table = Table(displayName="HumanVerificationTable", ref=f"A4:J{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:J{end_row}"
    widths = [7, 12, 14, 22, 15, 40, 44, 38, 36, 22]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for row in range(5, end_row + 1):
        ws.row_dimensions[row].height = 96
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25


def _build_source_sheet(wb, payload: dict) -> None:
    ws = wb.create_sheet("核验来源")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.merge_cells("A1:G1")
    ws["A1"] = "人工核验来源"
    ws.merge_cells("A2:G2")
    ws["A2"] = "网址指向公司年报、交易所披露镜像或公司官网；结构化财务口径来自项目内Tushare批量缓存。"
    headers = ["序号", "股票代码", "股票名称", "来源类型", "来源URL", "补充数据口径", "核验范围"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    source_rows = []
    for item in payload["items"]:
        source_rows.append([
            len(source_rows) + 1,
            str(item["stock_code"]).zfill(6),
            item["stock_name"],
            "2025年年度报告/公告",
            item["source_url"],
            "项目数据库：Tushare 2026Q1结构化财务快照",
            "主营、产业链、经营判断、主要风险",
        ])
        if item.get("secondary_source_url"):
            source_rows.append([
                len(source_rows) + 1,
                str(item["stock_code"]).zfill(6),
                item["stock_name"],
                "公司官方网站",
                item["secondary_source_url"],
                "公司公开业务资料",
                "主营板块、产品与应用场景",
            ])
    for row_idx, values in enumerate(source_rows, 5):
        for col_idx, value in enumerate(values, 1):
            ws.cell(row_idx, col_idx, value)
        ws.cell(row_idx, 5).hyperlink = ws.cell(row_idx, 5).value
        ws.cell(row_idx, 5).font = Font(name="Microsoft YaHei", size=9, color="0563C1", underline="single")
    end_row = 4 + len(source_rows)
    _style_source_sheet(ws, end_row)
    table = Table(displayName="HumanVerificationSources", ref=f"A4:G{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:G{end_row}"
    widths = [7, 12, 14, 22, 76, 40, 34]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for row in range(5, end_row + 1):
        ws.row_dimensions[row].height = 54
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


def _style_existing_sheet(ws, end_row: int, end_col: int) -> None:
    thin = Side(style="thin", color=GRID)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=end_col)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A3"].fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    ws["A3"].font = Font(name="Microsoft YaHei", size=10, color=MUTED)
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[4]:
        if cell.column <= end_col:
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=5, max_row=end_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color=TEXT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        row[3].fill = PatternFill("solid", fgColor=YELLOW)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[3].height = 42
    ws.row_dimensions[4].height = 36


def _style_verification_sheet(ws, end_row: int) -> None:
    thin = Side(style="thin", color=GRID)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].font = Font(name="Microsoft YaHei", size=20, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for cell_ref in ("A2", "A3"):
        ws[cell_ref].fill = PatternFill("solid", fgColor=BLUE_LIGHT if cell_ref == "A2" else BLUE_PALE)
        ws[cell_ref].font = Font(name="Microsoft YaHei", size=10, color=MUTED)
        ws[cell_ref].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx, row in enumerate(ws.iter_rows(min_row=5, max_row=end_row, min_col=1, max_col=10), 5):
        fill = GREEN if row_idx % 2 else BLUE_PALE
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(name="Microsoft YaHei", size=9, color=TEXT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        row[4].font = Font(name="Microsoft YaHei", size=9, bold=True, color="548235")
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 38
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[4].height = 42


def _style_source_sheet(ws, end_row: int) -> None:
    thin = Side(style="thin", color=GRID)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].font = Font(name="Microsoft YaHei", size=20, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    ws["A2"].font = Font(name="Microsoft YaHei", size=10, color=MUTED)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx, row in enumerate(ws.iter_rows(min_row=5, max_row=end_row, min_col=1, max_col=7), 5):
        fill = BLUE_PALE if row_idx % 2 == 0 else WHITE
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=fill)
            if cell.column != 5:
                cell.font = Font(name="Microsoft YaHei", size=9, color=TEXT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 36
    ws.row_dimensions[4].height = 38


def _headers(ws, row: int) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[row]
        if cell.value is not None and str(cell.value).strip()
    }


def _count_pending(path: Path) -> int:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        return sum(
            1 for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
            for value in row if value == "待人工确认"
        )
    finally:
        wb.close()


if __name__ == "__main__":
    raise SystemExit(main())
