from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from database.models.stock import StockMaster
from database.session import get_session


DETAIL_SHEET = "全A成交量明细"
SUMMARY_SHEET = "市场汇总"
NOTES_SHEET = "口径说明"


def build_volume_workbook(first_date: str, second_date: str, output: Path) -> dict:
    rows_by_date = {
        day: _read_daily_cache(day)
        for day in (first_date, second_date)
    }
    records_by_date = {
        day: {str(row["ts_code"]): row for row in rows if row.get("ts_code")}
        for day, rows in rows_by_date.items()
    }
    session = get_session()
    try:
        masters = {row.code: row for row in session.scalars(select(StockMaster))}
    finally:
        session.close()

    first_records = records_by_date[first_date]
    second_records = records_by_date[second_date]
    codes = sorted(
        set(first_records) | set(second_records),
        key=lambda code: (
            -_number((second_records.get(code) or {}).get("vol")),
            -_number((first_records.get(code) or {}).get("vol")),
            code,
        ),
    )

    workbook = Workbook()
    detail = workbook.active
    detail.title = DETAIL_SHEET
    summary = workbook.create_sheet(SUMMARY_SHEET)
    notes = workbook.create_sheet(NOTES_SHEET)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    _build_detail(detail, codes, first_records, second_records, masters, first_date, second_date)
    _build_summary(summary, len(codes), first_date, second_date)
    _build_notes(notes, first_date, second_date, len(codes))
    for sheet in workbook.worksheets:
        _configure_print(sheet)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    verification = _verify(output, len(codes))
    verification["source_reconciliation"] = {
        day: _source_summary(records)
        for day, records in records_by_date.items()
    }
    return verification


def _read_daily_cache(trade_date: str) -> list[dict]:
    path = ROOT_DIR / "data" / "cache" / "tushare" / "trade_date" / "daily" / f"{trade_date}.json"
    rows = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(rows, list) or not rows:
        raise ValueError(f"DAILY_CACHE_EMPTY:{trade_date}")
    actual_dates = {str(row.get("trade_date") or "") for row in rows}
    if actual_dates != {trade_date}:
        raise ValueError(f"DAILY_CACHE_DATE_MISMATCH:{trade_date}")
    return rows


def _build_detail(sheet, codes, first, second, masters, first_date: str, second_date: str) -> None:
    first_label = _date_label(first_date)
    second_label = _date_label(second_date)
    headers = [
        "序号", "股票代码", "股票名称", "交易所", "行业",
        f"{first_label} 成交量(手)", f"{first_label} 成交量权重",
        f"{first_label} 成交额(千元)", f"{first_label} 成交额权重", f"{first_label} 涨跌幅",
        f"{second_label} 成交量(手)", f"{second_label} 成交量权重",
        f"{second_label} 成交额(千元)", f"{second_label} 成交额权重", f"{second_label} 涨跌幅",
        "成交量增减(手)", "成交量变化率", "成交量权重变化",
        "成交额增减(千元)", "成交额变化率", "成交额权重变化",
        f"{first_label}量权涨跌贡献", f"{second_label}量权涨跌贡献", "量权贡献变化", "数据状态",
    ]
    styles = _styles()
    last_row = 3 + len(codes)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet["A1"] = f"全A股票成交量及权重对比（{_iso_date(first_date)} / {_iso_date(second_date)}）"
    _style_title(sheet["A1"], styles)
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet["A2"] = "权重口径：个股成交量 ÷ 当日全A成交量；成交额权重同理。成交量单位沿用 Tushare daily.vol（手）。"
    sheet["A2"].font = Font(name="Microsoft YaHei", size=10, color="475569")
    sheet["A2"].fill = PatternFill("solid", fgColor="F4F7F9")
    sheet["A2"].alignment = styles["left"]
    sheet.row_dimensions[2].height = 28

    for column, header in enumerate(headers, 1):
        cell = sheet.cell(3, column, header)
        _style_header(cell, styles)
    sheet.row_dimensions[3].height = 42

    for row_number, code in enumerate(codes, 4):
        first_row = first.get(code)
        second_row = second.get(code)
        master = masters.get(code)
        values = [
            row_number - 3,
            code,
            master.name if master else code,
            master.market if master and master.market else code.split(".")[-1],
            master.industry if master else "",
            _number(first_row.get("vol")) if first_row else None,
            None,
            _number(first_row.get("amount")) if first_row else None,
            None,
            _number(first_row.get("pct_chg")) / 100 if first_row else None,
            _number(second_row.get("vol")) if second_row else None,
            None,
            _number(second_row.get("amount")) if second_row else None,
            None,
            _number(second_row.get("pct_chg")) / 100 if second_row else None,
            None, None, None, None, None, None, None, None, None,
            "两日均有" if first_row and second_row else (f"仅{first_label}" if first_row else f"仅{second_label}"),
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)

        sheet.cell(row_number, 7, f"=IFERROR(F{row_number}/'{SUMMARY_SHEET}'!$B$4,0)")
        sheet.cell(row_number, 9, f"=IFERROR(H{row_number}/'{SUMMARY_SHEET}'!$B$5,0)")
        sheet.cell(row_number, 12, f"=IFERROR(K{row_number}/'{SUMMARY_SHEET}'!$C$4,0)")
        sheet.cell(row_number, 14, f"=IFERROR(M{row_number}/'{SUMMARY_SHEET}'!$C$5,0)")
        sheet.cell(row_number, 16, f'=IF(OR(F{row_number}="",K{row_number}=""),"",K{row_number}-F{row_number})')
        sheet.cell(row_number, 17, f'=IFERROR(K{row_number}/F{row_number}-1,"")')
        sheet.cell(row_number, 18, f"=L{row_number}-G{row_number}")
        sheet.cell(row_number, 19, f'=IF(OR(H{row_number}="",M{row_number}=""),"",M{row_number}-H{row_number})')
        sheet.cell(row_number, 20, f'=IFERROR(M{row_number}/H{row_number}-1,"")')
        sheet.cell(row_number, 21, f"=N{row_number}-I{row_number}")
        sheet.cell(row_number, 22, f'=IF(J{row_number}="","",G{row_number}*J{row_number})')
        sheet.cell(row_number, 23, f'=IF(O{row_number}="","",L{row_number}*O{row_number})')
        sheet.cell(row_number, 24, f"=W{row_number}-V{row_number}")

        for cell in sheet[row_number]:
            cell.alignment = styles["center"]
            cell.border = styles["border"]
            cell.font = styles["body_font"]
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
        sheet.cell(row_number, 2).number_format = "@"
        for column in (6, 8, 11, 13, 16, 19):
            sheet.cell(row_number, column).number_format = "#,##0.00"
        for column in (7, 9, 10, 12, 14, 15, 17, 18, 20, 21, 22, 23, 24):
            sheet.cell(row_number, column).number_format = "0.0000%"
        sheet.row_dimensions[row_number].height = 22

    sheet.freeze_panes = "F4"
    sheet.auto_filter.ref = f"A3:Y{last_row}"
    sheet.sheet_view.showGridLines = False
    widths = [8, 14, 14, 10, 18, 18, 17, 20, 17, 15, 18, 17, 20, 17, 15, 18, 15, 17, 20, 15, 17, 19, 19, 17, 13]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(3, column).column_letter].width = width
    for column in ("J", "O", "Q", "R", "T", "U", "V", "W", "X"):
        sheet.conditional_formatting.add(
            f"{column}4:{column}{last_row}",
            CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="C00000")),
        )
        sheet.conditional_formatting.add(
            f"{column}4:{column}{last_row}",
            CellIsRule(operator="lessThan", formula=["0"], font=Font(color="008000")),
        )
    sheet.conditional_formatting.add(
        f"Y4:Y{last_row}",
        FormulaRule(formula=['$Y4<>"两日均有"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )


def _build_summary(sheet, detail_count: int, first_date: str, second_date: str) -> None:
    styles = _styles()
    last_row = 3 + detail_count
    sheet.merge_cells("A1:E1")
    sheet["A1"] = "全A成交量市场汇总"
    _style_title(sheet["A1"], styles)
    for column, value in enumerate(["指标", _iso_date(first_date), _iso_date(second_date), "增减", "变化率"], 1):
        _style_header(sheet.cell(2, column, value), styles)
    metrics = [
        "股票记录数", "全A成交量(手)", "全A成交额(千元)",
        "成交量加权平均涨跌幅", "成交额加权平均涨跌幅", "平均涨跌幅", "中位涨跌幅",
        "上涨家数", "下跌家数", "平盘家数",
    ]
    for row, metric in enumerate(metrics, 3):
        sheet.cell(row, 1, metric)
    formulas = {
        3: (f"=COUNT('{DETAIL_SHEET}'!F4:F{last_row})", f"=COUNT('{DETAIL_SHEET}'!K4:K{last_row})"),
        4: (f"=SUM('{DETAIL_SHEET}'!F4:F{last_row})", f"=SUM('{DETAIL_SHEET}'!K4:K{last_row})"),
        5: (f"=SUM('{DETAIL_SHEET}'!H4:H{last_row})", f"=SUM('{DETAIL_SHEET}'!M4:M{last_row})"),
        6: (f"=SUM('{DETAIL_SHEET}'!V4:V{last_row})", f"=SUM('{DETAIL_SHEET}'!W4:W{last_row})"),
        7: (f"=SUMPRODUCT('{DETAIL_SHEET}'!I4:I{last_row},'{DETAIL_SHEET}'!J4:J{last_row})", f"=SUMPRODUCT('{DETAIL_SHEET}'!N4:N{last_row},'{DETAIL_SHEET}'!O4:O{last_row})"),
        8: (f"=AVERAGE('{DETAIL_SHEET}'!J4:J{last_row})", f"=AVERAGE('{DETAIL_SHEET}'!O4:O{last_row})"),
        9: (f"=MEDIAN('{DETAIL_SHEET}'!J4:J{last_row})", f"=MEDIAN('{DETAIL_SHEET}'!O4:O{last_row})"),
        10: (f'=COUNTIF(\'{DETAIL_SHEET}\'!J4:J{last_row},">0")', f'=COUNTIF(\'{DETAIL_SHEET}\'!O4:O{last_row},">0")'),
        11: (f'=COUNTIF(\'{DETAIL_SHEET}\'!J4:J{last_row},"<0")', f'=COUNTIF(\'{DETAIL_SHEET}\'!O4:O{last_row},"<0")'),
        12: (f'=COUNTIF(\'{DETAIL_SHEET}\'!J4:J{last_row},"=0")', f'=COUNTIF(\'{DETAIL_SHEET}\'!O4:O{last_row},"=0")'),
    }
    for row, (first_formula, second_formula) in formulas.items():
        sheet.cell(row, 2, first_formula)
        sheet.cell(row, 3, second_formula)
        sheet.cell(row, 4, f"=C{row}-B{row}")
        sheet.cell(row, 5, "" if row in {6, 7, 8, 9} else f'=IFERROR(D{row}/B{row},"")')
    for row in sheet.iter_rows(min_row=2, max_row=12, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = styles["center"]
            cell.border = styles["border"]
            if cell.row >= 3:
                cell.font = styles["body_font"]
            if cell.row >= 3:
                cell.fill = PatternFill("solid", fgColor="FFFFFF")
    for row in (6, 7, 8, 9):
        for column in range(2, 6):
            sheet.cell(row, column).number_format = "0.0000%"
    for row in (4, 5):
        for column in range(2, 6):
            sheet.cell(row, column).number_format = "#,##0.00"
    for row in (3, 10, 11, 12):
        for column in range(2, 5):
            sheet.cell(row, column).number_format = "#,##0"
    for row in (3, 4, 5, 10, 11, 12):
        sheet.cell(row, 5).number_format = "0.00%"
    sheet.column_dimensions["A"].width = 26
    for column in "BCDE":
        sheet.column_dimensions[column].width = 20
    sheet.freeze_panes = "B3"
    sheet.auto_filter.ref = "A2:E12"
    sheet.sheet_view.showGridLines = False


def _build_notes(sheet, first_date: str, second_date: str, count: int) -> None:
    styles = _styles()
    notes = [
        ("项目", "说明"),
        ("数据源", "Tushare daily 批量交易日缓存；未进行逐股 API 调用。"),
        ("日期", f"{_iso_date(first_date)}、{_iso_date(second_date)}。"),
        ("成交量单位", "Tushare daily.vol，单位为手。"),
        ("成交额单位", "Tushare daily.amount，单位为千元。"),
        ("成交量权重", "个股当日成交量 / 当日全A成交量。"),
        ("成交额权重", "个股当日成交额 / 当日全A成交额。"),
        ("量权涨跌贡献", "成交量权重 × 个股当日涨跌幅；汇总后为成交量加权平均涨跌幅。"),
        ("缺失处理", f"两日并集共 {count:,} 只；单日无记录时保留股票并标明数据状态。"),
        ("排序", f"按 {_iso_date(second_date)} 成交量降序，其次按 {_iso_date(first_date)} 成交量降序。"),
        ("安全边界", "数据表不包含 Token、API Key 或认证响应；iFinD 调用数为 0。"),
    ]
    for row_number, row in enumerate(notes, 1):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(row_number, column, value)
            cell.alignment = styles["center"]
            cell.border = styles["border"]
            cell.font = styles["body_font"]
            if row_number > 1:
                cell.fill = PatternFill("solid", fgColor="FFFFFF")
    for cell in sheet[1]:
        _style_header(cell, styles)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 90
    sheet.sheet_view.showGridLines = False


def _styles() -> dict:
    thin = Side(style="thin", color="CBD5E1")
    return {
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "left": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "body_font": Font(name="Microsoft YaHei", size=10),
    }


def _style_title(cell, styles: dict) -> None:
    cell.font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17365D")
    cell.alignment = styles["left"]


def _style_header(cell, styles: dict) -> None:
    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17365D")
    cell.alignment = styles["center"]
    cell.border = styles["border"]


def _configure_print(sheet) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_orientation = "landscape"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5


def _verify(output: Path, detail_count: int) -> dict:
    workbook = load_workbook(output, read_only=False, data_only=False)
    if workbook.sheetnames != [DETAIL_SHEET, SUMMARY_SHEET, NOTES_SHEET]:
        raise ValueError("VOLUME_WORKBOOK_SHEET_MISMATCH")
    detail = workbook[DETAIL_SHEET]
    expected_last_row = 3 + detail_count
    if detail.max_row != expected_last_row or detail.max_column != 25:
        raise ValueError("VOLUME_WORKBOOK_DIMENSION_MISMATCH")
    if detail.freeze_panes != "F4" or detail.auto_filter.ref != f"A3:Y{expected_last_row}":
        raise ValueError("VOLUME_WORKBOOK_NAVIGATION_MISMATCH")
    table_start_rows = {DETAIL_SHEET: 3, SUMMARY_SHEET: 2, NOTES_SHEET: 1}
    for sheet in workbook.worksheets:
        if sheet.tables:
            raise ValueError(f"VOLUME_WORKBOOK_TABLE_OBJECT_FOUND:{sheet.title}")
        for row_number in range(1, sheet.max_row + 1):
            if not any(
                not isinstance(cell, MergedCell) and cell.value is not None
                for cell in sheet[row_number]
            ):
                raise ValueError(f"VOLUME_WORKBOOK_BLANK_ROW:{sheet.title}:{row_number}")
        for row in sheet.iter_rows(min_row=table_start_rows[sheet.title]):
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                alignment = cell.alignment
                if alignment.horizontal != "center" or alignment.vertical != "center":
                    raise ValueError(
                        f"VOLUME_WORKBOOK_ALIGNMENT_MISMATCH:{sheet.title}:{cell.coordinate}"
                    )
    if detail["G4"].data_type != "f" or detail["W4"].data_type != "f":
        raise ValueError("VOLUME_WORKBOOK_FORMULA_MISSING")
    for row in range(1, min(expected_last_row, 100) + 1):
        for column in range(1, 26):
            value = str(detail.cell(row, column).value or "")
            if any(error in value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                raise ValueError("VOLUME_WORKBOOK_FORMULA_ERROR")
    return {
        "output": str(output),
        "size_bytes": output.stat().st_size,
        "sheets": workbook.sheetnames,
        "detail_rows": detail_count,
        "center_alignment_verified": True,
        "formula_cells_verified": True,
    }


def _source_summary(records: dict[str, dict]) -> dict:
    total_volume = sum(_number(row.get("vol")) for row in records.values())
    total_amount = sum(_number(row.get("amount")) for row in records.values())
    weighted_return = (
        sum(_number(row.get("vol")) * _number(row.get("pct_chg")) / 100 for row in records.values())
        / total_volume
        if total_volume
        else 0
    )
    return {
        "row_count": len(records),
        "total_volume_lots": total_volume,
        "total_amount_thousand_yuan": total_amount,
        "volume_weighted_return": weighted_return,
    }


def _number(value) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _date_label(value: str) -> str:
    return f"{int(value[4:6])}月{int(value[6:8])}日"


def _iso_date(value: str) -> str:
    return f"{value[:4]}-{value[4:6]}-{value[6:8]}"


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a centered full-A volume and weight workbook from Tushare trade-date caches.")
    parser.add_argument("--first-date", default="20260713")
    parser.add_argument("--second-date", default="20260714")
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT_DIR / "outputs" / "2026-07-14" / "全A成交量及权重_2026-07-13_2026-07-14.xlsx",
    )
    args = parser.parse_args()
    result = build_volume_workbook(args.first_date, args.second_date, args.output.resolve())
    print(json.dumps(result, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
