from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from openpyxl import load_workbook
from sqlalchemy import select


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from database.models.quant_run import QuantRun
from database.models.validation import ModelValidationRun, ProResumeRun
from database.session import get_session, init_db
from scripts.run_daily_routine import STAGES, _assert_advisory_only, _run_close


SHANGHAI = ZoneInfo("Asia/Shanghai")


def main() -> int:
    parser = argparse.ArgumentParser(description="One-click advisory-only daily close pipeline.")
    parser.add_argument("--trade-date", type=date.fromisoformat)
    parser.add_argument("--manual-pool", type=Path)
    parser.add_argument("--start-stage", choices=STAGES, default="data")
    parser.add_argument("--refresh-data", action="store_true")
    parser.add_argument("--market-supplement", type=Path)
    args = parser.parse_args()

    load_dotenv(ROOT / ".env", override=False)
    _assert_advisory_only()
    init_db()
    trade_date = args.trade_date or datetime.now(SHANGHAI).date()
    routine_args = SimpleNamespace(
        trade_date=trade_date,
        manual_pool=args.manual_pool,
        start_stage=args.start_stage,
        refresh_data=args.refresh_data,
        confirm_llm_budget=True,
        market_supplement=args.market_supplement,
    )
    try:
        report = _run_close(routine_args)
        completion = _completion_check(trade_date)
        payload = {
            "status": "SUCCESS" if completion["ready"] else "INCOMPLETE",
            "trade_date": trade_date.isoformat(),
            "routine": report,
            "completion": completion,
            "real_trading_enabled": False,
        }
    except Exception as exc:
        payload = {
            "status": "FAILED",
            "trade_date": trade_date.isoformat(),
            "error_category": type(exc).__name__,
            "error": str(exc),
            "resume_hint": "Run the same command again; completed checkpoints will be reused.",
            "real_trading_enabled": False,
        }
    audit = _write_audit(trade_date, payload)
    payload["audit"] = str(audit)
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
    return 0 if payload["status"] == "SUCCESS" else 2


def _completion_check(trade_date: date) -> dict:
    session = get_session()
    try:
        quant = session.scalar(select(QuantRun).where(
            QuantRun.base_market_trade_date == trade_date,
            QuantRun.status == "COMPLETED",
            QuantRun.actionable.is_(True),
        ).order_by(QuantRun.created_at.desc()))
        flash = session.scalar(select(ModelValidationRun).where(
            ModelValidationRun.base_market_trade_date == trade_date,
            ModelValidationRun.status.in_(["SUCCESS", "PARTIAL_SUCCESS"]),
        ).order_by(ModelValidationRun.created_at.desc()))
        pro = session.scalar(select(ProResumeRun).where(
            ProResumeRun.base_trade_date == trade_date,
            ProResumeRun.status == "COMPLETED",
        ).order_by(ProResumeRun.created_at.desc()))
    finally:
        session.close()

    root = ROOT / "outputs" / trade_date.isoformat()
    daily = root / f"智能交易助手_{trade_date.isoformat()}.xlsx"
    market = root / "复盘" / f"大盘复盘_{trade_date.isoformat()}.xlsx"
    today_recommendation_review = root / "复盘" / f"今日推荐复盘_截至{trade_date.isoformat()}.xlsx"
    key_candidates_review = root / "复盘" / f"重点候选复盘_截至{trade_date.isoformat()}.xlsx"
    issue_count = _workbook_issue_count(daily) if daily.is_file() else None
    quant_coverage = (
        int(quant.scored_count or 0) / int(quant.filtered_count or 1)
        if quant is not None and int(quant.filtered_count or 0) > 0
        else 0.0
    )
    checks = {
        "quant_coverage_passed": quant_coverage >= 0.95,
        "flash_completed": flash is not None,
        "pro_completed": pro is not None,
        "daily_workbook": daily.is_file(),
        "market_workbook": market.is_file(),
        "today_recommendation_review_workbook": today_recommendation_review.is_file(),
        "key_candidates_review_workbook": key_candidates_review.is_file(),
        "unresolved_issue_count": issue_count,
    }
    ready = all(value is True for key, value in checks.items() if key != "unresolved_issue_count") and issue_count == 0
    return {
        "ready": ready,
        "checks": checks,
        "quant_coverage_ratio": round(quant_coverage, 6),
        "quant_run_id": quant.run_id if quant else None,
        "flash_run_id": flash.run_id if flash else None,
        "pro_run_id": pro.run_id if pro else None,
    }


def _workbook_issue_count(path: Path) -> int:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[-1]
        return sum(
            1
            for row in worksheet.iter_rows(min_row=5, values_only=True)
            if re.fullmatch(r"\d{6}(?:\.(?:SH|SZ|BJ))?", str(row[0] or ""))
        )
    finally:
        workbook.close()


def _write_audit(trade_date: date, payload: dict) -> Path:
    path = ROOT / "outputs" / trade_date.isoformat() / "审计" / "一键收盘运行报告.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return path


if __name__ == "__main__":
    raise SystemExit(main())
