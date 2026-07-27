from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import uuid
import zipfile
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment
from sqlalchemy import select


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env", override=False)

from database.session import get_session, init_db
from database.models.quant_run import QuantRun
from database.models.validation import ProResumeRun
from backend.core.config_manager import ConfigManager
from scripts.run_guarded_llm_excel_validation import _secret_scan
from scripts.run_trader_demo_excel import _serialize_readback
from trader_demo.service import TraderDemoService
from reporting.workbook_standard import validate_trading_assistant_workbook
from reporting.workbook_style import WorkbookStyleService
from reporting.immutable_workbook import (
    safe_artifact_token,
    versioned_workbook_path,
    workbook_content_style_hashes,
)


DECISIONS = {
    "ADVANCE": "优先复核",
    "HOLD": "继续观察",
    "WATCH_ONLY": "普通观察",
    "REJECT": "暂不考虑",
    "BLOCK": "规则阻断",
}
SOURCES = {"LLM_TOP20": "模型筛选", "MANUAL": "人工关注", "BOTH": "共同入选"}
PRIORITIES = {"HIGH": "高", "MEDIUM": "中", "LOW": "低", "REVIEW_ONLY": "仅复核"}
FINANCIAL = {
    "HEALTHY": "稳健",
    "STABLE": "稳定",
    "PRESSURED": "承压",
    "HIGH_RISK": "高风险",
    "BLOCKED": "高风险",
    "NORMAL": "正常",
    "UNKNOWN": "信息不足",
}
CHAIN_POSITIONS = {
    "UPSTREAM": "上游", "MIDSTREAM": "中游", "DOWNSTREAM": "下游",
    "SERVICE_PLATFORM": "服务平台", "MULTI_SEGMENT": "多环节", "UNKNOWN": "信息不足",
}
ORDER_STATUS = {"DRAFT": "已生成", "BLOCKED": "已阻断", "NEEDS_REVIEW": "需复核"}
POSITION_STATUS = {"NON_ACTIONABLE": "模拟建议", "BLOCKED": "已阻断", "DRAFT": "已生成"}


def main() -> int:
    parser = argparse.ArgumentParser(description="生成按交易日归档的中文人工阅读版工作簿")
    parser.add_argument("--date", default=date.today().isoformat())
    parser.add_argument("--validation-run", default="")
    parser.add_argument("--use-existing", action="store_true")
    parser.add_argument("--replace-existing", action="store_true")
    parser.add_argument(
        "--repair-existing",
        action="store_true",
        help="Archive and atomically replace the same run artifact after an audited correction.",
    )
    args = parser.parse_args()
    if args.use_existing and args.repair_existing:
        raise ValueError("USE_EXISTING_AND_REPAIR_EXISTING_CONFLICT")
    trade_date = date.fromisoformat(args.date).isoformat()
    _assert_safe_runtime()

    validation_run_id = args.validation_run.strip()
    if not validation_run_id:
        checkpoint_path = _find_checkpoint(trade_date)
        checkpoint = json.loads(checkpoint_path.read_text(encoding="utf-8"))
        validation_run_id = str(checkpoint.get("flash_v4_run_id") or "")
    if not validation_run_id:
        raise ValueError("FLASH_VALIDATION_RUN_ID_REQUIRED")

    init_db()
    session = get_session()
    try:
        raw = _serialize_readback(TraderDemoService(session).readback(validation_run_id))
        minimum_recommendation_score = float(
            ConfigManager(session=session)
            .get_effective_config()["values"]
            .get("selection_performance.minimum_recommendation_score", 60)
        )
        pro_run = session.scalar(
            select(ProResumeRun)
            .where(
                ProResumeRun.flash_validation_run_id == validation_run_id,
                ProResumeRun.status == "COMPLETED",
            )
            .order_by(ProResumeRun.created_at.desc())
        )
        quant_run = (
            session.scalar(
                select(QuantRun).where(QuantRun.run_id == pro_run.quant_run_id)
            )
            if pro_run
            else None
        )
    finally:
        session.close()
    verifications = _load_verification_overlays(trade_date, validation_run_id)
    payload = build_human_payload(
        raw,
        trade_date,
        verifications=verifications,
        minimum_recommendation_score=minimum_recommendation_score,
    )
    if not _secret_scan(json.dumps(payload, ensure_ascii=False, default=str)):
        raise ValueError("HUMAN_OUTPUT_SECRET_SCAN_FAILED")

    daily_root = (ROOT / "outputs" / trade_date).resolve()
    output_root = (ROOT / "outputs").resolve()
    if output_root not in daily_root.parents:
        raise ValueError("DAILY_OUTPUT_PATH_OUTSIDE_OUTPUTS")
    daily_root.mkdir(parents=True, exist_ok=True)
    artifact_run_id = (
        str(pro_run.pipeline_run_id)
        if pro_run and pro_run.pipeline_run_id
        else validation_run_id
    )
    factor_version = (
        str(quant_run.factor_version or "LEGACY_UNCALIBRATED")
        if quant_run
        else "LEGACY_UNCALIBRATED"
    )
    output_path = versioned_workbook_path(
        daily_root,
        stem="智能交易助手",
        trade_date=trade_date,
        run_id=artifact_run_id,
        factor_version=factor_version,
    )
    if args.replace_existing:
        raise ValueError("IMMUTABLE_WORKBOOK_REPLACEMENT_FORBIDDEN")
    preview_dir = daily_root / "预览" / safe_artifact_token(artifact_run_id)
    reused_existing = output_path.exists()
    if args.use_existing and not reused_existing:
        raise FileNotFoundError(f"OUTPUT_NOT_FOUND:{output_path.name}")
    repaired_existing = bool(reused_existing and args.repair_existing)
    backup_path: Path | None = None
    if repaired_existing:
        history = daily_root / "历史版本"
        history.mkdir(parents=True, exist_ok=True)
        old_hash = workbook_content_style_hashes(output_path)["file_sha256"][:8]
        backup_path = history / f"{output_path.stem}_修复前_{old_hash}.xlsx"
        if not backup_path.exists():
            shutil.copy2(output_path, backup_path)
        candidate = output_path.with_name(
            f".{output_path.stem}_{uuid.uuid4().hex[:8]}_repair.xlsx"
        )
        try:
            _build_workbook(payload, candidate, preview_dir)
            _polish_workbook(candidate)
            validation = _validate_workbook(candidate)
            os.replace(candidate, output_path)
        finally:
            candidate.unlink(missing_ok=True)
            Path(f"{candidate}.inspect.ndjson").unlink(missing_ok=True)
    elif reused_existing:
        validation = _validate_workbook(output_path)
    else:
        candidate = output_path.with_name(f".{output_path.stem}_{uuid.uuid4().hex[:8]}_candidate.xlsx")
        try:
            _build_workbook(payload, candidate, preview_dir)
            _polish_workbook(candidate)
            validation = _validate_workbook(candidate)
            candidate.replace(output_path)
        finally:
            candidate.unlink(missing_ok=True)
            Path(f"{candidate}.inspect.ndjson").unlink(missing_ok=True)
    workbook_hashes = workbook_content_style_hashes(output_path)
    workbook_hash = workbook_hashes["file_sha256"]

    audit_dir = daily_root / "审计"
    audit_dir.mkdir(parents=True, exist_ok=True)
    manifest = {
        "run_id": artifact_run_id,
        "factor_version": factor_version,
        "input_hash": quant_run.request_hash if quant_run else None,
        "output_path": str(output_path),
        "content_hash": workbook_hashes["content_hash"],
        "style_hash": workbook_hashes["style_hash"],
        "file_sha256": workbook_hash,
        "reused_existing": reused_existing,
        "repaired_existing": repaired_existing,
        "repair_backup": str(backup_path) if backup_path else None,
        "交易日": trade_date,
        "生成状态": "完成",
        "数据来源运行": validation_run_id,
        "最终复核运行": pro_run.run_id if pro_run else None,
        "流水线运行": pro_run.pipeline_run_id if pro_run else None,
        "候选集摘要": pro_run.candidate_set_hash if pro_run else None,
        "工作簿": str(output_path),
        "工作簿摘要": workbook_hash,
        "候选数量": len(payload["candidates"]),
        "今日推荐数量": len(payload["recommendations"]),
        "今日推荐最低最终复核分": payload["minimum_recommendation_score"],
        "非零仓位数量": sum(int(row["建议股数"] or 0) > 0 for row in payload["orders"]),
        "当前问题数量": len(payload["issues"]),
        "检查结果": validation,
        "外部模型调用": 0,
        "外部行情调用": 0,
    }
    manifest_path = (
        audit_dir
        / f"人工阅读版_生成记录_{safe_artifact_token(artifact_run_id)}.json"
    )
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


def build_human_payload(
    raw: dict[str, Any],
    trade_date: str,
    *,
    verifications: dict[str, dict[str, Any]] | None = None,
    minimum_recommendation_score: float = 60,
) -> dict[str, Any]:
    verifications = verifications or {}
    llm_by_code = {row["stock_code"]: row for row in raw["llm_rows"]}
    fundamental_by_code = {row["stock_code"]: row for row in raw["fundamental_rows"]}
    candidates = []
    orders = []
    fundamentals = []
    for order in raw["order_rows"]:
        code = order["stock_code"]
        llm = llm_by_code.get(code, {})
        fundamental = fundamental_by_code.get(code, {})
        common = {
            "股票代码": code,
            "股票名称": order["stock_name"],
            "入选来源": SOURCES.get(order["selection_source"], _human_text(order["selection_source"])),
            "量化排名": order["quant_rank"],
            "量化得分": order["quant_score"],
            "二筛得分": order["llm_score"],
            "二筛结论": DECISIONS.get(order["llm_decision"], _human_text(order["llm_decision"])),
            "深度复核分": order["pro_score"],
            "深度复核排名": order["pro_rank"],
            "复核优先级": PRIORITIES.get(order["pro_priority"], _human_text(order["pro_priority"])),
        }
        candidates.append({
            **common,
            "一级行业": _human_text(order["level_one_sector"]),
            "产业链": _human_text(fundamental.get("industry_chain")),
            "财务状态": FINANCIAL.get(order["financial_status"], _human_text(order["financial_status"])),
            "建议仓位": order["position_percent"],
            "建议股数": order["quantity"],
            "参考价": order["recommended_price"],
            "止损价": order["stop_loss_price"],
            "第二目标价": order["take_profit_2"],
            "风险收益比": order["active_risk_reward"],
            "核心逻辑": _human_text(
                fundamental.get("pro_summary") or fundamental.get("investment_logic")
            ),
            "主要风险": _human_risk_summary(fundamental, llm),
            "当前状态": ORDER_STATUS.get(order["order_status"], _human_text(order["order_status"])),
        })
        orders.append({
            **common,
            "保守价": order["conservative_price"], "均衡价": order["balanced_price"],
            "积极价": order["aggressive_price"], "参考价": order["recommended_price"],
            "最高接受价": order["max_acceptable_price"], "止损价": order["stop_loss_price"],
            "第一目标价": order["take_profit_1"], "第二目标价": order["take_profit_2"],
            "第一目标风险收益比": order["risk_reward_to_tp1"],
            "第二目标风险收益比": order["risk_reward_to_tp2"],
            "当前风险收益比": order["active_risk_reward"],
            "建议仓位": order["position_percent"], "建议资金": order["capital_amount"],
            "建议股数": order["quantity"], "预计最大损失": order["max_loss"],
            "挂单状态": ORDER_STATUS.get(order["order_status"], _human_text(order["order_status"])),
            "仓位状态": POSITION_STATUS.get(order["position_status"], _human_text(order["position_status"])),
            "说明": _human_warning(order["order_warning"], order["position_warning"]),
        })
        fundamentals.append({
            **common,
            "一级行业": _human_text(fundamental.get("level_one_sector")),
            "产业链": _human_text(fundamental.get("industry_chain")),
            "链条位置": CHAIN_POSITIONS.get(fundamental.get("chain_position"), _human_text(fundamental.get("chain_position"))),
            "主营业务": _human_text(fundamental.get("main_business")),
            "核心产品": _human_text(fundamental.get("core_products")),
            "概念标签": _compact_concept_tags(fundamental.get("concept_tags")),
            "结构性方向": _human_text(fundamental.get("structural_theme_fit")),
            "潜在优势": _human_text(fundamental.get("competitive_advantage")),
            "行业趋势": _human_text(fundamental.get("industry_trend")),
            "核心逻辑": _human_text(fundamental.get("investment_logic")),
            "失效条件": _human_text(fundamental.get("logic_invalidation")),
            "财务状态": FINANCIAL.get(fundamental.get("financial_status"), _human_text(fundamental.get("financial_status"))),
            "财务说明": _human_text(fundamental.get("financial_status_reason")),
            "人工复核": _review_status(fundamental.get("manual_review")),
            "核验来源": "",
        })
        verification = verifications.get(str(code)[:6])
        if verification:
            _apply_verification_overlay(candidates[-1], fundamentals[-1], verification)
    candidates.sort(key=lambda row: (row["深度复核排名"] is None, row["深度复核排名"] or 9999, row["股票代码"]))
    orders.sort(key=lambda row: (row["深度复核排名"] is None, row["深度复核排名"] or 9999, row["股票代码"]))
    fundamentals.sort(key=lambda row: (row["深度复核排名"] is None, row["深度复核排名"] or 9999, row["股票代码"]))
    recommendations = [
        row for row in candidates
        if row["深度复核分"] is not None
        and float(row["深度复核分"]) >= float(minimum_recommendation_score)
    ]

    quant_top100 = [{
        "量化排名": row["rank"], "股票代码": row["stock_code"], "股票名称": row["stock_name"],
        "一级行业": _human_text(row["level_one_sector"]), "量化总分": row["total_score"],
        "技术得分": row["technical_score"], "资金得分": row["capital_score"],
        "情绪得分": row["emotion_score"], "动量得分": row["momentum_score"], "风险得分": row["risk_score"],
        "进入二筛": _yes_no(row["llm_evaluated"]), "进入重点候选": _yes_no(row["selection_source"]),
    } for row in raw["quant_rows"] if int(row["rank"]) <= 100]
    verified_codes = set(verifications)
    issues = [
        _human_issue(row, llm_by_code, order=False)
        for row in raw.get("errors") or []
        if str(row.get("stock_code") or "")[:6] not in verified_codes
    ]
    issues.extend(
        _human_issue(row, llm_by_code, order=True)
        for row in raw.get("warnings") or []
        if _is_reportable_order_warning(row, verified_codes)
    )
    resolved_current_issues = sum(
        str(row.get("stock_code") or "")[:6] in verified_codes
        for row in raw.get("errors") or []
    )
    run = raw["run"]
    return {
        "title": f"{trade_date} A股短线观察清单",
        "status_line": "供人工复核，最终交易决策由你确认。",
        "trade_date": trade_date,
        "target_date": run.get("target_market_trade_date") or run.get("target_trade_date") or "",
        "summary": {
            "量化股票数": len(raw["quant_rows"]), "二筛股票数": len(raw["llm_rows"]),
            "重点候选数": len(candidates), "非零仓位数": sum(int(row["建议股数"] or 0) > 0 for row in orders),
            "今日推荐数": len(recommendations),
            "当前问题数": len(issues),
            "历史已解决问题数": resolved_current_issues + sum(
                row.get("schema_status") == "RESOLVED_HISTORY" for row in raw.get("audits") or []
            ),
            "本次人工已解决问题数": resolved_current_issues,
        },
        "minimum_recommendation_score": float(minimum_recommendation_score),
        "recommendations": recommendations,
        "top10": candidates[:10], "candidates": candidates, "orders": orders,
        "fundamentals": fundamentals, "quant_top100": quant_top100, "issues": issues,
    }


def _load_verification_overlays(
    trade_date: str, validation_run_id: str
) -> dict[str, dict[str, Any]]:
    overlays: dict[str, dict[str, Any]] = {}
    outputs = ROOT / "outputs"
    if not outputs.is_dir():
        return overlays
    for path in sorted(outputs.glob("*/**/*.json")):
        parent_date = next(
            (part for part in path.parts if re.fullmatch(r"\d{4}-\d{2}-\d{2}", part)),
            "",
        )
        if parent_date != trade_date:
            continue
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError):
            continue
        items = payload.get("items") if isinstance(payload, dict) else None
        if (
            not isinstance(items, list)
            or payload.get("artifact_type") != "HUMAN_VERIFICATION_OVERLAY_V1"
            or payload.get("trade_date") != trade_date
            or payload.get("validation_run_id") != validation_run_id
            or not payload.get("reviewed_at")
        ):
            continue
        for item in items:
            if (
                not isinstance(item, dict)
                or not str(item.get("status") or "").startswith("已")
                or not str(item.get("source_url") or "").startswith(("http://", "https://"))
            ):
                continue
            code = re.sub(r"\D", "", str(item.get("stock_code") or ""))[:6]
            if len(code) == 6:
                overlays[code] = item
    return overlays


def _apply_verification_overlay(
    candidate: dict[str, Any],
    fundamental: dict[str, Any],
    verification: dict[str, Any],
) -> None:
    business = str(verification.get("business") or "").strip()
    industry_chain = str(verification.get("industry_chain") or "").strip()
    conclusion = str(verification.get("conclusion") or "").strip()
    risks = str(verification.get("risks") or "").strip()
    if industry_chain:
        candidate["产业链"] = industry_chain
        fundamental["产业链"] = industry_chain
        if "上游" in industry_chain:
            fundamental["链条位置"] = "上游"
        elif "中游" in industry_chain:
            fundamental["链条位置"] = "中游"
        elif "下游" in industry_chain:
            fundamental["链条位置"] = "下游"
    if business:
        fundamental["主营业务"] = business
        products = business.removeprefix("主营").split("，", 1)[0].rstrip("。")
        if products:
            fundamental["核心产品"] = products
    if conclusion:
        candidate["核心逻辑"] = conclusion
        fundamental["潜在优势"] = conclusion
        fundamental["行业趋势"] = conclusion
        fundamental["核心逻辑"] = conclusion
        fundamental["财务说明"] = conclusion
    if risks:
        candidate["主要风险"] = risks
        fundamental["失效条件"] = risks
    sources = [
        str(verification.get(key) or "").strip()
        for key in ("source_url", "secondary_source_url")
    ]
    fundamental["人工复核"] = "已联网核验"
    fundamental["核验来源"] = "；".join(value for value in sources if value)


def _is_reportable_order_warning(row: dict[str, Any], verified_codes: set[str]) -> bool:
    code = str(row.get("stock_code") or "")[:6]
    if code in verified_codes:
        return False
    reason = str(row.get("reason") or "")
    expected_manual_block = "人工选择，但LLM未入选" in reason and "人工选择，但LLM分析失败" not in reason
    return not expected_manual_block


def _human_issue(row: dict[str, Any], llm_by_code: dict[str, dict], *, order: bool) -> dict[str, Any]:
    code = str(row.get("stock_code") or "")
    llm = llm_by_code.get(code, {})
    reason = str(row.get("reason") or "")
    if "Flash component scores copied the prompt example" in reason or "quant_consistency_score" in reason:
        reason = "二筛评分结构与模板示例过于一致，结果已标记为待人工确认"
    elif "no_unsupported_customer" in reason:
        reason = "内容包含未核验的客户信息，需要人工确认"
    elif "no_unsupported_leadership_claim" in reason:
        reason = "内容包含缺少依据的行业领先表述，需要人工确认"
    elif "$.inferred_concept_tags" in reason and "too_long" in reason:
        reason = "模型返回的推断概念标签超过结构化上限，结果已阻断并等待自动压缩修复"
    elif "$.evidence_fields" in reason and "too_long" in reason:
        reason = "模型返回的证据字段超过结构化上限，结果已阻断并等待自动压缩修复"
    elif order:
        reason = (
            "基本面已完成外部资料核验；原挂单与仓位结果保持阻断，等待下游重新评估"
            if "EXTERNAL_FUNDAMENTAL_COMPLETED_DOWNSTREAM_REVIEW_PENDING" in reason
            else "基本面补全未通过，挂单和仓位建议已按规则阻断"
        )
    return {
        "股票代码": code, "股票名称": llm.get("stock_name") or "",
        "问题类型": "规则提醒" if order else "基本面补全",
        "当前状态": "已阻断" if order else "待人工确认",
        "说明": _human_text(reason),
    }


def _human_warning(*values: Any) -> str:
    text = "；".join(str(value or "") for value in values)
    replacements = {
        "仅供模型验证，不可作为正式交易仓位建议": "",
        "LLM_UNVERIFIED_POSITION_DISCOUNT_APPLIED": "未核验信息已按规则降权",
        "TARGET_DAY_LIMIT_RULE_ESTIMATED": "目标日涨跌停价为规则估算",
        "UNVERIFIED_FUNDAMENTAL_RESEARCH": "基本面信息含未核验推断",
        "BELOW_ONE_TRADING_LOT": "建议金额不足一手",
        "RISK_GATE_BLOCKED_OR_NO_RECOMMENDED_PRICE": "风险规则阻断或缺少有效参考价",
        "人工选择，但LLM未入选": "人工关注股，未进入模型重点名单",
        "人工选择，但LLM分析失败": "基本面补全未通过",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    parts = [part.strip() for part in re.split(r"[；;]", text) if part.strip()]
    return "；".join(dict.fromkeys(parts))


_CONCEPT_NOISE = re.compile(
    r"(?:同花顺|全A|沪深|上证指数|成份股|样本股|主板|股通|陆股通|"
    r"融资融券|QFII|机构重仓|重仓股|减持新规|回购增持|昨日|"
    r"高市盈率|低市盈率|高市净率|低市净率|高股息|破净股|"
    r"\(A股\)|（A股）)",
    re.I,
)


def _compact_concept_tags(value: Any, *, limit: int = 8) -> str:
    text = _human_text(value)
    selected: list[str] = []
    seen: set[str] = set()
    for raw in re.split(r"[；;,\n]+", text):
        item = raw.strip()
        if not item or _CONCEPT_NOISE.search(item):
            continue
        key = item.rstrip("*").strip().casefold()
        if not key or key in seen:
            continue
        seen.add(key)
        selected.append(item)
        if len(selected) >= limit:
            break
    return "；".join(selected) or "信息不足"


def _human_risk_summary(
    fundamental: dict[str, Any], llm: dict[str, Any]
) -> str:
    pro_risks = _human_text(fundamental.get("pro_risks"))
    if pro_risks and pro_risks not in {"信息不足", "信息不足*"}:
        return pro_risks
    text = _human_text(llm.get("risk_note"))
    replacements = {
        "receivable_risk": "应收账款风险",
        "inventory_risk": "存货风险",
        "goodwill_risk": "商誉风险",
        "shareholder_action_risk": "股东行为风险",
        "operating cash flow": "经营现金流",
        "Financial status": "财务状态",
        "fundamental inference": "基本面推断",
        "Missing data": "缺失数据",
        "unknown": "尚未核验",
        "low confidence": "置信度较低",
        "requires manual review": "需要人工复核",
        "require manual review": "需要人工复核",
        " and ": "；",
    }
    for source, target in replacements.items():
        text = re.sub(re.escape(source), target, text, flags=re.I)
    if re.search(r"[A-Za-z]{4,}", text):
        return "基本面与风险字段仍含未核验信息，需要结合公告原文复核。"
    return text or "未发现结构化硬风险，但仍需结合公告原文复核。"


def _review_status(value: Any) -> str:
    text = str(value or "").strip()
    if text.startswith("已"):
        return text
    if value in {True, "是", "true", "TRUE", 1}:
        return "需要"
    if value in {False, None, "", "否", "false", "FALSE", 0}:
        return "不需要"
    return "需要"


def _human_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    replacements = {
        "LLM_UNVERIFIED": "模型未核验",
        "MODEL_UNVERIFIED": "模型未核验",
        "DEEPSEEK_UNVERIFIED": "模型未核验",
        "NEGATIVE_OPERATING_CASH_FLOW": "经营现金流为负",
        "EXCESSIVE_LEVERAGE": "杠杆偏高",
        "HEALTHY": "稳健", "HIGH_RISK": "高风险",
        "UNKNOWN": "信息不足", "UNCLEAR": "尚不清晰",
        "INSUFFICIENT_DATA": "信息不足", "STABLE": "稳定",
        "PRESSURED": "承压", "NORMAL_WATCH": "普通观察", "KEY_WATCH": "重点观察",
        "WATCH_ONLY": "普通观察", "ADVANCE": "优先复核", "HOLD": "继续观察",
        "REJECT": "暂不考虑", "BLOCK": "规则阻断",
        "LLM": "模型", "MODEL_VALIDATION": "", "NON_ACTIONABLE": "",
        "MULTI_SEGMENT": "多环节", "MIDSTREAM": "中游", "UPSTREAM": "上游",
        "DOWNSTREAM": "下游", "SERVICE_PLATFORM": "服务平台",
        "Fundamental inference confidence low; industry position unclear; requires manual review.": "基本面推断置信度较低，行业位置尚不清晰，需要人工复核。",
        "Financial pressure and low confidence in fundamental inference require manual review.": "财务表现承压且基本面推断置信度较低，需要人工复核。",
        "Chip Probing": "晶圆测试", "FinalTest": "成品测试", "SoC": "系统级芯片",
        "GPU": "图形处理芯片", "PVC": "聚氯乙烯", "ROE": "净资产收益率",
        "LCD": "液晶显示面板", "LCM/TFT": "液晶显示模组/薄膜晶体管",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"^\d{6}：", "", text)
    text = re.sub(r"：P：stock_company：main_business$", "", text)
    return text.strip(" ；;")


def _yes_no(value: Any) -> str:
    if value in {True, "是", "true", "TRUE", 1}:
        return "是"
    if value in {False, None, "", "否", "false", "FALSE", 0}:
        return "否"
    return "是"


def _find_checkpoint(trade_date: str) -> Path:
    candidates = [
        ROOT / "outputs" / trade_date / "审计" / "flash_v4_checkpoint.json",
        ROOT / "outputs" / "daily_full_test" / "flash_v4_checkpoint.json",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("FLASH_V4_CHECKPOINT_NOT_FOUND")


def _build_workbook(payload: dict[str, Any], output_path: Path, preview_dir: Path) -> None:
    build_dir = output_path.parent / f".human-build-{uuid.uuid4().hex}"
    build_dir.mkdir(parents=True, exist_ok=False)
    link = build_dir / "node_modules"
    try:
        builder = build_dir / "build_human_daily_excel.mjs"
        shutil.copy2(ROOT / "scripts" / "build_human_daily_excel.mjs", builder)
        shutil.copy2(ROOT / "scripts" / "excel_alignment.mjs", build_dir / "excel_alignment.mjs")
        payload_path = build_dir / "payload.json"
        payload_path.write_text(json.dumps(payload, ensure_ascii=False, default=str), encoding="utf-8")
        dependency_root = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules"
        if not (dependency_root / "@oai" / "artifact-tool").exists():
            raise RuntimeError("ARTIFACT_TOOL_DEPENDENCY_NOT_FOUND")
        result = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(dependency_root)],
            capture_output=True, text=True, check=False,
        )
        if result.returncode != 0:
            raise RuntimeError("ARTIFACT_TOOL_JUNCTION_FAILED")
        completed = subprocess.run(
            ["node", str(builder), str(payload_path), str(output_path), str(preview_dir)],
            cwd=build_dir,
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if completed.returncode != 0 and not output_path.exists():
            raise RuntimeError(f"HUMAN_WORKBOOK_EXPORT_FAILED:{completed.returncode}")
    finally:
        if link.exists():
            os.rmdir(link)
        shutil.rmtree(build_dir, ignore_errors=True)


def _validate_workbook(path: Path) -> dict[str, Any]:
    if not path.exists() or path.stat().st_size == 0:
        raise ValueError("HUMAN_WORKBOOK_MISSING")
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("HUMAN_WORKBOOK_ZIP_ERROR")
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8", errors="replace")
        sheet_count = len(re.findall(r"<(?:\w+:)?sheet\b", workbook_xml))
        if sheet_count not in {6, 7}:
            raise ValueError("HUMAN_WORKBOOK_SHEET_COUNT_ERROR")
        xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist() if name.endswith(".xml")
        )
    if not _secret_scan(xml):
        raise ValueError("HUMAN_WORKBOOK_SECRET_SCAN_FAILED")
    forbidden = ("MODEL_VALIDATION", "NON_ACTIONABLE", "WATCH_ONLY", "ADVANCE", "request_hash", "reasoning_content")
    if any(value.lower() in xml.lower() for value in forbidden):
        raise ValueError("HUMAN_WORKBOOK_TECHNICAL_TEXT_LEAK")
    standard = validate_trading_assistant_workbook(path)
    excel_compatibility = WorkbookStyleService.validate_excel_compatibility(path)
    return {
        "状态": "通过",
        "工作表数量": sheet_count,
        "压缩包完整性": "通过",
        "敏感信息检查": "通过",
        "统一制表规范": standard,
        "Excel兼容性": excel_compatibility,
    }


def _polish_workbook(path: Path) -> None:
    """Apply deterministic table alignment and freeze panes after export."""
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        WorkbookStyleService.set_freeze_panes_safely(
            worksheet,
            "A9" if worksheet.title == "今日概览" else "D5" if worksheet.title in {
                "今日推荐", "重点候选", "挂单与仓位", "价格与权重", "基本面摘要", "复核依据"
            } else "A5",
        )
        worksheet.sheet_view.showGridLines = False
        table_start_row = 8 if worksheet.title == "今日概览" else 4
        for row in worksheet.iter_rows(min_row=table_start_row, max_row=worksheet.max_row):
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    try:
        trade_date = date.fromisoformat(path.parent.name)
        reference = WorkbookStyleService.resolve_recent_successful_reference(
            ROOT / "outputs",
            start=trade_date - timedelta(days=7),
            end=trade_date - timedelta(days=1),
        )
        WorkbookStyleService(reference).align_existing_workbook(workbook)
    except (FileNotFoundError, ValueError):
        pass
    workbook.save(path)


def _assert_safe_runtime() -> None:
    if os.getenv("ENABLE_REAL_TRADING", "false").strip().lower() != "false":
        raise RuntimeError("UNSAFE_RUNTIME_SWITCHES")


if __name__ == "__main__":
    raise SystemExit(main())
