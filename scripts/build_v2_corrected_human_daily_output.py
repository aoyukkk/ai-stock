from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import shutil
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Mapping

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from sqlalchemy import select
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from database.models.workbench import ManualSelectionRecord
from database.session import get_session
from reporting.immutable_workbook import workbook_content_style_hashes
from reporting.web_result_publish import publish_workbook_payload
from scripts.build_human_daily_output import (
    _build_workbook,
    _polish_workbook,
    _validate_workbook,
)

FACTOR_VERSION = "TUSHARE_QUANT_V2_CORRECTED_SHADOW"
DEFAULT_SOURCE_ROOT = ROOT / "outputs" / "quant_v2_validation"


def _code(value: Any) -> str:
    digits = "".join(char for char in str(value or "") if char.isdigit())
    return digits[:6]


def _number(value: Any) -> float | None:
    if value in {None, ""}:
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _integer(value: Any) -> int | None:
    number = _number(value)
    return int(number) if number is not None else None


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _text_list(value: Any, *, limit: int = 8) -> str:
    if isinstance(value, str):
        return _clean_text(value)
    if not isinstance(value, (list, tuple)):
        return ""
    items = [_clean_text(item) for item in value if _clean_text(item)]
    return "；".join(items[:limit])


def _fundamental_status(value: Any) -> str:
    status = str(value or "").strip().upper()
    return {
        "STABLE": "稳健",
        "NORMAL": "正常",
        "PRESSURED": "承压",
        "RISKY": "高风险",
        "INSUFFICIENT": "信息不足",
        "UNKNOWN": "待核验",
    }.get(status, "待核验")


def _decode_gbk_mojibake(value: str) -> str:
    try:
        decoded = value.encode("latin-1").decode("gbk")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value
    return decoded


def _repair_export_encoding(path: Path) -> None:
    """Repair the legacy builder's GBK-as-Latin-1 text without changing styles."""
    workbook = load_workbook(path)
    original_titles = [worksheet.title for worksheet in workbook.worksheets]
    repaired_titles = [_decode_gbk_mojibake(value) for value in original_titles]
    if len(set(repaired_titles)) != len(repaired_titles):
        workbook.close()
        raise RuntimeError("REPAIRED_SHEET_TITLES_NOT_UNIQUE")

    for worksheet, repaired_title in zip(workbook.worksheets, repaired_titles):
        worksheet.title = repaired_title
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _decode_gbk_mojibake(cell.value)
                if cell.comment and isinstance(cell.comment.text, str):
                    cell.comment.text = _decode_gbk_mojibake(cell.comment.text)
        for rules in worksheet.conditional_formatting._cf_rules.values():
            for rule in rules:
                if rule.formula:
                    rule.formula = [
                        _decode_gbk_mojibake(value) for value in rule.formula
                    ]
                if getattr(rule, "text", None):
                    rule.text = _decode_gbk_mojibake(rule.text)
        for table in worksheet.tables.values():
            min_col, min_row, max_col, _ = range_boundaries(table.ref)
            if table.tableColumns:
                for offset, column in enumerate(table.tableColumns):
                    if min_col + offset > max_col:
                        break
                    header = worksheet.cell(min_row, min_col + offset).value
                    if header is not None:
                        column.name = str(header)
    workbook.save(path)
    workbook.close()


def _manual_rows(trade_date: date) -> list[dict[str, Any]]:
    session = get_session()
    try:
        rows = list(
            session.scalars(
                select(ManualSelectionRecord)
                .where(ManualSelectionRecord.trade_date == trade_date)
                .order_by(
                    ManualSelectionRecord.created_at,
                    ManualSelectionRecord.stock_code,
                )
            )
        )
        return [
            {
                "stock_code": _code(row.stock_code),
                "reason": row.reason,
                "priority": row.priority,
            }
            for row in rows
        ]
    finally:
        session.close()


def _priority(value: Any) -> str:
    return {
        "HIGH": "高",
        "MEDIUM": "中",
        "LOW": "低",
        "REVIEW_ONLY": "人工复核",
    }.get(str(value or "").upper(), "人工复核")


def _flash_decision(value: Any) -> str:
    return {
        "ADVANCE": "优先复核",
        "HOLD": "继续观察",
        "WATCH_ONLY": "继续观察",
        "BLOCK": "规则阻断",
        "REJECT": "规则阻断",
    }.get(str(value or "").upper(), "待人工复核")


def _clean_text(value: Any) -> str:
    text = str(value or "").strip()
    replacements = {
        "MODEL_VALIDATION": "模型验证",
        "NON_ACTIONABLE": "仅供参考",
        "WATCH_ONLY": "仅观察",
        "ADVANCE": "优先复核",
        "request_hash": "请求摘要",
        "reasoning_content": "模型分析内容",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def _source(code: str, model_codes: set[str], manual_codes: set[str]) -> str:
    if code in model_codes and code in manual_codes:
        return "人工关注+模型"
    if code in manual_codes:
        return "人工关注"
    return "模型筛选"


def _candidate_status(
    code: str,
    *,
    active_codes: set[str],
    watch_codes: set[str],
    model_codes: set[str],
    hard_gate: bool,
) -> str:
    if hard_gate:
        return "规则阻断"
    if code in active_codes:
        return "可执行影子候选"
    if code in watch_codes:
        return "观察池"
    if code in model_codes:
        return "模型复核"
    return "人工关注"


def _candidate_prices(plan: Mapping[str, Any] | None) -> dict[str, Any]:
    if not plan:
        return {
            "保守价": None,
            "均衡价": None,
            "积极价": None,
            "参考价": None,
            "最高接受价": None,
            "止损价": None,
            "第一目标价": None,
            "第二目标价": None,
            "第一目标收益比": None,
            "第二目标收益比": None,
            "当前收益比": None,
        }
    low = _number(plan.get("price_range_low"))
    high = _number(plan.get("price_range_high"))
    candidates = {
        str(row.get("price_type") or "").upper(): _number(row.get("price"))
        for row in (plan.get("candidates") or [])
    }

    def valid(value: float | None) -> float | None:
        if value is None:
            return None
        if low is not None and value < low - 0.01:
            return None
        if high is not None and value > high + 0.01:
            return None
        return value

    return {
        "保守价": valid(candidates.get("CONSERVATIVE")),
        "均衡价": valid(candidates.get("BALANCED")),
        "积极价": valid(candidates.get("AGGRESSIVE")),
        "参考价": _number(plan.get("recommended_price")),
        "最高接受价": _number(plan.get("max_acceptable_price")),
        "止损价": _number(plan.get("stop_loss_price")),
        "第一目标价": _number(plan.get("take_profit_1_price")),
        "第二目标价": _number(plan.get("take_profit_2_price")),
        "第一目标收益比": _number(plan.get("risk_reward_to_tp1")),
        "第二目标收益比": _number(plan.get("risk_reward_to_tp2")),
        "当前收益比": _number(plan.get("active_risk_reward")),
    }


def _payload(
    trade_date: date,
    full_rows: list[dict[str, Any]],
    final_audit: Mapping[str, Any],
    manual_rows: list[dict[str, Any]],
    fundamental_rows: list[dict[str, Any]] | None = None,
    web_verifications: Mapping[str, Any] | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    expected_prefix = f"v2-recovery-{trade_date:%Y%m%d}-"
    if not str(final_audit.get("base_run_id") or "").startswith(expected_prefix):
        raise RuntimeError("V2_SOURCE_RUN_MISMATCH")
    if final_audit.get("trade_date") != trade_date.isoformat():
        raise RuntimeError("V2_TRADE_DATE_MISMATCH")
    if final_audit.get("final_status") not in {
        "MONDAY_V2_SHADOW_FINALIZED",
        "MONDAY_WATCH_POOL_ONLY",
    }:
        raise RuntimeError("V2_FINALIZATION_NOT_READY")

    for row in full_rows:
        if row.get("factor_version") != FACTOR_VERSION:
            raise RuntimeError("V2_FACTOR_VERSION_MISMATCH")
    v2_by_code = {_code(row.get("stock_code")): row for row in full_rows}
    top100 = sorted(full_rows, key=lambda row: int(row["rank"]))[:100]
    emotion_values = [
        value
        for value in (_number(row.get("emotion_score")) for row in top100)
        if value is not None
    ]
    if len({round(value, 6) for value in emotion_values}) <= 1:
        raise RuntimeError("V2_EMOTION_SCORE_NOT_DIFFERENTIATED")
    if emotion_values and all(abs(value - 50.0) < 1e-9 for value in emotion_values):
        raise RuntimeError("V2_EMOTION_SCORE_LEGACY_CONSTANT")

    flash_rows = list((final_audit.get("flash") or {}).get("top20") or [])
    pro_rows = list((final_audit.get("pro") or {}).get("results") or [])
    if len(pro_rows) != 20:
        raise RuntimeError("V2_PRO_TOP20_INCOMPLETE")
    flash_by_code = {_code(row.get("stock_code")): row for row in flash_rows}
    pro_by_code = {_code(row.get("stock_code")): row for row in pro_rows}
    model_codes = set(pro_by_code)
    manual_by_code = {
        row["stock_code"]: row for row in manual_rows if row.get("stock_code")
    }
    manual_codes = set(manual_by_code)
    active_by_code = {
        _code(row.get("stock_code")): row
        for row in (final_audit.get("active_shadow") or [])
    }
    watch_by_code = {
        _code(row.get("stock_code")): row
        for row in (final_audit.get("watch_pool") or [])
    }
    active_codes = set(active_by_code)
    watch_codes = set(watch_by_code)
    plans_by_code = {
        _code(row.get("stock_code")): row
        for row in (final_audit.get("order_plans") or [])
    }
    strict_fundamental_overlay = fundamental_rows is not None
    fundamental_rows = fundamental_rows or []
    web_verifications = web_verifications or {}
    fundamental_by_code = {
        _code(row.get("stock_code")): row
        for row in fundamental_rows
        if row.get("status") == "COMPLETE"
    }
    expected_codes = model_codes | manual_codes
    if strict_fundamental_overlay and expected_codes - set(fundamental_by_code):
        missing = sorted(expected_codes - set(fundamental_by_code))
        raise RuntimeError(f"FUNDAMENTAL_LLM_RESULT_MISSING:{','.join(missing)}")

    ordered_codes = [
        _code(row.get("stock_code"))
        for row in sorted(pro_rows, key=lambda row: int(row["pro_rank"]))
    ]
    ordered_codes.extend(code for code in manual_by_code if code not in model_codes)

    candidates: list[dict[str, Any]] = []
    orders: list[dict[str, Any]] = []
    fundamentals: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for candidate_index, code in enumerate(ordered_codes, 1):
        quant = v2_by_code.get(code) or {}
        pro = pro_by_code.get(code) or {}
        flash = flash_by_code.get(code) or {}
        active = active_by_code.get(code) or {}
        manual = manual_by_code.get(code) or {}
        fundamental = fundamental_by_code.get(code) or {}
        profile = fundamental.get("profile") or {}
        inference = fundamental.get("inference") or {}
        web_entry = dict(web_verifications.get(code) or {})
        web_override = dict(web_entry.get("overrides") or {})
        hard_gate = not bool(quant) or str(quant.get("hard_gate")).lower() == "true"
        source = _source(code, model_codes, manual_codes)
        stock_name = (
            quant.get("stock_name")
            or pro.get("stock_name")
            or active.get("stock_name")
            or profile.get("stock_name")
            or code
        )
        sector = (
            quant.get("level_one_sector")
            or active.get("industry")
            or profile.get("level_one_sector")
            or "信息不足"
        )
        pro_score = _number(pro.get("pro_score"))
        flash_score = _number(pro.get("flash_score") or flash.get("llm_score"))
        summary = _clean_text(
            web_override.get("investment_logic")
            or inference.get("investment_logic")
            or pro.get("final_summary")
            or manual.get("reason")
            or "人工关注，未进入模型Top20，保留人工复核。"
        )
        key_risks = (
            web_override.get("invalidation_conditions")
            or inference.get("invalidation_conditions")
            or pro.get("key_risks")
            or []
        )
        risk_text = _text_list(key_risks, limit=8)
        if hard_gate:
            risk_text = "ST或数据硬门禁；未进入V2可评分Universe。"
        elif not risk_text:
            risk_text = "人工关注候选，需完成基本面和交易风险复核。"
        priority = _priority(pro.get("priority") or manual.get("priority"))
        plan = plans_by_code.get(code)
        prices = _candidate_prices(plan)
        position = 0.10 if code in active_codes else 0.0
        reference_price = prices["参考价"]
        shares = (
            math.floor(10000.0 / reference_price / 100.0) * 100
            if position > 0 and reference_price
            else 0
        )
        max_loss = (
            shares * max(0.0, reference_price - (prices["止损价"] or reference_price))
            if shares and reference_price
            else 0.0
        )
        current_status = _candidate_status(
            code,
            active_codes=active_codes,
            watch_codes=watch_codes,
            model_codes=model_codes,
            hard_gate=hard_gate,
        )
        candidate = {
            "深度复核排名": pro.get("pro_rank") or candidate_index,
            "股票代码": code,
            "股票名称": stock_name,
            "入选来源": source,
            "量化排名": _integer(quant.get("rank")),
            "量化得分": _number(quant.get("total_score")),
            "二筛得分": flash_score,
            "二筛结论": _flash_decision(
                pro.get("flash_decision") or flash.get("screening_decision")
            ),
            "深度复核分": pro_score,
            "复核优先级": priority,
            "一级行业": sector,
            "产业链": sector,
            "财务状态": _fundamental_status(
                web_override.get("financial_status")
                or inference.get("financial_status")
                or (profile.get("financial_status") or {}).get("status")
            ),
            "建议仓位": position,
            "建议股数": shares,
            "参考价": reference_price,
            "止损价": prices["止损价"],
            "第二目标价": prices["第二目标价"],
            "风险收益比": prices["当前收益比"],
            "核心逻辑": summary,
            "主要风险": risk_text,
            "当前状态": current_status,
        }
        candidates.append(candidate)

        note = (
            "RISK_OFF部署上限内的ACTIVE_SHADOW；仅供人工复核，不创建订单。"
            if code in active_codes
            else "RISK_OFF，仅观察，不配置仓位。"
            if code in watch_codes
            else "未进入ACTIVE_SHADOW，不配置仓位。"
            if code in model_codes
            else "人工关注候选；未通过模型Top20，不配置仓位。"
        )
        if hard_gate:
            note = "ST或数据硬门禁，禁止生成收益和仓位建议。"
        if plan and prices["均衡价"] is None:
            note += " Order Price Engine的均衡价超出合法价格区间，已留空而未写入异常值。"
        orders.append(
            {
                "深度复核排名": candidate["深度复核排名"],
                "股票代码": code,
                "股票名称": stock_name,
                "入选来源": source,
                "量化得分": candidate["量化得分"],
                "二筛得分": flash_score,
                "二筛结论": candidate["二筛结论"],
                "深度复核分": pro_score,
                **prices,
                "建议仓位": position,
                "建议资金": 10000.0 if position > 0 else 0.0,
                "建议股数": shares,
                "预计最大损失": max_loss,
                "说明": note,
            }
        )
        fundamentals.append(
            {
                "深度复核排名": candidate["深度复核排名"],
                "股票代码": code,
                "股票名称": stock_name,
                "入选来源": source,
                "一级行业": sector,
                "产业链": web_override.get("industry_chain")
                or inference.get("industry_chain")
                or sector,
                "链条位置": web_override.get("chain_position")
                or inference.get("chain_position")
                or "待核验",
                "主营业务": web_override.get("main_business")
                or inference.get("main_business")
                or _text_list(profile.get("main_business"), limit=4)
                or "待核验",
                "核心产品": _text_list(
                    web_override.get("core_products")
                    or inference.get("core_products")
                    or profile.get("core_products"),
                    limit=8,
                )
                or "待核验",
                "概念标签": _text_list(
                    web_override.get("concept_tags")
                    or inference.get("concept_tags")
                    or active.get("all_theme_clusters")
                    or profile.get("normalized_concept_tags"),
                    limit=8,
                )
                or "待核验",
                "结构性方向": web_override.get("structural_direction")
                or inference.get("structural_direction")
                or active.get("binding_cluster")
                or sector,
                "潜在优势": web_override.get("competitive_advantage")
                or inference.get("competitive_advantage")
                or "；".join(
                    _clean_text(item) for item in (pro.get("key_strengths") or [])
                )
                or "待核验",
                "行业趋势": web_override.get("industry_trend")
                or inference.get("industry_trend")
                or "待核验",
                "核心逻辑": web_override.get("investment_logic")
                or inference.get("investment_logic")
                or summary,
                "失效条件": _text_list(
                    web_override.get("invalidation_conditions")
                    or inference.get("invalidation_conditions"),
                    limit=8,
                )
                or risk_text,
                "财务状态": _fundamental_status(
                    web_override.get("financial_status")
                    or inference.get("financial_status")
                    or (profile.get("financial_status") or {}).get("status")
                ),
                "财务说明": web_override.get("financial_summary")
                or inference.get("financial_summary")
                or "结构化财务数据仍不足，保留待核验。",
                "人工复核": "已联网核验"
                if web_entry
                else "LLM补全",
                "复核优先级": priority,
                "核验来源": "；".join(web_entry.get("source_urls") or [])
                or (
                    "Tushare结构化公司与财务数据；"
                    "DeepSeek结构化归纳（workbook_fundamental_refill_v1）"
                ),
            }
        )
        if hard_gate:
            issues.append(
                {
                    "股票代码": code,
                    "股票名称": stock_name,
                    "问题类型": "V2硬门禁",
                    "当前状态": "已阻断",
                    "说明": "ST或数据硬门禁；保留人工关注记录，但不计算V2分数、收益或仓位。",
                }
            )

    recommendations = [
        row
        for row in candidates
        if row["深度复核分"] is not None and row["深度复核分"] >= 60.0
    ]
    quant_top100 = [
        {
            "量化排名": int(row["rank"]),
            "股票代码": _code(row["stock_code"]),
            "股票名称": row["stock_name"],
            "一级行业": row.get("level_one_sector") or "信息不足",
            "量化总分": _number(row.get("total_score")),
            "技术得分": _number(row.get("technical_score")),
            "资金得分": _number(row.get("capital_score")),
            "情绪得分": _number(row.get("emotion_score")),
            "动量得分": _number(row.get("momentum_score")),
            "风险得分": _number(row.get("risk_score")),
            "进入二筛": "是",
            "进入重点候选": "是" if _code(row["stock_code"]) in model_codes else "否",
        }
        for row in top100
    ]
    payload = {
        "title": f"{trade_date.isoformat()} A股短线观察清单",
        "status_line": (
            f"V2正式日线；因子版本：{FACTOR_VERSION}；"
            f"源运行：{final_audit['base_run_id']}；供人工复核，最终交易决策由你确认。"
        ),
        "trade_date": trade_date.isoformat(),
        "target_date": str(final_audit.get("target_trade_date") or ""),
        "minimum_recommendation_score": 60.0,
        "summary": {
            "量化股票数": len(full_rows),
            "二筛股票数": int((final_audit.get("flash") or {}).get("business_inputs") or 0),
            "重点候选数": len(candidates),
            "非零仓位数": len(active_codes),
            "非零仓位数量": len(active_codes),
            "今日推荐数": len(recommendations),
            "当前问题数": len(issues),
            "历史已解决问题数": 0,
        },
        "recommendations": recommendations,
        "top10": candidates[:10],
        "candidates": candidates,
        "orders": orders,
        "fundamentals": fundamentals,
        "quant_top100": quant_top100,
        "issues": issues,
    }
    audit = {
        "trade_date": trade_date.isoformat(),
        "factor_version": FACTOR_VERSION,
        "source_run_id": final_audit["base_run_id"],
        "finalization_status": final_audit["final_status"],
        "full_universe_count": len(full_rows),
        "top100_emotion_unique_count": len(
            {round(value, 6) for value in emotion_values}
        ),
        "top100_emotion_min": min(emotion_values),
        "top100_emotion_max": max(emotion_values),
        "top100_emotion_all_50": all(
            abs(value - 50.0) < 1e-9 for value in emotion_values
        ),
        "model_top20_count": len(model_codes),
        "manual_count": len(manual_codes),
        "candidate_count": len(candidates),
        "recommendation_count": len(recommendations),
        "watch_pool_count": len(watch_codes),
        "active_shadow_count": len(active_codes),
        "real_orders": 0,
        "virtual_orders": 0,
        "scheduler": False,
        "production_quant_promoted": True,
        "fundamental_llm_completed": len(fundamental_by_code),
        "fundamental_llm_actual_network_calls": sum(
            1
            for row in fundamental_rows
            if (row.get("usage") or {}).get("cached") is False
        ),
        "web_verified_count": len(web_verifications),
    }
    return payload, audit


def build(
    trade_date: date,
    *,
    overwrite: bool,
    fundamental_overlay: Path,
    web_overlay: Path | None,
) -> dict[str, Any]:
    source_dir = DEFAULT_SOURCE_ROOT / trade_date.isoformat()
    full_path = source_dir / "v2_full_universe.csv"
    final_path = source_dir / "monday_v2_candidate_audit.json"
    if not full_path.exists() or not final_path.exists():
        raise FileNotFoundError("V2_SOURCE_ARTIFACT_MISSING")
    full_rows = _read_csv(full_path)
    final_audit = _read_json(final_path)
    manual_rows = _manual_rows(trade_date)
    fundamental_audit = _read_json(fundamental_overlay)
    if fundamental_audit.get("trade_date") != trade_date.isoformat():
        raise RuntimeError("FUNDAMENTAL_OVERLAY_TRADE_DATE_MISMATCH")
    if fundamental_audit.get("status") != "COMPLETE":
        raise RuntimeError("FUNDAMENTAL_OVERLAY_NOT_COMPLETE")
    web_audit = (
        _read_json(web_overlay)
        if web_overlay is not None
        else {"trade_date": trade_date.isoformat(), "rows": []}
    )
    if web_audit.get("trade_date") != trade_date.isoformat():
        raise RuntimeError("WEB_OVERLAY_TRADE_DATE_MISMATCH")
    web_verifications = {
        _code(row.get("stock_code")): row
        for row in (web_audit.get("rows") or [])
    }
    fundamental_rows = list(fundamental_audit.get("rows") or [])
    payload, audit = _payload(
        trade_date,
        full_rows,
        final_audit,
        manual_rows,
        fundamental_rows,
        web_verifications,
    )

    day_dir = ROOT / "outputs" / trade_date.isoformat()
    formal_dir = day_dir / "正式日线"
    audit_dir = day_dir / "审计"
    formal_dir.mkdir(parents=True, exist_ok=True)
    audit_dir.mkdir(parents=True, exist_ok=True)
    filename = (
        f"智能交易助手_{trade_date.isoformat()}_"
        f"TUSHARE_QUANT_V2_CORRECTED_{final_audit['base_run_id']}.xlsx"
    )
    formal_target = formal_dir / filename
    root_target = day_dir / filename
    if (formal_target.exists() or root_target.exists()) and not overwrite:
        raise FileExistsError("OUTPUT_EXISTS_USE_OVERWRITE")

    timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    build_root = Path(
        tempfile.mkdtemp(prefix="v2-corrected-workbook-", dir=audit_dir)
    )
    generated = build_root / filename
    preview_dir = audit_dir / f"v2_corrected_visual_qa_{timestamp}"
    _build_workbook(payload, generated, preview_dir)
    _repair_export_encoding(generated)
    _polish_workbook(generated)
    validation = _validate_workbook(generated)
    generated_hashes = workbook_content_style_hashes(generated)

    backups: list[str] = []
    for target, label in ((formal_target, "formal"), (root_target, "root")):
        if target.exists():
            backup = audit_dir / f"替换前错误Legacy表_{label}_{timestamp}.xlsx"
            shutil.copy2(target, backup)
            backups.append(str(backup))
        replacement = target.with_suffix(target.suffix + ".tmp")
        shutil.copy2(generated, replacement)
        os.replace(replacement, target)

    formal_hashes = workbook_content_style_hashes(formal_target)
    root_hashes = workbook_content_style_hashes(root_target)
    if formal_hashes["file_sha256"] != root_hashes["file_sha256"]:
        raise RuntimeError("FORMAL_ROOT_WORKBOOK_HASH_MISMATCH")
    if formal_hashes["file_sha256"] != generated_hashes["file_sha256"]:
        raise RuntimeError("WORKBOOK_REPLACEMENT_HASH_MISMATCH")

    shutil.rmtree(build_root, ignore_errors=True)
    target_trade_date = date.fromisoformat(
        str(final_audit.get("target_trade_date") or trade_date.isoformat())
    )
    web_publish = publish_workbook_payload(
        trade_date=trade_date,
        target_trade_date=target_trade_date,
        payload=payload,
        full_quant_rows=full_rows,
        final_audit=final_audit,
        fundamental_rows=fundamental_rows,
        workbook_path=root_target,
        workbook_sha256=formal_hashes["file_sha256"],
        workbook_content_hash=formal_hashes["content_hash"],
        factor_version=FACTOR_VERSION,
        production_promoted=True,
    )
    deleted_superseded: list[str] = []
    for directory in (day_dir, formal_dir):
        for path in directory.glob(f"智能交易助手_{trade_date.isoformat()}_*.xlsx"):
            if path.resolve() == (root_target if directory == day_dir else formal_target).resolve():
                continue
            path.unlink()
            deleted_superseded.append(str(path))
    audit.update(
        {
            "status": "V2_CORRECTED_HUMAN_WORKBOOK_AND_WEB_READY",
            "formal_workbook": str(formal_target),
            "root_workbook": str(root_target),
            "workbook_hashes": formal_hashes,
            "web_publish": web_publish,
            "validation": validation,
            "visual_qa_dir": str(preview_dir),
            "backups": backups,
            "fundamental_overlay": str(fundamental_overlay),
            "web_overlay": str(web_overlay),
            "deleted_superseded_v1_workbooks": deleted_superseded,
        }
    )
    audit_path = audit_dir / f"智能交易助手_V2修正版_{timestamp}_audit.json"
    audit_path.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    audit["audit_path"] = str(audit_path)
    return audit


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--trade-date", default="2026-07-24")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--fundamental-overlay", type=Path, required=True)
    parser.add_argument("--web-overlay", type=Path)
    args = parser.parse_args()
    report = build(
        date.fromisoformat(args.trade_date),
        overwrite=args.overwrite,
        fundamental_overlay=args.fundamental_overlay.resolve(),
        web_overlay=args.web_overlay.resolve() if args.web_overlay else None,
    )
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
