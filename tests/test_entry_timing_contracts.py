from __future__ import annotations

from openpyxl import Workbook, load_workbook

from backend.application.excel_export import _write_sheet
from backend.api.entry_timing import router as entry_timing_router
from entry_timing.historical_v2 import EntryTimingV2Excel, _sheet


def test_entry_timing_api_is_registered_and_shadow_named() -> None:
    paths = {route.path for route in entry_timing_router.routes}
    assert "/api/workbench/entry-timing/run-shadow" in paths
    assert "/api/workbench/entry-timing/latest" in paths
    assert "/api/workbench/entry-timing/results" in paths
    assert "/api/workbench/entry-timing/methodology" in paths
    assert "/api/workbench/entry-timing/v2/run-shadow" in paths
    assert "/api/workbench/entry-timing/v2/latest" in paths
    assert "/api/workbench/entry-timing/v2/results" in paths
    assert "/api/workbench/entry-timing/v2/methodology" in paths


def test_entry_timing_excel_is_centered_and_stock_code_is_text(tmp_path) -> None:
    output = tmp_path / "entry.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "08_买入准入分析", [{"股票代码": "000001", "准入状态": "通过", "时机总分": 72.5}])
    workbook.save(output)
    workbook.close()
    loaded = load_workbook(output)
    sheet = loaded["08_买入准入分析"]
    assert sheet["A2"].value == "000001"
    assert sheet["A2"].number_format == "@"
    assert all(cell.alignment.horizontal == "center" and cell.alignment.vertical == "center" for row in sheet.iter_rows() for cell in row if cell.value is not None)
    loaded.close()


def test_frontend_exposes_entry_timing_shadow_page() -> None:
    from pathlib import Path

    root = Path(__file__).resolve().parents[1]
    view = (root / "frontend/src/views/EntryTimingView.vue").read_text(encoding="utf-8")
    api = (root / "frontend/src/api/entryTiming.ts").read_text(encoding="utf-8")
    assert "买入准入分析" in view
    assert "运行影子分析" in view
    assert "confirm_shadow: true" in api
    assert "strategy_id" in view
    assert "market_emotion_state" in view
    assert "admission_status_v2" in view
    assert "/api/workbench/entry-timing/v2/run-shadow" in api


def test_v2_excel_contract_has_required_order_alignment_and_text_code(tmp_path) -> None:
    assert EntryTimingV2Excel.SHEETS == (
        "01_总体对比", "02_逐股V1_V2明细", "03_策略分组", "04_市场情绪分组",
        "05_被过滤股票", "06_保留股票", "07_人工挑战池", "08_口径与异常",
    )
    output = tmp_path / "v2.xlsx"
    workbook = Workbook()
    _sheet(workbook, "V2", [{"股票代码": "000001", "Admission V2": "PASS", "D+1收益": .01}])
    workbook.remove(workbook["Sheet"])
    workbook.save(output)
    workbook.close()
    loaded = load_workbook(output, data_only=False)
    sheet = loaded["V2"]
    assert sheet["A2"].number_format == "@"
    assert sheet["C2"].number_format == "[Red]0.00%;[Green]-0.00%;-"
    assert all(cell.alignment.horizontal == "center" and cell.alignment.vertical == "center" and cell.alignment.wrap_text for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert not any(isinstance(cell.value, str) and cell.value.startswith("#") for row in sheet.iter_rows() for cell in row)
    loaded.close()
