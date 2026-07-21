from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table

from backend.core.config import AppConfig
from database.session import get_session, init_db
from datasource.ifind.http.errors import IFindHttpError, IFindHttpErrorCategory
from midday.coordinator import MiddayAppCoordinator
from midday.core import MiddayTimeGate
from midday.human_excel import _humanize_note
from midday.provider import MiddayIFindCollector
from midday.scoring import MiddayScoringEngine
from midday.service import MiddayRecommendationService, _equal_weights
from reporting.workbook_standard import _validate_table


SHANGHAI = ZoneInfo("Asia/Shanghai")


def _config(tmp_path: Path) -> AppConfig:
    return AppConfig(
        root_dir=tmp_path,
        config_dir=tmp_path,
        env={"ENABLE_REAL_TRADING": "false"},
        config_files={
            "system": {"system": {}, "safety": {"enable_real_trading": False}},
            "virtual_trading": {"virtual_trading": {"enabled": False, "real_trading_enabled": False}},
            "midday_recommendation": {"midday_recommendation": {
                "enabled": True, "auto_run_midday": False,
                "start_after": "11:32", "latest_start_time": "12:50",
                "valid_from": "13:00", "valid_until": "14:45",
                "source_mode": "PREVIOUS_DAY_TUSHARE_WITH_MORNING_IFIND_SHADOW",
                "pool": {"base_top_n": 100},
                "safety": {"actionable": False, "create_orders": False},
            }},
        },
    )


def test_midday_time_gate_boundaries() -> None:
    gate = MiddayTimeGate(start_after="11:32", latest="12:50")
    assert gate.evaluate(date(2026, 7, 15), datetime(2026, 7, 15, 11, 32, tzinfo=SHANGHAI))["passed"] is True
    assert gate.evaluate(date(2026, 7, 15), datetime(2026, 7, 15, 12, 50, tzinfo=SHANGHAI))["passed"] is True
    assert gate.evaluate(date(2026, 7, 15), datetime(2026, 7, 15, 12, 51, tzinfo=SHANGHAI))["status"] == "MISSED_MIDDAY_WINDOW"
    assert gate.evaluate(date(2026, 7, 15), datetime(2026, 7, 15, 10, 0, tzinfo=SHANGHAI))["status"] == "MIDDAY_SESSION_REQUIRED"


def test_missed_window_persists_without_external_or_llm(tmp_path: Path) -> None:
    from database.session import get_engine

    engine = get_engine("sqlite:///:memory:")
    init_db(engine)
    session = get_session(engine)

    class Forbidden:
        def __getattr__(self, name):
            raise AssertionError(f"external service called: {name}")

    try:
        result = MiddayRecommendationService(session, _config(tmp_path), collector=Forbidden(), reviewer=Forbidden()).run(
            date(2026, 7, 15),
            decision_time=datetime(2026, 7, 15, 12, 51, tzinfo=SHANGHAI),
        )
        assert result["status"] == "MISSED_MIDDAY_WINDOW"
        assert result["checkpoint"]["external_calls"] == 0
        assert result["checkpoint"]["llm_calls"] == 0
        assert result["orders_created"] == 0
    finally:
        session.close()


def test_scoring_separates_baseline_snapshot_and_minute_scopes() -> None:
    engine = MiddayScoringEngine({"overlay": {"scale": 0.12, "minimum_minute_completeness": 0.95}})
    item = {"base_quant_score": 70, "base_quant_rank": 5, "position_status": "HELD"}
    baseline = engine.score(item, None, [], index_return=None)
    assert baseline.feature_scope == "BASELINE_ONLY"
    assert baseline.delta == 0
    assert baseline.held_action == "DATA_INSUFFICIENT"
    snapshot = {"latest": 10.5, "pre_close": 10, "open": 10, "high": 10.6, "low": 9.9, "limit_up": 11, "limit_down": 9}
    partial = engine.score(item, snapshot, [], index_return=0.01)
    assert partial.feature_scope == "MORNING_SNAPSHOT_ONLY"
    bars = [{"close": 10 + index * 0.01, "volume": 1000 + index} for index in range(61)]
    full = engine.score(item, snapshot, bars, index_return=0.01)
    assert full.feature_scope == "MORNING_FULL_MINUTE"
    assert abs(full.delta) <= 6


def test_auto_coordinator_stays_disabled(tmp_path: Path) -> None:
    class Service:
        def run(self, *_args, **_kwargs):
            raise AssertionError("run must stay disabled")

    result = MiddayAppCoordinator(Service(), {"auto_run_midday": False}).tick(
        date(2026, 7, 15), now=datetime(2026, 7, 15, 11, 45, tzinfo=SHANGHAI),
    )
    assert result == {"status": "AUTO_RUN_DISABLED", "started": False}


def test_equal_reference_weights_sum_exactly_to_one() -> None:
    weights = _equal_weights([f"{index:06d}.SZ" for index in range(18)])
    assert sum(weights.values()) == 1


def test_midday_frontend_and_api_contracts_exist() -> None:
    root = Path(__file__).resolve().parents[1]
    api = (root / "backend/api/midday.py").read_text(encoding="utf-8")
    view = (root / "frontend/src/views/MiddayRecommendationView.vue").read_text(encoding="utf-8")
    router = (root / "frontend/src/router/index.ts").read_text(encoding="utf-8")
    for endpoint in ("/run", "/status", "/results", "/history", "/recheck", "/export", "/methodology"):
        assert f'\"/api/workbench/midday{endpoint}\"' in api
    assert "CenteredDataTable" in view
    assert "historical_validation: false" in (root / "frontend/src/api/midday.ts").read_text(encoding="utf-8")
    assert "midday-recommendation" in router
    assert "MIDDAY_AFTERNOON_RECHECK" in api
    assert "RequestPriority.P0" in api
    assert "collect_snapshots" in api
    assert "recheck(body.run_id, snapshots=snapshots)" in api


def test_midday_human_output_translates_model_notes() -> None:
    assert _humanize_note("High close quality (85.8)") == "收盘位置较强（85.8）"
    assert _humanize_note("Hard gate status PASS") == "硬性门槛通过"
    assert _humanize_note("Low liquidity confirmation (35.4)") == "量能确认偏弱（35.4）"


def test_human_excel_overview_can_use_weight_count() -> None:
    root = Path(__file__).resolve().parents[1]
    builder = (root / "scripts/build_human_daily_excel.mjs").read_text(encoding="utf-8")
    assert 'data.summary["非零仓位数量"]' in builder


def test_workbook_gate_rejects_table_header_metadata_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "mismatch.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测试"
    sheet.append(["股票代码", "股票名称"])
    sheet.append(["000001", "测试股票"])
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.add_table(Table(displayName="MismatchTable", ref="A1:B2"))
    workbook.save(path)
    workbook.close()

    workbook = load_workbook(path)
    sheet = workbook["测试"]
    sheet["A1"] = "证券代码"
    with pytest.raises(ValueError, match="WORKBOOK_TABLE_HEADER_METADATA_MISMATCH"):
        _validate_table(sheet, next(iter(sheet.tables.values())))
    workbook.close()


def test_midday_fast_path_stays_within_ifind_hard_call_limit() -> None:
    root = Path(__file__).resolve().parents[1]
    config = yaml.safe_load((root / "config" / "midday_recommendation.yaml").read_text(encoding="utf-8"))["midday_recommendation"]
    ifind = config["ifind"]
    assert ifind["fast_minutes_only"] is True
    assert ifind["max_external_calls"] == 30
    assert ifind["minute_max_stocks"] <= ifind["max_external_calls"]
    assert (root / "run_midday_once.cmd").is_file()
    assert (root / "scripts" / "run_midday_once.py").is_file()


def test_midday_minute_collection_isolates_recoverable_provider_failure(tmp_path: Path) -> None:
    from database.session import get_engine

    engine = get_engine("sqlite:///:memory:")
    init_db(engine)
    session = get_session(engine)

    class Provider:
        def get_minute_bars(self, *_args, **_kwargs):
            raise IFindHttpError(IFindHttpErrorCategory.TIMEOUT, "test timeout")

    try:
        collector = MiddayIFindCollector(
            session,
            _config(tmp_path),
            {"minute_max_stocks": 30, "max_external_calls": 30},
            provider=Provider(),
        )
        result = collector.collect_minutes(date(2026, 7, 15), ["000001.SZ"])
        assert result == {"000001.SZ": []}
        assert collector.summary()["minute_failures"] == 1
        assert collector.summary()["failures"] == [{"stock_code": "000001.SZ", "category": "TIMEOUT"}]
    finally:
        session.close()
