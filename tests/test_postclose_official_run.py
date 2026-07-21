from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.views import Selection

import post_close.official_run as official
import scripts.run_postclose_official_once as reusable_postclose
from post_close.official_run import DATASETS, PostCloseOfficialRunner, TushareTemporalGate
from post_close.seven_day_comparison import SevenDayComparisonService
from reporting.workbook_standard import DAILY_SHEETS
from reporting.workbook_style import WorkbookStyleService


def _dataset_row(dataset: str, code: str, trade_date: str) -> dict:
    common = {"ts_code": code, "trade_date": trade_date}
    return common | {
        "daily": {"open": 10, "high": 11, "low": 9, "close": 10.5},
        "daily_basic": {"close": 10.5, "turnover_rate": 2, "total_mv": 1000},
        "moneyflow": {},
        "stk_limit": {"up_limit": 11.5, "down_limit": 9.5},
        "adj_factor": {"adj_factor": 1.2},
    }[dataset]


def _write_dataset(root: Path, dataset: str, day: date, count: int) -> None:
    target = root / "data" / "cache" / "tushare" / "trade_date" / dataset / f"{day:%Y%m%d}.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps([
        _dataset_row(dataset, f"{index:06d}.SZ", f"{day:%Y%m%d}")
        for index in range(count)
    ]), encoding="utf-8")


def _reference(path: Path) -> Path:
    workbook = Workbook()
    workbook.active.title = DAILY_SHEETS[0]
    for name in DAILY_SHEETS[1:]:
        workbook.create_sheet(name)
    for worksheet in workbook.worksheets:
        worksheet["A1"] = worksheet.title
        worksheet["A4"] = "股票代码"
    workbook.save(path)
    return path


def test_temporal_gate_uses_batch_cache_and_coverage_ratios(tmp_path: Path) -> None:
    previous = date(2026, 7, 17)
    today = date(2026, 7, 20)
    for dataset in DATASETS:
        _write_dataset(tmp_path, dataset, previous, 100)
        _write_dataset(tmp_path, dataset, today, 98)

    result = TushareTemporalGate(tmp_path, today, previous).check()

    assert result["passed"] is True
    assert all(item["trade_dates"] == ["20260720"] for item in result["datasets"].values())
    assert result["datasets"]["daily"]["eligible_coverage"] == pytest.approx(0.98)
    assert all(item["duplicate_count"] == 0 for item in result["datasets"].values())


def test_pre_1700_guard_prevents_provider_and_llm(monkeypatch, tmp_path: Path) -> None:
    called = {"provider": 0}

    def forbidden_provider(**_kwargs):
        called["provider"] += 1
        raise AssertionError("provider must not run before the temporal guard")

    monkeypatch.setattr(official, "run_prewarm", forbidden_provider)
    runner = PostCloseOfficialRunner(tmp_path, date(2099, 7, 20), tmp_path / "unused.xlsx")
    with pytest.raises(RuntimeError, match="POSTCLOSE_CANNOT_RUN_BEFORE_1700"):
        runner.run()
    assert called == {"provider": 0}


def test_seven_trading_days_are_resolved_from_calendar(tmp_path: Path) -> None:
    reference = _reference(tmp_path / "reference.xlsx")
    cache = tmp_path / "data" / "cache" / "tushare"
    cache.mkdir(parents=True)
    calendar = ["20260709", "20260710", "20260713", "20260714", "20260715", "20260716", "20260717", "20260720"]
    (cache / "trade_cal_fixture.json").write_text(
        json.dumps([{"cal_date": value, "is_open": 1} for value in calendar]),
        encoding="utf-8",
    )

    class EmptySession:
        @staticmethod
        def scalars(_query):
            return []

    service = SevenDayComparisonService(EmptySession(), tmp_path, reference)
    assert service.trading_days(date(2026, 7, 20)) == [
        date(2026, 7, 10), date(2026, 7, 13), date(2026, 7, 14),
        date(2026, 7, 15), date(2026, 7, 16), date(2026, 7, 17), date(2026, 7, 20),
    ]


def test_shared_style_centers_wraps_and_preserves_stock_codes(tmp_path: Path) -> None:
    reference = _reference(tmp_path / "reference.xlsx")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "测试"
    worksheet.append(["股票代码", "长文本"])
    worksheet.append(["000001", "需要自动换行并保持居中的长文本"])
    WorkbookStyleService(reference).align_existing_workbook(workbook)
    output = tmp_path / "styled.xlsx"
    workbook.save(output)
    workbook.close()

    check = load_workbook(output)
    try:
        cell = check["测试"]["A2"]
        assert cell.value == "000001"
        assert cell.number_format == "@"
        assert cell.alignment.horizontal == "center"
        assert cell.alignment.vertical == "center"
        assert cell.alignment.wrap_text is True
    finally:
        check.close()


def test_shared_style_reapply_keeps_excel_compatible_views_and_filters(tmp_path: Path) -> None:
    reference = _reference(tmp_path / "reference.xlsx")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "测试"
    worksheet.append(["股票代码", "指标", "说明", "值"])
    worksheet.append(["000001", 1, "内容", 2])
    worksheet.add_table(Table(displayName="CompatibilityTable", ref="A1:D2"))
    worksheet.auto_filter.ref = "A1:D2"
    worksheet.freeze_panes = "D2"

    service = WorkbookStyleService(reference)
    service.align_existing_workbook(workbook)
    service.align_existing_workbook(workbook)
    output = tmp_path / "excel-compatible.xlsx"
    workbook.save(output)
    workbook.close()

    result = WorkbookStyleService.validate_excel_compatibility(output)
    assert result["status"] == "PASS"
    check = load_workbook(output)
    try:
        sheet = check["测试"]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is None
        assert len(sheet.tables) == 1
        assert [selection.pane for selection in sheet.sheet_view.selection] == ["bottomLeft"]
    finally:
        check.close()


def test_excel_compatibility_gate_rejects_known_excel_repair_patterns(tmp_path: Path) -> None:
    filter_book = Workbook()
    filter_sheet = filter_book.active
    filter_sheet.append(["股票代码", "值"])
    filter_sheet.append(["000001", 1])
    filter_sheet.add_table(Table(displayName="DuplicateFilterTable", ref="A1:B2"))
    filter_sheet.auto_filter.ref = "A1:B2"
    filter_path = tmp_path / "overlapping-filter.xlsx"
    filter_book.save(filter_path)
    filter_book.close()
    with pytest.raises(ValueError, match="OVERLAPPING_TABLE_AND_SHEET_FILTER"):
        WorkbookStyleService.validate_excel_compatibility(filter_path)

    view_book = Workbook()
    view_sheet = view_book.active
    view_sheet["A1"] = "测试"
    view_sheet.sheet_view.selection = [
        Selection(pane="bottomLeft", activeCell="A1", sqref="A1"),
        Selection(pane="bottomLeft", activeCell="A1", sqref="A1"),
    ]
    view_path = tmp_path / "duplicate-view.xlsx"
    view_book.save(view_path)
    view_book.close()
    with pytest.raises(ValueError, match="DUPLICATE_SHEET_VIEW_SELECTION"):
        WorkbookStyleService.validate_excel_compatibility(view_path)


def test_official_report_names_follow_runner_trade_date(tmp_path: Path) -> None:
    reference = _reference(tmp_path / "reference.xlsx")
    runner = PostCloseOfficialRunner(tmp_path, date(2026, 7, 21), reference)
    runner.output_dir.mkdir(parents=True)

    json_path, markdown_path, audit_path = runner._write_reports({
        "final_status": "POSTCLOSE_FULL_A_SUCCESS",
        "quant": {"scored_count": 10},
        "final_candidate_count": 2,
    })

    assert json_path.name.startswith("postclose_official_2026-07-21_")
    assert markdown_path.read_text(encoding="utf-8").startswith("# 2026-07-21 盘后全A正式运行")
    assert audit_path.name.startswith("postclose_official_audit_2026-07-21_")


def test_reference_resolver_finds_nested_official_workbook(tmp_path: Path) -> None:
    output_root = tmp_path / "outputs"
    nested = output_root / "2026-07-20" / "正式日线" / "智能交易助手_2026-07-20.xlsx"
    nested.parent.mkdir(parents=True)
    _reference(nested)

    result = WorkbookStyleService.resolve_recent_successful_reference(
        output_root,
        start=date(2026, 7, 13),
        end=date(2026, 7, 20),
    )

    assert result == nested.resolve()


def test_reusable_postclose_calendar_gate_distinguishes_open_and_closed_days(
    monkeypatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setattr(reusable_postclose, "ROOT", tmp_path)
    cache = tmp_path / "data" / "cache" / "tushare"
    cache.mkdir(parents=True)
    (cache / "trade_cal_fixture.json").write_text(json.dumps([
        {"cal_date": "20260720", "is_open": 1},
        {"cal_date": "20260721", "is_open": 0},
    ]), encoding="utf-8")

    assert reusable_postclose._local_trade_day_status(date(2026, 7, 20)) is True
    assert reusable_postclose._local_trade_day_status(date(2026, 7, 21)) is False
    assert reusable_postclose._calendar_horizon_end() == date(2026, 7, 21)
