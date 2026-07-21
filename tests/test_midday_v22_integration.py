from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
import pytest

from datasource.ifind.http.provider import IFindHttpP0Provider
from datasource.ifind.http.schemas import IFindHttpResponse
from midday.v22_export import SHEETS, export_midday_v22
from midday.v22_service import _reported_regime_confidence
from midday.v22_service import MiddayV22OneShotService
from database.models.midday_v22 import MiddayV22Result, MiddayV22Run


class BatchClient:
    def __init__(self) -> None:
        self.calls = []

    def post(self, endpoint_name, payload):
        self.calls.append((endpoint_name, payload))
        tables = []
        for code, close in (("000001.SZ", 10.1), ("600000.SH", 11.2)):
            tables.append({
                "thscode": code,
                "time": ["2026-07-20 11:30:00"],
                "table": {
                    "open": [close], "high": [close], "low": [close], "close": [close],
                    "volume": [100.0], "amount": [close * 100.0],
                },
            })
        return IFindHttpResponse(
            endpoint=endpoint_name, status_code=200, latency_ms=1, provider_code=0,
            provider_message="", schema_hash="batch-test", payload={"tables": tables},
        )


def test_ifind_minute_batch_uses_one_call_and_splits_canonical_codes():
    client = BatchClient()
    provider = IFindHttpP0Provider(client, enabled=True)
    result = provider.get_minute_bars_batch(["000001", "600000"], "2026-07-20 09:30:00", "2026-07-20 11:30:00")
    assert set(result) == {"000001.SZ", "600000.SH"}
    assert result["600000.SH"][0].close == 11.2
    assert len(client.calls) == 1
    assert client.calls[0][1]["codes"] == "000001.SZ,600000.SH"


def test_ifind_minute_batch_keeps_missing_requested_code_explicit():
    client = BatchClient()
    provider = IFindHttpP0Provider(client, enabled=True)
    result = provider.get_minute_bars_batch(["000001", "600000", "000002"], "2026-07-20 09:30:00", "2026-07-20 11:30:00")
    assert result["000002.SZ"] == []


def test_limited_top100_scope_reduces_reported_regime_confidence():
    confidence, basis = _reported_regime_confidence(1, "QUANT_TOP100_CROSS_SECTION")
    assert confidence == 0.75
    assert basis == "LIMITED_QUANT_TOP100_CROSS_SECTION"
    assert _reported_regime_confidence(0.8, "FULL_A_CROSS_SECTION") == (0.8, "FULL_SCOPE")


def test_v22_export_is_centered_and_stock_codes_are_text(tmp_path: Path):
    run = SimpleNamespace(run_id="midday-v22-export-test", trade_date=date(2026, 7, 20), status="SUCCESS")
    report = {
        "run_id": run.run_id, "cutoff": "2026-07-20T11:30:00+08:00", "status": "SUCCESS",
        "previous_regime": "CRASH", "midday_regime": "REPAIR", "regime_confidence": 0.75,
        "transition_reason": ["FIRST_SESSION_AFTER_CRASH"], "cooldown": 1,
        "live_breadth_scope": "QUANT_TOP100_CROSS_SECTION", "counts": {},
        "results": [{"stock_code": "000001.SZ", "stock_name": "测试", "pool_type": "AI_POOL", "result_layer": "BUY_READY", "afternoon_recheck": {"recheck_status": "NOT_REQUIRED"}}],
        "index_quality": {}, "provider_audit": [], "hashes": {},
    }
    paths = export_midday_v22(tmp_path, run, report)
    workbook = load_workbook(paths["excel"], data_only=False)
    try:
        assert workbook.sheetnames == list(SHEETS)
        focus = workbook[SHEETS[1]]
        code_column = next(cell.column for cell in focus[1] if cell.value == "stock_code")
        assert focus.cell(2, code_column).value == "000001.SZ"
        assert focus.cell(2, code_column).number_format == "@"
        assert focus.cell(2, code_column).alignment.horizontal == "center"
        assert focus.freeze_panes == "A2"
    finally:
        workbook.close()


def test_v22_run_is_committed_before_preflight_and_visible_to_new_session():
    engine = create_engine("sqlite:///:memory:")
    MiddayV22Run.__table__.create(engine)
    with Session(engine, expire_on_commit=False) as first:
        service = object.__new__(MiddayV22OneShotService)
        service.session = first
        service.v22 = {"test": True}
        run = service._start_run(date(2026, 7, 20), datetime.combine(date(2026, 7, 20), time(11, 30)))
        run_id = run.run_id
    with Session(engine) as second:
        persisted = second.scalar(select(MiddayV22Run).where(MiddayV22Run.run_id == run_id))
        assert persisted is not None and persisted.status == "RUNNING" and persisted.current_stage == "PREFLIGHT"


def test_v22_result_rows_are_immutable_after_insert():
    engine = create_engine("sqlite:///:memory:")
    MiddayV22Result.__table__.create(engine)
    with Session(engine, expire_on_commit=False) as session:
        row = MiddayV22Result(run_id="run", stock_code="000001.SZ", stock_name="测试", pool_type="AI_POOL", result_layer="BLOCKED", payload_json={})
        session.add(row); session.commit(); row.result_layer = "FOCUS"
        with pytest.raises(ValueError, match="IMMUTABLE_MIDDAY_V22_RESULT"):
            session.commit()
