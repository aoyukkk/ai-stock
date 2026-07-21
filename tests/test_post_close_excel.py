from __future__ import annotations

from openpyxl import load_workbook

from post_close.excel import _is_confirmed_nonheld

from backend.application.excel_export import write_tabular_workbook


def test_post_close_excel_contract_is_centered_wrapped_and_stock_code_is_text(tmp_path):
    output = tmp_path / "actions.xlsx"
    result = write_tabular_workbook(output, {
        "01_持仓操作建议": [{"股票代码": "000001.SZ", "当前采用建议": "建议减仓", "浮盈亏": .1234, "主要依据": "第一行\n第二行"}],
        "02_未持仓候选建议": [],
        "03_原方案与iFinD影子对比": [],
        "04_规则与异常": [],
        "05_扩展": [],
    })
    assert result["sheet_count"] == 5
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["01_持仓操作建议", "02_未持仓候选建议", "03_原方案与iFinD影子对比", "04_规则与异常", "05_扩展"]
    sheet = workbook["01_持仓操作建议"]
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.alignment.horizontal == "center"
            assert cell.alignment.vertical == "center"
            assert cell.alignment.wrap_text is True
    assert sheet["A2"].value == "000001.SZ"
    assert sheet["A2"].number_format == "@"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref
    workbook.close()


def test_missing_position_data_is_not_exported_as_confirmed_non_held():
    assert _is_confirmed_nonheld("SELECTED_NOT_HELD") is True
    assert _is_confirmed_nonheld("POSITION_DATA_MISSING") is False
