from openpyxl import Workbook, load_workbook

from scripts.build_human_daily_output import (
    _compact_concept_tags,
    _human_issue,
    _human_risk_summary,
    _human_text,
    _human_warning,
    _polish_workbook,
    _review_status,
    _yes_no,
)
from scripts.organize_outputs_by_date import _date_from_name


def test_human_output_translates_machine_statuses():
    assert _human_text("WATCH_ONLY") == "普通观察"
    assert _human_text("STABLE") == "稳定"
    assert _human_text("MULTI_SEGMENT") == "多环节"
    assert _human_text("HEALTHY") == "稳健"
    assert _human_text("HIGH_RISK") == "高风险"


def test_human_output_removes_repeated_disclaimer_and_translates_warnings():
    value = _human_warning(
        "LLM_UNVERIFIED_POSITION_DISCOUNT_APPLIED；仅供模型验证，不可作为正式交易仓位建议",
        "BELOW_ONE_TRADING_LOT；仅供模型验证，不可作为正式交易仓位建议",
    )
    assert value == "未核验信息已按规则降权；建议金额不足一手"
    assert "LLM" not in value
    assert "仅供模型验证" not in value


def test_human_output_yes_no_handles_chinese_false_value():
    assert _yes_no("是") == "是"
    assert _yes_no("否") == "否"
    assert _yes_no("") == "否"


def test_human_output_translates_technical_terms_and_removes_provenance():
    value = _human_text("603726：智能SoC、GPU、PVC：P：stock_company：main_business")
    assert value == "智能系统级芯片、图形处理芯片、聚氯乙烯"


def test_output_organizer_extracts_date_from_artifact_name():
    assert _date_from_name("ai_trader_demo_20260709.xlsx") == "2026-07-09"
    assert _date_from_name("flash_v4_checkpoint.json") is None


def test_human_output_translates_flash_template_validation_error():
    result = _human_issue(
        {
            "stock_code": "301277",
            "reason": "$.quant_consistency_score: Flash component scores copied the prompt example.",
        },
        {"301277": {"stock_name": "新天地"}},
        order=False,
    )
    assert result["说明"] == "二筛评分结构与模板示例过于一致，结果已标记为待人工确认"
    assert "$" not in result["说明"]
    assert "Flash" not in result["说明"]


def test_human_output_compacts_concepts_and_translates_machine_risk():
    concepts = "同花顺全A；沪深300；人工智能；人工智能；算力；融资融券；数据中心；云计算；边缘计算；软件"
    compacted = _compact_concept_tags(concepts, limit=4)

    assert compacted == "人工智能；算力；数据中心；云计算"
    assert _human_risk_summary(
        {},
        {"risk_note": "receivable_risk and goodwill_risk require manual review"},
    ) == "应收账款风险；商誉风险 需要人工复核"
    assert _review_status(True) == "需要"
    assert _review_status(False) == "不需要"


def test_human_output_polish_centers_tables_and_freezes_panes(tmp_path):
    path = tmp_path / "human.xlsx"
    workbook = Workbook()
    workbook.active.title = "今日概览"
    for name in ("今日推荐", "重点候选", "挂单与仓位", "基本面摘要", "量化前100", "当前问题"):
        workbook.create_sheet(name)
    for worksheet in workbook.worksheets:
        worksheet["A4"] = "表头"
        worksheet["A5"] = "内容"
    workbook.save(path)

    _polish_workbook(path)

    result = load_workbook(path)
    assert [sheet.freeze_panes for sheet in result.worksheets] == ["A9", "D5", "D5", "D5", "D5", "A5", "A5"]
    assert result.worksheets[1]["A4"].alignment.horizontal == "center"
    assert result.worksheets[1]["A4"].alignment.vertical == "center"
    assert result.worksheets[1]["A4"].alignment.wrap_text is True
    for sheet in result.worksheets:
        panes = [selection.pane for selection in sheet.sheet_view.selection if selection.pane]
        assert len(panes) == len(set(panes))
