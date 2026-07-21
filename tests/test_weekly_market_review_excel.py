from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from market_review.weekly_excel import WeeklyMarketReviewExcelExporter


def test_weekly_market_review_splits_daily_industries_and_shows_capital_flow(tmp_path: Path) -> None:
    bundles = [_bundle(date(2026, 7, 13), 0.0), _bundle(date(2026, 7, 14), 0.002)]
    flows = {
        date(2026, 7, 13): {
            "values": {"行业00": 12.5, "行业24": -9.25},
            "coverage": 0.98,
        },
        date(2026, 7, 14): {
            "values": {"行业01": 18.0, "行业23": -11.5},
            "coverage": 0.99,
        },
    }
    output = tmp_path / "本周大盘复盘.xlsx"
    result = WeeklyMarketReviewExcelExporter().export(output, bundles, industry_moneyflow=flows)

    assert result["blank_rows"] == 0
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == [
        "本周概览", "每日复盘", "7月13日行业", "7月14日行业", "行业资金流向", "全部行业",
    ]
    daily = workbook["7月13日行业"]
    assert daily["A1"].value == "7月13日行业前十后十"
    assert daily.max_row == 23
    assert daily["H4"].value == 12.5
    flow = workbook["行业资金流向"]
    assert flow["C4"].value == "行业24"
    assert flow["D4"].value == -9.25
    assert flow["F4"].value == "行业00"
    assert flow["G4"].value == 12.5
    assert "流出" in flow["I4"].value
    workbook.close()


def _bundle(trade_day: date, amount_shift: float) -> dict:
    industries = []
    for index in range(25):
        industries.append({
            "rank": index + 1,
            "sector_code": f"I{index:03d}",
            "sector_name": f"行业{index:02d}",
            "change_percent": (12 - index) / 100,
            "advancing_ratio": max(0.0, min(1.0, (25 - index) / 25)),
            "limit_up_count": index % 3,
            "amount": 100_000_000 * (index + 1) * (1 + amount_shift),
            "member_count": 10 + index,
        })
    return {
        "run": {"trade_date": trade_day.isoformat()},
        "snapshot": {
            "breadth": {
                "equal_weight_return": 0.01,
                "median_return": 0.008,
                "advancing_count": 3000,
                "declining_count": 2000,
                "advancing_ratio": 0.6,
            },
            "turnover": {"total_amount": 2_000_000_000_000, "change_ratio": 0.03},
            "limit_structure": {"limit_up_count": 50, "limit_down_count": 5},
            "data_quality_score": 99,
            "industries": industries,
        },
        "regime": {"market_regime": "MIXED_ROTATION"},
        "review": {
            "market_summary": "市场震荡轮动",
            "breadth_summary": "涨多跌少",
            "turnover_summary": "成交温和放大",
        },
        "drivers": [{"explanation": "成交与市场宽度共同改善。"}],
    }
