from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.api.entry_timing import router
from entry_timing.historical_v22 import V22Excel


def test_v22_routes_are_explicitly_shadow_only() -> None:
    paths = {route.path for route in router.routes}
    assert "/api/workbench/entry-timing/v22/historical-shadow" in paths
    assert "/api/workbench/entry-timing/v22/latest" in paths
    assert "/api/workbench/entry-timing/v22/results" in paths
    assert "/api/workbench/entry-timing/v22/methodology" in paths


def test_v22_frontend_has_required_filters_and_safety_copy() -> None:
    root = Path(__file__).resolve().parents[1]
    view = (root / "frontend/src/views/EntryTimingView.vue").read_text(encoding="utf-8")
    api = (root / "frontend/src/api/entryTiming.ts").read_text(encoding="utf-8")
    for field in ("regime_state", "deployment_status", "crowding_status", "trigger_status", "pool_type"):
        assert field in view
    assert "不创建真实或虚拟订单" in view
    assert "confirm_shadow: true" in api
    assert "/api/workbench/entry-timing/v22/historical-shadow" in api


def test_v22_excel_contract_is_centered_and_stock_code_is_text(tmp_path) -> None:
    output = tmp_path / "v22.xlsx"
    report = {
        "versions": {"V2_2_DEPLOYMENT": {"absolute_win_rate": 0.5, "relative_win_rate": 0.6}},
        "timeline": [{"trade_date": "2026-07-13", "current_state": "CRASH"}],
        "industry_concentration_before_after": [{"industry": "医药", "before_concentration": 4, "after_concentration": 2}],
        "target_before_stop_note": "Not evaluated",
        "promotion_recommendation": "KEEP_V2_2_SHADOW",
        "reproducible_hash": "a" * 64,
    }
    rows = [{"股票代码": "000001", "trigger_status": "DATA_INSUFFICIENT", "crowding_status": "RETAINED"}]
    V22Excel().export(output, rows, report)
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.sheetnames == list(V22Excel.SHEETS)
        assert workbook["03_逐股部署明细"]["A2"].number_format == "@"
        assert all(
            cell.alignment.horizontal == "center" and cell.alignment.vertical == "center" and cell.alignment.wrap_text
            for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None
        )
    finally:
        workbook.close()
