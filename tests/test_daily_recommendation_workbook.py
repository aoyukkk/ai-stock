from pathlib import Path

from openpyxl import load_workbook

from scripts.build_human_daily_output import (
    _build_workbook,
    _polish_workbook,
    _validate_workbook,
    build_human_payload,
)


def test_today_recommendations_use_inclusive_final_score_threshold_for_all_sources() -> None:
    raw = _raw_payload([
        ("000001", "模型一", "LLM_TOP20", 60),
        ("000002", "人工一", "MANUAL", 61),
        ("000003", "共同一", "BOTH", 59.99),
    ])
    payload = build_human_payload(raw, "2026-07-17", minimum_recommendation_score=60)

    assert [row["股票代码"] for row in payload["candidates"]] == ["000001", "000002", "000003"]
    assert [(row["股票代码"], row["入选来源"]) for row in payload["recommendations"]] == [
        ("000001", "模型筛选"),
        ("000002", "人工关注"),
    ]
    assert payload["summary"]["今日推荐数"] == 2


def test_today_recommendation_sheet_is_valid_when_no_stock_reaches_threshold(tmp_path: Path) -> None:
    payload = build_human_payload(
        _raw_payload([("000001", "模型一", "LLM_TOP20", 59.99)]),
        "2026-07-17",
        minimum_recommendation_score=60,
    )
    output = tmp_path / "智能交易助手_2026-07-17.xlsx"
    _build_workbook(payload, output, tmp_path / "预览")
    _polish_workbook(output)
    result = _validate_workbook(output)

    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.sheetnames == [
            "今日概览", "今日推荐", "重点候选", "挂单与仓位", "基本面摘要", "量化前100", "当前问题",
        ]
        assert workbook["今日推荐"]["A5"].value == "暂无推荐"
        assert len(workbook["今日推荐"].tables) == 1
        assert result["工作表数量"] == 7
    finally:
        workbook.close()


def _raw_payload(rows: list[tuple[str, str, str, float]]) -> dict:
    quant_rows = []
    llm_rows = []
    order_rows = []
    fundamental_rows = []
    for rank, (code, name, source, pro_score) in enumerate(rows, start=1):
        quant_rows.append({
            "rank": rank, "stock_code": code, "stock_name": name, "level_one_sector": "测试行业",
            "total_score": 80 - rank, "technical_score": 70, "capital_score": 60,
            "emotion_score": 50, "momentum_score": 75, "risk_score": 65,
            "llm_evaluated": True, "selection_source": source,
        })
        llm_rows.append({"stock_code": code, "risk_note": "风险待跟踪"})
        order_rows.append({
            "stock_code": code, "stock_name": name, "selection_source": source,
            "quant_rank": rank, "quant_score": 80 - rank, "llm_score": 70 - rank,
            "llm_decision": "ADVANCE", "pro_score": pro_score, "pro_rank": rank,
            "pro_priority": "MEDIUM", "level_one_sector": "测试行业", "financial_status": "STABLE",
            "position_percent": 0.05, "quantity": 100, "recommended_price": 10,
            "stop_loss_price": 9.5, "take_profit_2": 11, "active_risk_reward": 2,
            "order_status": "DRAFT", "conservative_price": 9.8, "balanced_price": 10,
            "aggressive_price": 10.2, "max_acceptable_price": 10.5, "take_profit_1": 10.5,
            "risk_reward_to_tp1": 1, "risk_reward_to_tp2": 2, "capital_amount": 1000,
            "max_loss": 50, "position_status": "DRAFT", "order_warning": "", "position_warning": "",
        })
        fundamental_rows.append({
            "stock_code": code, "level_one_sector": "测试行业", "industry_chain": "测试产业链",
            "chain_position": "MIDSTREAM", "main_business": "测试主营", "core_products": "测试产品",
            "concept_tags": "测试概念", "structural_theme_fit": "测试方向",
            "competitive_advantage": "测试优势", "industry_trend": "测试趋势",
            "investment_logic": "测试逻辑", "logic_invalidation": "测试失效条件",
            "financial_status": "STABLE", "financial_status_reason": "财务稳定", "manual_review": False,
        })
    return {
        "run": {"target_trade_date": "2026-07-20"},
        "quant_rows": quant_rows,
        "llm_rows": llm_rows,
        "order_rows": order_rows,
        "fundamental_rows": fundamental_rows,
        "errors": [], "warnings": [], "audits": [],
    }
