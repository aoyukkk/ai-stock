from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from database.base import Base
from database.models import ForwardOutcome, StrategyTimingContract
from review.forward_shadow import BackfillOptions, ForwardShadowBackfillService, GATE_SCOPES, LocalDailyCache, sample_status
from review.forward_shadow_report import ForwardShadowWorkbookExporter, SHEETS


def _session():
    engine = create_engine("sqlite:///:memory:", future=True)
    import database.models  # noqa: F401
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)()


def _cache(root: Path, *, limit_up: bool = False, missing_entry: bool = False) -> LocalDailyCache:
    for kind in ("daily", "stk_limit", "daily_basic"):
        (root / kind).mkdir(parents=True)
    rows = {
        "20260721": [] if missing_entry else [{"ts_code": "000001.SZ", "open": 10.0, "high": 10.8, "low": 9.8, "close": 10.5, "vol": 10000}],
        "20260722": [{"ts_code": "000001.SZ", "open": 10.6, "high": 11.0, "low": 10.2, "close": 10.8, "vol": 10000}],
        "20260723": [{"ts_code": "000001.SZ", "open": 10.9, "high": 11.2, "low": 10.4, "close": 11.0, "vol": 10000}],
    }
    for day, values in rows.items():
        (root / "daily" / f"{day}.json").write_text(json.dumps(values), encoding="utf-8")
        limits = [{"ts_code": "000001.SZ", "up_limit": 10.0 if limit_up and day == "20260721" else 11.0, "down_limit": 9.0}]
        (root / "stk_limit" / f"{day}.json").write_text(json.dumps(limits), encoding="utf-8")
        (root / "daily_basic" / f"{day}.json").write_text("[]", encoding="utf-8")
    return LocalDailyCache(root)


def _contract(session) -> StrategyTimingContract:
    row = StrategyTimingContract(stock_code="000001.SZ", trade_date=date(2026, 7, 20),
        observation_end_ts=datetime(2026, 7, 20, 15), available_at_ts=datetime(2026, 7, 20, 15, 5),
        signal_generated_at=datetime(2026, 7, 20, 15, 10), order_eligible_at=datetime(2026, 7, 21, 9, 30),
        execution_policy="NEXT_OPEN", feature_version="test", data_snapshot_id="data", universe_snapshot_id="universe")
    session.add(row); session.commit(); return row


def test_next_open_no_same_day_fill_and_horizon_uses_trading_days(tmp_path):
    session = _session(); contract = _contract(session)
    ForwardShadowBackfillService(session, _cache(tmp_path)).run(BackfillOptions(as_of_date=date(2026, 7, 22)))
    row = session.scalar(select(ForwardOutcome))
    assert row.entry_trade_date == date(2026, 7, 21)
    assert float(row.entry_price) == 10.0
    assert float(row.return_d1) == .05
    assert row.return_d3 is None and row.horizon_status_json["d3"] == "PENDING"
    assert contract.trade_date != row.entry_trade_date


def test_slippage_sensitivity_and_missing_returns_are_null(tmp_path):
    session = _session(); _contract(session)
    ForwardShadowBackfillService(session, _cache(tmp_path)).run(BackfillOptions(as_of_date=date(2026, 7, 21), buy_slippage_bps=20))
    row = session.scalar(select(ForwardOutcome))
    assert round(float(row.entry_price), 3) == 10.02
    assert row.sensitivity_json["OPEN_PLUS_30BP"]["entry_price"] == 10.03
    assert row.return_d3 is None and row.return_d5 is None and row.return_d10 is None


def test_limit_up_and_suspended_do_not_enter_strategy_return(tmp_path):
    session = _session(); _contract(session)
    ForwardShadowBackfillService(session, _cache(tmp_path, limit_up=True)).run(BackfillOptions(as_of_date=date(2026, 7, 22)))
    row = session.scalar(select(ForwardOutcome))
    assert row.entry_status == "PATH_AMBIGUOUS"
    assert row.return_d1 is None and row.horizon_status_json["d1"] == "NOT_TRADABLE"


def test_max_acceptable_price_blocks_fill(tmp_path):
    session = _session(); contract = _contract(session)
    service = ForwardShadowBackfillService(session, _cache(tmp_path))
    values = service._evaluate(contract, {"source_run_id": "x", "route": "BASELINE", "baseline_status": "WATCH",
                                          "v2_2_status": None, "v3_status": None, "v3": None}, None,
                               BackfillOptions(as_of_date=date(2026, 7, 22)), max_price=9.99)
    assert values["entry_status"] == "NOT_FILLED_PRICE_TOO_HIGH"
    assert values.get("return_d1") is None
    assert values["non_fill_quality"] == "MISSED_GAIN" and values["missed_opportunity"] > 0


def test_gate_scope_sample_status_and_offline_guards(tmp_path):
    assert GATE_SCOPES["MARKET_REGIME"] == "GLOBAL_MARKET"
    assert GATE_SCOPES["CONCENTRATION"] == "PORTFOLIO"
    assert sample_status(29) == "INSUFFICIENT_SAMPLE"
    assert sample_status(30) == "PRELIMINARY"
    assert sample_status(100) == "USABLE" and sample_status(500) == "STABLE"
    session = _session(); _contract(session)
    try:
        ForwardShadowBackfillService(session, _cache(tmp_path)).run(BackfillOptions(as_of_date=date(2026, 7, 22), no_llm=False))
    except ValueError as exc:
        assert str(exc) == "SHADOW_BACKFILL_REQUIRES_OFFLINE_READ_ONLY_MODE"
    else:
        raise AssertionError("offline guard did not fire")


def test_binding_gate_and_overlap_are_ordered_without_duplicate_attribution():
    class V3:
        risk_penalties = {"MARKET_RED": {"score_penalty": 10}, "HIGH_POSITION_RISK": {"score_penalty": 4}}
    gates, binding = ForwardShadowBackfillService._gates({"v3": V3(), "v2_2_status": None}, True)
    assert gates == ["MARKET_REGIME", "HIGH_POSITION_RISK"]
    assert binding == "MARKET_REGIME"


def test_workbook_has_13_centered_sheets_text_codes_and_no_formula_errors(tmp_path):
    session = _session(); _contract(session)
    ForwardShadowBackfillService(session, _cache(tmp_path / "cache")).run(BackfillOptions(as_of_date=date(2026, 7, 22)))
    path, _ = ForwardShadowWorkbookExporter(session).export(date(2026, 7, 22), tmp_path / "out")
    wb = load_workbook(path, data_only=False)
    assert tuple(wb.sheetnames) == SHEETS
    code_cell = wb["03_逐股结果"]["D2"]
    assert code_cell.value == "000001" and code_cell.number_format == "@"
    for ws in wb.worksheets:
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref
        for row in ws.iter_rows():
            for cell in row:
                assert cell.alignment.horizontal == "center" and cell.alignment.vertical == "center"
                assert cell.alignment.wrap_text is True
                assert not (isinstance(cell.value, str) and "#REF!" in cell.value)
    wb.close()


def test_forward_shadow_api_is_get_only():
    from backend.api.forward_shadow import router
    routes = {route.path: route.methods for route in router.routes}
    assert set(routes) == {"/api/forward-shadow/summary", "/api/forward-shadow/outcomes",
                           "/api/forward-shadow/model-comparison", "/api/forward-shadow/gates",
                           "/api/forward-shadow/factors", "/api/forward-shadow/pending"}
    assert all(methods == {"GET"} for methods in routes.values())
