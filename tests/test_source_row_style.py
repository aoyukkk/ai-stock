from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from reporting.source_row_style import (
    apply_selection_source_rows,
    disable_table_banding,
)


def test_source_row_style_treats_combined_selection_as_manual() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["股票代码", "入选来源"])
    sheet.append(["000001", "模型筛选"])
    sheet.append(["000002", "人工关注"])
    sheet.append(["000003", "前日量化前100、人工关注"])
    sheet.append(["000004", "共同入选"])

    assert apply_selection_source_rows(sheet, header_row=1) is True
    assert sheet["A2"].fill.fgColor.rgb == "00DCEAF5"
    assert sheet["A3"].fill.fgColor.rgb == "00FFF2CC"
    assert sheet["A4"].fill.fgColor.rgb == "00FFF2CC"
    assert sheet["A5"].fill.fgColor.rgb == "00FFF2CC"


def test_disable_table_banding_removes_row_and_column_stripes() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["股票代码", "入选来源"])
    sheet.append(["000001", "模型筛选"])
    table = Table(displayName="Candidates", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showColumnStripes=True,
    )
    sheet.add_table(table)

    disable_table_banding(sheet)

    assert table.tableStyleInfo.showRowStripes is False
    assert table.tableStyleInfo.showColumnStripes is False
