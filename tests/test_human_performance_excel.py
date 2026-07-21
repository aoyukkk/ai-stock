from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from review.human_performance_excel import HumanPerformanceExcelExporter, WeeklyStockPerformanceExcelExporter


def test_human_performance_workbook_pivots_daily_returns(tmp_path: Path) -> None:
    payload = {
        "run": {"evaluation_end_date": date(2026, 7, 14)},
        "cohorts": [
            {"selection_trade_date": date(2026, 7, 10), "baseline_trade_date": date(2026, 7, 13), "coverage_ratio": 1},
            {"selection_trade_date": date(2026, 7, 13), "baseline_trade_date": date(2026, 7, 14), "coverage_ratio": 1},
        ],
        "daily": [
            {"selection_trade_date": date(2026, 7, 10), "evaluation_trade_date": date(2026, 7, 13), "daily_return": 0.02},
            {"selection_trade_date": date(2026, 7, 10), "evaluation_trade_date": date(2026, 7, 14), "daily_return": -0.01},
            {"selection_trade_date": date(2026, 7, 13), "evaluation_trade_date": date(2026, 7, 14), "daily_return": 0.03},
        ],
        "stocks": [
            _stock(date(2026, 7, 10), date(2026, 7, 13), "000001.SZ", 1, 0.02, 0.02),
            _stock(date(2026, 7, 10), date(2026, 7, 14), "000001.SZ", 1, -0.01, 0.0098),
            _stock(date(2026, 7, 13), date(2026, 7, 14), "000002.SZ", 2, 0.03, 0.03),
        ],
    }
    output = tmp_path / "选股复盘.xlsx"
    result = HumanPerformanceExcelExporter().export(output, payload)
    assert result["blank_rows"] == 0
    assert result["table_objects"] == 0
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["近三日汇总", "7月10日选股复盘", "7月13日选股复盘"]
    first = workbook["7月10日选股复盘"]
    assert [cell.value for cell in first[4]] == [
        "原排名", "股票代码", "股票名称", "来源", "7月13日当日涨跌", "7月14日当日涨跌", "截至7月14日总涨跌",
    ]
    assert first["G5"].value == 0.0098
    assert first["B5"].number_format == "@"
    assert first.tables == {}
    workbook.close()


def test_review_workbook_uses_scope_label_and_supports_empty_recommendation(tmp_path: Path) -> None:
    payload = {
        "run": {
            "evaluation_end_date": date(2026, 7, 14),
            "selection_scope": "FINAL_CANDIDATES",
            "lookback_value": 7,
        },
        "cohorts": [{
            "selection_trade_date": date(2026, 7, 14),
            "baseline_trade_date": date(2026, 7, 14),
            "coverage_ratio": 0,
        }],
        "daily": [],
        "stocks": [],
        "members": [],
    }
    output = tmp_path / "今日推荐复盘.xlsx"
    HumanPerformanceExcelExporter().export(output, payload)
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["近七日汇总", "7月14日选股复盘"]
    assert workbook["近七日汇总"]["A1"].value == "近七日今日推荐复盘"
    assert workbook["7月14日选股复盘"]["A1"].value == "7月14日今日推荐复盘"
    assert workbook["7月14日选股复盘"]["A5"].value == "暂无推荐"
    workbook.close()


def test_review_workbook_uses_key_candidate_label(tmp_path: Path) -> None:
    payload = {
        "run": {
            "evaluation_end_date": date(2026, 7, 14),
            "selection_scope": "KEY_CANDIDATES",
            "lookback_value": 7,
        },
        "cohorts": [{
            "selection_trade_date": date(2026, 7, 14),
            "baseline_trade_date": date(2026, 7, 14),
            "coverage_ratio": 0,
        }],
        "daily": [],
        "stocks": [],
        "members": [],
    }
    output = tmp_path / "重点候选复盘.xlsx"
    HumanPerformanceExcelExporter().export(output, payload)
    workbook = load_workbook(output, data_only=False)
    assert workbook["近七日汇总"]["A1"].value == "近七日重点候选复盘"
    assert workbook["7月14日选股复盘"]["A5"].value == "暂无重点候选"
    workbook.close()


def test_weekly_stock_comparison_is_one_vertical_table(tmp_path: Path) -> None:
    payload = {
        "run": {"evaluation_end_date": date(2026, 7, 17)},
        "members": [
            {
                "selection_trade_date": date(2026, 7, 13),
                "stock_code": "000001.SZ",
                "stock_name": "测试一",
                "selection_source": "LLM",
                "pro_rank": 1,
                "pro_score": 65.5,
            },
            {
                "selection_trade_date": date(2026, 7, 17),
                "stock_code": "000002.SZ",
                "stock_name": "测试二",
                "selection_source": "LLM",
                "pro_rank": 2,
                "pro_score": 62.0,
            },
        ],
        "stocks": [
            _stock(date(2026, 7, 13), date(2026, 7, 14), "000001.SZ", 1, 0.02, 0.02),
            _stock(date(2026, 7, 13), date(2026, 7, 15), "000001.SZ", 1, -0.01, 0.0098),
        ],
    }
    output = tmp_path / "本周模型选股对比.xlsx"
    WeeklyStockPerformanceExcelExporter().export(output, payload)
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["本周模型选股对比"]
    sheet = workbook.active
    assert [cell.value for cell in sheet[3]] == [
        "推荐日期", "观测起始", "原排名", "股票代码", "股票名称", "来源", "复核分",
        "7月14日当日涨跌", "7月15日当日涨跌", "截至7月17日总涨跌",
    ]
    assert sheet["D4"].value == "000001"
    assert sheet["H4"].value == 0.02
    assert sheet["I4"].value == -0.01
    assert sheet["J4"].value == 0.0098
    assert sheet["J5"].value == "待观测"
    workbook.close()


def test_weekly_key_candidates_use_same_flat_layout_and_keep_all_sources(tmp_path: Path) -> None:
    payload = {
        "run": {
            "evaluation_end_date": date(2026, 7, 17),
            "selection_scope": "KEY_CANDIDATES",
        },
        "members": [
            {
                "selection_trade_date": date(2026, 7, 16),
                "stock_code": "000001.SZ",
                "stock_name": "模型候选",
                "selection_source": "LLM",
                "pro_rank": 1,
                "pro_score": 58.0,
            },
            {
                "selection_trade_date": date(2026, 7, 16),
                "stock_code": "000002.SZ",
                "stock_name": "人工候选",
                "selection_source": "MANUAL",
                "pro_rank": 2,
                "pro_score": 52.0,
            },
            {
                "selection_trade_date": date(2026, 7, 16),
                "stock_code": "000003.SZ",
                "stock_name": "共同候选",
                "selection_source": "BOTH",
                "pro_rank": 3,
                "pro_score": 49.0,
            },
        ],
        "stocks": [
            _stock(date(2026, 7, 16), date(2026, 7, 17), "000001.SZ", 1, 0.01, 0.01),
            _stock(date(2026, 7, 16), date(2026, 7, 17), "000002.SZ", 2, -0.02, -0.02),
            _stock(date(2026, 7, 16), date(2026, 7, 17), "000003.SZ", 3, 0.03, 0.03),
        ],
    }
    output = tmp_path / "本周重点候选对比.xlsx"
    WeeklyStockPerformanceExcelExporter().export(output, payload)
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["本周重点候选对比"]
    sheet = workbook.active
    assert sheet["A1"].value == "本周重点候选表现对比"
    assert [sheet.cell(row, 6).value for row in range(4, 7)] == ["模型选股", "人工选股", "模型+人工"]
    assert [sheet.cell(row, 7).value for row in range(4, 7)] == [58.0, 52.0, 49.0]
    workbook.close()


def test_weekly_key_candidates_deduplicate_by_code_and_keep_earliest_selection(tmp_path: Path) -> None:
    payload = {
        "run": {
            "evaluation_end_date": date(2026, 7, 17),
            "selection_scope": "KEY_CANDIDATES",
        },
        "members": [
            {
                "selection_trade_date": date(2026, 7, 15),
                "stock_code": "000001.SZ",
                "stock_name": "重复候选",
                "selection_source": "MANUAL",
                "pro_rank": 8,
                "pro_score": 55.0,
            },
            {
                "selection_trade_date": date(2026, 7, 13),
                "stock_code": "000001.SZ",
                "stock_name": "重复候选",
                "selection_source": "LLM",
                "pro_rank": 3,
                "pro_score": 63.0,
            },
        ],
        "stocks": [
            _stock(date(2026, 7, 13), date(2026, 7, 14), "000001.SZ", 3, 0.02, 0.02),
            _stock(date(2026, 7, 15), date(2026, 7, 16), "000001.SZ", 8, -0.01, -0.01),
        ],
    }
    output = tmp_path / "本周重点候选去重.xlsx"
    WeeklyStockPerformanceExcelExporter().export(output, payload)
    workbook = load_workbook(output, data_only=False)
    sheet = workbook.active
    assert sheet.max_row == 4
    assert sheet["A4"].value.date() == date(2026, 7, 13)
    assert sheet["C4"].value == 3
    assert sheet["F4"].value == "模型选股"
    assert sheet["G4"].value == 63.0
    workbook.close()


def _stock(
    selection_day: date,
    evaluation_day: date,
    code: str,
    rank: int,
    daily_return: float,
    cumulative_return: float,
) -> dict:
    return {
        "selection_trade_date": selection_day,
        "evaluation_trade_date": evaluation_day,
        "stock_code": code,
        "stock_name": "测试股票",
        "selection_source": "LLM",
        "quant_rank": rank,
        "pro_rank": rank,
        "daily_return": daily_return,
        "cumulative_return": cumulative_return,
    }
