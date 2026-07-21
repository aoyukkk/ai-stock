from __future__ import annotations

import json
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker

from backend.api import market_review as market_review_api
from backend.main import create_app
from database.base import Base
from database.models.market_review import MarketDailySnapshot, MarketReviewRun
from database.models.stock import StockMaster
from database.models.system import LLMUsage
from database.session import create_engine_from_url
from datasource.search_provider import MockMarketSearchProvider, TavilySearchProvider
import market_review.human_excel as human_market_excel
from market_review.evidence import MarketEvidenceSearchService
from market_review.excel import DISCLAIMER, _finish, add_market_review_sheet
from market_review.human_excel import HumanMarketReviewExcelExporter
from market_review.rules import MarketOutlookRuleEngine, MarketRegimeEngine
from market_review.schemas import SearchResult
from market_review.service import MarketReviewService
from market_review.snapshot import MarketDailySnapshotService, _rank_sectors


TRADE_DATE = date(2026, 7, 13)


def _factory(tmp_path: Path):
    engine = create_engine_from_url(f"sqlite:///{(tmp_path / 'market_review.db').as_posix()}")
    import database.models  # noqa: F401

    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _config() -> dict:
    return {
        "enabled": True,
        "schema_version": "market_daily_snapshot_v1",
        "regime_version": "market_regime_v1",
        "outlook_rule_version": "market_outlook_rule_v1",
        "prompt_version": "market_daily_review_prompt_v1",
        "contract_version": "market_daily_review_wire_v1",
        "amount_multiplier": 1000,
        "flat_return_epsilon": 0.0001,
        "search": {"enabled": True, "provider": "MOCK", "max_queries": 4, "max_results_per_query": 5, "max_evidence_items": 10, "max_age_hours": 36, "evidence_min_confidence": 0.55, "multi_source_confirmation_count": 2},
        "pro": {"enabled": False, "model_alias": "controller-high-capability", "max_tokens": 8000},
        "indices": [{"index_code": "000001.SH", "index_name": "上证综指"}],
    }


def _cache(tmp_path: Path) -> Path:
    root = tmp_path / "tushare"
    rows = [
        {"ts_code": "000001.SZ", "trade_date": "20260713", "open": 10, "high": 10.5, "low": 9.9, "close": 10.4, "pre_close": 10, "pct_chg": 4, "vol": 100, "amount": 1000},
        {"ts_code": "000002.SZ", "trade_date": "20260713", "open": 20, "high": 20.2, "low": 19.5, "close": 19.6, "pre_close": 20, "pct_chg": -2, "vol": 100, "amount": 2000},
        {"ts_code": "600000.SH", "trade_date": "20260713", "open": 8, "high": 8.2, "low": 7.9, "close": 8.08, "pre_close": 8, "pct_chg": 1, "vol": 100, "amount": 1500},
        {"ts_code": "600001.SH", "trade_date": "20260713", "open": 12, "high": 12.1, "low": 11.8, "close": 11.88, "pre_close": 12, "pct_chg": -1, "vol": 100, "amount": 1800},
    ]
    _write(root / "trade_date" / "daily" / "20260713.json", rows)
    _write(root / "trade_date" / "daily" / "20260710.json", [{**row, "trade_date": "20260710", "amount": row["amount"] * 0.9} for row in rows])
    _write(root / "trade_date" / "daily_basic" / "20260713.json", [{"ts_code": row["ts_code"], "trade_date": "20260713", "circ_mv": 10000 + index * 1000} for index, row in enumerate(rows)])
    _write(root / "trade_date" / "stk_limit" / "20260713.json", [{"ts_code": row["ts_code"], "trade_date": "20260713", "up_limit": row["pre_close"] * 1.1, "down_limit": row["pre_close"] * 0.9} for row in rows])
    _write(root / "trade_date" / "moneyflow" / "20260713.json", [{"ts_code": row["ts_code"], "trade_date": "20260713", "net_mf_amount": index * 10} for index, row in enumerate(rows)])
    _write(root.parent / "market_review" / "index" / "20260713.json", [{"index_code": "000001.SH", "trade_date": "20260713", "open": 3500, "high": 3530, "low": 3490, "close": 3520, "pre_close": 3500, "pct_chg": 0.5714, "amount": 500000}])
    return root


def _write(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def _snapshot(tmp_path: Path):
    factory = _factory(tmp_path)
    session = factory()
    session.add_all([
        StockMaster(code="000001.SZ", name="平安银行", industry="银行"),
        StockMaster(code="000002.SZ", name="万科A", industry="房地产"),
        StockMaster(code="600000.SH", name="浦发银行", industry="银行"),
        StockMaster(code="600001.SH", name="样本", industry="工业"),
    ])
    session.commit()
    snapshot = MarketDailySnapshotService(session, cache_root=_cache(tmp_path), config=_config()).build(
        TRADE_DATE, datetime(2026, 7, 13, 15, 30, tzinfo=timezone.utc)
    )
    return session, snapshot


def test_snapshot_rules_are_deterministic_and_probabilities_sum_to_100(tmp_path: Path) -> None:
    session, snapshot = _snapshot(tmp_path)
    assert snapshot.trade_date == TRADE_DATE
    assert snapshot.breadth["valid_count"] == 4
    assert snapshot.indices[0].data_status == "AVAILABLE"
    assert snapshot.indices[0].change_percent == pytest.approx(0.005714, abs=1e-6)
    regime = MarketRegimeEngine().evaluate(snapshot)
    outlook = MarketOutlookRuleEngine().evaluate(snapshot, regime)
    assert outlook.base_case_probability + outlook.bull_case_probability + outlook.bear_case_probability == 100
    assert outlook == MarketOutlookRuleEngine().evaluate(snapshot, regime)
    session.close()


def test_evidence_is_normalized_deduplicated_and_audited(tmp_path: Path) -> None:
    session, snapshot = _snapshot(tmp_path)
    result = SearchResult(
        title="交易所发布当日市场运行信息", url="https://www.sse.com.cn/news/item?utm_source=test",
        domain="sse.com.cn", source_name="上海证券交易所",
        publish_time=datetime(2026, 7, 13, 8, tzinfo=timezone.utc), fetched_at=datetime(2026, 7, 13, 9, tzinfo=timezone.utc),
        snippet="交易所发布当日公开市场运行信息。", provider="MOCK", provider_result_id="official-1",
    )
    provider = MockMarketSearchProvider([result, result])
    collected = MarketEvidenceSearchService(provider, _config()["search"]).collect(snapshot, enabled=True)
    assert provider.call_count > 0
    assert collected["external_call_occurred"] is True
    assert collected["search_status"] == "VERIFIED"
    assert collected["stats"]["valid_evidence_count"] == 1
    assert "utm_source" not in str(collected["evidence"][0].url)
    session.close()


def test_tavily_requires_key_and_explicit_runtime_approval(monkeypatch) -> None:
    calls = []
    provider = TavilySearchProvider(api_key="test-only-key", allow_external_calls=False, transport=lambda body, headers: calls.append((body, headers)) or {"results": []})
    monkeypatch.delenv("SEARCH_REAL_CALLS_ENABLED", raising=False)
    with pytest.raises(RuntimeError, match="REAL_SEARCH_CALLS_NOT_APPROVED"):
        provider.search("测试", datetime.now(timezone.utc), datetime.now(timezone.utc), "zh-CN", 5)
    assert calls == []


def test_service_persists_cache_hit_and_never_calls_llm_in_data_only_mode(tmp_path: Path) -> None:
    factory = _factory(tmp_path)
    session = factory()
    cache = _cache(tmp_path)
    service = MarketReviewService(session, config=_config(), cache_root=cache, search_provider=MockMarketSearchProvider())
    first = service.run(TRADE_DATE, mode="DATA_ONLY")
    second = service.run(TRADE_DATE, mode="DATA_ONLY")
    assert first["run"]["status"] == "DATA_ONLY"
    assert second["run"]["cache_status"] == "SUCCESS_CACHE_HIT"
    assert session.scalar(select(func.count()).select_from(MarketDailySnapshot)) == 1
    assert session.scalar(select(func.count()).select_from(MarketReviewRun)) == 1
    assert session.scalar(select(func.count()).select_from(LLMUsage)) == 0
    session.close()
    fresh = factory()
    assert fresh.scalar(select(func.count()).select_from(MarketReviewRun)) == 1
    fresh.close()


def test_market_review_api_returns_persisted_bundle(tmp_path: Path, monkeypatch) -> None:
    factory = _factory(tmp_path)
    session = factory()
    MarketReviewService(session, config=_config(), cache_root=_cache(tmp_path), search_provider=MockMarketSearchProvider()).run(TRADE_DATE, mode="DATA_ONLY")
    session.close()
    monkeypatch.setattr(market_review_api, "_session", lambda: factory())
    client = TestClient(create_app())
    response = client.get("/api/workbench/market-review/latest", params={"trade_date": TRADE_DATE.isoformat()})
    assert response.status_code == 200
    assert response.json()["data"]["run"]["trade_date"] == TRADE_DATE.isoformat()
    assert client.get("/api/workbench/market-review/methodology").json()["data"]["probability_owner"] == "DETERMINISTIC_RULE_ENGINE"


def test_market_review_excel_sheet_is_chinese_centered_and_wrapped(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_market_review_sheet(workbook, None)
    output = tmp_path / "review.xlsx"
    workbook.save(output)
    workbook.close()
    loaded = load_workbook(output)
    assert loaded.sheetnames == ["06_大盘复盘"]
    sheet = loaded["06_大盘复盘"]
    assert DISCLAIMER in " ".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
    assert all(cell.alignment.horizontal == "center" and cell.alignment.vertical == "center" and cell.alignment.wrap_text for row in sheet.iter_rows() for cell in row if cell.value is not None)
    loaded.close()


def test_market_review_excel_formats_all_percentage_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["行业", "涨跌幅"])
    sheet.append(["银行", 0.01])
    sheet.append(["工业", -0.02])
    sheet.append(["科技", 0.035])
    _finish(sheet)
    assert [sheet.cell(row=row, column=2).number_format for row in range(2, 5)] == ["0.00%"] * 3
    workbook.close()


def test_industry_ranking_keeps_all_industries_for_full_review() -> None:
    grouped = {
        f"行业{index:02d}": [{
            "ts_code": f"{index:06d}.SZ",
            "pct_chg": index - 12,
            "amount": 1000 + index,
        }]
        for index in range(25)
    }
    full = _rank_sectors(grouped, "INDUSTRY", set(), 1000, include_all=True)
    summary = _rank_sectors(grouped, "CONCEPT", set(), 1000)
    assert len(full) == 25
    assert [item.rank for item in full] == list(range(1, 26))
    assert full[0].change_percent >= full[-1].change_percent
    assert len(summary) == 20


def test_human_market_review_exports_summary_and_mobile_safe_full_industries(tmp_path: Path) -> None:
    industries = [
        {
            "rank": index + 1,
            "sector_code": f"I{index:03d}",
            "sector_name": f"行业{index:03d}",
            "change_percent": (12 - index) / 100,
            "advancing_ratio": max(0, min(1, (25 - index) / 25)),
            "limit_up_count": index % 3,
            "member_count": 10 + index,
            "amount": 100_000_000 + index,
        }
        for index in range(25)
    ]
    bundle = {
        "run": {
            "trade_date": "2026-07-15",
            "confidence": 0.8,
            "base_case_probability": 57,
            "bull_case_probability": 23,
            "bear_case_probability": 20,
        },
        "snapshot": {
            "breadth": {"advancing_count": 10, "declining_count": 8, "flat_count": 2, "advancing_ratio": 0.5, "equal_weight_return": 0.01},
            "turnover": {"total_amount": 1_000_000_000, "change_ratio": -0.02, "relative_to_5d": 0.95},
            "limit_structure": {"limit_up_count": 2, "limit_down_count": 1, "failed_limit_up_ratio": 0.1},
            "capital": {"net_main_inflow": -100_000_000},
            "data_quality_score": 90,
            "industries": industries,
        },
        "outlook": {},
    }
    supplement = {
        "indices": [
            {"name": "上涨指数", "close": 100, "change": 0.01, "feature": "上涨", "source": "测试", "date": "2026-07-15"},
            {"name": "下跌指数", "close": 100, "change": -0.01, "feature": "下跌", "source": "测试", "date": "2026-07-15"},
        ],
        "causes": [],
    }
    output = tmp_path / "human_market_review.xlsx"
    HumanMarketReviewExcelExporter().export(output, bundle, supplement)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["大盘概览", "行业轮动", "全部行业", "盘面原因", "次日观察"]
    assert workbook["行业轮动"].max_row == 23
    assert workbook["全部行业"].max_row == 28
    assert workbook["全部行业"]["D4"].value == pytest.approx(0.12)
    assert workbook["全部行业"]["D28"].value == pytest.approx(-0.12)
    for coordinate, expected_color in (("D4", "C00000"), ("D28", "008000")):
        cell = workbook["全部行业"][coordinate]
        assert cell.number_format == "0.00%"
        assert cell.font.color.type == "rgb"
        assert cell.font.color.rgb.startswith("FF")
        assert cell.font.color.rgb[-6:] == expected_color
        assert cell.fill.fill_type == "solid"
        assert cell.fill.fgColor.rgb.startswith("FF")
        assert cell.value is not None
    title = workbook["行业轮动"]["A1"]
    header = workbook["行业轮动"]["A3"]
    assert title.fill.fgColor.rgb.startswith("FF") and title.font.color.rgb.startswith("FF")
    assert header.fill.fgColor.rgb.startswith("FF") and header.font.color.rgb.startswith("FF")
    assert title.font.color.rgb == "FF17365D"
    assert header.font.color.rgb == "FF17365D"
    assert all("[Red]" not in cell.number_format and "[Green]" not in cell.number_format for sheet in workbook for row in sheet.iter_rows() for cell in row)
    workbook.close()


def test_human_market_review_uses_update_file_when_canonical_is_locked(tmp_path: Path, monkeypatch) -> None:
    output = tmp_path / "大盘复盘.xlsx"
    bundle = {
        "run": {"trade_date": "2026-07-15", "base_case_probability": 57, "bull_case_probability": 23, "bear_case_probability": 20},
        "snapshot": {"breadth": {}, "turnover": {}, "limit_structure": {}, "capital": {}, "industries": []},
        "outlook": {},
    }
    real_replace = human_market_excel.os.replace
    attempts = []

    def replace_with_first_target_locked(source, target):
        attempts.append(Path(target))
        if Path(target) == output:
            raise PermissionError("test lock")
        return real_replace(source, target)

    monkeypatch.setattr(human_market_excel.os, "replace", replace_with_first_target_locked)
    result = HumanMarketReviewExcelExporter().export(output, bundle, {"indices": [], "causes": []})
    actual = Path(result["output"])
    assert result["canonical_replaced"] is False
    assert actual.name == "大盘复盘_更新版.xlsx"
    assert actual.is_file()
    assert attempts[0] == output
