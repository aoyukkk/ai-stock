from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from database.models.midday_v22 import MiddayV22AfternoonRun
from midday.v22_afternoon_export import SHEETS, export_v22_afternoon
from midday.v22_afternoon_service import MiddayV22AfternoonRecheckService, _afternoon_cutoff, _five_minute_structure, _resolve_recheck_layer, _session_vwap

SH = ZoneInfo("Asia/Shanghai")


def _bars():
    rows = []
    for minute in range(10):
        rows.append({"bar_time": f"2026-07-20T13:{minute:02d}:00+08:00", "open": 10, "high": 10.2, "low": 9.9, "close": 10 + minute / 100, "volume": 100, "amount": (10 + minute / 100) * 100})
    return rows


def test_afternoon_vwap_and_five_minute_structure_are_derived_from_afternoon_only():
    bars = [{"bar_time": "2026-07-20T11:30:00+08:00", "close": 99, "volume": 100, "amount": 9900}] + _bars()
    assert _session_vwap(bars, datetime.strptime("13:00", "%H:%M").time()) < 11
    structure = _five_minute_structure(bars)
    assert structure["bar_count"] == 2 and structure["trend"] == "UP"


def test_afternoon_cutoff_is_second_precision_for_ifind_contract():
    value=_afternoon_cutoff(datetime(2026,7,20,13,11,22,987654,tzinfo=SH))
    assert value.isoformat()=="13:11:22"


def test_watch_upgrades_only_when_every_condition_and_trigger_pass():
    conditions = {key: True for key in ("regime_stable", "price_acceptable", "volume_confirmed", "sector_relative_confirmed", "structure_valid", "no_new_hard_risk", "llm_not_blocked")}
    assert _resolve_recheck_layer("AFTERNOON_WATCH", "ENTRY_TRIGGERED", conditions, current_price=10, stop_loss=9) == ("BUY_READY", "UPGRADED")
    conditions["volume_confirmed"] = False
    assert _resolve_recheck_layer("AFTERNOON_WATCH", "ENTRY_TRIGGERED", conditions, current_price=10, stop_loss=9) == ("AFTERNOON_WATCH", "WATCH")


def test_watch_rejects_on_hard_trigger_or_stop_break():
    conditions = {"ok": True}
    assert _resolve_recheck_layer("AFTERNOON_WATCH", "TRIGGER_REJECTED", conditions, current_price=10, stop_loss=9) == ("BLOCKED", "REJECTED")
    assert _resolve_recheck_layer("AFTERNOON_WATCH", "WAITING_TRIGGER", conditions, current_price=8, stop_loss=9) == ("BLOCKED", "REJECTED")


def test_before_1300_fails_and_persists_run(tmp_path: Path):
    engine = create_engine("sqlite:///:memory:")
    MiddayV22AfternoonRun.__table__.create(engine)
    with Session(engine, expire_on_commit=False) as session:
        result = MiddayV22AfternoonRecheckService(session, output_root=tmp_path).run(date(2026, 7, 20), ["600101.SH"], now=datetime(2026, 7, 20, 12, 30, tzinfo=SH))
        assert result["status"] == "FAILED" and result["error_code"] == "AFTERNOON_SESSION_NOT_STARTED"
        persisted = session.scalar(select(MiddayV22AfternoonRun).where(MiddayV22AfternoonRun.run_id == result["run_id"]))
        assert persisted is not None and persisted.status == "FAILED" and persisted.real_orders == 0


def test_afternoon_excel_reads_final_database_status(tmp_path: Path):
    run = SimpleNamespace(run_id="recheck-test-12345678", trade_date=date(2026, 7, 20), status="AFTERNOON_RECHECK_COMPLETE")
    report = {"run_id": run.run_id, "midday_run_id": "midday", "recheck_time": "2026-07-20T13:30:00+08:00", "status": run.status, "previous_midday_regime": "REPAIR", "afternoon_regime": "REPAIR", "results": [], "counts": {"result_layers": {}}, "provider_audit": [], "real_orders": 0, "virtual_orders": 0}
    paths = export_v22_afternoon(tmp_path, run, report)
    workbook = load_workbook(paths["excel"])
    try:
        assert workbook.sheetnames == list(SHEETS)
        summary = workbook[SHEETS[0]]
        status_column = next(cell.column for cell in summary[1] if cell.value == "最终状态")
        assert summary.cell(2, status_column).value == "AFTERNOON_RECHECK_COMPLETE"
    finally: workbook.close()
