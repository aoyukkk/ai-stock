from pathlib import Path

from openpyxl import load_workbook

from review.performance_excel import PerformanceExcelExporter


def test_performance_excel_has_four_centered_wrapped_sheets(tmp_path: Path) -> None:
    output = tmp_path / "selection_performance_2026-07-10_5td.xlsx"
    row = {"selection_trade_date": "2026-07-03", "stock_code": "000001", "daily_return": 0.01, "cumulative_return": 0.021}
    result = PerformanceExcelExporter().export(output, {"cohorts": [row], "daily": [row], "stocks": [row], "methodology": {"statement": "本统计反映选股后的市场价格表现，不代表实际成交收益或交易建议。"}})
    assert result["sheet_count"] == 4
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["01_选股日汇总", "02_组合每日涨跌", "03_个股收益明细", "04_口径与异常"]
    assert workbook.active.title == "01_选股日汇总"
    for sheet in workbook.worksheets:
        for row_cells in sheet.iter_rows(min_row=4):
            for cell in row_cells:
                assert cell.alignment.horizontal == "center"
                assert cell.alignment.vertical == "center"
                assert cell.alignment.wrap_text is True
    stock_sheet = workbook["03_个股收益明细"]
    headers = {cell.value: cell.column for cell in stock_sheet[4]}
    assert stock_sheet.cell(5, headers["stock_code"]).number_format == "@"
    assert stock_sheet.cell(5, headers["daily_return"]).number_format == "0.00%"
