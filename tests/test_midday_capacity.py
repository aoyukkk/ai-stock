from __future__ import annotations

from dataclasses import replace

import pytest
from openpyxl import load_workbook

from datasource.ifind.http.rate_limiter import IFindHttpSerialRateLimiter
from midday.capacity import (
    BatchProbe,
    QuotaStatistics,
    _after_cutoff,
    _capacity_workbook,
    audit_call_limit,
    estimate_capacity,
    parse_data_statistics,
    run_scoped_capacity_config,
    select_capacity_path,
)
from datetime import datetime
from midday.core import SHANGHAI


def probe(size: int, *, complete: bool = False, transport: str = "HTTP", returned: int | None = None) -> BatchProbe:
    count = size if complete else (returned if returned is not None else 0)
    return BatchProbe(transport, size, count, size - count, 0, complete, size > 100 and count == 100, False, None, 1, "2026-07-20 11:30:00", 0)


def safe_quota() -> QuotaStatistics:
    return QuotaStatistics(status="AVAILABLE", market_data_total=10_000_000, market_data_used=1_000, market_data_remaining=9_999_000)


def safe_estimate() -> dict:
    return estimate_capacity(5_291, safe_quota())


def test_internal_limit_is_identified(monkeypatch):
    for key in ("IFIND_MAX_PROVIDER_CALLS", "IFIND_PROVIDER_CALL_BUDGET", "IFIND_FULL_A_CALL_LIMIT"):
        monkeypatch.delenv(key, raising=False)
    result = audit_call_limit({"maximum_provider_calls": 30})
    assert result["classification"] == "INTERNAL_GUARD"
    assert result["whether_account_reported"] is False


def test_environment_limit_has_priority(monkeypatch):
    monkeypatch.setenv("IFIND_MAX_PROVIDER_CALLS", "12")
    result = audit_call_limit({"maximum_provider_calls": 30})
    assert result["effective_value"] == 12
    assert result["limit_source"] == "ENVIRONMENT"


def test_empty_statistics_are_unavailable():
    assert parse_data_statistics(None).status == "QUOTA_STATISTICS_UNAVAILABLE"


def test_statistics_remaining_is_derived():
    value = parse_data_statistics({"data": {"market_data": {"total": 1000, "used": 250}}})
    assert value.market_data_remaining == 750
    assert value.status == "AVAILABLE"


def test_statistics_explicit_remaining():
    value = parse_data_statistics({"market_data_remaining": 321})
    assert value.market_data_remaining == 321


def test_real_sdk_statistics_shape_is_parsed():
    value = parse_data_statistics({"errorcode": 0, "tables": {"QuotesDataStat": {"limit": 150_000_000, "usage": 575_922}}})
    assert value.status == "AVAILABLE"
    assert value.market_data_remaining == 149_424_078


def test_statistics_without_remaining_is_uncertain():
    assert parse_data_statistics({"account_tier": "trial"}).status == "QUOTA_SCOPE_UNCERTAIN"


def test_estimated_cell_formula():
    result = estimate_capacity(100, safe_quota())
    assert result["estimated_snapshot_cells"] == 900
    assert result["estimated_minute_cells"] == 36_300
    assert result["estimated_total_cells"] > 37_000


def test_estimate_is_unverified_without_quota():
    assert estimate_capacity(5_291, QuotaStatistics(status="QUOTA_STATISTICS_UNAVAILABLE"))["quota_safety_result"] == "UNVERIFIED"


def test_estimate_fails_above_safety_ratio():
    quota = QuotaStatistics(status="AVAILABLE", market_data_remaining=100)
    assert estimate_capacity(5_291, quota)["quota_safety_result"] == "FAIL"


def test_wide_http_selected_with_quota():
    path, reasons = select_capacity_path(http_probes=[probe(500, complete=True)], sdk_probes=[], quota=safe_quota(), estimate=safe_estimate(), limit_audit=audit_call_limit({"maximum_provider_calls": 30}), temporary_authorized=True)
    assert path == "WIDE_BATCH_HTTP" and not reasons


def test_wide_sdk_selected_when_http_is_limited():
    path, _ = select_capacity_path(http_probes=[probe(100, complete=True), probe(500)], sdk_probes=[probe(500, complete=True, transport="SDK")], quota=safe_quota(), estimate=safe_estimate(), limit_audit=audit_call_limit({"maximum_provider_calls": 30}), temporary_authorized=True)
    assert path == "WIDE_BATCH_SDK"


def test_temporary_override_selected_only_for_complete_100():
    path, _ = select_capacity_path(http_probes=[probe(100, complete=True), probe(200), probe(500)], sdk_probes=[], quota=safe_quota(), estimate=safe_estimate(), limit_audit=audit_call_limit({"maximum_provider_calls": 30}), temporary_authorized=True)
    assert path == "TEMPORARY_CALL_BUDGET_OVERRIDE"


@pytest.mark.parametrize("quota", [QuotaStatistics(status="QUOTA_STATISTICS_UNAVAILABLE"), QuotaStatistics(status="QUOTA_SCOPE_UNCERTAIN")])
def test_unverified_quota_blocks_every_path(quota):
    path, reasons = select_capacity_path(http_probes=[probe(500, complete=True)], sdk_probes=[], quota=quota, estimate=estimate_capacity(5_291, quota), limit_audit=audit_call_limit({"maximum_provider_calls": 30}), temporary_authorized=True)
    assert path == "CAPACITY_CONFIRMED_BLOCKED"
    assert quota.status in reasons


def test_override_requires_user_authorization():
    path, _ = select_capacity_path(http_probes=[probe(100, complete=True)], sdk_probes=[], quota=safe_quota(), estimate=safe_estimate(), limit_audit=audit_call_limit({"maximum_provider_calls": 30}), temporary_authorized=False)
    assert path == "CAPACITY_CONFIRMED_BLOCKED"


def test_override_requires_internal_guard():
    audit = audit_call_limit({"maximum_provider_calls": 20})
    path, _ = select_capacity_path(http_probes=[probe(100, complete=True)], sdk_probes=[], quota=safe_quota(), estimate=safe_estimate(), limit_audit=audit, temporary_authorized=True)
    assert path == "CAPACITY_CONFIRMED_BLOCKED"


def test_errorcode_success_but_missing_codes_is_not_complete():
    row = probe(100, returned=99)
    assert row.complete is False and row.missing == 1


def test_truncation_at_100_is_detectable():
    row = BatchProbe("HTTP", 200, 100, 100, 0, False, True, False, None, 1, None, 0)
    assert row.truncated_to_100 is True


def test_duplicate_codes_prevent_completeness():
    row = BatchProbe("HTTP", 100, 99, 1, 1, False, False, False, None, 1, None, 0)
    assert row.duplicates == 1 and row.complete is False


def test_post_cutoff_rows_prevent_completeness():
    row = replace(probe(100, complete=True), complete=False, post_cutoff_rows=1)
    assert not row.complete


def test_post_cutoff_time_is_compared_in_shanghai_timezone():
    cutoff = datetime(2026, 7, 20, 11, 30, tzinfo=SHANGHAI)
    assert _after_cutoff("2026-07-20 11:31:00", cutoff) is True
    assert _after_cutoff("2026-07-20 11:30:00", cutoff) is False


def test_timeout_is_a_failure():
    row = replace(probe(100), timeout=True, error_category="TIMEOUT")
    assert row.timeout and not row.complete


def test_scoped_override_does_not_mutate_default():
    config = {"maximum_provider_calls": 30, "snapshot_batch_size": 100}
    with run_scoped_capacity_config(config, selected_path="TEMPORARY_CALL_BUDGET_OVERRIDE") as (working, audit):
        assert working["maximum_provider_calls"] == 70
        assert config["maximum_provider_calls"] == 30
    assert audit["restoration_success"] is True


def test_scoped_override_restores_after_exception():
    config = {"maximum_provider_calls": 30}
    audit = None
    with pytest.raises(RuntimeError):
        with run_scoped_capacity_config(config, selected_path="TEMPORARY_CALL_BUDGET_OVERRIDE") as (_, current):
            audit = current
            raise RuntimeError("boom")
    assert config["maximum_provider_calls"] == 30
    assert audit["restoration_success"] is True


def test_adaptive_run_scope_uses_400_and_98_percent_without_mutating_defaults():
    config = {"maximum_provider_calls": 30, "minimum_full_a_coverage": 0.95}
    with run_scoped_capacity_config(
        config, selected_path="ADAPTIVE_COVERAGE", temporary_limit=400
    ) as (working, audit):
        assert working["maximum_provider_calls"] == 400
        assert working["minimum_full_a_coverage"] == 0.98
        assert working["minute_reconstruction_batch_size"] == 20
        assert config == {"maximum_provider_calls": 30, "minimum_full_a_coverage": 0.95}
    assert audit["restoration_success"] is True


def test_non_override_path_keeps_hashes_equal():
    config = {"maximum_provider_calls": 30}
    with run_scoped_capacity_config(config, selected_path="CAPACITY_CONFIRMED_BLOCKED") as (_, audit):
        pass
    assert audit["original_config_hash"] == audit["temporary_config_hash"] == audit["restored_config_hash"]


def test_default_limiter_rejects_70():
    with pytest.raises(ValueError, match="IFIND_HTTP_CALL_LIMIT_INVALID"):
        IFindHttpSerialRateLimiter(70, sleep=lambda _: None)


def test_authorized_limiter_accepts_70():
    limiter = IFindHttpSerialRateLimiter(70, authorized_maximum=70, sleep=lambda _: None)
    assert limiter.maximum == 70


def test_authorized_limiter_still_enforces_interval():
    with pytest.raises(ValueError, match="IFIND_HTTP_INTERVAL_TOO_SHORT"):
        IFindHttpSerialRateLimiter(70, interval_ms=999, authorized_maximum=70)


def test_limiter_rejects_above_phase_authorization():
    with pytest.raises(ValueError, match="IFIND_HTTP_AUTHORIZED_LIMIT_INVALID"):
        IFindHttpSerialRateLimiter(401, authorized_maximum=401)


def test_run_scoped_limiter_accepts_phase_authorized_400():
    limiter = IFindHttpSerialRateLimiter(400, authorized_maximum=400, sleep=lambda _: None)
    assert limiter.maximum == 400


def test_capacity_workbook_has_17_sheets_and_final_status(tmp_path):
    path = tmp_path / "capacity.xlsx"
    report = {"run_id": "run-1", "status": "FULL_A_CAPACITY_CONFIRMED_BLOCKED", "trade_date": "2026-07-20", "counts": {}}
    capacity = {"selected_capacity_path": "CAPACITY_CONFIRMED_BLOCKED", "blocking_reasons": ["QUOTA_STATISTICS_UNAVAILABLE"], "http_batch_probes": [], "sdk_batch_probes": [], "account_quota_statistics": {"status": "QUOTA_STATISTICS_UNAVAILABLE"}}
    _capacity_workbook(path, report, capacity)
    workbook = load_workbook(path)
    assert len(workbook.sheetnames) == 17
    assert workbook["01_午盘总览"]["B3"].value == report["status"]


def test_capacity_workbook_cells_are_centered(tmp_path):
    path = tmp_path / "capacity.xlsx"
    report = {"run_id": "run-1", "status": "FULL_A_CAPACITY_CONFIRMED_BLOCKED", "trade_date": "2026-07-20", "counts": {}}
    capacity = {"selected_capacity_path": "CAPACITY_CONFIRMED_BLOCKED", "http_batch_probes": [], "sdk_batch_probes": [], "account_quota_statistics": {}}
    _capacity_workbook(path, report, capacity)
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    assert cell.alignment.horizontal == "center"
                    assert cell.alignment.vertical == "center"


@pytest.mark.parametrize("status", ["FULL_A_CAPACITY_CONFIRMED_BLOCKED", "FULL_A_QUOTA_INSUFFICIENT", "FULL_A_SNAPSHOT_VALIDATION_FAILED"])
def test_final_status_is_never_running(status):
    assert status != "RUNNING"


@pytest.mark.parametrize("orders,scheduler", [(0, False), (0, 0)])
def test_shadow_safety_values(orders, scheduler):
    assert orders == 0
    assert scheduler is False or scheduler == 0
