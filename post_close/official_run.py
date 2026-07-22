from __future__ import annotations

import hashlib
import json
import os
import shutil
import time as time_module
import uuid
from argparse import Namespace
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Callable
from zoneinfo import ZoneInfo

from sqlalchemy import func, select

from backend.workbench.service import WorkbenchService
from database.models.postclose_official import PostCloseOfficialRun
from database.models.quant_run import QuantRankResult, QuantRun
from database.models.system import LLMUsage
from database.models.validation import (
    ModelValidationLLMAudit,
    ModelValidationRun,
    ModelValidationSample,
    ProCandidateReview,
    ProResumeRun,
)
from database.models.workbench import ManualSelectionRecord
from database.session import get_session, init_db
from post_close.seven_day_comparison import SevenDayComparisonService
from post_close.shadow_integration import PostCloseShadowIntegrationService
from reporting.workbook_standard import validate_trading_assistant_workbook
from reporting.workbook_style import WorkbookStyleService
from scripts.prewarm_tushare_trade_date_cache import run_prewarm
from scripts.run_daily_routine import (
    _build_monitor,
    _ensure_flash,
    _ensure_manual,
    _ensure_market,
    _ensure_pro,
    _ensure_quant,
    _export_daily,
)


SHANGHAI = ZoneInfo("Asia/Shanghai")
DATASETS = ("daily", "daily_basic", "moneyflow", "stk_limit", "adj_factor")
OFFICIAL_SUCCESS = {"POSTCLOSE_FULL_A_SUCCESS", "POSTCLOSE_FULL_A_PARTIAL_SUCCESS"}
ALLOWED_STATUSES = OFFICIAL_SUCCESS | {
    "TUSHARE_TEMPORAL_GATE_FAILED",
    "POSTCLOSE_QUANT_FAILED",
    "POSTCLOSE_FLASH_FAILED",
    "POSTCLOSE_PRO_FAILED",
    "POSTCLOSE_EXPORT_FAILED",
    "POSTCLOSE_INTEGRATION_FAILED",
}


class TushareTemporalGate:
    def __init__(self, root: Path | str, trade_date: date, previous_date: date) -> None:
        self.root = Path(root).resolve()
        self.cache = self.root / "data" / "cache" / "tushare" / "trade_date"
        self.trade_date = trade_date
        self.previous_date = previous_date
        self.previous_daily_codes = set(self._codes("daily", previous_date))

    def check(self) -> dict[str, Any]:
        datasets = {name: self._check_dataset(name) for name in DATASETS}
        return {
            "checked_at": datetime.now(SHANGHAI).isoformat(),
            "trade_date": self.trade_date.isoformat(),
            "previous_trade_date": self.previous_date.isoformat(),
            "datasets": datasets,
            "passed": all(value["passed"] for value in datasets.values()),
            "missing_or_failed": [name for name, value in datasets.items() if not value["passed"]],
        }

    def _check_dataset(self, dataset: str) -> dict[str, Any]:
        records = _records(self._path(dataset, self.trade_date))
        previous = _records(self._path(dataset, self.previous_date))
        codes = [str(row.get("ts_code") or "") for row in records if row.get("ts_code")]
        unique = set(codes)
        previous_codes = {str(row.get("ts_code")) for row in previous if row.get("ts_code")}
        dates = {str(row.get("trade_date") or "") for row in records}
        duplicate_count = max(0, len(codes) - len(unique))
        eligible_base = self.previous_daily_codes
        eligible_coverage = len(unique & eligible_base) / len(eligible_base) if eligible_base else 0.0
        previous_ratio = len(unique) / len(previous_codes) if previous_codes else 0.0
        required_fields = {
            "daily": ("ts_code", "trade_date", "open", "high", "low", "close"),
            "daily_basic": ("ts_code", "trade_date", "close", "turnover_rate", "total_mv"),
            "moneyflow": ("ts_code", "trade_date"),
            "stk_limit": ("ts_code", "trade_date", "up_limit", "down_limit"),
            "adj_factor": ("ts_code", "trade_date", "adj_factor"),
        }[dataset]
        invalid_rows = sum(
            any(row.get(field) is None or row.get(field) == "" for field in required_fields)
            for row in records
        )
        field_completeness = 1 - invalid_rows / len(records) if records else 0.0
        coverage_threshold = 0.97 if dataset in {"daily", "daily_basic", "adj_factor"} else 0.80
        previous_threshold = 0.97 if dataset in {"daily", "daily_basic", "adj_factor", "stk_limit"} else 0.80
        eligible_required = dataset in {"daily", "daily_basic", "adj_factor"}
        passed = bool(
            records
            and dates == {f"{self.trade_date:%Y%m%d}"}
            and duplicate_count == 0
            and field_completeness >= 0.99
            and previous_ratio >= previous_threshold
            and (eligible_coverage >= coverage_threshold if eligible_required else True)
        )
        return {
            "dataset": dataset,
            "passed": passed,
            "path": str(self._path(dataset, self.trade_date)),
            "row_count": len(records),
            "unique_code_count": len(unique),
            "previous_unique_code_count": len(previous_codes),
            "eligible_reference_count": len(eligible_base),
            "eligible_coverage": eligible_coverage,
            "previous_day_ratio": previous_ratio,
            "duplicate_count": duplicate_count,
            "invalid_required_field_rows": invalid_rows,
            "field_completeness": field_completeness,
            "trade_dates": sorted(dates),
            "coverage_threshold": coverage_threshold if eligible_required else None,
            "previous_day_threshold": previous_threshold,
        }

    def _path(self, dataset: str, day: date) -> Path:
        return self.cache / dataset / f"{day:%Y%m%d}.json"

    def _codes(self, dataset: str, day: date) -> list[str]:
        return [str(row.get("ts_code")) for row in _records(self._path(dataset, day)) if row.get("ts_code")]


class PostCloseOfficialRunner:
    def __init__(self, root: Path | str, trade_date: date, reference_path: Path | str) -> None:
        self.root = Path(root).resolve()
        self.trade_date = trade_date
        self.reference_path = Path(reference_path).resolve()
        self.output_dir = self.root / "outputs" / trade_date.isoformat() / "正式日线"
        self.audit_dir = self.output_dir
        self.log_dir = self.root / "logs"
        self.run_id = f"postclose-full-a-{uuid.uuid4().hex[:20]}"
        self.provider_calls = 0
        self.cache_hits = 0
        self.temporal_attempts: list[dict[str, Any]] = []

    def dry_run(self) -> dict[str, Any]:
        session = get_session()
        try:
            settings = WorkbenchService(session).settings()
            comparison = SevenDayComparisonService(session, self.root, self.reference_path)
            days = comparison.trading_days(self.trade_date)
            gate = TushareTemporalGate(self.root, self.trade_date, days[-2]).check()
            payload = comparison.build_payload(days)
        finally:
            session.close()
        weights = _quant_weights(self.root)
        if weights != {"technical": 0.25, "capital": 0.25, "emotion": 0.20, "momentum": 0.15, "risk": 0.15}:
            raise RuntimeError("QUANT_WEIGHTS_CHANGED")
        return {
            "status": "DRY_RUN_READY",
            "trade_date": self.trade_date.isoformat(),
            "resolved_active_config": _active_config(settings),
            "reference_workbook": str(self.reference_path),
            "reference_workbook_hash": hashlib.sha256(self.reference_path.read_bytes()).hexdigest(),
            "expected_sheet_structure": list(load_sheet_names(self.reference_path)),
            "seven_trading_days": [value.isoformat() for value in days],
            "historical_missing": payload["missing_historical_runs"],
            "temporal_gate_local_fixture": gate,
            "per_stock_api_call_count": 0,
            "no_external_calls": True,
            "no_llm_calls": True,
            "scheduler": False,
            "real_orders": 0,
            "virtual_orders": 0,
        }

    def run(self) -> dict[str, Any]:
        now = datetime.now(SHANGHAI)
        if now < datetime.combine(self.trade_date, time(17, 0), SHANGHAI):
            raise RuntimeError("POSTCLOSE_CANNOT_RUN_BEFORE_1700")
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.log_dir.mkdir(parents=True, exist_ok=True)
        init_db()
        started = now
        config_hash_before = _config_hash(self.root)
        llm_before = self._llm_usage()
        self._create_run(started)
        report: dict[str, Any] = {
            "phase": f"{self.trade_date.isoformat()} Post-Close Full-A Official Run",
            "run_id": self.run_id,
            "trade_date": self.trade_date.isoformat(),
            "execution_start": started.isoformat(),
            "reference_workbook": str(self.reference_path),
            "template_hash": hashlib.sha256(self.reference_path.read_bytes()).hexdigest(),
            "real_orders": 0,
            "virtual_orders": 0,
            "scheduler": False,
            "production_config_changed": False,
            "ifind_full_a_calls": 0,
            "warnings": [],
            "failures": [],
        }
        stage = "TEMPORAL_GATE"
        try:
            calendar_session = get_session()
            try:
                days = SevenDayComparisonService(
                    calendar_session, self.root, self.reference_path
                ).trading_days(self.trade_date)
            finally:
                calendar_session.close()
            previous = days[-2]
            gate = self._run_temporal_gate(previous)
            report["temporal_gate"] = gate
            if not gate["passed"]:
                raise RuntimeError("TUSHARE_TEMPORAL_GATE_FAILED")
            stage = "QUANT"
            quant_stage = _ensure_quant(self.trade_date)
            stage = "MANUAL"
            manual_stage = self._today_manual_pool()
            stage = "FLASH"
            flash_stage = _ensure_flash(self.trade_date)
            stage = "PRO"
            pro_stage = _ensure_pro(self.trade_date)
            stage = "MARKET"
            market_stage = _ensure_market(self.trade_date)
            report["stages"] = {
                "quant": quant_stage,
                "manual": manual_stage,
                "flash": flash_stage,
                "pro": pro_stage,
                "market": market_stage,
            }
            report.update(self._business_readback())
            stage = "SHADOW"
            shadow_session = get_session()
            try:
                report["shadow"] = PostCloseShadowIntegrationService(shadow_session).run(
                    trade_date=self.trade_date,
                    quant_run_id=str((report.get("quant") or {}).get("run_id") or ""),
                    available_at_ts=datetime.fromisoformat(str(gate["passed_at"])),
                    baseline_snapshot={
                        "quant": report.get("quant"),
                        "flash": report.get("flash"),
                        "pro": report.get("pro"),
                        "final_candidate_count": report.get("final_candidate_count"),
                        "final_candidate_list": report.get("final_candidate_list"),
                    },
                )
            except Exception as shadow_exc:
                shadow_session.rollback()
                report["shadow"] = {
                    "status": "FAILED_CLOSED",
                    "mode": "PARALLEL_READ_ONLY_SHADOW",
                    "formal_result_affected": False,
                    "error_code": type(shadow_exc).__name__,
                    "error_message": _safe_error(shadow_exc),
                }
                report["warnings"].append(f"SHADOW_FAILED_CLOSED:{type(shadow_exc).__name__}")
            finally:
                shadow_session.close()
            has_partial_failures = bool(
                (report.get("flash") or {}).get("failure")
                or (report.get("pro") or {}).get("failure")
            )
            final_status = (
                "POSTCLOSE_FULL_A_PARTIAL_SUCCESS"
                if has_partial_failures
                else "POSTCLOSE_FULL_A_SUCCESS"
            )
            if not report.get("final_candidate_count"):
                report["warnings"].append("FINAL_CANDIDATE_EMPTY_BY_RULES")
            self._persist_final_before_export(final_status, report)
            reread = self._read_run_status()
            if reread != final_status:
                raise RuntimeError("FINAL_STATUS_REREAD_MISMATCH")
            stage = "EXPORT"
            export = _export_daily(self.trade_date, None)
            generated_workbook = Path(export["daily_workbook"])
            main_workbook = self.output_dir / generated_workbook.name
            if main_workbook.exists():
                raise FileExistsError(f"OUTPUT_EXISTS:{main_workbook}")
            shutil.copy2(generated_workbook, main_workbook)
            self.style.annotate_official_status(main_workbook, final_status)
            structure = self.style.compare_official_structure(self.reference_path, main_workbook)
            if not all(structure[key] for key in ("sheet_names_match", "sheet_order_match", "headers_match")):
                raise RuntimeError("OFFICIAL_WORKBOOK_TEMPLATE_MISMATCH")
            formatting = validate_trading_assistant_workbook(main_workbook)
            stage = "SEVEN_DAY"
            session = get_session()
            try:
                seven_day = SevenDayComparisonService(session, self.root, self.reference_path).export(self.trade_date, self.output_dir)
            finally:
                session.close()
            stage = "MONITOR"
            monitor = _build_monitor(self.trade_date)
            outputs = {
                "main_workbook": str(main_workbook),
                "main_workbook_sha256": hashlib.sha256(main_workbook.read_bytes()).hexdigest(),
                "seven_day_workbook": seven_day["workbook"],
                "seven_day_json": seven_day["json"],
                "seven_day_markdown": seven_day["markdown"],
                "market_workbook": export.get("market_workbook"),
                "monitor": monitor,
            }
            report.update({
                "main_workbook_sheet_comparison": structure,
                "formatting_validation": formatting,
                "seven_day": seven_day,
                "output_paths": outputs,
            })
            report["final_status"] = final_status
            report["execution_end"] = datetime.now(SHANGHAI).isoformat()
            report["execution_duration_seconds"] = (datetime.now(SHANGHAI) - started).total_seconds()
            report["config_hash_before"] = config_hash_before
            report["config_hash_after"] = _config_hash(self.root)
            report["production_config_changed"] = report["config_hash_before"] != report["config_hash_after"]
            report["provider_calls"] = self.provider_calls
            report["cache_hits"] = self.cache_hits
            report.update(self._llm_delta(llm_before))
            json_path, md_path, audit_path = self._write_reports(report)
            report["output_paths"].update({"json": str(json_path), "markdown": str(md_path), "audit": str(audit_path)})
            self._write_reports(report)
            self._complete_run(report, outputs=report["output_paths"])
            return report
        except Exception as exc:
            status = _failure_status(stage, exc)
            report.update({
                "final_status": status,
                "failure_stage": stage,
                "failures": [status],
                "error_code": type(exc).__name__,
                "error_message": _safe_error(exc),
                "execution_end": datetime.now(SHANGHAI).isoformat(),
                "execution_duration_seconds": (datetime.now(SHANGHAI) - started).total_seconds(),
                "provider_calls": self.provider_calls,
                "cache_hits": self.cache_hits,
                "config_hash_before": config_hash_before,
                "config_hash_after": _config_hash(self.root),
            })
            report["production_config_changed"] = report["config_hash_before"] != report["config_hash_after"]
            report.update(self._llm_delta(llm_before))
            json_path, md_path, audit_path = self._write_reports(report)
            report["output_paths"] = {"json": str(json_path), "markdown": str(md_path), "audit": str(audit_path)}
            self._write_reports(report)
            self._fail_run(report)
            return report

    @property
    def style(self) -> WorkbookStyleService:
        return WorkbookStyleService(self.reference_path)

    def _today_manual_pool(self) -> dict[str, Any]:
        """Read only today's explicit manual pool; an empty pool is valid and is never fabricated."""
        session = get_session()
        try:
            count = int(session.scalar(
                select(func.count()).select_from(ManualSelectionRecord).where(
                    ManualSelectionRecord.trade_date == self.trade_date
                )
            ) or 0)
            return {
                "status": "READY" if count else "EMPTY",
                "count": count,
                "source": "DATABASE_TODAY_ONLY",
                "carried_forward": False,
            }
        finally:
            session.close()

    def _run_temporal_gate(self, previous: date) -> dict[str, Any]:
        gate = TushareTemporalGate(self.root, self.trade_date, previous)
        deadline = datetime.combine(self.trade_date, time(19, 0), SHANGHAI)
        first_check = None
        passed_datasets: set[str] = set()
        while True:
            audit = gate.check()
            first_check = first_check or audit["checked_at"]
            passed_datasets.update(name for name, value in audit["datasets"].items() if value["passed"])
            missing = [name for name in DATASETS if name not in passed_datasets]
            attempt = {
                "attempt": len(self.temporal_attempts) + 1,
                "checked_at": audit["checked_at"],
                "datasets": audit["datasets"],
                "passed_before_fetch": sorted(passed_datasets),
                "requested_datasets": [],
            }
            if not missing:
                attempt["result"] = "PASS"
                self.temporal_attempts.append(attempt)
                self.cache_hits += len(passed_datasets) if len(self.temporal_attempts) == 1 else 0
                return {
                    "passed": True,
                    "first_check_time": first_check,
                    "passed_at": audit["checked_at"],
                    "attempts": self.temporal_attempts,
                    "latest_trade_date": self.trade_date.isoformat(),
                    "datasets": audit["datasets"],
                }
            if datetime.now(SHANGHAI) >= deadline:
                attempt["result"] = "DEADLINE_FAILED"
                self.temporal_attempts.append(attempt)
                return {
                    "passed": False,
                    "first_check_time": first_check,
                    "passed_at": None,
                    "attempts": self.temporal_attempts,
                    "latest_trade_date": max((self.trade_date if value["trade_dates"] == [f"{self.trade_date:%Y%m%d}"] else previous) for value in audit["datasets"].values()).isoformat(),
                    "datasets": audit["datasets"],
                }
            attempt["requested_datasets"] = missing
            report_path = self.output_dir / f"tushare_gate_attempt_{attempt['attempt']:02d}.json"
            fetch = run_prewarm(
                start_date=self.trade_date.isoformat(),
                end_date=self.trade_date.isoformat(),
                interfaces=missing,
                include_reference=False,
                refresh_cache=True,
                output=report_path,
                progress=False,
            )
            self.provider_calls += len(missing)
            attempt["provider_report"] = {
                "path": str(report_path),
                "row_count": fetch.get("row_count"),
                "error_count": fetch.get("error_count"),
            }
            attempt["result"] = "FETCHED_RECHECK_REQUIRED"
            post_fetch = gate.check()
            attempt["post_fetch_check"] = post_fetch
            passed_datasets.update(
                name for name, value in post_fetch["datasets"].items() if value["passed"]
            )
            if len(passed_datasets) == len(DATASETS):
                attempt["result"] = "PASS_AFTER_FETCH"
                self.temporal_attempts.append(attempt)
                return {
                    "passed": True,
                    "first_check_time": first_check,
                    "passed_at": post_fetch["checked_at"],
                    "attempts": self.temporal_attempts,
                    "latest_trade_date": self.trade_date.isoformat(),
                    "datasets": post_fetch["datasets"],
                }
            self.temporal_attempts.append(attempt)
            time_module.sleep(
                min(300, max(0, (deadline - datetime.now(SHANGHAI)).total_seconds()))
            )

    def _create_run(self, started: datetime) -> None:
        session = get_session()
        try:
            session.add(PostCloseOfficialRun(
                run_id=self.run_id,
                trade_date=self.trade_date,
                status="RUNNING",
                stage="TEMPORAL_GATE",
                started_at=started,
                report_json={},
                output_paths_json={},
                real_orders=0,
                virtual_orders=0,
                scheduler_enabled=False,
                production_config_changed=False,
            ))
            session.commit()
        finally:
            session.close()

    def _persist_final_before_export(self, status: str, report: dict[str, Any]) -> None:
        if status not in ALLOWED_STATUSES:
            raise ValueError("INVALID_POSTCLOSE_STATUS")
        session = get_session()
        try:
            row = session.scalar(select(PostCloseOfficialRun).where(PostCloseOfficialRun.run_id == self.run_id))
            row.status = status
            row.stage = "BUSINESS_COMPLETED"
            row.report_json = report
            session.commit()
        finally:
            session.close()

    def _read_run_status(self) -> str:
        session = get_session()
        try:
            return str(session.scalar(select(PostCloseOfficialRun.status).where(PostCloseOfficialRun.run_id == self.run_id)))
        finally:
            session.close()

    def _complete_run(self, report: dict[str, Any], outputs: dict[str, Any]) -> None:
        session = get_session()
        try:
            row = session.scalar(select(PostCloseOfficialRun).where(PostCloseOfficialRun.run_id == self.run_id))
            row.status = report["final_status"]
            row.stage = "COMPLETED"
            row.completed_at = datetime.now(SHANGHAI)
            row.report_json = report
            row.output_paths_json = outputs
            row.provider_calls = int(report.get("provider_calls") or 0)
            row.llm_calls = int(report.get("llm_calls") or 0)
            row.total_tokens = int(report.get("total_token_usage") or 0)
            row.per_stock_api_calls = int((report.get("quant") or {}).get("per_stock_api_call_count") or 0)
            row.production_config_changed = bool(report.get("production_config_changed"))
            session.commit()
        finally:
            session.close()

    def _fail_run(self, report: dict[str, Any]) -> None:
        session = get_session()
        try:
            row = session.scalar(select(PostCloseOfficialRun).where(PostCloseOfficialRun.run_id == self.run_id))
            row.status = report["final_status"]
            row.stage = "FAILED"
            row.completed_at = datetime.now(SHANGHAI)
            row.report_json = report
            row.output_paths_json = report.get("output_paths") or {}
            row.error_code = report.get("error_code")
            row.error_message = report.get("error_message")
            row.provider_calls = int(report.get("provider_calls") or 0)
            row.llm_calls = int(report.get("llm_calls") or 0)
            row.total_tokens = int(report.get("total_token_usage") or 0)
            session.commit()
        finally:
            session.close()

    def _business_readback(self) -> dict[str, Any]:
        session = get_session()
        try:
            settings = WorkbenchService(session).settings()
            quant = session.scalar(select(QuantRun).where(
                QuantRun.base_market_trade_date == self.trade_date,
                QuantRun.status == "COMPLETED",
            ).order_by(QuantRun.created_at.desc()))
            flash = session.scalar(select(ModelValidationRun).where(
                ModelValidationRun.base_market_trade_date == self.trade_date,
                ModelValidationRun.real_llm.is_(True),
            ).order_by(ModelValidationRun.created_at.desc()))
            pro = session.scalar(select(ProResumeRun).where(
                ProResumeRun.base_trade_date == self.trade_date,
                ProResumeRun.status == "COMPLETED",
            ).order_by(ProResumeRun.created_at.desc()))
            ranks = list(session.scalars(select(QuantRankResult).where(
                QuantRankResult.quant_run_id == quant.run_id
            ).order_by(QuantRankResult.rank))) if quant else []
            samples = list(session.scalars(select(ModelValidationSample).where(
                ModelValidationSample.validation_run_id == flash.run_id
            ).order_by(ModelValidationSample.rank))) if flash else []
            reviews = list(session.scalars(select(ProCandidateReview).where(
                ProCandidateReview.pro_resume_run_id == pro.run_id
            ).order_by(ProCandidateReview.pro_rank))) if pro else []
            return {
                "active_config": _active_config(settings),
                "quant": {
                    "run_id": quant.run_id if quant else None,
                    "universe_count": quant.universe_count if quant else 0,
                    "eligible_count": quant.filtered_count if quant else 0,
                    "scored_count": quant.scored_count if quant else 0,
                    "top_q": quant.top_count if quant else 0,
                    "per_stock_api_call_count": quant.per_stock_api_call_count if quant else None,
                    "no_llm_call_verified": quant.no_llm_call_verified if quant else None,
                    "hash": _hash_rows([[row.stock_code, row.rank, str(row.total_score)] for row in ranks]),
                    "top20": [{"rank": row.rank, "stock_code": row.stock_code, "score": float(row.total_score)} for row in ranks[:20]],
                },
                "flash": {
                    "run_id": flash.run_id if flash else None,
                    "status": flash.status if flash else None,
                    "input": len(samples),
                    "success": sum(str(((row.screening_result or {}).get("_trader_demo") or {}).get("execution_status") or "").upper() == "SUCCESS" for row in samples),
                    "failure": sum(str(((row.screening_result or {}).get("_trader_demo") or {}).get("execution_status") or "").upper() != "SUCCESS" for row in samples),
                    "knowledge_mode": flash.knowledge_mode if flash else None,
                    "hash": _hash_rows([[row.stock_code, row.screening_result] for row in samples]),
                },
                "pro": {
                    "run_id": pro.run_id if pro else None,
                    "input": pro.candidate_count if pro else 0,
                    "success": len(reviews),
                    "failure": max(0, int(pro.candidate_count or 0) - len(reviews)) if pro else 0,
                    "hash": pro.candidate_set_hash if pro else None,
                    "results": [{"rank": row.pro_rank, "stock_code": row.stock_code, "score": float(row.pro_score), "status": row.review_status} for row in reviews],
                },
                "final_candidate_count": len(reviews),
                "final_candidate_list": [row.stock_code for row in reviews],
            }
        finally:
            session.close()

    def _llm_usage(self) -> dict[str, int]:
        session = get_session()
        try:
            query = select(
                func.count(ModelValidationLLMAudit.id),
                func.coalesce(func.sum(ModelValidationLLMAudit.input_tokens + ModelValidationLLMAudit.output_tokens), 0),
            ).join(ModelValidationRun, ModelValidationRun.run_id == ModelValidationLLMAudit.validation_run_id).where(
                ModelValidationRun.base_market_trade_date == self.trade_date
            )
            flash_calls, flash_tokens = session.execute(query).one()
            pro_query = select(
                func.count(LLMUsage.id),
                func.coalesce(func.sum(LLMUsage.total_tokens), 0),
            ).join(
                ProResumeRun, ProResumeRun.run_id == LLMUsage.pro_resume_run_id
            ).where(
                ProResumeRun.base_trade_date == self.trade_date,
                LLMUsage.usage_source == "CURRENT_CALL",
            )
            pro_calls, pro_tokens = session.execute(pro_query).one()
            return {
                "calls": int(flash_calls or 0) + int(pro_calls or 0),
                "tokens": int(flash_tokens or 0) + int(pro_tokens or 0),
                "flash_calls": int(flash_calls or 0),
                "flash_tokens": int(flash_tokens or 0),
                "pro_calls": int(pro_calls or 0),
                "pro_tokens": int(pro_tokens or 0),
            }
        finally:
            session.close()

    def _llm_delta(self, before: dict[str, int]) -> dict[str, int]:
        after = self._llm_usage()
        return {
            "llm_calls": max(0, after["calls"] - before["calls"]),
            "total_token_usage": max(0, after["tokens"] - before["tokens"]),
            "flash_llm_calls": max(0, after["flash_calls"] - before.get("flash_calls", 0)),
            "flash_token_usage": max(0, after["flash_tokens"] - before.get("flash_tokens", 0)),
            "pro_llm_calls": max(0, after["pro_calls"] - before.get("pro_calls", 0)),
            "pro_token_usage": max(0, after["pro_tokens"] - before.get("pro_tokens", 0)),
        }

    def _write_reports(self, report: dict[str, Any]) -> tuple[Path, Path, Path]:
        suffix = self.run_id[-8:]
        trade_date = self.trade_date.isoformat()
        json_path = self.output_dir / f"postclose_official_{trade_date}_{suffix}.json"
        md_path = self.output_dir / f"postclose_official_{trade_date}_{suffix}.md"
        audit_path = self.output_dir / f"postclose_official_audit_{trade_date}_{suffix}.json"
        json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        audit_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        md_path.write_text("\n".join([
            f"# {trade_date} 盘后全A正式运行",
            "",
            f"- Run ID：{self.run_id}",
            f"- 最终状态：{report.get('final_status')}",
            f"- 全A量化：{(report.get('quant') or {}).get('scored_count', 0)}",
            f"- 最终候选：{report.get('final_candidate_count', 0)}",
            "- 真实订单：0",
            "- 项目Scheduler：关闭",
            "",
        ]), encoding="utf-8")
        return json_path, md_path, audit_path


def _failure_status(stage: str, exc: Exception) -> str:
    if "TUSHARE_TEMPORAL_GATE_FAILED" in str(exc) or stage == "TEMPORAL_GATE":
        return "TUSHARE_TEMPORAL_GATE_FAILED"
    return {
        "QUANT": "POSTCLOSE_QUANT_FAILED",
        "MANUAL": "POSTCLOSE_INTEGRATION_FAILED",
        "FLASH": "POSTCLOSE_FLASH_FAILED",
        "PRO": "POSTCLOSE_PRO_FAILED",
        "MARKET": "POSTCLOSE_INTEGRATION_FAILED",
        "EXPORT": "POSTCLOSE_EXPORT_FAILED",
        "SEVEN_DAY": "POSTCLOSE_EXPORT_FAILED",
        "MONITOR": "POSTCLOSE_FULL_A_PARTIAL_SUCCESS",
    }.get(stage, "POSTCLOSE_INTEGRATION_FAILED")


def _active_config(settings: dict[str, Any]) -> dict[str, Any]:
    return {
        "Q": int(settings["quant_top_n"]),
        "L": int(settings["llm_analysis_n"]),
        "N": int(settings["llm_top_n"]),
        "R": int(settings["final_display_n"]),
        "W": int(settings["manual_soft_limit"]),
        "daily_token_limit": int(settings["daily_token_limit"]),
        "token_warning": int(float(settings["daily_token_limit"]) * float(settings["token_warning_ratio"])),
    }


def _quant_weights(root: Path) -> dict[str, float]:
    import yaml
    value = yaml.safe_load((root / "config" / "quant_factor.yaml").read_text(encoding="utf-8"))
    return {key: float(number) for key, number in value["quant_factor"]["weights"].items()}


def _config_hash(root: Path) -> str:
    paths = [
        root / "config" / name
        for name in ("quant_factor.yaml", "models.yaml", "stock_scan.yaml", "order_price.yaml", "position_sizing.yaml", "schedule.yaml")
    ]
    return _hash_rows([[path.name, hashlib.sha256(path.read_bytes()).hexdigest()] for path in paths])


def _hash_rows(rows: list[Any]) -> str:
    return hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")).hexdigest()


def load_sheet_names(path: Path | str) -> list[str]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _records(path: Path) -> list[dict[str, Any]]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    if isinstance(value, dict):
        value = value.get("records", [])
    return [row for row in value if isinstance(row, dict)] if isinstance(value, list) else []


def _safe_error(exc: Exception) -> str:
    text = str(exc).replace("\r", " ").replace("\n", " ")[:1000]
    for name in ("TUSHARE_TOKEN", "DEEPSEEK_API_KEY", "OPENAI_API_KEY", "IFIND_ACCESS_TOKEN", "IFIND_REFRESH_TOKEN"):
        secret = os.getenv(name, "")
        if secret:
            text = text.replace(secret, "[REDACTED]")
    return text
