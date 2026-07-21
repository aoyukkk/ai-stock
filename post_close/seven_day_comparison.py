from __future__ import annotations

import hashlib
import json
import statistics
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import func, select

from database.models.market_review import MarketDailySnapshot
from database.models.postclose_official import PostCloseOfficialRun
from database.models.quant_run import QuantRankResult, QuantRun
from database.models.stock import StockMaster
from database.models.system import LLMUsage
from database.models.validation import (
    ModelValidationFailureAudit,
    ModelValidationLLMAudit,
    ModelValidationRun,
    ModelValidationSample,
    ProCandidateReview,
    ProResumeRun,
)
from database.models.workbench import ManualSelectionRecord
from reporting.workbook_style import WorkbookStyleService


SHEET_NAMES = (
    "01_七日总览",
    "02_市场行情对比",
    "03_全A量化对比",
    "04_LLM推荐对比",
    "05_候选股票明细",
    "06_候选重合度",
    "07_行业分布对比",
    "08_市场状态对比",
    "09_候选后续表现",
    "10_数据口径与缺失",
)


class SevenDayComparisonService:
    def __init__(self, session, root: Path | str, reference_path: Path | str) -> None:
        self.session = session
        self.root = Path(root).resolve()
        self.cache_root = self.root / "data" / "cache" / "tushare"
        self.style = WorkbookStyleService(reference_path)
        self.stock_master = {
            row.code: row for row in self.session.scalars(select(StockMaster))
        }

    def trading_days(self, end_date: date, count: int = 7) -> list[date]:
        values: set[date] = set()
        for path in self.cache_root.glob("trade_cal_*.json"):
            for row in _records(path):
                text = str(row.get("cal_date") or "")
                if row.get("is_open") in (1, "1", True) and len(text) == 8 and text.isdigit():
                    value = datetime.strptime(text, "%Y%m%d").date()
                    if value <= end_date:
                        values.add(value)
        if end_date not in values and self._dataset_path("daily", end_date).is_file():
            values.add(end_date)
        selected = sorted(values)[-count:]
        if len(selected) != count or selected[-1] != end_date:
            raise ValueError("SEVEN_TRADING_DAY_CALENDAR_INCOMPLETE")
        return selected

    def export(self, end_date: date, output_dir: Path | str) -> dict[str, Any]:
        days = self.trading_days(end_date)
        output_dir = Path(output_dir).resolve()
        output_dir.mkdir(parents=True, exist_ok=True)
        stem = f"近7交易日数据对比_{days[0].isoformat()}_至_{end_date.isoformat()}"
        xlsx = _available_path(output_dir / f"{stem}.xlsx")
        json_path = xlsx.with_suffix(".json")
        markdown = xlsx.with_suffix(".md")
        payload = self.build_payload(days)
        workbook = self._workbook(payload)
        self.style.align_existing_workbook(workbook)
        workbook.save(xlsx)
        workbook.close()
        validation = self.validate(xlsx, days)
        payload["validation"] = validation
        payload["output_paths"] = {
            "workbook": str(xlsx),
            "json": str(json_path),
            "markdown": str(markdown),
        }
        json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        markdown.write_text(self._markdown(payload), encoding="utf-8")
        return {
            "days": [value.isoformat() for value in days],
            "workbook": str(xlsx),
            "json": str(json_path),
            "markdown": str(markdown),
            "sha256": hashlib.sha256(xlsx.read_bytes()).hexdigest(),
            "validation": validation,
            "missing_historical_runs": payload["missing_historical_runs"],
        }

    def build_payload(self, days: list[date]) -> dict[str, Any]:
        summaries: list[dict[str, Any]] = []
        quant_rows: list[dict[str, Any]] = []
        llm_rows: list[dict[str, Any]] = []
        candidate_rows: list[dict[str, Any]] = []
        industry_rows: list[dict[str, Any]] = []
        regime_rows: list[dict[str, Any]] = []
        missing_rows: list[dict[str, Any]] = []
        top20_by_day: dict[date, list[str]] = {}
        candidates_by_day: dict[date, list[str]] = {}
        daily_cache = {day: self._daily_map(day) for day in self._all_cached_days(days[-1])}
        for day in days:
            daily = list(daily_cache.get(day, {}).values())
            market = _market_metrics(daily)
            quant = self._latest_quant(day)
            ranks = list(self.session.scalars(select(QuantRankResult).where(
                QuantRankResult.quant_run_id == quant.run_id
            ).order_by(QuantRankResult.rank))) if quant else []
            flash = self._latest_flash(day, quant.run_id if quant else None)
            samples = list(self.session.scalars(select(ModelValidationSample).where(
                ModelValidationSample.validation_run_id == flash.run_id
            ).order_by(ModelValidationSample.rank))) if flash else []
            pro = self._latest_pro(day, flash.run_id if flash else None)
            reviews = list(self.session.scalars(select(ProCandidateReview).where(
                ProCandidateReview.pro_resume_run_id == pro.run_id
            ).order_by(ProCandidateReview.pro_rank))) if pro else []
            manual_count = int(self.session.scalar(select(func.count()).select_from(ManualSelectionRecord).where(
                ManualSelectionRecord.trade_date == day
            )) or 0)
            llm_audits = list(self.session.scalars(select(ModelValidationLLMAudit).where(
                ModelValidationLLMAudit.validation_run_id == flash.run_id
            ))) if flash else []
            pro_usage = list(self.session.scalars(select(LLMUsage).where(
                LLMUsage.pro_resume_run_id == pro.run_id,
                LLMUsage.usage_source == "CURRENT_CALL",
            ))) if pro else []
            failure_count = int(self.session.scalar(select(func.count()).select_from(ModelValidationFailureAudit).where(
                ModelValidationFailureAudit.validation_run_id == flash.run_id
            )) or 0) if flash else 0
            token_total = (
                sum(int(row.input_tokens or 0) + int(row.output_tokens or 0) for row in llm_audits)
                + sum(int(row.total_tokens or 0) for row in pro_usage)
            )
            top20 = [row.stock_code for row in ranks[:20]]
            candidates = [row.stock_code for row in reviews]
            top20_by_day[day] = top20
            candidates_by_day[day] = candidates
            scores = [_float(row.total_score) for row in ranks]
            factor_values = {
                key: [_float(getattr(row, key)) for row in ranks]
                for key in ("technical_score", "capital_score", "emotion_score", "momentum_score", "risk_score")
            }
            snapshot = self.session.scalar(select(MarketDailySnapshot).where(
                MarketDailySnapshot.trade_date == day
            ).order_by(MarketDailySnapshot.created_at.desc()))
            final_status = "HISTORICAL_RUN_MISSING"
            if quant and flash and pro:
                final_status = "COMPLETE"
            elif quant:
                final_status = "PARTIAL_HISTORY"
            official = self.session.scalar(select(PostCloseOfficialRun).where(
                PostCloseOfficialRun.trade_date == day,
                PostCloseOfficialRun.status.in_([
                    "POSTCLOSE_FULL_A_SUCCESS", "POSTCLOSE_FULL_A_PARTIAL_SUCCESS"
                ]),
            ).order_by(PostCloseOfficialRun.created_at.desc()))
            if official is not None:
                final_status = official.status
            index_metrics = {
                "上证涨跌幅": _index_value(snapshot, "shanghai"),
                "深证涨跌幅": _index_value(snapshot, "shenzhen"),
                "创业板涨跌幅": _index_value(snapshot, "chinext"),
                "科创相关指数涨跌幅": _index_value(snapshot, "star"),
            }
            if any(value is None for value in index_metrics.values()):
                missing_rows.append({
                    "交易日期": day,
                    "缺失类型": "INDEX_DATA_NOT_AVAILABLE",
                    "Quant": bool(quant),
                    "Flash": bool(flash),
                    "Pro": bool(pro),
                    "处理方式": "指数本地快照不可用；保持空值，不以全A横截面收益冒充指数收益",
                })
            summaries.append({
                "交易日期": day,
                "市场状态": snapshot.market_regime if snapshot else None,
                **index_metrics,
                "全市场成交额": market["amount"],
                "上涨家数": market["positive"],
                "下跌家数": market["negative"],
                "平盘家数": market["flat"],
                "全市场中位涨跌幅": market["median_return"],
                "涨停数量": market["limit_up"],
                "跌停数量": market["limit_down"],
                "Quant Universe": quant.universe_count if quant else None,
                "Quant Scored": quant.scored_count if quant else None,
                "Quant Top Q": quant.top_count if quant else None,
                "Flash输入数": len(samples),
                "Flash成功数": sum(_sample_success(row) for row in samples),
                "Flash失败数": failure_count,
                "正式推荐数": len(reviews),
                "人工池数量": manual_count,
                "最终候选数": len(reviews),
                "Pro成功数": sum(str(row.review_status).upper() in {"SUCCESS", "COMPLETED", "PASS"} for row in reviews),
                "Token使用": token_total,
                "最终运行状态": final_status,
            })
            quant_rows.append({
                "交易日期": day,
                "Quant Run ID": quant.run_id if quant else None,
                "Universe": quant.universe_count if quant else None,
                "Scored": quant.scored_count if quant else None,
                "Top Q": quant.top_count if quant else None,
                "技术平均分": _mean(factor_values["technical_score"]),
                "技术中位数": _median(factor_values["technical_score"]),
                "资金平均分": _mean(factor_values["capital_score"]),
                "资金中位数": _median(factor_values["capital_score"]),
                "情绪平均分": _mean(factor_values["emotion_score"]),
                "情绪中位数": _median(factor_values["emotion_score"]),
                "动量平均分": _mean(factor_values["momentum_score"]),
                "动量中位数": _median(factor_values["momentum_score"]),
                "风险平均分": _mean(factor_values["risk_score"]),
                "风险中位数": _median(factor_values["risk_score"]),
                "Top Q最低分": scores[min(len(scores), int(quant.top_count or 0)) - 1] if quant and scores and quant.top_count else None,
                "Top20平均分": _mean(scores[:20]),
                "每日Top20清单": "、".join(top20),
            })
            for sample in samples:
                screening = sample.screening_result or {}
                llm_rows.append({
                    "交易日期": day,
                    "股票代码": sample.stock_code,
                    "股票名称": sample.stock_name,
                    "Flash分": screening.get("llm_score"),
                    "Flash结论": screening.get("screening_decision"),
                    "风险等级": screening.get("risk_note"),
                    "主要行业": getattr(self.stock_master.get(sample.stock_code), "industry", None),
                    "调用状态": (screening.get("_trader_demo") or {}).get("execution_status"),
                })
            for review in reviews:
                master = self.stock_master.get(review.stock_code)
                screening = next((row.screening_result for row in samples if row.stock_code == review.stock_code), {}) or {}
                candidate_rows.append({
                    "交易日期": day,
                    "股票代码": review.stock_code,
                    "股票名称": getattr(master, "name", None),
                    "行业": getattr(master, "industry", None),
                    "Flash分": screening.get("llm_score"),
                    "Pro分": _float(review.pro_score),
                    "Pro排名": review.pro_rank,
                    "Pro结论": review.review_status,
                    "风险等级": review.priority,
                    "摘要": review.final_summary,
                })
            counts = Counter(getattr(self.stock_master.get(code), "industry", None) or "未知行业" for code in candidates)
            for industry, count in counts.most_common():
                industry_rows.append({
                    "交易日期": day,
                    "行业": industry,
                    "候选数量": count,
                    "当日候选占比": count / len(candidates) if candidates else None,
                })
            regime_rows.append({
                "交易日期": day,
                "市场方向": snapshot.market_direction if snapshot else None,
                "市场状态": snapshot.market_regime if snapshot else None,
                "状态分": snapshot.regime_score if snapshot else None,
                "置信度": snapshot.regime_confidence if snapshot else None,
                "数据质量": snapshot.data_quality_score if snapshot else None,
            })
            if not quant or not flash or not pro:
                missing_rows.append({
                    "交易日期": day,
                    "缺失类型": "HISTORICAL_RUN_MISSING",
                    "Quant": bool(quant),
                    "Flash": bool(flash),
                    "Pro": bool(pro),
                    "处理方式": "仅使用现有持久化数据；未重新调用历史LLM",
                })
        overlap = _overlap_rows(days, top20_by_day, candidates_by_day)
        forward = self._forward_rows(candidate_rows, days, daily_cache)
        missing_historical = sorted({
            row["交易日期"].isoformat()
            for row in missing_rows
            if row.get("缺失类型") == "HISTORICAL_RUN_MISSING"
        })
        return {
            "trading_days": [value.isoformat() for value in days],
            "overview": summaries,
            "market": [{"交易日期": row["交易日期"], **{key: row[key] for key in (
                "全市场成交额", "上涨家数", "下跌家数", "平盘家数", "全市场中位涨跌幅", "涨停数量", "跌停数量"
            )}} for row in summaries],
            "quant": quant_rows,
            "llm": llm_rows,
            "candidates": candidate_rows,
            "overlap": overlap,
            "industry": industry_rows,
            "regime": regime_rows,
            "forward": forward,
            "missing": missing_rows or [{"交易日期": days[-1], "缺失类型": "NONE", "处理方式": "无历史运行缺失"}],
            "missing_historical_runs": missing_historical,
        }

    def _workbook(self, payload: dict[str, Any]):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets = (
            (SHEET_NAMES[0], payload["overview"], "最近7个交易日总览"),
            (SHEET_NAMES[1], payload["market"], "全A市场行情对比"),
            (SHEET_NAMES[2], payload["quant"], "全A量化因子与Top Q对比"),
            (SHEET_NAMES[3], payload["llm"], "真实Flash与Pro历史结果对比"),
            (SHEET_NAMES[4], payload["candidates"], "正式候选股票明细"),
            (SHEET_NAMES[5], payload["overlap"], "相邻交易日候选重合度"),
            (SHEET_NAMES[6], payload["industry"], "候选行业分布与集中度"),
            (SHEET_NAMES[7], payload["regime"], "市场状态对比"),
            (SHEET_NAMES[8], payload["forward"], "候选后续表现；未来数据为空值"),
            (SHEET_NAMES[9], payload["missing"], "数据口径、有效样本与缺失说明"),
        )
        for index, (name, rows, title) in enumerate(sheets, 1):
            _write_sheet(workbook, name, rows, title, index)
        return workbook

    def validate(self, path: Path, days: list[date]) -> dict[str, Any]:
        excel_compatibility = WorkbookStyleService.validate_excel_compatibility(path)
        workbook = load_workbook(path, data_only=False)
        try:
            if tuple(workbook.sheetnames) != SHEET_NAMES:
                raise ValueError("SEVEN_DAY_SHEET_STRUCTURE_MISMATCH")
            observed_days = [
                _as_date(row[0].value)
                for row in workbook[SHEET_NAMES[0]].iter_rows(min_row=5, min_col=1, max_col=1)
            ]
            if observed_days != days:
                raise ValueError("SEVEN_DAY_CALENDAR_MISMATCH")
            formula_errors = 0
            stock_cells = 0
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        if cell.alignment.horizontal != "center" or cell.alignment.vertical != "center" or not cell.alignment.wrap_text:
                            raise ValueError(f"SEVEN_DAY_ALIGNMENT_ERROR:{worksheet.title}!{cell.coordinate}")
                        if isinstance(cell.value, str) and any(error in cell.value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                            formula_errors += 1
                header_row = 4
                code_columns = [cell.column for cell in worksheet[header_row] if "股票代码" in str(cell.value or "")]
                for column in code_columns:
                    for row in range(header_row + 1, worksheet.max_row + 1):
                        cell = worksheet.cell(row, column)
                        if cell.value not in (None, ""):
                            stock_cells += 1
                            if cell.number_format != "@" or not isinstance(cell.value, str):
                                raise ValueError(f"SEVEN_DAY_STOCK_CODE_FORMAT_ERROR:{worksheet.title}!{cell.coordinate}")
            if formula_errors:
                raise ValueError(f"SEVEN_DAY_FORMULA_ERRORS:{formula_errors}")
            forward_sheet = workbook[SHEET_NAMES[8]]
            headers = {str(cell.value or ""): cell.column - 1 for cell in forward_sheet[4]}
            today_rows = [
                row for row in forward_sheet.iter_rows(min_row=5, values_only=True)
                if row and _as_date(row[headers["交易日期"]]) == days[-1]
            ]
            return_columns = [headers[name] for name in ("D+1收益", "D+3收益", "D+5收益")]
            if not today_rows or any(any(row[column] is not None for column in return_columns) for row in today_rows):
                raise ValueError("TODAY_FORWARD_RETURN_MUST_BE_BLANK")
            if any(row[headers["数据状态"]] != "PENDING_FUTURE_DATA" for row in today_rows):
                raise ValueError("TODAY_FORWARD_STATUS_ERROR")
            return {
                "status": "PASS",
                "sheet_count": len(workbook.sheetnames),
                "sheet_names": list(workbook.sheetnames),
                "formula_errors": 0,
                "stock_code_cells": stock_cells,
                "centered": True,
                "wrapped": True,
                "today_forward_status": "PENDING_FUTURE_DATA",
                "excel_compatibility": excel_compatibility,
            }
        finally:
            workbook.close()

    def _latest_quant(self, day: date):
        return self.session.scalar(select(QuantRun).where(
            QuantRun.base_market_trade_date == day,
            QuantRun.status == "COMPLETED",
        ).order_by(QuantRun.created_at.desc()))

    def _latest_flash(self, day: date, quant_run_id: str | None):
        query = select(ModelValidationRun).where(
            ModelValidationRun.base_market_trade_date == day,
            ModelValidationRun.status.in_(["COMPLETED", "SUCCESS", "PARTIAL_SUCCESS"]),
        )
        if quant_run_id:
            query = query.where(ModelValidationRun.quant_run_id == quant_run_id)
        return self.session.scalar(query.order_by(ModelValidationRun.created_at.desc()))

    def _latest_pro(self, day: date, flash_run_id: str | None):
        query = select(ProResumeRun).where(
            ProResumeRun.base_trade_date == day,
            ProResumeRun.status == "COMPLETED",
        )
        if flash_run_id:
            query = query.where(ProResumeRun.flash_validation_run_id == flash_run_id)
        return self.session.scalar(query.order_by(ProResumeRun.created_at.desc()))

    def _dataset_path(self, dataset: str, day: date) -> Path:
        return self.cache_root / "trade_date" / dataset / f"{day:%Y%m%d}.json"

    def _daily_map(self, day: date) -> dict[str, dict[str, Any]]:
        return {str(row.get("ts_code")): row for row in _records(self._dataset_path("daily", day)) if row.get("ts_code")}

    def _all_cached_days(self, end: date) -> list[date]:
        values = []
        for path in (self.cache_root / "trade_date" / "daily").glob("*.json"):
            if len(path.stem) == 8 and path.stem.isdigit():
                value = datetime.strptime(path.stem, "%Y%m%d").date()
                if value <= end:
                    values.append(value)
        return sorted(values)

    def _forward_rows(self, candidates: list[dict[str, Any]], days: list[date], daily_cache) -> list[dict[str, Any]]:
        all_days = sorted(daily_cache)
        rows = []
        for candidate in candidates:
            selected = candidate["交易日期"]
            code = candidate["股票代码"]
            base = daily_cache.get(selected, {}).get(code)
            future = [value for value in all_days if value > selected]
            result = {
                "交易日期": selected,
                "记录类型": "候选明细",
                "股票代码": code,
                "股票名称": candidate.get("股票名称"),
                "行业": candidate.get("行业"),
                "D+1收益": _forward_return(base, daily_cache.get(future[0], {}).get(code)) if len(future) >= 1 else None,
                "D+3收益": _forward_return(base, daily_cache.get(future[2], {}).get(code)) if len(future) >= 3 else None,
                "D+5收益": _forward_return(base, daily_cache.get(future[4], {}).get(code)) if len(future) >= 5 else None,
                "MAE": None,
                "MFE": None,
                "有效样本数": None,
                "正收益比例": None,
                "平均收益": None,
                "中位收益": None,
                "最大盈利": None,
                "最大亏损": None,
                "Profit Factor": None,
                "数据状态": "PENDING_FUTURE_DATA" if selected == days[-1] else "AVAILABLE_RANGE_ONLY",
            }
            if base and future:
                base_close = _float(base.get("close"))
                path_rows = [daily_cache.get(day, {}).get(code) for day in future[:5]]
                lows = [_float(row.get("low")) for row in path_rows if row and row.get("low") is not None]
                highs = [_float(row.get("high")) for row in path_rows if row and row.get("high") is not None]
                if base_close:
                    result["MAE"] = min((value / base_close - 1 for value in lows), default=None)
                    result["MFE"] = max((value / base_close - 1 for value in highs), default=None)
            rows.append(result)
        for selected in days:
            selected_rows = [row for row in rows if row["交易日期"] == selected and row["记录类型"] == "候选明细"]
            valid = [row["D+1收益"] for row in selected_rows if row["D+1收益"] is not None]
            gains = [value for value in valid if value > 0]
            losses = [value for value in valid if value < 0]
            rows.append({
                "交易日期": selected,
                "记录类型": "日汇总",
                "股票代码": "",
                "股票名称": "日汇总" if selected_rows else "无正式候选",
                "行业": "",
                "D+1收益": None,
                "D+3收益": None,
                "D+5收益": None,
                "MAE": None,
                "MFE": None,
                "有效样本数": len(valid),
                "正收益比例": len(gains) / len(valid) if valid else None,
                "平均收益": statistics.mean(valid) if valid else None,
                "中位收益": statistics.median(valid) if valid else None,
                "最大盈利": max(valid) if valid else None,
                "最大亏损": min(valid) if valid else None,
                "Profit Factor": sum(gains) / abs(sum(losses)) if losses else None,
                "数据状态": "PENDING_FUTURE_DATA" if selected == days[-1] else (
                    "AVAILABLE" if valid else "HISTORICAL_RUN_MISSING"
                ),
            })
        return rows

    @staticmethod
    def _markdown(payload: dict[str, Any]) -> str:
        days = payload["trading_days"]
        missing = payload["missing_historical_runs"]
        return "\n".join([
            "# 近7交易日数据对比口径说明",
            "",
            f"- 交易日范围：{days[0]} 至 {days[-1]}",
            "- 交易日由本地持久化Tushare交易日历动态解析。",
            "- 历史LLM结果只读取数据库，不重新调用模型。",
            f"- {days[-1]}的D+1/D+3/D+5为空，并标记PENDING_FUTURE_DATA。",
            f"- 历史运行缺失：{'、'.join(missing) if missing else '无'}",
            "",
        ])


def _write_sheet(workbook, name: str, rows: Iterable[dict[str, Any]], title: str, index: int) -> None:
    values = list(rows)
    headers = list(dict.fromkeys(key for row in values for key in row)) or ["状态"]
    if not values:
        values = [{headers[0]: "NO_DATA"}]
    worksheet = workbook.create_sheet(name)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(headers)))
    worksheet.cell(1, 1, title)
    worksheet.cell(3, 1, "全部数据来自本地缓存和数据库持久化结果；不为历史日期重新调用LLM。")
    for column, header in enumerate(headers, 1):
        worksheet.cell(4, column, header)
    for row_index, row in enumerate(values, 5):
        for column, header in enumerate(headers, 1):
            value = row.get(header)
            worksheet.cell(row_index, column, _cell(value))
            if "股票代码" in header:
                worksheet.cell(row_index, column).number_format = "@"
            elif "日期" in header and isinstance(value, date):
                worksheet.cell(row_index, column).number_format = "yyyy-mm-dd"
            elif any(token in header for token in ("涨跌幅", "收益", "占比", "换手率", "置信度", "MAE", "MFE")) and isinstance(value, (int, float, Decimal)):
                worksheet.cell(row_index, column).number_format = "0.00%"
            elif isinstance(value, (int, float, Decimal)):
                worksheet.cell(row_index, column).number_format = "#,##0.00"
    table = Table(displayName=f"SevenDayTable{index:02d}", ref=f"A4:{_column_letter(len(headers))}{4 + len(values)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    worksheet.add_table(table)
    worksheet.freeze_panes = "A5"
    worksheet.sheet_view.showGridLines = False
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _market_metrics(rows: list[dict[str, Any]]) -> dict[str, Any]:
    changes = [_float(row.get("pct_chg")) / 100 for row in rows if row.get("pct_chg") is not None]
    return {
        "amount": sum(_float(row.get("amount")) for row in rows),
        "positive": sum(value > 0 for value in changes),
        "negative": sum(value < 0 for value in changes),
        "flat": sum(value == 0 for value in changes),
        "median_return": _median(changes),
        "limit_up": sum(value >= 0.095 for value in changes),
        "limit_down": sum(value <= -0.095 for value in changes),
    }


def _overlap_rows(days, top20_by_day, candidates_by_day) -> list[dict[str, Any]]:
    rows = []
    for previous, current in zip(days, days[1:]):
        top_previous, top_current = set(top20_by_day[previous]), set(top20_by_day[current])
        candidate_previous, candidate_current = set(candidates_by_day[previous]), set(candidates_by_day[current])
        rows.append({
            "上一交易日": previous,
            "当前交易日": current,
            "Quant Top20重合数": len(top_previous & top_current),
            "Quant Top20换手率": 1 - len(top_previous & top_current) / max(1, len(top_previous | top_current)),
            "候选重合数": len(candidate_previous & candidate_current),
            "候选换手率": 1 - len(candidate_previous & candidate_current) / max(1, len(candidate_previous | candidate_current)),
            "共同候选": "、".join(sorted(candidate_previous & candidate_current)),
        })
    return rows


def _records(path: Path) -> list[dict[str, Any]]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    if isinstance(value, dict):
        value = value.get("records", [])
    return [row for row in value if isinstance(row, dict)] if isinstance(value, list) else []


def _index_value(snapshot, key: str):
    if snapshot is None:
        return None
    values = snapshot.index_summary_json or {}
    for name, payload in values.items() if isinstance(values, dict) else []:
        if key in str(name).lower() and isinstance(payload, dict):
            value = payload.get("pct_chg") or payload.get("change_pct") or payload.get("return")
            if value is not None:
                numeric = _float(value)
                return numeric / 100 if abs(numeric) > 1 else numeric
    return None


def _sample_success(sample) -> bool:
    return str(((sample.screening_result or {}).get("_trader_demo") or {}).get("execution_status") or "").upper() == "SUCCESS"


def _forward_return(base, future):
    if not base or not future:
        return None
    left, right = _float(base.get("close")), _float(future.get("close"))
    return right / left - 1 if left else None


def _cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if value is None or isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return json.dumps(value, ensure_ascii=False, default=str, separators=(",", ":"))


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def _float(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _mean(values: list[float]):
    return statistics.fmean(values) if values else None


def _median(values: list[float]):
    return statistics.median(values) if values else None


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _available_path(path: Path) -> Path:
    if not path.exists():
        return path
    return path.with_name(f"{path.stem}_{datetime.now():%H%M%S}{path.suffix}")
