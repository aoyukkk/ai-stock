from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy import select, update

from database.models import StockMaster
from database.models.post_close import PositionTruthConfirmation, TraderPositionImportBatch, TraderPositionSnapshot
from stock_codes import normalize_ts_code


ACCOUNT_SCOPES = {"HUMAN_REFERENCE", "AI_SIMULATION"}
HEADER_ALIASES = {
    "stock_code": {"stock_code", "code", "股票代码", "证券代码"},
    "stock_name": {"stock_name", "name", "股票名称", "证券名称"},
    "quantity": {"quantity", "持仓数量", "当前数量"},
    "available_quantity": {"available_quantity", "可卖数量", "当前可卖数量"},
    "cost_price": {"cost_price", "成本价", "持仓成本"},
    "buy_date": {"buy_date", "买入日期", "建仓日期"},
    "market_value": {"market_value", "市值", "持仓市值"},
    "position_percent": {"position_percent", "仓位比例", "持仓比例"},
    "account_scope": {"account_scope", "账户类型", "账户范围"},
    "note": {"note", "备注"},
}


class PositionImportService:
    def __init__(self, session) -> None:
        self.session = session

    def preview(self, *, filename: str, content_base64: str, account_scope: str) -> dict[str, Any]:
        scope = _scope(account_scope)
        content = base64.b64decode(content_base64, validate=True)
        if len(content) > 5 * 1024 * 1024:
            raise ValueError("POSITION_FILE_TOO_LARGE")
        raw_rows = _read_rows(filename, content)
        rows, errors = _normalize_rows(raw_rows)
        errors.extend(self._known_code_errors(rows))
        for index, row in enumerate(rows, start=2):
            if row.get("account_scope") and row["account_scope"] != scope:
                errors.append({"row": index, "error": "POSITION_ACCOUNT_SCOPE_CONFLICT"})
        payload = json.dumps(rows, sort_keys=True, ensure_ascii=True, separators=(",", ":"))
        preview_id = f"position-preview-{uuid.uuid4().hex[:20]}"
        batch = TraderPositionImportBatch(
            preview_id=preview_id,
            account_scope=scope,
            source_filename=Path(filename).name,
            payload_hash=hashlib.sha256(payload.encode()).hexdigest(),
            rows_json=rows,
            status="VALID" if not errors else "INVALID",
            error_json=errors,
        )
        self.session.add(batch)
        self.session.commit()
        return {"preview_id": preview_id, "account_scope": scope, "status": batch.status, "row_count": len(rows), "items": rows, "errors": errors, "payload_hash": batch.payload_hash}

    def confirm(self, preview_id: str, *, confirmed_by: str = "LOCAL_TRADER") -> dict[str, Any]:
        batch = self.session.scalar(select(TraderPositionImportBatch).where(TraderPositionImportBatch.preview_id == preview_id))
        if batch is None:
            raise ValueError("POSITION_PREVIEW_NOT_FOUND")
        if batch.status == "CONFIRMED":
            version = _version(preview_id)
            truth = self.session.scalar(select(PositionTruthConfirmation).where(
                PositionTruthConfirmation.account_scope == batch.account_scope,
                PositionTruthConfirmation.snapshot_version == version,
            ))
            if truth is None:
                self._record_truth(
                    account_scope=batch.account_scope,
                    trade_date_value=max((date.fromisoformat(row["trade_date"]) for row in batch.rows_json), default=(batch.confirmed_at or datetime.now(timezone.utc)).date()),
                    position_count=len(batch.rows_json),
                    confirmed_by=confirmed_by,
                    version=version,
                )
                self.session.commit()
            return {"preview_id": preview_id, "status": "CONFIRMED", "version": version, "row_count": len(batch.rows_json)}
        if batch.status != "VALID":
            raise ValueError("POSITION_PREVIEW_INVALID")
        now = datetime.now(timezone.utc)
        version = _version(preview_id)
        self.session.execute(update(TraderPositionSnapshot).where(
            TraderPositionSnapshot.account_scope == batch.account_scope,
            TraderPositionSnapshot.is_current.is_(True),
        ).values(is_current=False))
        for row in batch.rows_json:
            self.session.add(TraderPositionSnapshot(
                account_scope=batch.account_scope,
                snapshot_time=now,
                trade_date=date.fromisoformat(row.get("trade_date") or now.date().isoformat()),
                stock_code=row["stock_code"],
                stock_name_snapshot=row.get("stock_name"),
                quantity=row["quantity"],
                available_quantity=row["available_quantity"],
                cost_price=Decimal(str(row["cost_price"])),
                buy_date=date.fromisoformat(row["buy_date"]) if row.get("buy_date") else None,
                market_value=Decimal(str(row["market_value"])) if row.get("market_value") is not None else None,
                position_percent=Decimal(str(row["position_percent"])) if row.get("position_percent") is not None else None,
                source="CONFIRMED_IMPORT",
                version=version,
                is_current=True,
            ))
        batch.status = "CONFIRMED"
        batch.confirmed_at = now
        self._record_truth(
            account_scope=batch.account_scope,
            trade_date_value=max((date.fromisoformat(row["trade_date"]) for row in batch.rows_json), default=now.date()),
            position_count=len(batch.rows_json),
            confirmed_by=confirmed_by,
            version=version,
        )
        self.session.commit()
        return {"preview_id": preview_id, "status": "CONFIRMED", "version": version, "row_count": len(batch.rows_json)}

    def current(self, account_scope: str | None = None) -> list[TraderPositionSnapshot]:
        query = select(TraderPositionSnapshot).where(TraderPositionSnapshot.is_current.is_(True))
        if account_scope:
            query = query.where(TraderPositionSnapshot.account_scope == _scope(account_scope))
        return list(self.session.scalars(query.order_by(TraderPositionSnapshot.account_scope, TraderPositionSnapshot.stock_code)))

    def confirm_empty(self, *, account_scope: str, trade_date_value: date, confirmed_by: str = "LOCAL_TRADER") -> dict[str, Any]:
        scope = _scope(account_scope)
        now = datetime.now(timezone.utc)
        seed = f"{scope}|{trade_date_value.isoformat()}|{now.isoformat()}"
        version = f"position-empty-v1-{hashlib.sha256(seed.encode()).hexdigest()[:16]}"
        self.session.execute(update(TraderPositionSnapshot).where(
            TraderPositionSnapshot.account_scope == scope,
            TraderPositionSnapshot.is_current.is_(True),
        ).values(is_current=False))
        confirmation = self._record_truth(
            account_scope=scope,
            trade_date_value=trade_date_value,
            position_count=0,
            confirmed_by=confirmed_by,
            version=version,
        )
        self.session.commit()
        return {
            "status": "CONFIRMED_EMPTY",
            "account_scope": scope,
            "trade_date": trade_date_value,
            "snapshot_time": confirmation.snapshot_time,
            "position_count": 0,
            "version": version,
        }

    def _record_truth(self, *, account_scope: str, trade_date_value: date, position_count: int, confirmed_by: str, version: str) -> PositionTruthConfirmation:
        now = datetime.now(timezone.utc)
        payload = f"{account_scope}|{trade_date_value.isoformat()}|{position_count}|{version}"
        digest = hashlib.sha256(payload.encode()).hexdigest()
        existing = self.session.scalar(select(PositionTruthConfirmation).where(PositionTruthConfirmation.snapshot_hash == digest))
        if existing:
            return existing
        self.session.execute(update(PositionTruthConfirmation).where(
            PositionTruthConfirmation.account_scope == account_scope,
            PositionTruthConfirmation.is_current.is_(True),
        ).values(is_current=False))
        row = PositionTruthConfirmation(
            account_scope=account_scope,
            trade_date=trade_date_value,
            snapshot_time=now,
            position_count=position_count,
            confirmation_status="CONFIRMED_POSITIONS" if position_count else "CONFIRMED_EMPTY",
            confirmed_by=str(confirmed_by)[:64] or "LOCAL_TRADER",
            snapshot_version=version,
            snapshot_hash=digest,
            is_current=True,
        )
        self.session.add(row)
        return row

    def _known_code_errors(self, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        known = {normalize_ts_code(value) for value in self.session.scalars(select(StockMaster.code)).all()}
        if not known:
            return []
        return [
            {"row": index, "error": "POSITION_STOCK_CODE_UNKNOWN"}
            for index, row in enumerate(rows, start=2)
            if row["stock_code"] not in known
        ]


def _read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        return [dict(row) for row in csv.DictReader(io.StringIO(text))]
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return []
            headers = [str(value or "").strip() for value in values[0]]
            return [dict(zip(headers, row)) for row in values[1:] if any(value not in (None, "") for value in row)]
        finally:
            workbook.close()
    raise ValueError("POSITION_FILE_TYPE_UNSUPPORTED")


def _normalize_rows(values: Iterable[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows, errors = [], []
    seen: set[str] = set()
    for index, raw in enumerate(values, start=2):
        mapped = {_canonical_header(str(key or "")): value for key, value in raw.items() if _canonical_header(str(key or ""))}
        try:
            code = normalize_ts_code(str(mapped.get("stock_code") or ""))
            if not code or code in seen:
                raise ValueError("STOCK_CODE_DUPLICATE_OR_MISSING")
            quantity = _integer(mapped.get("quantity"))
            available = _integer(mapped.get("available_quantity", quantity))
            cost = _decimal(mapped.get("cost_price"))
            if quantity < 0 or available < 0 or available > quantity:
                raise ValueError("POSITION_QUANTITY_INVALID")
            if cost <= 0:
                raise ValueError("POSITION_COST_PRICE_INVALID")
            buy_date = _date_value(mapped.get("buy_date"))
            if buy_date and buy_date > date.today():
                raise ValueError("POSITION_BUY_DATE_IN_FUTURE")
            account_scope = str(mapped.get("account_scope") or "").strip().upper() or None
            if account_scope and account_scope not in ACCOUNT_SCOPES:
                raise ValueError("POSITION_ACCOUNT_SCOPE_INVALID")
            rows.append({
                "stock_code": code,
                "stock_name": str(mapped.get("stock_name") or "").strip() or None,
                "quantity": quantity,
                "available_quantity": available,
                "cost_price": float(cost),
                "buy_date": buy_date.isoformat() if buy_date else None,
                "market_value": _optional_float(mapped.get("market_value")),
                "position_percent": _optional_float(mapped.get("position_percent")),
                "trade_date": date.today().isoformat(),
                "account_scope": account_scope,
                "note": str(mapped.get("note") or "").strip() or None,
            })
            seen.add(code)
        except (ValueError, InvalidOperation) as exc:
            errors.append({"row": index, "error": str(exc)})
    return rows, errors


def _canonical_header(value: str) -> str | None:
    normalized = value.strip().lower()
    return next((canonical for canonical, aliases in HEADER_ALIASES.items() if normalized in {item.lower() for item in aliases}), None)


def _integer(value: Any) -> int:
    number = Decimal(str(value))
    if number != number.to_integral_value():
        raise ValueError("POSITION_QUANTITY_NOT_INTEGER")
    return int(number)


def _decimal(value: Any) -> Decimal:
    return Decimal(str(value))


def _optional_float(value: Any) -> float | None:
    return None if value in (None, "") else float(value)


def _date_value(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _scope(value: str) -> str:
    scope = str(value).upper()
    if scope not in ACCOUNT_SCOPES:
        raise ValueError("POSITION_ACCOUNT_SCOPE_INVALID")
    return scope


def _version(preview_id: str) -> str:
    return f"position-v1-{hashlib.sha256(preview_id.encode()).hexdigest()[:16]}"
