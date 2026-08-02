from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from event_overlay.hashing import file_hash


REQUIRED_SHEETS = [
    "V3事件初筛Top20",
    "V2与V3对照",
    "Top100事件搜索状态",
    "事件证据明细",
    "Event Score拆解",
    "事件风险与Hard Gate",
    "Pro复核结果",
    "Market Regime与部署限制",
    "主题集中度",
    "数据时效与来源",
    "Checkpoint复用审计",
    "版本与Hash",
    "数据质量问题",
]


def export_run(
    output_dir: Path,
    *,
    trade_date: str,
    run_id: str,
    manifest: dict[str, Any],
    items: list[dict[str, Any]],
    snapshots: list[dict[str, Any]],
    checkpoint_audit: list[dict[str, Any]],
    data_quality: list[dict[str, Any]],
    comparison: list[dict[str, Any]],
    pro_results: list[dict[str, Any]] | None = None,
    market_regime: dict[str, Any] | None = None,
    theme_concentration: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=False)
    evidence_items = [
        {**event, "snapshot_id": snapshot["snapshot_id"], "search_status": snapshot["search_status"]}
        for snapshot in snapshots
        for event in snapshot.get("items", [])
    ]
    json_path = output_dir / "event_evidence_snapshot.json"
    json_path.write_text(json.dumps(snapshots, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    _write_csv(output_dir / "event_evidence_items.csv", evidence_items)
    _write_csv(output_dir / "v3_screening_top20.csv", [row for row in items if row["selected_top20"]])
    _write_csv(output_dir / "v2_v3_comparison.csv", comparison)
    (output_dir / "checkpoint_reuse_audit.json").write_text(
        json.dumps(checkpoint_audit, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    _write_csv(output_dir / "data_quality.csv", data_quality)
    (output_dir / "run_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )

    workbook_path = output_dir / f"V3事件覆盖层Shadow_{trade_date}_{run_id}.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet_rows = _sheet_rows(
        items=items,
        snapshots=snapshots,
        evidence_items=evidence_items,
        checkpoint_audit=checkpoint_audit,
        data_quality=data_quality,
        comparison=comparison,
        manifest=manifest,
        pro_results=pro_results or [],
        market_regime=market_regime or {},
        theme_concentration=theme_concentration or [],
    )
    for sheet_name in REQUIRED_SHEETS:
        _add_sheet(workbook, sheet_name, sheet_rows[sheet_name])
    workbook.save(workbook_path)
    validation = validate_workbook(workbook_path)
    validation_path = output_dir / "validation_report.md"
    validation_path.write_text(
        "\n".join([
            "# V3 Event Overlay Shadow Validation",
            "",
            f"- workbook: `{workbook_path.name}`",
            f"- workbook_sha256: `{file_hash(workbook_path)}`",
            f"- required_sheets: `{validation['required_sheets']}`",
            f"- stock_code_text_format: `{validation['stock_code_text_format']}`",
            f"- evidence_source_visible: `{validation['evidence_source_visible']}`",
            f"- search_status_visible: `{validation['search_status_visible']}`",
            f"- immutable_output_directory: `true`",
            f"- real_orders: `{manifest['real_orders']}`",
            f"- virtual_orders: `{manifest['virtual_orders']}`",
            f"- scheduler: `{str(manifest['scheduler']).lower()}`",
        ]),
        encoding="utf-8",
    )
    return {
        "output_dir": str(output_dir),
        "workbook": str(workbook_path),
        "workbook_sha256": file_hash(workbook_path),
        "validation": validation,
        "files": sorted(path.name for path in output_dir.iterdir()),
    }


def validate_workbook(path: Path) -> dict[str, bool]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    required = workbook.sheetnames == REQUIRED_SHEETS
    code_format = True
    for sheet in workbook.worksheets:
        headers = [cell.value for cell in sheet[1]]
        for index, header in enumerate(headers, start=1):
            if header in {"股票代码", "stock_code"}:
                code_format = code_format and all(
                    cell.number_format == "@" for cell in list(sheet.columns)[index - 1][1:] if cell.value not in {None, ""}
                )
    evidence_headers = [cell.value for cell in workbook["事件证据明细"][1]]
    status_headers = [cell.value for cell in workbook["Top100事件搜索状态"][1]]
    workbook.close()
    return {
        "required_sheets": required,
        "stock_code_text_format": code_format,
        "evidence_source_visible": "来源" in evidence_headers and "URL" in evidence_headers,
        "search_status_visible": "搜索状态" in status_headers,
    }


def _sheet_rows(**payload: Any) -> dict[str, list[dict[str, Any]]]:
    items = payload["items"]
    snapshots = payload["snapshots"]
    snapshot_by_id = {row["snapshot_id"]: row for row in snapshots}
    top20 = [row for row in items if row["selected_top20"]]
    search_rows = [{
        "股票代码": row["stock_code"],
        "股票名称": row["stock_name"],
        "Quant排名": row["quant_rank"],
        "搜索状态": row["search_status"],
        "Evidence Confidence": row["evidence_confidence"],
        "事件数": len(snapshot_by_id[row["event_snapshot_id"]].get("items", [])),
        "直搜降级": snapshot_by_id[row["event_snapshot_id"]]["direct_search_used"],
        "来源Provider": snapshot_by_id[row["event_snapshot_id"]]["provider"],
    } for row in items]
    evidence_rows = [{
        "股票代码": row["stock_code"],
        "事件类型": row["event_type"],
        "事件方向": row["event_direction"],
        "标题": row["title"],
        "摘要": row["summary"],
        "来源": row.get("domain") or row["provider"],
        "来源等级": row["source_tier"],
        "发布时间": row.get("published_at"),
        "评分资格": row.get("score_eligible", True),
        "不合格原因": "；".join(row.get("score_exclusion_reasons") or []),
        "时效状态": row.get("temporal_status"),
        "时间衰减": row.get("time_decay"),
        "URL": row.get("url"),
        "搜索状态": row["search_status"],
        "事件簇": row["event_cluster_id"],
    } for row in payload["evidence_items"]]
    score_rows = [{
        "股票代码": row["stock_code"],
        "Quant分": row["quant_score"],
        "Event Opportunity Score": row["event_opportunity_score"],
        "Evidence Confidence": row["evidence_confidence"],
        "Evidence Breadth": row["evidence_breadth"],
        "Risk Action": row["risk_action"],
        "V3初筛分": row["v3_screening_score"],
        "V3排名": row["v3_rank"],
    } for row in items]
    risk_rows = [{
        "股票代码": row["stock_code"],
        "Risk Action": row["risk_action"],
        "Hard Gate原因": "；".join(row.get("hard_gate_reasons") or []),
        "进入Top20": row["selected_top20"],
        "证据快照": row["event_snapshot_id"],
    } for row in items]
    versions = [{"字段": key, "值": value} for key, value in payload["manifest"].items() if "version" in key or "hash" in key or key.endswith("_id")]
    freshness = [{
        "股票代码": row["stock_code"],
        "决策时点": row["decision_as_of_time"],
        "搜索状态": row["search_status"],
        "Provider": row["provider"],
        "Provider已验证": row["provider_verified"],
        "直搜降级": row["direct_search_used"],
        "生产可用": row["production_eligible"],
        "Shadow可用": row["shadow_eligible"],
        "证据总数": len(row.get("items") or []),
        "评分合格证据": sum(
            1 for item in (row.get("items") or [])
            if item.get("score_eligible", True)
        ),
        "过期或未知证据": sum(
            1 for item in (row.get("items") or [])
            if not item.get("score_eligible", True)
        ),
    } for row in snapshots]
    return {
        "V3事件初筛Top20": [_display_item(row) for row in top20],
        "V2与V3对照": payload["comparison"],
        "Top100事件搜索状态": search_rows,
        "事件证据明细": evidence_rows,
        "Event Score拆解": score_rows,
        "事件风险与Hard Gate": risk_rows,
        "Pro复核结果": payload["pro_results"] or [{"状态": "SKIPPED", "说明": "Shadow默认关闭Pro"}],
        "Market Regime与部署限制": [payload["market_regime"]] if payload["market_regime"] else [{"状态": "NOT_AVAILABLE", "说明": "未改变V2部署结果"}],
        "主题集中度": payload["theme_concentration"] or [{"状态": "PRESERVED", "说明": "沿用现有V2主题集中度限制，不在本模块重写"}],
        "数据时效与来源": freshness,
        "Checkpoint复用审计": payload["checkpoint_audit"],
        "版本与Hash": versions,
        "数据质量问题": payload["data_quality"] or [{"issue_code": "", "detail": ""}],
    }


def _display_item(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "股票代码": row["stock_code"],
        "股票名称": row["stock_name"],
        "Quant排名": row["quant_rank"],
        "Quant分": row["quant_score"],
        "V3排名": row["v3_rank"],
        "V3初筛分": row["v3_screening_score"],
        "Event Opportunity Score": row["event_opportunity_score"],
        "Evidence Confidence": row["evidence_confidence"],
        "Risk Action": row["risk_action"],
        "搜索状态": row["search_status"],
    }


def _add_sheet(workbook: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(name)
    rows = rows or [{"状态": "EMPTY"}]
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header)) for header in headers])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, header in enumerate(headers, start=1):
        values = [str(row.get(header) or "") for row in rows]
        width = min(48, max(10, len(str(header)) + 2, max((len(value) for value in values), default=0) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width
        if header in {"股票代码", "stock_code"}:
            for cell in list(sheet.columns)[index - 1][1:]:
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value).zfill(6)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        if isinstance(value, list):
            value = "；".join(str(item) for item in value)
        else:
            value = "；".join(f"{key}={item}" for key, item in value.items())
    if isinstance(value, str):
        probe = value.lstrip()
        if probe.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    headers = list(rows[0]) if rows else ["status"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for row in rows or [{"status": "EMPTY"}]:
            writer.writerow({key: _cell_value(row.get(key)) for key in headers})
