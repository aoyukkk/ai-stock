from __future__ import annotations

import json
import os
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reporting.workbook_style import WorkbookStyleService
from reporting.source_row_style import apply_selection_source_rows


NAVY="193B63";BLUE="D9EAF7";PALE="EDF4F8";RED="C00000";GREEN="16833A";WHITE="FFFFFF";GOLD="FFF2CC";GRID="C9D6E2"


def export_full_a_midday(output_root:Path,run,report:dict[str,Any])->dict[str,str]:
    directory=Path(output_root)/report["trade_date"]/"午盘推荐_全A";directory.mkdir(parents=True,exist_ok=True);suffix=run.run_id[-8:];stem=f"全A午盘推荐_V2_2_{report['trade_date']}_{suffix}"
    json_path=directory/f"{stem}.json";md_path=directory/f"{stem}.md";audit_path=directory/f"全A午盘推荐_V2_2_{report['trade_date']}_audit_{suffix}.json";capacity_path=directory/f"全A午盘推荐_V2_2_{report['trade_date']}_capacity_{suffix}.json";xlsx_path=directory/f"{stem}.xlsx"
    capacity=report.get("capacity_audit") or {};_write_json(json_path,report);_write_json(capacity_path,capacity);_write_json(audit_path,{"run_id":run.run_id,"status":report["status"],"provider_audit":report.get("provider_audit",[]),"snapshot_semantic_validation":report.get("snapshot_semantic_validation"),"capacity_audit":capacity,"source_hashes_before":report.get("source_hashes_before"),"source_hashes_after":report.get("source_hashes_after"),"quant_hash_unchanged":report.get("quant_hash_unchanged"),"flash_hash_unchanged":report.get("flash_hash_unchanged"),"pro_hash_unchanged":report.get("pro_hash_unchanged"),"real_orders":0,"virtual_orders":0,"scheduler":False,"production_config_changed":False});md_path.write_text(_markdown(report),encoding="utf-8")
    workbook=_workbook(report);_apply_shared_style(workbook,Path(output_root),report);temporary=xlsx_path.with_suffix(".tmp.xlsx");workbook.save(temporary);os.replace(temporary,xlsx_path);_verify(xlsx_path,report["status"])
    return {"excel":str(xlsx_path),"json":str(json_path),"markdown":str(md_path),"capacity":str(capacity_path),"audit":str(audit_path)}


def export_full_a_failure(output_root:Path,run,report:dict[str,Any])->dict[str,str]:
    directory=Path(output_root)/report["trade_date"]/"午盘推荐_全A";directory.mkdir(parents=True,exist_ok=True);suffix=run.run_id[-8:];stem=f"全A午盘推荐_V2_2_{report['trade_date']}_失败_{suffix}"
    json_path=directory/f"{stem}.json";md_path=directory/f"{stem}.md";audit_path=directory/f"{stem}_audit.json";_write_json(json_path,report);_write_json(audit_path,report);md_path.write_text(_markdown(report),encoding="utf-8");return {"json":str(json_path),"markdown":str(md_path),"audit":str(audit_path)}


def _workbook(report):
    wb=Workbook();wb.remove(wb.active);results=report.get("results",[]);radar=report.get("radar_top200",[])
    overview=[["项目","结果"],["运行编号",report["run_id"]],["最终状态",report["status"]],["交易日",report["trade_date"]],["截止时间",report["cutoff"]],["市场状态",f"{report.get('previous_regime')} → {report.get('midday_regime')}"],["全A有效数量",report["counts"]["full_a_eligible_count"]],["全A覆盖率",report["counts"]["full_a_coverage"]],["雷达Top200",report["counts"]["radar_top200_count"]],["Admission PASS",report["counts"]["admission_distribution"].get("PASS",0)],["BUY_READY",report["counts"]["buy_ready"]],["下午观察",report["counts"]["afternoon_watch"]],["真实订单",0],["虚拟订单",0],["调度器","关闭"]]
    _sheet(wb,"01_午盘总览",overview,title="全A午盘推荐 V2.2")
    overview_ws=wb["01_午盘总览"]
    overview_ws.column_dimensions["B"].width=54
    overview_ws.row_dimensions[8].height=54
    _sheet(wb,"02_全A雷达Top200",_radar_rows(radar),title="全A午盘雷达 Top200")
    _sheet(wb,"03_Admission_PASS",_result_rows([r for r in results if r.get("admission_status_v2")=="PASS"]),title="Admission PASS")
    _sheet(wb,"04_BUY_READY",_result_rows([r for r in results if r.get("result_layer")=="BUY_READY"]),title="BUY_READY")
    _sheet(wb,"05_下午观察池",_result_rows([r for r in results if r.get("result_layer")=="AFTERNOON_WATCH"]),title="下午观察池")
    _sheet(wb,"06_高分市场阻断",_result_rows([r for r in results if r.get("result_layer")=="REGIME_BLOCKED_HIGH_SCORE"]),title="高分市场阻断")
    _sheet(wb,"07_行业集中复核",_result_rows([r for r in results if r.get("result_layer")=="CONCENTRATION_REVIEW"]),title="行业集中复核")
    _sheet(wb,"08_REVIEW池",_result_rows([r for r in results if r.get("admission_status_v2")=="REVIEW"]),title="REVIEW 池")
    _sheet(wb,"09_BLOCK池",_result_rows([r for r in results if r.get("admission_status_v2")=="BLOCK"]),title="BLOCK 池")
    _sheet(wb,"10_人工挑战池",_result_rows([r for r in results if r.get("pool_type")=="MANUAL_CHALLENGE_POOL"]),title="人工挑战池")
    breadth=report.get("breadth",{});_sheet(wb,"11_市场状态",[["指标","数值"],["上一状态",report.get("previous_regime")],["午盘状态",report.get("midday_regime")],["状态置信度",report.get("regime_confidence")],["状态原因","；".join(report.get("regime_reason") or [])]],title="全A市场状态")
    market_ws=wb["11_市场状态"];market_ws.column_dimensions["B"].width=54;market_ws.row_dimensions[5].height=48;market_ws.row_dimensions[7].height=54
    _sheet(wb,"12_行业状态",_industry_rows(report.get("industries",[])),title="全行业状态")
    _sheet(wb,"13_全A市场宽度",[["指标","数值"]]+[[key,value] for key,value in breadth.items()],title="FULL_A_MARKET_BREADTH")
    validation=report.get("snapshot_semantic_validation") or {};coverage=(report.get("capacity_audit") or {}).get("coverage_reconciliation") or {};unresolved=coverage.get("unresolved_codes") or {};unresolved_items=list(unresolved.items());unresolved_summary="；".join(f"{code}:{reason}" for code,reason in unresolved_items[:5]);unresolved_summary += f"；其余{len(unresolved_items)-5}条详见capacity JSON" if len(unresolved_items)>5 else "";quality=[["项目","数值"],["解析方法",report.get("asof_resolution_method")],["请求数量",report["counts"]["full_a_requested"]],["返回数量",report["counts"]["full_a_returned"]],["覆盖率",report["counts"]["full_a_coverage"]],["验证样本",validation.get("sample_count")],["一致样本",validation.get("consistent_count")],["验证结论",validation.get("passed")],["排除11:30后记录",report.get("post_cutoff_rows_excluded",0)],["未解决代码",len(unresolved)],["未解决明细",unresolved_summary]];_sheet(wb,"14_数据质量",quality,title="数据质量与覆盖审计");quality_ws=wb["14_数据质量"];quality_ws.column_dimensions["B"].width=70;quality_ws.row_dimensions[4].height=48;quality_ws.row_dimensions[13].height=54
    capacity=report.get("capacity_audit") or {};before=capacity.get("usage_before") or {};after=capacity.get("usage_after") or {};capacity_rows=[["项目","数值"],["容量模式",capacity.get("capacity_mode")],["原始调用上限",capacity.get("original_call_limit")],["运行级调用预算",capacity.get("run_scoped_call_budget")],["配置已恢复",capacity.get("call_budget_restored")],["可靠完整批量",coverage.get("reliable_complete_batch_size")],["按市场可靠批量",json.dumps(coverage.get("reliable_complete_batch_size_by_market") or {},ensure_ascii=False)],["自适应拆分次数",coverage.get("adaptive_split_count")],["单代码回退次数",coverage.get("single_code_fallback_count")],["SDK回退次数",coverage.get("sdk_fallback_count")],["分钟线回退次数",coverage.get("minute_fallback_count")],["配额前已用",before.get("market_data_used")],["配额后已用",after.get("market_data_used")],["实际配额增量",capacity.get("actual_usage_delta")],["配额安全状态",capacity.get("quota_safety_status")]];_sheet(wb,"15_容量审计",capacity_rows,title="运行级容量与配额审计");capacity_ws=wb["15_容量审计"];capacity_ws.column_dimensions["B"].width=70;capacity_ws.row_dimensions[9].height=54
    audit=[["能力","用途","市场","请求代码数","解析代码数","返回行数","延迟毫秒","错误类别","并发"]]+[[row.get("capability"),row.get("purpose"),row.get("market"),row.get("requested_code_count"),row.get("normalized_output_code_count"),row.get("returned_rows"),row.get("latency_ms"),row.get("error_category"),row.get("concurrency")] for row in report.get("provider_audit",[])];_sheet(wb,"16_调用审计",audit,title="外部调用审计")
    weights=report.get("radar_weight_hash");algorithm=[["项目","说明"],["雷达版本",report.get("radar_version")],["权重哈希",weights],["综合公式","40%正式Quant + 20%上午相对强度 + 15%量价 + 15%行业共振 + 10%开盘风险质量"],["基础数据","2026-07-17正式TUSHARE_BASELINE_V1全量Quant"],["时点数据","严格使用2026-07-20 11:30及以前的iFinD真实数据"],["排序规则","雷达分、相对强度、行业共振、Quant分、风险质量、股票代码"],["安全","Shadow只读；无真实订单；无虚拟订单；调度器关闭"]];_sheet(wb,"15_算法说明",algorithm,title="算法说明");algorithm_ws=wb["15_算法说明"];algorithm_ws.column_dimensions["B"].width=70;algorithm_ws.row_dimensions[5].height=48
    wb["15_算法说明"].title="17_算法说明"
    return wb


def _radar_rows(rows):
    header=["午盘排名","股票代码","股票名称","行业","Quant分","雷达分","上午相对强度","上午量价","行业共振","开盘风险质量","上午涨跌幅","上午振幅","行业内排名","数据质量","风险标签"]
    return [header]+[[r.get("midday_rank"),r.get("stock_code"),r.get("stock_name"),r.get("industry"),r.get("baseline_quant_score"),r.get("midday_radar_score"),r.get("morning_relative_strength"),r.get("morning_volume_price"),r.get("sector_resonance"),r.get("opening_risk_quality"),r.get("change_pct_to_cutoff"),r.get("amplitude_to_cutoff"),r.get("industry_rank"),r.get("data_quality"),"、".join(r.get("risk_flags") or [])] for r in rows]


def _result_rows(rows):
    header=["午盘排名","股票代码","股票名称","来源","行业","策略","Base Strategy","Live Strategy","Quant分","Radar分","Strategy Fit","Entry Timing V2.1","Admission分","Admission","上午涨跌幅","Price vs VWAP","市场状态","部署状态","集中门禁","Trigger","Flash","Pro","保守观察价","均衡观察价","最高可接受价","止损","止盈1","止盈2","升级条件","主要风险","数据质量"]
    data=[header]
    for r in rows:
        plan=r.get("price_plan") or {};flash=r.get("flash") or {};pro=r.get("pro") or {}
        data.append([r.get("midday_rank"),r.get("stock_code"),r.get("stock_name"),r.get("pool_type"),r.get("industry"),r.get("strategy_id"),r.get("base_strategy_id"),r.get("live_strategy_status"),r.get("quant_score"),r.get("midday_radar_score"),r.get("strategy_fit_score"),r.get("entry_timing_v2_score"),r.get("admission_ranking_score_v2"),r.get("admission_status_v2"),r.get("change_pct_to_cutoff"),_ratio(r.get("close_at_cutoff"),r.get("vwap_to_cutoff")),r.get("market_regime"),r.get("deployment_status"),r.get("concentration_status"),r.get("trigger_status"),flash.get("decision"),pro.get("decision"),plan.get("conservative_watch_price"),plan.get("balanced_watch_price"),plan.get("maximum_acceptable_price"),plan.get("stop_loss"),plan.get("take_profit_1"),plan.get("take_profit_2"),"、".join(plan.get("afternoon_upgrade_conditions") or []),"、".join(r.get("risk_flags") or r.get("block_reasons") or []),r.get("data_quality")])
    return data


def _industry_rows(rows):
    header=["行业排名","行业","股票数","有效数","等权涨幅","中位涨幅","上涨比例","成交强度","领涨股涨幅","离散度","行业强度分"]
    return [header]+[[r.get("industry_rank"),r.get("sector_name"),r.get("industry_stock_count"),r.get("industry_valid_count"),r.get("industry_return_equal_weight"),r.get("industry_median_return"),r.get("industry_positive_ratio"),r.get("industry_turnover_strength"),r.get("industry_top_stock_return"),r.get("industry_dispersion"),r.get("industry_strength_score")] for r in rows]


def _sheet(wb,name,rows,title):
    ws=wb.create_sheet(name);width=max(2,max((len(row) for row in rows),default=2));ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=width);cell=ws.cell(1,1,title);cell.fill=PatternFill("solid",fgColor=NAVY);cell.font=Font(color=WHITE,bold=True,size=15);cell.alignment=Alignment(horizontal="center",vertical="center");ws.row_dimensions[1].height=30
    if not rows:rows=[["状态"],["无符合条件记录"]]
    for rindex,row in enumerate(rows,3):
        for cindex,value in enumerate(row,1):
            cell=ws.cell(rindex,cindex,value);cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True);cell.border=Border(bottom=Side(style="thin",color=GRID));cell.fill=PatternFill("solid",fgColor=NAVY if rindex==3 else BLUE if rindex%2==0 else WHITE);cell.font=Font(color=WHITE if rindex==3 else "000000",bold=rindex==3)
            if rindex>3 and isinstance(value,float) and ("率" in str(ws.cell(3,cindex).value) or "涨跌幅" in str(ws.cell(3,cindex).value) or "振幅" in str(ws.cell(3,cindex).value) or "收益" in str(ws.cell(3,cindex).value) or "覆盖" in str(ws.cell(3,cindex).value)):cell.number_format="0.00%"
            if rindex>3 and str(ws.cell(3,cindex).value) in {"股票代码"}:cell.number_format="@";cell.value=str(value or "")
        ws.row_dimensions[rindex].height=34 if rindex>3 else 28
    ws.freeze_panes="A4";ws.auto_filter.ref=f"A3:{get_column_letter(width)}{max(3,ws.max_row)}";ws.sheet_view.showGridLines=False
    for col in range(1,width+1):
        values=[str(ws.cell(row,col).value or "") for row in range(3,min(ws.max_row,80)+1)];ws.column_dimensions[get_column_letter(col)].width=min(28,max(10,max((len(value) for value in values),default=8)+2))
    apply_selection_source_rows(ws,header_row=3,first_data_row=4)


def _verify(path,status):
    wb=load_workbook(path,data_only=False);assert wb["01_午盘总览"]["B5"].value==status
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.coordinate!="A1":assert cell.alignment.horizontal=="center" and cell.alignment.vertical=="center"
                if isinstance(cell.value,str):assert not any(token in cell.value for token in ("#REF!","#DIV/0!","#VALUE!","#NAME?","#N/A"))


def _markdown(report):
    if report.get("status") not in {"FULL_A_MIDDAY_SUCCESS","FULL_A_MIDDAY_EMPTY_BUY_READY","FULL_A_RULE_ONLY_INCOMPLETE","FULL_A_LLM_STAGE_FAILED","FULL_A_PARTIAL_SUCCESS"}:return f"# 全A午盘运行未完成\n\n- 运行编号：{report.get('run_id')}\n- 状态：{report.get('status')}\n- 原因：{report.get('error_code')} {report.get('error_message')}\n"
    focus=[row for row in report.get("results",[]) if row.get("result_layer")=="BUY_READY"];watch=[row for row in report.get("results",[]) if row.get("result_layer")=="AFTERNOON_WATCH"]
    lines=["# 全A午盘推荐 V2.2","",f"- 运行编号：{report['run_id']}",f"- 状态：{report['status']}",f"- 市场状态：{report.get('midday_regime')}",f"- 全A覆盖：{report['counts']['full_a_returned']} / {report['counts']['full_a_requested']}",f"- BUY_READY：{len(focus)}",f"- 下午观察：{len(watch)}","","## BUY_READY"]
    lines += [f"- {row['stock_code']} {row.get('stock_name') or ''}" for row in focus] or ["- 无"]
    lines += ["","## 下午观察"]+[f"- {row['stock_code']} {row.get('stock_name') or ''}：{'、'.join(row.get('trigger_reasons') or [])}" for row in watch] or ["- 无"]
    return "\n".join(lines)+"\n"


def _write_json(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
def _ratio(left,right):return float(left)/float(right)-1 if left is not None and right not in (None,0) else None


def _apply_shared_style(workbook,output_root:Path,report:dict[str,Any])->None:
    trade_date=date.fromisoformat(report["trade_date"])
    try:
        reference=WorkbookStyleService.resolve_recent_successful_reference(
            output_root,start=trade_date-timedelta(days=7),end=trade_date-timedelta(days=1)
        )
    except FileNotFoundError:
        return
    status=str(report["status"])
    complete=status in {"FULL_A_MIDDAY_SUCCESS","FULL_A_MIDDAY_EMPTY_BUY_READY","FULL_A_PARTIAL_SUCCESS"}
    nature="完整全A午盘推荐" if complete else "仅完成全A数据覆盖和Radar Top200，不是完整投资推荐"
    WorkbookStyleService(reference).align_midday_workbook(workbook,final_status=status,result_nature=nature)
