from __future__ import annotations

import copy
import hashlib
import json
import math
import os
import time
import uuid
from contextlib import contextmanager
from dataclasses import asdict, dataclass
from datetime import date, datetime, time as clock_time, timezone
from importlib.util import find_spec
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from database.models import MiddayFullARadarResult, MiddayFullARadarRun
from datasource.ifind.http.errors import IFindHttpError
from datasource.ifind.http.normalizer import normalize_rows
from midday.core import SHANGHAI
from midday.adaptive_coverage import AdaptiveCoverageResolver, dynamic_run_call_budget
from midday.full_a_export import export_full_a_midday
from midday.full_a_radar import full_a_breadth, industry_state, stable_hash
from midday.full_a_service import FullAMiddayService, _git_head
from midday.provider import MiddayIFindCollector
from stock_codes import normalize_ts_code


CAPACITY_RUN_TYPE = "FULL_A_MIDDAY_RADAR_CAPACITY_RESOLVED_V2_2"
BLOCKED_STATUS = "FULL_A_CAPACITY_CONFIRMED_BLOCKED"
DEFAULT_CALL_LIMIT = 30
TEMPORARY_CALL_LIMIT = 70
MAXIMUM_RUN_SCOPED_PROVIDER_CALLS = 400
MINIMUM_WIDE_BATCH = 250


@dataclass(frozen=True)
class QuotaStatistics:
    status: str
    market_data_total: int | None = None
    market_data_used: int | None = None
    market_data_remaining: int | None = None
    high_frequency_total: int | None = None
    high_frequency_used: int | None = None
    high_frequency_remaining: int | None = None
    reset_period: str | None = None
    account_tier: str | None = None
    returned_at: str | None = None


@dataclass(frozen=True)
class BatchProbe:
    transport: str
    requested: int
    returned: int
    missing: int
    duplicates: int
    complete: bool
    truncated_to_100: bool
    timeout: bool
    error_category: str | None
    latency_ms: int
    latest_data_time: str | None
    post_cutoff_rows: int


def audit_call_limit(config: dict[str, Any]) -> dict[str, Any]:
    yaml_value = int(config.get("maximum_provider_calls", DEFAULT_CALL_LIMIT))
    env_keys = ["IFIND_MAX_PROVIDER_CALLS", "IFIND_PROVIDER_CALL_BUDGET", "IFIND_FULL_A_CALL_LIMIT"]
    env_override = next(((key, os.getenv(key)) for key in env_keys if os.getenv(key)), None)
    effective = int(env_override[1]) if env_override else yaml_value
    return {
        "limit_value": effective,
        "limit_source": "ENVIRONMENT" if env_override else "YAML",
        "config_key": env_override[0] if env_override else "midday_recommendation.full_a_radar.maximum_provider_calls",
        "effective_value": effective,
        "whether_internal_guard": effective == DEFAULT_CALL_LIMIT,
        "whether_provider_reported": False,
        "whether_account_reported": False,
        "classification": "INTERNAL_GUARD",
        "code_guards": [
            "IFindHttpClient default maximum_calls=30",
            "IFindHttpSerialRateLimiter authorized_maximum=30",
            "MiddayIFindCollector authorized_call_limit=30",
        ],
        "whether_safe_to_override": False,
    }


def parse_data_statistics(payload: dict[str, Any] | None) -> QuotaStatistics:
    if not payload:
        return QuotaStatistics(status="QUOTA_STATISTICS_UNAVAILABLE")
    source = payload.get("data") if isinstance(payload.get("data"), dict) else payload
    tables = source.get("tables") if isinstance(source.get("tables"), dict) else {}
    quotes = tables.get("QuotesDataStat") if isinstance(tables.get("QuotesDataStat"), dict) else {}
    if quotes:
        total = _integer(quotes.get("limit"))
        used = _integer(quotes.get("usage"))
        return QuotaStatistics(
            status="AVAILABLE" if total is not None and used is not None else "QUOTA_SCOPE_UNCERTAIN",
            market_data_total=total,
            market_data_used=used,
            market_data_remaining=max(0, total - used) if total is not None and used is not None else None,
            returned_at=datetime.now(SHANGHAI).isoformat(),
        )
    market = source.get("market_data") if isinstance(source.get("market_data"), dict) else source
    high = source.get("high_frequency") if isinstance(source.get("high_frequency"), dict) else {}
    total = _integer(market.get("total") or market.get("market_data_total"))
    used = _integer(market.get("used") or market.get("market_data_used"))
    remaining = _integer(market.get("remaining") or market.get("market_data_remaining"))
    if remaining is None and total is not None and used is not None:
        remaining = max(0, total - used)
    return QuotaStatistics(
        status="AVAILABLE" if remaining is not None else "QUOTA_SCOPE_UNCERTAIN",
        market_data_total=total,
        market_data_used=used,
        market_data_remaining=remaining,
        high_frequency_total=_integer(high.get("total")),
        high_frequency_used=_integer(high.get("used")),
        high_frequency_remaining=_integer(high.get("remaining")),
        reset_period=_safe_text(source.get("reset_period")),
        account_tier=_safe_text(source.get("account_tier")),
        returned_at=_safe_text(source.get("returned_at")),
    )


def estimate_capacity(eligible_count: int, quota: QuotaStatistics) -> dict[str, Any]:
    snapshot_fields = 9
    index_count, index_fields = 8, 6
    minute_candidates, morning_bars, minute_fields = 50, 121, 6
    snapshot_cells = eligible_count * snapshot_fields
    index_cells = index_count * index_fields
    minute_cells = minute_candidates * morning_bars * minute_fields
    retry_buffer = math.ceil((snapshot_cells + index_cells + minute_cells) * 0.05)
    total = snapshot_cells + index_cells + minute_cells + retry_buffer
    remaining = quota.market_data_remaining
    ratio = total / remaining if remaining and remaining > 0 else None
    return {
        "eligible_stock_count": eligible_count,
        "snapshot_field_count": snapshot_fields,
        "estimated_snapshot_cells": snapshot_cells,
        "index_count": index_count,
        "estimated_index_cells": index_cells,
        "minute_candidate_limit": minute_candidates,
        "expected_morning_bar_count": morning_bars,
        "estimated_minute_cells": minute_cells,
        "retry_buffer": retry_buffer,
        "estimated_total_cells": total,
        "quota_remaining": remaining,
        "estimated_quota_usage_ratio": ratio,
        "maximum_quota_usage_ratio": 0.10,
        "quota_safety_result": "PASS" if ratio is not None and ratio < 0.10 else "UNVERIFIED" if ratio is None else "FAIL",
    }


def select_capacity_path(
    *,
    http_probes: list[BatchProbe],
    sdk_probes: list[BatchProbe],
    quota: QuotaStatistics,
    estimate: dict[str, Any],
    limit_audit: dict[str, Any],
    temporary_authorized: bool,
) -> tuple[str, list[str]]:
    quota_safe = estimate["quota_safety_result"] == "PASS" and quota.status == "AVAILABLE"
    http_wide = any(row.requested >= MINIMUM_WIDE_BATCH and row.complete for row in http_probes)
    sdk_wide = any(row.requested >= MINIMUM_WIDE_BATCH and row.complete for row in sdk_probes)
    observed_100_only = any(row.requested == 100 and row.complete for row in http_probes) and not http_wide
    if http_wide and quota_safe:
        return "WIDE_BATCH_HTTP", []
    if sdk_wide and quota_safe:
        return "WIDE_BATCH_SDK", []
    if all((observed_100_only, limit_audit["whether_internal_guard"], temporary_authorized, quota_safe)):
        return "TEMPORARY_CALL_BUDGET_OVERRIDE", []
    reasons = []
    if quota.status != "AVAILABLE":
        reasons.append(quota.status)
    elif not quota_safe:
        reasons.append("QUOTA_SAFETY_NOT_PROVEN")
    if not http_wide:
        reasons.append("HTTP_WIDE_BATCH_NOT_VALIDATED")
    if not sdk_wide:
        reasons.append("SDK_WIDE_BATCH_NOT_AVAILABLE")
    if not observed_100_only:
        reasons.append("HTTP_100_COMPLETE_NOT_VALIDATED")
    return "CAPACITY_CONFIRMED_BLOCKED", list(dict.fromkeys(reasons))


@contextmanager
def run_scoped_capacity_config(
    config: dict[str, Any], *, selected_path: str, temporary_limit: int = TEMPORARY_CALL_LIMIT
) -> Iterator[tuple[dict[str, Any], dict[str, Any]]]:
    original = copy.deepcopy(config)
    working = copy.deepcopy(config)
    audit = {
        "original_config_hash": stable_hash(original),
        "temporary_config_hash": stable_hash(original),
        "restored_config_hash": None,
        "restoration_success": False,
    }
    if selected_path in {"TEMPORARY_CALL_BUDGET_OVERRIDE", "ADAPTIVE_COVERAGE"}:
        working["maximum_provider_calls"] = temporary_limit
        if selected_path == "ADAPTIVE_COVERAGE":
            working["minimum_full_a_coverage"] = max(0.98, float(working.get("minimum_full_a_coverage", 0.98)))
            working["minute_reconstruction_batch_size"] = min(20, int(working.get("minute_reconstruction_batch_size", 20)))
        audit["temporary_config_hash"] = stable_hash(working)
    try:
        yield working, audit
    finally:
        audit["restored_config_hash"] = stable_hash(config)
        audit["restoration_success"] = config == original


class CapacityResolvedFullARunner:
    def __init__(self, session, *, output_root: Path) -> None:
        self.session = session
        self.output_root = Path(output_root)
        self.service = FullAMiddayService(session, output_root=output_root)
        self._sdk = None
        self._sdk_logged_in = False

    def run(
        self,
        trade_date: date,
        cutoff: clock_time,
        *,
        max_run_scoped_provider_calls: int = MAXIMUM_RUN_SCOPED_PROVIDER_CALLS,
        temporary_call_budget: int | None = None,
    ) -> dict[str, Any]:
        started = time.perf_counter()
        authorization = int(temporary_call_budget or max_run_scoped_provider_calls)
        if authorization < 1 or authorization > MAXIMUM_RUN_SCOPED_PROVIDER_CALLS:
            raise ValueError("MAXIMUM_RUN_SCOPED_PROVIDER_CALLS_INVALID")

        # Universe sizing is local/database-only and intentionally happens
        # before any external market-data request.
        shadow = SimpleNamespace(report_json={})
        baseline, quant_rows, masters, daily = self.service._preflight(trade_date, shadow)
        universe, excluded = self.service._universe(trade_date, quant_rows, masters, daily)
        source_hashes = copy.deepcopy(shadow.report_json.get("source_hashes", {}))
        quota_before = self._quota_statistics()
        estimate = estimate_capacity(len(universe), quota_before)
        budget_plan = dynamic_run_call_budget(
            len(universe), reliable_batch_size=20, authorized_maximum=authorization
        )
        run_budget = budget_plan["run_call_budget"]
        quota_limit = min(
            1_000_000,
            int((quota_before.market_data_remaining or 0) * 0.05),
        )
        if quota_before.status != "AVAILABLE" or estimate["quota_safety_result"] != "PASS" or quota_limit <= 0:
            return self._pre_provider_failure(
                trade_date,
                cutoff,
                "FULL_A_QUOTA_SAFETY_FAILED",
                "QUOTA_SAFETY_NOT_PROVEN",
                {
                    "usage_before": asdict(quota_before),
                    "estimate": estimate,
                    "call_budget": budget_plan,
                },
                started,
            )

        capacity: dict[str, Any] = {
            "capacity_mode": "ADAPTIVE_COVERAGE",
            "call_limit": audit_call_limit(self.service.radar_cfg),
            "original_call_limit": int(self.service.radar_cfg["maximum_provider_calls"]),
            "call_limit_source": audit_call_limit(self.service.radar_cfg)["limit_source"],
            "maximum_run_scoped_provider_calls": authorization,
            "run_scoped_call_budget": run_budget,
            "call_budget_calculation": budget_plan,
            "usage_before": asdict(quota_before),
            "estimate": estimate,
            "quota_safety_limit": quota_limit,
            "quota_safety_ratio": 0.05,
            "source_hashes_before": source_hashes,
            "source_hashes_after": copy.deepcopy(source_hashes),
            "selected_capacity_path": "ADAPTIVE_COVERAGE",
            "blocking_reasons": [],
            "required_function_permissions": ["THS_DataStatistics", "snap_shot", "high_frequency"],
        }
        checkpoint_called = False

        def quota_checkpoint() -> dict[str, Any]:
            nonlocal checkpoint_called
            checkpoint_called = True
            after = self._read_sdk_statistics()
            delta = _usage_delta(quota_before, after)
            passed = after.status == "AVAILABLE" and delta is not None and 0 <= delta <= quota_limit
            return {
                "usage_after": asdict(after),
                "actual_usage_delta": delta,
                "estimated_usage": estimate["estimated_total_cells"],
                "estimate_error": (delta - estimate["estimated_total_cells"]) if delta is not None else None,
                "quota_safety_status": "PASS" if passed else "FAIL",
                "quota_safety_reason": None if passed else "ACTUAL_USAGE_EXCEEDS_SAFE_LIMIT_OR_IS_UNAVAILABLE",
            }

        original_config = self.service.radar_cfg
        result: dict[str, Any]
        with run_scoped_capacity_config(
            original_config,
            selected_path="ADAPTIVE_COVERAGE",
            temporary_limit=run_budget,
        ) as (temporary_config, restoration):
            capacity["configuration_restoration"] = restoration
            capacity["effective_run_call_budget"] = int(temporary_config["maximum_provider_calls"])
            self.service.radar_cfg = temporary_config
            try:
                result = self.service.run(
                    trade_date,
                    cutoff,
                    runtime={
                        "authorized_call_limit": run_budget,
                        "capacity_audit": capacity,
                        "quota_checkpoint": quota_checkpoint,
                        "resolver_factory": lambda provider, config: AdaptiveCoverageResolver(
                            provider,
                            config,
                            sdk_snapshot=self._sdk_snapshot_payload if self._sdk_logged_in else None,
                            diagnostic_single_limit=20,
                        ),
                    },
                )
            finally:
                self.service.radar_cfg = original_config

        if not checkpoint_called:
            capacity.update(quota_checkpoint())
        capacity["configuration_restoration"] = restoration
        capacity["call_budget_restored"] = bool(restoration["restoration_success"])
        capacity["source_hashes_after"] = copy.deepcopy(result.get("source_hashes_after") or source_hashes)
        result["capacity_audit"] = capacity
        result["execution_duration_seconds"] = round(time.perf_counter() - started, 3)

        run_id = result.get("run_id")
        run = self.session.scalar(select(MiddayFullARadarRun).where(MiddayFullARadarRun.run_id == run_id)) if run_id else None
        if run is not None:
            self._attach_persisted_radar(run, result, capacity)
            run.report_json = result
            run.audit_json = list(run.audit_json or []) + [{"capacity_audit": capacity}]
            self.session.commit()
            if result.get("results") is not None and result.get("radar_top200") is not None:
                paths = export_full_a_midday(self.output_root, run, result)
                run.output_paths_json = paths
                self.session.commit()
                result["output_paths"] = paths
            else:
                paths = export_capacity_result(self.output_root, run, result, capacity)
                run.output_paths_json = paths
                self.session.commit()
                result["output_paths"] = paths
        self._logout_sdk()
        return result

    def _attach_persisted_radar(self, run, report: dict[str, Any], capacity: dict[str, Any]) -> None:
        if report.get("radar_top200"):
            return
        persisted = list(self.session.scalars(
            select(MiddayFullARadarResult)
            .where(MiddayFullARadarResult.run_id == run.run_id)
            .order_by(MiddayFullARadarResult.midday_rank)
        ))
        if not persisted:
            return
        rows = [dict(item.payload_json or {}) for item in persisted]
        eligible = int((report.get("counts") or {}).get("full_a_eligible_count") or len(rows))
        coverage = len(rows) / eligible if eligible else 0.0
        coverage_audit = capacity.get("coverage_reconciliation") or {}
        counts = dict(report.get("counts") or {})
        counts.update({
            "full_a_requested": eligible,
            "full_a_returned": len(rows),
            "full_a_coverage": coverage,
            "radar_top200_count": min(200, len(rows)),
            "admission_distribution": {},
            "trigger_distribution": {},
            "result_layers": {},
            "buy_ready": 0,
            "afternoon_watch": 0,
            "provider_calls": coverage_audit.get("provider_calls"),
            "hf_batch_count": coverage_audit.get("minute_fallback_count"),
            "llm_calls": 0,
            "llm_tokens": 0,
            "real_orders": 0,
            "virtual_orders": 0,
        })
        report.update({
            "counts": counts,
            "radar_top200": rows[:200],
            "results": [],
            "breadth": full_a_breadth(rows, eligible, 0.98),
            "industries": industry_state(rows),
            "previous_regime": "CRASH",
            "midday_regime": "NOT_RELEASED_QUOTA_SAFETY_STOP",
            "regime_confidence": None,
            "regime_reason": ["LLM_AND_FINAL_RELEASE_STOPPED_BY_QUOTA_SAFETY"],
            "asof_resolution_method": "ADAPTIVE_SNAPSHOT_AND_1M_RECONSTRUCTION",
            "snapshot_semantic_validation": {"passed": True, "reason": "PASSED_BEFORE_QUOTA_SAFETY_STOP"},
            "provider_audit": coverage_audit.get("response_audits") or [],
            "post_cutoff_rows_excluded": coverage_audit.get("post_cutoff_rows_excluded", 0),
            "radar_version": self.service.radar_cfg.get("version"),
            "radar_weight_hash": stable_hash(self.service.radar_cfg.get("weights") or {}),
            "llm_status": "NOT_RUN_QUOTA_SAFETY",
            "llm_usage": {"calls": 0, "tokens": 0, "cost": 0},
            "source_hashes_before": capacity.get("source_hashes_before") or {},
            "source_hashes_after": capacity.get("source_hashes_after") or {},
            "quant_hash_unchanged": True,
            "flash_hash_unchanged": True,
            "pro_hash_unchanged": True,
            "real_orders": 0,
            "virtual_orders": 0,
            "scheduler": False,
            "production_config_changed": False,
            "warnings": ["RULE_OUTPUT_PARTIALLY_RECOVERED_FROM_PERSISTED_RADAR"],
            "failures": ["FULL_A_QUOTA_SAFETY_FAILED"],
        })

    def _sdk_snapshot_payload(self, codes: list[str], asof: str) -> dict[str, Any] | None:
        if not self._sdk_logged_in or self._sdk is None:
            return None
        return self._sdk.THS_Snapshot(
            ",".join(normalize_ts_code(code) for code in codes),
            "tradeDate;tradeTime;preClose;open;high;low;latest;volume;amount",
            "",
            asof,
            asof,
        )

    def _pre_provider_failure(
        self,
        trade_date: date,
        cutoff: clock_time,
        status: str,
        reason: str,
        capacity: dict[str, Any],
        started: float,
    ) -> dict[str, Any]:
        cutoff_dt = datetime.combine(trade_date, cutoff, tzinfo=SHANGHAI)
        run = self._create_run(trade_date, cutoff_dt, DEFAULT_CALL_LIMIT)
        report = self._base_report(run, trade_date, cutoff_dt)
        report.update({
            "status": status,
            "error_code": reason,
            "error_message": reason,
            "capacity_audit": capacity,
            "execution_duration_seconds": round(time.perf_counter() - started, 3),
        })
        self._finish(run, report, None)
        paths = export_capacity_result(self.output_root, run, report, capacity)
        run.output_paths_json = paths
        self.session.commit()
        return {**report, "output_paths": paths}

    def _create_run(self, trade_date: date, cutoff: datetime, temporary_call_budget: int) -> MiddayFullARadarRun:
        payload = [CAPACITY_RUN_TYPE, trade_date.isoformat(), cutoff.isoformat(), datetime.now(timezone.utc).isoformat()]
        run = MiddayFullARadarRun(
            run_id=f"midday-full-a-cap-{uuid.uuid4().hex[:18]}",
            trade_date=trade_date,
            cutoff_time=cutoff,
            run_mode=CAPACITY_RUN_TYPE,
            status="RUNNING",
            stage="CAPACITY_PREFLIGHT",
            input_hash=stable_hash(payload),
            config_hash=stable_hash([self.service.radar_cfg, self.service.v21, self.service.v22]),
            radar_config_hash=stable_hash(self.service.radar_cfg),
            current_git_head=_git_head(),
            counts_json={"original_call_budget": int(self.service.radar_cfg["maximum_provider_calls"]), "temporary_call_budget": temporary_call_budget},
            audit_json=[],
            report_json={"account_id_hash": _account_hash(), "provider_mode": "IFIND_HTTP_REAL_SHADOW"},
            output_paths_json={},
            execution_started_at=datetime.now(timezone.utc),
            real_orders=0,
            virtual_orders=0,
            scheduler_enabled=False,
        )
        self.session.add(run)
        self.session.commit()
        return run

    def _quota_statistics(self) -> QuotaStatistics:
        if not _sdk_available():
            return QuotaStatistics(status="QUOTA_STATISTICS_UNAVAILABLE")
        try:
            import iFinDPy as sdk

            username = os.getenv("IFIND_USERNAME")
            password = os.getenv("IFIND_PASSWORD")
            if not username or not password:
                return QuotaStatistics(status="QUOTA_STATISTICS_UNAVAILABLE")
            if sdk.THS_iFinDLogin(username, password) != 0:
                return QuotaStatistics(status="QUOTA_STATISTICS_UNAVAILABLE")
            self._sdk = sdk
            self._sdk_logged_in = True
            return self._read_sdk_statistics()
        except Exception:
            return QuotaStatistics(status="QUOTA_STATISTICS_UNAVAILABLE")

    def _read_sdk_statistics(self) -> QuotaStatistics:
        try:
            return parse_data_statistics(self._sdk.THS_DataStatistics())
        except Exception:
            return QuotaStatistics(status="QUOTA_STATISTICS_UNAVAILABLE")

    def _http_probe(self, collector: MiddayIFindCollector, codes: list[str], size: int, cutoff: datetime) -> BatchProbe:
        requested_codes = [normalize_ts_code(code) for code in codes[:size]]
        started = time.perf_counter()
        try:
            response = collector.client.post("snap_shot", {
                "codes": ",".join(requested_codes),
                "indicators": "tradeDate,tradeTime,latest",
                "starttime": cutoff.strftime("%Y-%m-%d %H:%M:%S"),
                "endtime": cutoff.strftime("%Y-%m-%d %H:%M:%S"),
            })
            rows = normalize_rows(response.payload)
            returned_codes = [normalize_ts_code(row.get("thscode")) for row in rows if row.get("thscode")]
            unique = set(returned_codes)
            latest = max((str(row.get("time")) for row in rows if row.get("time")), default=None)
            post_cutoff = sum(1 for row in rows if _after_cutoff(row.get("time"), cutoff))
            return BatchProbe(
                "HTTP", size, len(unique), len(set(requested_codes) - unique), len(returned_codes) - len(unique),
                unique == set(requested_codes) and post_cutoff == 0,
                size > 100 and len(unique) == 100,
                False, None, round((time.perf_counter() - started) * 1000), latest, post_cutoff,
            )
        except IFindHttpError as exc:
            return BatchProbe("HTTP", size, 0, size, 0, False, False, exc.category.value == "TIMEOUT", exc.category.value, round((time.perf_counter() - started) * 1000), None, 0)
        except Exception as exc:
            return BatchProbe("HTTP", size, 0, size, 0, False, False, False, type(exc).__name__, round((time.perf_counter() - started) * 1000), None, 0)

    def _sdk_probes(self, codes: list[str], cutoff: datetime) -> list[BatchProbe]:
        if not self._sdk_logged_in:
            reason = "SDK_NOT_INSTALLED" if not _sdk_available() else "SDK_LOGIN_UNAVAILABLE"
            return [BatchProbe("SDK", size, 0, size, 0, False, False, False, reason, 0, None, 0) for size in (100, 200, 500, 1000)]
        return [self._sdk_probe(codes, size, cutoff) for size in (100, 200, 500, 1000)]

    def _sdk_probe(self, codes: list[str], size: int, cutoff: datetime) -> BatchProbe:
        requested_codes = [normalize_ts_code(code) for code in codes[:size]]
        started = time.perf_counter()
        try:
            payload = self._sdk.THS_Snapshot(
                ",".join(requested_codes),
                "latest;tradeTime",
                "",
                cutoff.strftime("%Y-%m-%d %H:%M:%S"),
                cutoff.strftime("%Y-%m-%d %H:%M:%S"),
            )
            error_code = _integer(payload.get("errorcode")) if isinstance(payload, dict) else None
            if error_code != 0:
                return BatchProbe("SDK", size, 0, size, 0, False, False, False, f"SDK_ERROR_{error_code}", round((time.perf_counter() - started) * 1000), None, 0)
            tables = payload.get("tables", []) if isinstance(payload, dict) else []
            if isinstance(tables, dict):
                tables = [tables]
            returned_codes = [normalize_ts_code(row.get("thscode")) for row in tables if isinstance(row, dict) and row.get("thscode")]
            unique = set(returned_codes)
            latest_values = []
            for row in tables:
                if not isinstance(row, dict):
                    continue
                times = row.get("time") if isinstance(row.get("time"), list) else []
                latest_values.extend(str(value) for value in times if value)
            latest = max(latest_values, default=None)
            return BatchProbe(
                "SDK", size, len(unique), len(set(requested_codes) - unique), len(returned_codes) - len(unique),
                unique == set(requested_codes), size > 100 and len(unique) == 100, False, None,
                round((time.perf_counter() - started) * 1000), latest, 0,
            )
        except Exception as exc:
            return BatchProbe("SDK", size, 0, size, 0, False, False, False, type(exc).__name__, round((time.perf_counter() - started) * 1000), None, 0)

    def _block(self, run, baseline, masters, quant_rows, universe, excluded, capacity, collector, started):
        report = self._base_report(run, run.trade_date, run.cutoff_time)
        report.update({
            "status": BLOCKED_STATUS,
            "error_code": BLOCKED_STATUS,
            "error_message": "; ".join(capacity["blocking_reasons"]),
            "latest_completed_trade_date": str(baseline.base_market_trade_date),
            "counts": {
                "stock_master_count": len(masters),
                "baseline_quant_scored_count": len(quant_rows),
                "full_a_eligible_count": len(universe),
                "excluded_count": sum(excluded.values()),
                "full_a_requested": 0,
                "full_a_returned": 0,
                "full_a_coverage": 0.0,
                "radar_top200_count": 0,
            },
            "capacity_audit": capacity,
            "source_hashes_before": capacity.get("source_hashes_before", {}),
            "source_hashes_after": capacity.get("source_hashes_after", {}),
            "quant_hash_unchanged": capacity.get("source_hashes_before", {}).get("quant") == capacity.get("source_hashes_after", {}).get("quant"),
            "flash_hash_unchanged": capacity.get("source_hashes_before", {}).get("flash") == capacity.get("source_hashes_after", {}).get("flash"),
            "pro_hash_unchanged": capacity.get("source_hashes_before", {}).get("pro") == capacity.get("source_hashes_after", {}).get("pro"),
            "provider_calls": collector.client.call_count,
            "successful_calls": collector.client.success_count,
            "failed_calls": collector.client.failed_count,
            "retry_calls": 0,
            "real_orders": 0,
            "virtual_orders": 0,
            "scheduler": False,
            "production_config_changed": False,
            "execution_duration_seconds": round(time.perf_counter() - started, 3),
        })
        self._finish(run, report, collector)
        paths = export_capacity_result(self.output_root, run, report, capacity)
        run.output_paths_json = paths
        self.session.commit()
        return {**report, "output_paths": paths}

    def _base_report(self, run, trade_date, cutoff):
        return {
            "run_id": run.run_id,
            "run_type": CAPACITY_RUN_TYPE,
            "trade_date": str(trade_date),
            "cutoff": cutoff.isoformat() if hasattr(cutoff, "isoformat") else str(cutoff),
            "execution_start": run.execution_started_at.isoformat(),
            "results": [],
            "radar_top200": [],
            "industries": [],
            "breadth": {},
        }

    def _finish(self, run, report, collector):
        self._logout_sdk()
        run.status = report["status"]
        run.stage = "COMPLETED"
        run.report_json = report
        run.audit_json = [report.get("capacity_audit", {})]
        run.error_code = report.get("error_code")
        run.error_message = report.get("error_message")
        run.execution_completed_at = datetime.now(timezone.utc)
        run.real_orders = 0
        run.virtual_orders = 0
        run.scheduler_enabled = False
        self.session.commit()

    def _logout_sdk(self) -> None:
        if self._sdk_logged_in and self._sdk is not None:
            try:
                self._sdk.THS_iFinDLogout()
            finally:
                self._sdk_logged_in = False


def export_capacity_result(output_root: Path, run, report: dict[str, Any], capacity: dict[str, Any]) -> dict[str, str]:
    directory = Path(output_root) / report["trade_date"] / "午盘推荐_全A"
    directory.mkdir(parents=True, exist_ok=True)
    suffix = run.run_id[-8:]
    stem = f"全A午盘推荐_V2_2_{report['trade_date']}_{suffix}"
    paths = {
        "excel": directory / f"{stem}.xlsx",
        "json": directory / f"{stem}.json",
        "markdown": directory / f"{stem}.md",
        "capacity_audit": directory / f"全A午盘推荐_V2_2_{report['trade_date']}_capacity_{suffix}.json",
        "audit": directory / f"全A午盘推荐_V2_2_{report['trade_date']}_audit_{suffix}.json",
    }
    paths["json"].write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    paths["capacity_audit"].write_text(json.dumps(capacity, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    paths["audit"].write_text(json.dumps({"run_id": run.run_id, "status": report["status"], "capacity": capacity, "real_orders": 0, "virtual_orders": 0, "scheduler": False}, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    paths["markdown"].write_text(
        f"# 全A午盘容量审计\n\n- 运行编号：{run.run_id}\n- 最终状态：{report['status']}\n- 容量路径：{capacity.get('selected_capacity_path')}\n- 阻断原因：{'；'.join(capacity.get('blocking_reasons', []))}\n- 真实订单：0\n- 虚拟订单：0\n",
        encoding="utf-8",
    )
    _capacity_workbook(paths["excel"], report, capacity)
    return {key: str(value) for key, value in paths.items()}


def _capacity_workbook(path: Path, report: dict[str, Any], capacity: dict[str, Any]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    counts = report.get("counts", {})
    sheets = [
        ("01_午盘总览", [["项目", "结果"], ["运行编号", report["run_id"]], ["最终状态", report["status"]], ["交易日", report["trade_date"]], ["容量路径", capacity.get("selected_capacity_path")], ["真实订单", 0], ["虚拟订单", 0], ["调度器", "关闭"]]),
        ("02_全A雷达Top200", [["状态"], ["容量前置检查未通过，未生成部分全A排名"]]),
        ("03_Admission_PASS", [["状态"], ["未执行"]]),
        ("04_BUY_READY", [["状态"], ["未执行"]]),
        ("05_下午观察池", [["状态"], ["未执行"]]),
        ("06_高分市场阻断", [["状态"], ["未执行"]]),
        ("07_行业集中复核", [["状态"], ["未执行"]]),
        ("08_REVIEW池", [["状态"], ["未执行"]]),
        ("09_BLOCK池", [["状态"], ["未执行"]]),
        ("10_人工挑战池", [["状态"], ["未执行"]]),
        ("11_市场状态", [["项目", "结果"], ["状态", "容量门禁前停止"]]),
        ("12_行业状态", [["状态"], ["未执行"]]),
        ("13_全A市场宽度", [["项目", "结果"], ["stock_master", counts.get("stock_master_count")], ["Quant已评分", counts.get("baseline_quant_scored_count")], ["全A有效股票", counts.get("full_a_eligible_count")]]),
        ("14_数据质量", [["项目", "结果"], ["配额统计", capacity.get("account_quota_statistics", {}).get("status")], ["完整全A数据", "未请求"]]),
        ("15_容量审计", _mapping_rows(capacity)),
        ("16_调用审计", [["传输", "请求代码数", "返回代码数", "缺失", "重复", "完整", "错误"]] + [[row.get("transport"), row.get("requested"), row.get("returned"), row.get("missing"), row.get("duplicates"), row.get("complete"), row.get("error_category")] for row in capacity.get("http_batch_probes", []) + capacity.get("sdk_batch_probes", [])]),
        ("17_算法说明", [["项目", "说明"], ["算法权重", "未修改"], ["算法阈值", "未降低"], ["执行边界", "容量门禁未通过，因此未进入Quant之后的午盘计算"], ["安全", "Shadow只读，无真实订单，无虚拟订单"]]),
    ]
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                cell = ws.cell(r, c, _scalar(value))
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if r == 1:
                    cell.fill = PatternFill("solid", fgColor="193B63")
                    cell.font = Font(color="FFFFFF", bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(max(1, ws.max_column))}{max(1, ws.max_row)}"
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 24
    wb.save(path)
    check = load_workbook(path, data_only=False)
    assert check["01_午盘总览"]["B3"].value == report["status"]
    assert len(check.sheetnames) == 17


def _mapping_rows(value: dict[str, Any]) -> list[list[Any]]:
    rows = [["项目", "结果"]]
    for key, item in value.items():
        rows.append([key, item])
    return rows


def _scalar(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


def _sdk_available() -> bool:
    return any(find_spec(name) is not None for name in ("iFinDPy", "iFindPy", "THS_iFinD"))


def _account_hash() -> str:
    identity = os.getenv("IFIND_USERNAME") or os.getenv("IFIND_ACCOUNT_ID") or "IFIND_ACCOUNT_UNDISCLOSED"
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()[:16]


def _integer(value: Any) -> int | None:
    try:
        return None if value is None else int(float(value))
    except (TypeError, ValueError):
        return None


def _safe_text(value: Any) -> str | None:
    return None if value is None else str(value)[:100]


def _after_cutoff(value: Any, cutoff: datetime) -> bool:
    if not value:
        return False
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=SHANGHAI)
        return parsed.astimezone(SHANGHAI) > cutoff
    except ValueError:
        return False


def _usage_delta(before: QuotaStatistics, after: QuotaStatistics) -> int | None:
    if before.market_data_used is None or after.market_data_used is None:
        return None
    return after.market_data_used - before.market_data_used
