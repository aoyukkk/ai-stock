from __future__ import annotations
import hashlib,json
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from reporting.source_row_style import apply_selection_source_rows

CENTER=Alignment(horizontal="center",vertical="center",wrap_text=True)
SHEETS=("01_午盘结论","02_重点候选","03_下午观察池","04_被阻断候选","05_人工挑战池","06_市场状态","07_数据质量","08_调用审计")

def export_midday_v22(root:Path,run,report:dict)->dict:
    folder=Path(root)/run.trade_date.isoformat()/"午盘推荐";folder.mkdir(parents=True,exist_ok=True);suffix="";stem=f"午盘推荐_V2_2_{run.trade_date.isoformat()}"
    audit_name=f"{stem}_audit_v2.json"
    if (folder/audit_name).exists():suffix=f"_{run.run_id[-8:]}"
    xlsx=folder/f"{stem}{suffix}.xlsx";json_path=folder/f"{stem}{suffix}.json";md=folder/f"{stem}{suffix}.md";audit=folder/(f"{stem}_audit_v2{suffix}.json")
    results=report.get("results") or [];focus=[r for r in results if r.get("result_layer")=="BUY_READY"];watch=[r for r in results if r.get("result_layer")=="AFTERNOON_WATCH"];regime_blocked=[r for r in results if r.get("result_layer")=="REGIME_BLOCKED_HIGH_SCORE"];concentration=[r for r in results if r.get("result_layer")=="CONCENTRATION_REVIEW"];blocked=[r for r in results if r.get("result_layer")=="BLOCKED"][:10];manual=[r for r in results if r.get("pool_type")=="MANUAL_CHALLENGE_POOL"]
    blocked_display=(regime_blocked+concentration+blocked)[:20]
    conclusion=[{"运行编号":run.run_id,"交易日期":run.trade_date,"截止时间":str(report.get("cutoff")),"运行状态":report.get("status",run.status),"上一状态":report.get("previous_regime"),"午盘状态":report.get("midday_regime"),"可买准备":len(focus),"下午观察":len(watch),"状态门禁高分":len(regime_blocked),"集中度复核":len(concentration),"真实订单":0,"虚拟订单":0,"下午复核":"必须"}]
    market=[{"上一状态":report.get("previous_regime"),"当前状态":report.get("midday_regime"),"置信度":report.get("regime_confidence"),"转换原因":report.get("transition_reason"),"冷却期":report.get("cooldown"),"广度范围":report.get("live_breadth_scope")}]
    quality=[{"代码":code,"数据质量":row.get("data_quality"),"最后合法时间":row.get("last_eligible_bar_time"),"分钟数量":row.get("minute_bar_count"),"重建方法":row.get("reconstruction_method"),"截止合规":row.get("cutoff_compliant")} for code,row in (report.get("index_quality") or {}).items()]
    audits=report.get("provider_audit") or [{"状态":"未调用Provider"}]
    wb=Workbook();wb.remove(wb.active)
    for title,rows in zip(SHEETS,(conclusion,focus,watch,blocked_display,manual,market,quality,audits)):_sheet(wb,title,rows)
    wb.save(xlsx);wb.close();_verify(xlsx)
    serialized=json.dumps(report,ensure_ascii=False,indent=2,default=str);json_path.write_text(serialized,encoding="utf-8")
    md.write_text(_markdown(report,focus,watch,regime_blocked,concentration,blocked),encoding="utf-8")
    audit_payload={"run_id":run.run_id,"status":report.get("status",run.status),"trade_date":run.trade_date.isoformat(),"cutoff":report.get("cutoff"),"provider_audit":report.get("provider_audit",[]),"counts":report.get("counts",{}),"hashes":report.get("hashes",{}),"real_orders":0,"virtual_orders":0,"scheduler":False,"production_config_changed":False,"workbook_sha256":hashlib.sha256(xlsx.read_bytes()).hexdigest(),"formula_error_scan":"PASSED","centered":True,"stock_code_text":True}
    audit.write_text(json.dumps(audit_payload,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
    return {"excel":str(xlsx),"json":str(json_path),"markdown":str(md),"audit":str(audit)}

def _sheet(wb,title,rows):
    ws=wb.create_sheet(title);headers=list(rows[0]) if rows else ["暂无数据"];ws.append(headers)
    for row in rows:ws.append([_cell(row.get(h)) for h in headers])
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
    for c in ws[1]:c.fill=PatternFill("solid",fgColor="1F4E78");c.font=Font(name="Microsoft YaHei",bold=True,color="FFFFFF")
    for row in ws.iter_rows():
        for cell in row:cell.alignment=CENTER
        ws.row_dimensions[row[0].row].height=34 if any(isinstance(c.value,str) and len(c.value)>28 for c in row) else 25
    for i,h in enumerate(headers,1):
        if "股票代码" in str(h) or str(h) in {"代码","stock_code"}:
            for cell in ws[get_column_letter(i)]:cell.number_format="@"
        ws.column_dimensions[get_column_letter(i)].width=min(34,max(12,max(len(str(ws.cell(r,i).value or "")) for r in range(1,min(ws.max_row,120)+1))*1.4+2))
    apply_selection_source_rows(ws,header_row=1,first_data_row=2)
def _cell(v):return json.dumps(v,ensure_ascii=False,sort_keys=True,default=str) if isinstance(v,(dict,list)) else v
def _verify(path):
    wb=load_workbook(path,data_only=False)
    try:
        assert wb.sheetnames==list(SHEETS)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:assert cell.alignment.horizontal=="center" and cell.alignment.vertical=="center" and cell.alignment.wrap_text;assert not (isinstance(cell.value,str) and cell.value.startswith(("#REF!","#DIV/0!","#VALUE!","#NAME?")))
    finally:wb.close()
def _markdown(report,focus,watch,regime_blocked,concentration,blocked):
    lines=[f"# 2026-07-20 午盘推荐 V2.2","",f"- 运行编号：{report.get('run_id')}",f"- 最终状态：{report.get('status')}",f"- 截止时间：{report.get('cutoff')}",f"- 市场状态：{report.get('previous_regime')} → {report.get('midday_regime')}",f"- 可买准备：{len(focus)}",f"- 下午观察：{len(watch)}","","## BUY_READY"]
    lines.extend([f"- {r.get('stock_code')} {r.get('stock_name') or ''}：{r.get('strategy_id')}" for r in focus] or ["- 无"]);lines.extend(["","## AFTERNOON_WATCH"]);lines.extend([f"- {r.get('stock_code')} {r.get('stock_name') or ''}：未满足 {'、'.join(r.get('trigger_reasons') or [])}；观察条件 {_cell(r.get('afternoon_recheck'))}" for r in watch] or ["- 无"]);lines.extend(["","## REGIME_BLOCKED_HIGH_SCORE"]);lines.extend([f"- {r.get('stock_code')} {r.get('stock_name') or ''}" for r in regime_blocked] or ["- 无"]);lines.extend(["","## CONCENTRATION_REVIEW"]);lines.extend([f"- {r.get('stock_code')} {r.get('stock_name') or ''}" for r in concentration] or ["- 无"]);lines.extend(["","## BLOCKED"]);lines.extend([f"- {r.get('stock_code')} {r.get('stock_name') or ''}：{'、'.join(r.get('block_reasons') or r.get('trigger_reasons') or [])}" for r in blocked] or ["- 无"]);return "\n".join(lines)+"\n"
