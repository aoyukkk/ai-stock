from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reporting.source_row_style import apply_selection_source_rows

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
SHEETS = ("01_复核结论", "02_逐股复核", "03_条件观察计划", "04_市场状态", "05_调用审计")


def export_v22_afternoon(root: Path, run, report: dict) -> dict:
    folder = Path(root) / run.trade_date.isoformat() / "下午复核"; folder.mkdir(parents=True, exist_ok=True)
    stem = f"下午复核_V2_2_{run.trade_date.isoformat()}_{run.run_id[-8:]}"
    xlsx, json_path, md, audit = (folder / f"{stem}{suffix}" for suffix in (".xlsx", ".json", ".md", "_audit.json"))
    results = report.get("results") or []; counts = report.get("counts") or {}
    summary = [{"运行编号": run.run_id, "午盘运行": report.get("midday_run_id"), "复核时间": report.get("recheck_time"), "最终状态": report.get("status"), "午盘市场状态": report.get("previous_midday_regime"), "下午市场状态": report.get("afternoon_regime"), "复核股票": len(results), "BUY_READY": counts.get("result_layers", {}).get("BUY_READY", 0), "AFTERNOON_WATCH": counts.get("result_layers", {}).get("AFTERNOON_WATCH", 0), "真实订单": 0, "虚拟订单": 0}]
    plans = [{"stock_code": row["stock_code"], "stock_name": row.get("stock_name"), **(row.get("price_plan") or {})} for row in results]
    market = [{"午盘状态": report.get("previous_midday_regime"), "下午状态": report.get("afternoon_regime"), "转换原因": report.get("regime_reasons"), "横截面": report.get("breadth")}]
    wb = Workbook(); wb.remove(wb.active)
    for title, rows in zip(SHEETS, (summary, results, plans, market, report.get("provider_audit") or [])): _sheet(wb, title, rows)
    wb.save(xlsx); wb.close(); _verify(xlsx)
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    lines = [f"# {run.trade_date.isoformat()} 下午复核 V2.2", "", f"- 最终状态：{report.get('status')}", f"- 市场状态：{report.get('previous_midday_regime')} → {report.get('afternoon_regime')}", "", "## 逐股结果"]
    lines.extend(f"- {row['stock_code']} {row.get('stock_name') or ''}：{row['prior_layer']} → {row['result_layer']}；触发 {row['trigger_status']}；未满足 {'、'.join(row.get('trigger_reasons') or []) or '无'}" for row in results)
    md.write_text("\n".join(lines)+"\n", encoding="utf-8")
    audit_payload = {"run_id": run.run_id, "status": report.get("status"), "provider_audit": report.get("provider_audit"), "counts": counts, "real_orders": 0, "virtual_orders": 0, "scheduler": False, "workbook_sha256": hashlib.sha256(xlsx.read_bytes()).hexdigest(), "centered": True, "stock_code_text": True}
    audit.write_text(json.dumps(audit_payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return {"excel": str(xlsx), "json": str(json_path), "markdown": str(md), "audit": str(audit)}


def _sheet(wb, title, rows):
    ws = wb.create_sheet(title); headers = list(rows[0]) if rows else ["暂无数据"]; ws.append(headers)
    for row in rows: ws.append([json.dumps(row.get(key), ensure_ascii=False, default=str) if isinstance(row.get(key), (dict, list)) else row.get(key) for key in headers])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions; ws.sheet_view.showGridLines = False
    for cell in ws[1]: cell.fill = PatternFill("solid", fgColor="1F4E78"); cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    for row in ws.iter_rows():
        for cell in row: cell.alignment = CENTER
    for index, header in enumerate(headers, 1):
        if header in {"stock_code", "股票代码", "代码"}:
            for cell in ws[get_column_letter(index)]: cell.number_format = "@"
        ws.column_dimensions[get_column_letter(index)].width = min(36, max(12, max(len(str(ws.cell(row, index).value or "")) for row in range(1, ws.max_row+1))*1.3+2))
    apply_selection_source_rows(ws, header_row=1, first_data_row=2)


def _verify(path):
    wb = load_workbook(path, data_only=False)
    try:
        assert wb.sheetnames == list(SHEETS)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None: assert cell.alignment.horizontal == "center" and cell.alignment.vertical == "center" and cell.alignment.wrap_text
    finally: wb.close()
