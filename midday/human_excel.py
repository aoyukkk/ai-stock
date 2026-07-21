from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import uuid
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from sqlalchemy import select
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment

from database.models import MiddayRecommendationResult, MiddayRecommendationRun, StockMaster
from database.models.quant_run import QuantRankResult
from backend.core.config_manager import ConfigManager
from stock_codes import display_stock_code, normalize_ts_code


ACTION_LABELS = {
    "AFTERNOON_PREPARE_ENTRY": "下午准备介入",
    "WAIT_PULLBACK": "等待回落",
    "KEEP_WATCH": "继续观察",
    "REMOVE_FROM_POOL": "移出候选池",
    "DO_NOT_CHASE": "不追高",
    "MANUAL_REVIEW": "人工复核",
    "CONTINUE_HOLD": "继续持有",
    "HOLD_WITH_TIGHT_STOP": "持有并收紧止损",
    "REDUCE_IF_WEAKENS": "走弱时减仓",
    "EXIT_IF_TRIGGERED": "触发条件时退出",
    "T_PLUS_ONE_LOCKED": "T+1锁定",
    "DATA_INSUFFICIENT": "数据不足",
}

FLASH_LABELS = {
    "PASS": "通过",
    "WATCH": "观察",
    "REJECT": "暂不考虑",
}

FEATURE_LABELS = {
    "MORNING_FULL_MINUTE": "上午完整分钟线",
    "MORNING_SNAPSHOT_ONLY": "上午快照",
    "BASELINE_ONLY": "前日量化基线",
    "DATA_CONFLICTED": "数据冲突",
}

SOURCE_LABELS = {
    "BASE_TOP100": "前日量化前100",
    "MANUAL": "人工关注",
    "HUMAN_POSITION": "人工持仓",
    "AI_POSITION": "模拟持仓",
    "ACTIVE_ORDER_PLAN": "已有计划",
}


def build_midday_human_payload(
    session,
    run: MiddayRecommendationRun,
    rows: Iterable[MiddayRecommendationResult],
) -> dict[str, Any]:
    results = list(rows)
    final = sorted(
        (row for row in results if row.pro_rank is not None),
        key=lambda row: (row.pro_rank or 9999, row.stock_code),
    )
    quant_rows = list(
        session.scalars(
            select(QuantRankResult)
            .where(QuantRankResult.quant_run_id == run.baseline_quant_run_id)
            .order_by(QuantRankResult.rank)
            .limit(100)
        )
    )
    stock_codes = {normalize_ts_code(row.stock_code) for row in results}
    stock_codes.update(normalize_ts_code(row.stock_code) for row in quant_rows)
    masters = {
        normalize_ts_code(row.code): row
        for row in session.scalars(select(StockMaster).where(StockMaster.code.in_(stock_codes)))
    }
    result_by_code = {row.stock_code: row for row in results}
    result_by_display = {display_stock_code(row.stock_code): row for row in results}

    candidates = [_candidate_row(row, masters.get(row.stock_code)) for row in final]
    minimum_recommendation_score = float(
        ConfigManager(session=session)
        .get_effective_config()["values"]
        .get("selection_performance.minimum_recommendation_score", 60)
    )
    recommendations = [
        row for row in candidates
        if row["深度复核分"] is not None
        and float(row["深度复核分"]) >= minimum_recommendation_score
    ]
    orders = [_price_row(row) for row in final]
    fundamentals = [_context_row(row, masters.get(row.stock_code)) for row in final]
    quant_top100 = []
    for quant in quant_rows:
        canonical_code = normalize_ts_code(quant.stock_code)
        result = result_by_display.get(display_stock_code(canonical_code))
        master = masters.get(canonical_code)
        quant_top100.append({
            "量化排名": quant.rank,
            "股票代码": display_stock_code(quant.stock_code),
            "股票名称": (master.name if master else None) or (result.stock_name if result else ""),
            "一级行业": master.industry if master else "",
            "量化总分": _number(quant.total_score),
            "技术得分": _number(quant.technical_score),
            "资金得分": _number(quant.capital_score),
            "情绪得分": _number(quant.emotion_score),
            "动量得分": _number(quant.momentum_score),
            "风险得分": _number(quant.risk_score),
            "进入二筛": "是" if result and result.midday_flash_score is not None else "否",
            "进入重点候选": "是" if result and result.pro_rank is not None else "否",
        })

    failures = (run.checkpoint_json or {}).get("failures") or []
    issues = []
    for failure in failures:
        code = str(failure.get("stock_code") or "")
        row = result_by_code.get(code)
        issues.append({
            "股票代码": display_stock_code(code),
            "股票名称": row.stock_name if row else "",
            "问题类型": "深度复核未完成",
            "当前状态": "待人工确认",
            "说明": "结构化结果未通过校验，本次未进入最终推荐。",
        })

    held_count = sum(row.position_status == "HELD" for row in results)
    return {
        "output_mode": "midday",
        "title": f"{run.session_trade_date.isoformat()} A股午间观察清单",
        "status_line": "按上午11:30前数据整理，供下午盘中复核。",
        "trade_date": run.session_trade_date.isoformat(),
        "target_date": run.session_trade_date.isoformat(),
        "summary": {
            "量化股票数": len(quant_rows),
            "二筛股票数": run.flash_count,
            "重点候选数": len(final),
            "今日推荐数": len(recommendations),
            "非零仓位数量": sum((_number(row.suggested_weight) or 0) > 0 for row in final),
            "当前问题数量": len(issues),
            "历史已解决问题数": 0,
            "确认持仓数": held_count,
        },
        "minimum_recommendation_score": minimum_recommendation_score,
        "recommendations": recommendations,
        "top10": candidates[:10],
        "candidates": candidates,
        "orders": orders,
        "fundamentals": fundamentals,
        "quant_top100": quant_top100,
        "issues": issues,
    }


def write_midday_human_workbook(payload: dict[str, Any], output: Path, preview_dir: Path) -> None:
    output = output.resolve()
    preview_dir = preview_dir.resolve()
    build_dir = output.parent / f".midday-human-build-{uuid.uuid4().hex}"
    build_dir.mkdir(parents=True, exist_ok=False)
    link = build_dir / "node_modules"
    try:
        shutil.copy2(Path(__file__).resolve().parents[1] / "scripts" / "build_human_daily_excel.mjs", build_dir)
        shutil.copy2(Path(__file__).resolve().parents[1] / "scripts" / "excel_alignment.mjs", build_dir)
        shutil.copy2(Path(__file__).resolve().parents[1] / "scripts" / "render_existing_excel.mjs", build_dir)
        payload_path = (build_dir / "payload.json").resolve()
        payload_path.write_text(json.dumps(payload, ensure_ascii=False, default=str), encoding="utf-8")
        dependency_root = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules"
        if not (dependency_root / "@oai" / "artifact-tool").exists():
            raise RuntimeError("ARTIFACT_TOOL_DEPENDENCY_NOT_FOUND")
        linked = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(dependency_root)],
            capture_output=True,
            text=True,
            check=False,
        )
        if linked.returncode != 0:
            raise RuntimeError("ARTIFACT_TOOL_JUNCTION_FAILED")
        completed = subprocess.run(
            ["node", "build_human_daily_excel.mjs", str(payload_path), str(output), str(preview_dir)],
            cwd=build_dir,
            check=False,
        )
        if completed.returncode != 0 and not output.exists():
            raise RuntimeError(f"MIDDAY_HUMAN_WORKBOOK_EXPORT_FAILED:{completed.returncode}")
        _polish_workbook(output)
        final_preview_dir = preview_dir / "最终版"
        rendered = subprocess.run(
            ["node", "render_existing_excel.mjs", str(output), str(final_preview_dir)],
            cwd=build_dir,
            check=False,
        )
        if rendered.returncode != 0 and len(list(final_preview_dir.glob("*.png"))) < 7:
            raise RuntimeError(f"MIDDAY_HUMAN_WORKBOOK_RENDER_FAILED:{rendered.returncode}")
    finally:
        Path(f"{output}.inspect.ndjson").unlink(missing_ok=True)
        if link.exists():
            os.rmdir(link)
        shutil.rmtree(build_dir, ignore_errors=True)


def _candidate_row(row: MiddayRecommendationResult, master: StockMaster | None) -> dict[str, Any]:
    return {
        "深度复核排名": row.pro_rank,
        "股票代码": display_stock_code(row.stock_code),
        "股票名称": row.stock_name or (master.name if master else ""),
        "入选来源": _sources(row.pool_sources),
        "量化排名": row.base_quant_rank,
        "量化得分": _number(row.base_quant_score),
        "二筛得分": _number(row.midday_flash_score),
        "二筛结论": FLASH_LABELS.get(row.flash_decision or "", row.flash_decision or ""),
        "深度复核分": _number(row.midday_pro_score),
        "复核优先级": "高" if (row.pro_rank or 999) <= 5 else "中",
        "一级行业": master.industry if master else "",
        "产业链": "",
        "财务状态": "午间未更新",
        "建议仓位": _number(row.suggested_weight),
        "建议股数": None,
        "参考价": _number(row.recommended_price),
        "止损价": _number(row.stop_loss),
        "第二目标价": _number(row.take_profit_2),
        "风险收益比": _risk_reward(row.recommended_price, row.stop_loss, row.take_profit_2),
        "核心逻辑": _lines(row.key_reasons),
        "主要风险": _lines(row.key_risks),
        "当前状态": ACTION_LABELS.get(row.candidate_action, row.candidate_action),
    }


def _price_row(row: MiddayRecommendationResult) -> dict[str, Any]:
    rr1 = _risk_reward(row.recommended_price, row.stop_loss, row.take_profit_1)
    rr2 = _risk_reward(row.recommended_price, row.stop_loss, row.take_profit_2)
    return {
        "深度复核排名": row.pro_rank,
        "股票代码": display_stock_code(row.stock_code),
        "股票名称": row.stock_name or "",
        "入选来源": _sources(row.pool_sources),
        "量化得分": _number(row.base_quant_score),
        "二筛得分": _number(row.midday_flash_score),
        "二筛结论": FLASH_LABELS.get(row.flash_decision or "", row.flash_decision or ""),
        "深度复核分": _number(row.midday_pro_score),
        "保守价": None,
        "均衡价": _number(row.recommended_price),
        "积极价": _number(row.max_acceptable_price),
        "参考价": _number(row.recommended_price),
        "最高接受价": _number(row.max_acceptable_price),
        "止损价": _number(row.stop_loss),
        "第一目标价": _number(row.take_profit_1),
        "第二目标价": _number(row.take_profit_2),
        "第一目标风险收益比": rr1,
        "第二目标风险收益比": rr2,
        "当前风险收益比": rr2,
        "建议仓位": _number(row.suggested_weight),
        "建议资金": None,
        "建议股数": None,
        "预计最大损失": None,
        "说明": f"{ACTION_LABELS.get(row.candidate_action, row.candidate_action)}；{FEATURE_LABELS.get(row.feature_scope, row.feature_scope)}",
    }


def _context_row(row: MiddayRecommendationResult, master: StockMaster | None) -> dict[str, Any]:
    return {
        "深度复核排名": row.pro_rank,
        "股票代码": display_stock_code(row.stock_code),
        "股票名称": row.stock_name or (master.name if master else ""),
        "入选来源": _sources(row.pool_sources),
        "一级行业": master.industry if master else "",
        "产业链": "",
        "链条位置": "",
        "主营业务": "",
        "核心产品": "",
        "概念标签": "",
        "结构性方向": FEATURE_LABELS.get(row.feature_scope, row.feature_scope),
        "潜在优势": _lines(row.key_reasons),
        "行业趋势": "",
        "核心逻辑": _lines(row.key_reasons),
        "失效条件": _lines(row.key_risks),
        "财务状态": "午间未更新",
        "财务说明": "沿用前一交易日基础资料，本次仅更新上午行情与模型复核结果。",
        "人工复核": "是" if row.requires_manual_review else "否",
        "复核优先级": "高" if (row.pro_rank or 999) <= 5 else "中",
    }


def _sources(values: list[Any]) -> str:
    return "、".join(SOURCE_LABELS.get(str(value), str(value)) for value in values)


def _lines(values: list[Any]) -> str:
    return "；".join(_humanize_note(str(value).strip()) for value in values if str(value).strip())


def _humanize_note(value: str) -> str:
    text = value
    replacements = {
        "Hard gate status: PASS": "硬性门槛通过",
        "Hard gate status PASS": "硬性门槛通过",
        "Hard gate passed": "硬性门槛通过",
        "Hard gate pass": "硬性门槛通过",
        "Positive midday overlay score": "午间增强增量为正",
        "Positive midday delta": "午间增强增量为正",
        "High intraday stability": "盘中稳定性较高",
        "Strong intraday stability": "盘中稳定性较高",
        "Good intraday stability": "盘中稳定性较好",
        "Intraday stability high": "盘中稳定性较高",
        "Strong close quality": "收盘位置较强",
        "Good close quality": "收盘位置较好",
        "High close quality": "收盘位置较强",
        "Neutral close quality": "收盘位置中性",
        "Strong morning strength": "上午强度较高",
        "Good market regime fit": "市场环境匹配度较好",
        "Strong market regime fit": "市场环境匹配度较好",
        "Favorable market regime fit": "市场环境匹配度较好",
        "Market regime fit acceptable": "市场环境匹配度尚可",
        "Low liquidity confirmation": "量能确认偏弱",
        "Below average liquidity confirmation": "量能确认低于平均水平",
        "Below-average liquidity confirmation": "量能确认低于平均水平",
        "Liquidity confirmation borderline": "量能确认接近临界值",
        "Moderate relative strength": "相对强度一般",
        "Below-average relative strength": "相对强度低于平均水平",
        "Score is moderate": "综合得分处于中等水平",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    substitutions = (
        (r"Strong morning performance \(score ([\d.]+)\)", r"上午表现较强（评分 \1）"),
        (r"Strong morning strength \(([^)]+)\)", r"上午强度较高（\1）"),
        (r"Excellent morning strength \(([^)]+)\)", r"上午强度突出（\1）"),
        (r"Morning strength is perfect at ([\d.]+)", r"上午强度评分为 \1"),
        (r"Morning strength slightly low at ([\d.]+)", r"上午强度略低（\1）"),
        (r"Strong relative strength \(([^)]+)\)", r"相对强度较高（\1）"),
        (r"Good relative strength \(score ([\d.]+)\)", r"相对强度较好（评分 \1）"),
        (r"Healthy relative strength \(([^)]+)\)", r"相对强度健康（\1）"),
        (r"Relative strength is high at ([\d.]+)", r"相对强度较高（\1）"),
        (r"Relative strength moderate at ([\d.]+)", r"相对强度一般（\1）"),
        (r"Relative strength moderate \(([^)]+)\)", r"相对强度一般（\1）"),
        (r"Relative strength neutral: ([\d.]+)", r"相对强度中性（\1）"),
        (r"Neutral relative strength \(([^)]+)\)", r"相对强度中性（\1）"),
        (r"Low relative strength: ([\d.]+)", r"相对强度偏低（\1）"),
        (r"High close quality \(score ([\d.]+)\)", r"收盘位置较强（评分 \1）"),
        (r"High close quality \(([^)]+)\)", r"收盘位置较强（\1）"),
        (r"High close quality: ([\d.]+)", r"收盘位置较强（\1）"),
        (r"Close quality score ([\d.]+)", r"收盘位置评分 \1"),
        (r"Close quality moderate at ([\d.]+)", r"收盘位置一般（\1）"),
        (r"Close quality neutral: ([\d.]+)", r"收盘位置中性（\1）"),
        (r"Close quality average \(([^)]+)\)", r"收盘位置一般（\1）"),
        (r"Low close quality: ([\d.]+)", r"收盘位置偏弱（\1）"),
        (r"Intraday stability: ([\d.]+)", r"盘中稳定性评分 \1"),
        (r"Moderate intraday stability \(score ([\d.]+)\)", r"盘中稳定性一般（评分 \1）"),
        (r"Moderate intraday stability \(([^)]+)\)", r"盘中稳定性一般（\1）"),
        (r"Intraday stability is low at ([\d.]+)", r"盘中稳定性偏低（\1）"),
        (r"Intraday stability score: ([\d.]+) \(below 50\)", r"盘中稳定性偏低（\1）"),
        (r"Liquidity confirmation moderate at ([\d.]+)", r"量能确认一般（\1）"),
        (r"Liquidity confirmation below 50: ([\d.]+)", r"量能确认偏弱（\1）"),
        (r"Liquidity confirmation is low at ([\d.]+)", r"量能确认偏弱（\1）"),
        (r"Liquidity confirmation low \(([^)]+)\)", r"量能确认偏弱（\1）"),
        (r"Liquidity confirmation: ([\d.]+)", r"量能确认评分 \1"),
        (r"Liquidity confirmation score: ([\d.]+) \(below 50\)", r"量能确认偏弱（\1）"),
        (r"Market regime fit is neutral \(score ([\d.]+)\)", r"市场环境匹配度中性（评分 \1）"),
        (r"Good market regime fit \(([^)]+)\)", r"市场环境匹配度较好（\1）"),
        (r"Strong market regime fit \(([^)]+)\)", r"市场环境匹配度较好（\1）"),
        (r"Base quant rank top ([\d]+)", r"前日量化排名位于前 \1"),
        (r"Base quant rank ([\d]+) in top 100", r"前日量化排名第 \1"),
        (r"Base quant rank: ([\d]+) \(top 100\)", r"前日量化排名第 \1"),
        (r"Base quant rank: ([\d]+)", r"前日量化排名第 \1"),
        (r"Base quant rank ([\d]+)", r"前日量化排名第 \1"),
        (r"Base quant score below 60 \(([^)]+)\)", r"前日量化得分低于60（\1）"),
        (r"Midday enhanced score: ([\d.]+)", r"午间增强分 \1"),
        (r"Midday enhanced score ([\d.]+) above threshold", r"午间增强分 \1，高于筛选阈值"),
        (r"Midday enhanced score above 60", r"午间增强分高于60"),
        (r"Midday delta small \(([^)]+)\)", r"午间增量较小（\1）"),
        (r"Negative midday delta: ([\-\d.]+)", r"午间增量为负（\1）"),
    )
    for pattern, replacement in substitutions:
        text = re.sub(pattern, replacement, text)
    text = re.sub(r"\s*\(score ([\d.]+)\)", r"（评分 \1）", text)
    text = re.sub(r"\s*\(([\d.]+)\)", r"（\1）", text)
    return text


def _polish_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A9" if worksheet.title == "今日概览" else "D5" if worksheet.title in {
            "今日推荐", "重点候选", "价格与权重", "复核依据"
        } else "A5"
        worksheet.sheet_view.showGridLines = False
        table_start_row = 8 if worksheet.title == "今日概览" else 4
        for row in worksheet.iter_rows(min_row=table_start_row, max_row=worksheet.max_row):
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    workbook.save(path)


def _number(value: Any) -> float | None:
    if value is None:
        return None
    return float(value)


def _risk_reward(entry: Decimal | None, stop: Decimal | None, target: Decimal | None) -> float | None:
    if entry is None or stop is None or target is None:
        return None
    risk = Decimal(entry) - Decimal(stop)
    if risk <= 0:
        return None
    return float((Decimal(target) - Decimal(entry)) / risk)
